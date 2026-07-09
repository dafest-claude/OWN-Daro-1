#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Rev. Alt 7 sobre la Rev. Alt 6.
Incorpora las BANDEJAS PORTACABLES Y CANALIZACIONES de la especialidad ELÉCTRICA
(faltaban por completo), tomadas del listado de materiales ACAL-00102-LM-E-0018 Rev.0
(totales consolidados de todas las áreas). Usa rendimientos análogos a los de Instrumentación.
 - Anexo E: nueva subsección E.9 con tabla de cantidades y HH.
 - C.1: estándares EL de bandeja (tramo/accesorio 4,80 HH/u; tapa 1,44 HH/u; conduit 0,30 HH/m).
 - C.2 (EL): nueva línea 'Montaje de bandejas y canalizaciones (LM-E-0018)' = 7.167 HH;
   subtotal EL 45.018 -> 52.185.
 - C.4: frente EL de montaje incluye bandejas/canalizaciones (2.840 -> 10.007 HH).
 - TOTAL E&I 87.663 -> 94.830 HH (~539 persona-mes).
 - Cuadro ANEXOS y DOCUMENTOS DE REFERENCIA: se agrega LM-E-0018.
No borra revisiones anteriores.
Base:  Análisis generados/Pliego subcontrato E&I - Rev Alt 6 - 2026-07-07 (con Anexos).docx
Salida: Análisis generados/Pliego subcontrato E&I - Rev Alt 7 - 2026-07-07 (con Anexos).docx
"""
import os
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

GEN="/home/user/OWN-Daro-1/Constructibilidad/Análisis generados"
SRC=os.path.join(GEN,"Pliego subcontrato E&I - Rev Alt 6 - 2026-07-07 (con Anexos).docx")
OUT=os.path.join(GEN,"Pliego subcontrato E&I - Rev Alt 7 - 2026-07-07 (con Anexos).docx")
doc=Document(SRC)
TBL='Tabla Cuadro'
try: doc.styles[TBL]
except KeyError: TBL='Table Grid'

def set_cell_text(cell,new):
    p=cell.paragraphs[0]
    if p.runs:
        p.runs[0].text=new
        for r in list(p.runs)[1:]: r._element.getparent().remove(r._element)
    else:
        p.add_run(new)

def insert_row_before(table,ref_idx,values,size=8.5):
    new=table.add_row()
    for j,v in enumerate(values):
        c=new.cells[j]; c.text=''
        r=c.paragraphs[0].add_run(str(v)); r.font.size=Pt(size); r.font.name='Calibri'
        if j>0 and len(str(v))<22: c.paragraphs[0].alignment=WD_ALIGN_PARAGRAPH.CENTER
    table.rows[ref_idx]._tr.addprevious(new._tr)
    trPr=new._tr.get_or_add_trPr()
    if trPr.find(qn('w:cantSplit')) is None: trPr.append(OxmlElement('w:cantSplit'))
    return new

# ---------- C.1: estándares EL de canalizaciones (bloque EL, tras 'Montaje de motor') ----------
t17=doc.tables[17]
for i,r in enumerate(t17.rows):
    if r.cells[0].text.strip()=='IN' and 'Tendido cable simple' in r.cells[1].text:
        insert_row_before(t17,i,["EL","Montaje de bandeja portacable (tramo / accesorio)","4,80","HH/unidad","2","8"])
        insert_row_before(t17,i,["EL","Montaje de tapa de bandeja","1,44","HH/unidad","2","8"])
        insert_row_before(t17,i,["EL","Montaje de conduit / caño (RMC / PVC)","0,30","HH/m","2","8"])
        break

# ---------- C.2 (EL): nueva línea de bandejas + subtotal ----------
t18=doc.tables[18]
for i,r in enumerate(t18.rows):
    if 'Pruebas, energización y precomisionado' in r.cells[0].text:
        insert_row_before(t18,i,
          ["Montaje de bandejas y canalizaciones (LM-E-0018)","827+227+1.135 u + 1.575 m",
           "4,80 / 1,44 HH/u · 0,30 HH/m","7.167"])
        break
for r in t18.rows:
    if 'SUBTOTAL ELECTRICIDAD' in r.cells[0].text:
        set_cell_text(r.cells[3],"≈ 52.185 HH")

# ---------- C.4: frente EL de montaje incluye bandejas ----------
t20=doc.tables[20]
for r in t20.rows:
    if r.cells[0].text.strip()=='EL' and r.cells[1].text.strip()=='Montaje de equipos':
        set_cell_text(r.cells[1],"Montaje de equipos + bandejas/canalizaciones")
        set_cell_text(r.cells[2],"10.007")   # 2.840 + 7.167
        set_cell_text(r.cells[3],"6")
        set_cell_text(r.cells[4],"1.668")     # 10.007 / 6
        set_cell_text(r.cells[5],"~9,5")
        break

# ---------- TOTAL E&I ----------
for p in doc.paragraphs:
    if 'TOTAL mano de obra estimada' in p.text:
        for r in p.runs:
            r.text=r.text.replace('87.663','94.830').replace('498 persona-mes','539 persona-mes')

# ---------- cuadro ANEXOS + DOCUMENTOS DE REFERENCIA (tabla 1) ----------
for t in doc.tables:
    labs=[t.rows[r].cells[0].text.strip().upper() for r in range(len(t.rows))]
    if 'ANEXOS' in labs and 'OBJETIVO' in labs:
        for r in range(len(t.rows)):
            k=t.rows[r].cells[0].text.strip().upper()
            cell=t.rows[r].cells[1]
            if k=='ANEXOS':
                for p in cell.paragraphs:
                    for run in p.runs:
                        if 'SPCDA, iluminación)' in run.text:
                            run.text=run.text.replace('SPCDA, iluminación)','SPCDA, iluminación, bandejas/canalizaciones)')
            if k=='DOCUMENTOS DE REFERENCIA':
                tgt=None
                for p in cell.paragraphs:
                    if 'LM-E-0020' in p.text: tgt=p
                if tgt is not None:
                    npar=cell.add_paragraph()
                    rr=npar.add_run("   • ACAL-00102-LM-E-0018 Rev. 0 — Canalizaciones (bandejas y conduits, montaje eléctrico).")
                    rr.font.size=Pt(8.5)
                    tgt._p.addnext(npar._p)
        break

# ---------- Anexo E: nueva subsección E.9 + tabla (AL FINAL: mueve índices de tablas) ----------
anchor=None
for p in doc.paragraphs:
    if p.text.strip().startswith('Iluminación exterior: 144 proyectores'):
        anchor=p; break
h=doc.add_heading("E.9  Bandejas portacables y canalizaciones — Electricidad (LM-E-0018)",level=2)
h.paragraph_format.keep_with_next=True
intro=doc.add_paragraph()
intro.add_run("Cantidades consolidadas (todas las áreas) del listado de materiales "
    "ACAL-00102-LM-E-0018 Rev. 0 — Canalizaciones. El montaje de bandejas y accesorios de "
    "Electricidad no estaba computado en revisiones anteriores; se incorpora con rendimientos "
    "análogos a los de Instrumentación (Tabla C.1). Las grapas de fijación y cuplas de unión se "
    "consideran incluidas en el montaje de los tramos.")
tb=doc.add_table(rows=1,cols=5); tb.style=TBL
for j,x in enumerate(["Concepto","Cantidad","Unidad","Estándar (C.1)","HH estimadas"]):
    c=tb.rows[0].cells[j]; c.text=''; rr=c.paragraphs[0].add_run(x); rr.bold=True; rr.font.size=Pt(8.5)
E9=[
 ["Bandeja recta tipo escalera (tramos 6 m; anchos 300 / 450 / 600 mm)","827","u","4,80 HH/u","3.970"],
 ["Accesorios (curvas horiz./vert., tees, reductores)","227","u","4,80 HH/u","1.090"],
 ["Tapas de bandeja","1.135","u","1,44 HH/u","1.634"],
 ["Conduit RMC (metálico) + PVC rígido","1.575","m","0,30 HH/m","473"],
 ["Fijaciones (grapas 2.255 + cuplas de unión 2.255)","4.510","u","incluidas en montaje","—"],
 ["TOTAL canalizaciones — Electricidad","—","—","—","≈ 7.167 HH"],
]
for ri,row in enumerate(E9):
    cs=tb.add_row().cells
    for j,v in enumerate(row):
        cs[j].text=''; rr=cs[j].paragraphs[0].add_run(v); rr.font.size=Pt(8.5)
        if ri==len(E9)-1: rr.bold=True
        if j>0 and len(v)<20: cs[j].paragraphs[0].alignment=WD_ALIGN_PARAGRAPH.CENTER
for row in tb.rows:
    trPr=row._tr.get_or_add_trPr()
    if trPr.find(qn('w:cantSplit')) is None: trPr.append(OxmlElement('w:cantSplit'))
anchor._p.addnext(tb._tbl); anchor._p.addnext(intro._p); anchor._p.addnext(h._p)

# ---------- tipografía Calibri en runs nuevos ----------
def force_cal(run):
    rpr=run._element.get_or_add_rPr(); rf=rpr.find(qn('w:rFonts'))
    if rf is None: rf=OxmlElement('w:rFonts'); rpr.insert(0,rf)
    for a in ('w:asciiTheme','w:hAnsiTheme','w:cstheme'):
        if rf.get(qn(a)) is not None: del rf.attrib[qn(a)]
    rf.set(qn('w:ascii'),'Calibri'); rf.set(qn('w:hAnsi'),'Calibri'); rf.set(qn('w:cs'),'Calibri')
for p in doc.paragraphs:
    for r in p.runs: force_cal(r)
for t in doc.tables:
    for row in t.rows:
        for c in row.cells:
            for p in c.paragraphs:
                for r in p.runs: force_cal(r)

doc.save(OUT)
print("OK ->",OUT,"| tablas:",len(doc.tables))
