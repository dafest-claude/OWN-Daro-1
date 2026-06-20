#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Cronograma OPTIMIZADO CPF2 La Calera II - EL + IN  (escenario con palancas de aceleración)
Palancas respecto del escenario realista conservador:
 - Entrega PROGRESIVA de cables: el tendido arranca sin esperar el total (inicio c/20%).
 - Jornada 6 días/semana (se trabaja sábados).
 - Conexionado en salas eléctricas (SE#3/SE#4) y Sala INS a DOBLE TURNO.
 - Solapamiento PARCIAL precomisionado ↔ comisionado.
 - Dimensionado por PUNTAS (EL ~5.178 · IN 15.752).
Piso duro: entregas de equipo (PCS 24-jun, skids 20-jul-2027) no se aceleran con mano de obra.
Salidas: Cronograma_Optimizado_CPF2.pdf y .xlsx
"""
import os
from datetime import date
BASE=os.path.dirname(os.path.abspath(__file__))
OUT_PDF=os.path.join(BASE,"Cronograma_Optimizado_CPF2.pdf")
OUT_XLS=os.path.join(BASE,"Cronograma_Optimizado_CPF2.xlsx")
HOY=date(2026,6,19)
RFSU_BASE=date(2028,1,31)   # contractual
RFSU_OPT =date(2028,3,31)   # optimizado estimado
RFSU_CONS=date(2028,5,31)   # conservador (escenario anterior)
T0=date(2027,1,1); T1=date(2028,8,31)
def d(s): y,m,dd=map(int,s.split('-')); return date(y,m,dd)
C_PROC='#bf9000'; C_IN='#2e75b6'; C_EL='#c55a11'; C_TEST='#548235'; C_COM='#7030a0'; C_CRIT='#c00000'
T=[
 ('H','##','SUMINISTROS / ENTREGAS (Plan 18-06-26) — cables PROGRESIVOS',None,None,None,False,''),
 ('H','S0','▼ Cables IN/EL — llegada PROGRESIVA','2027-01-20',None,C_PROC,False,'no se espera el total'),
 ('H','S2','▼ SIS (HIMA)','2027-04-18',None,C_PROC,False,''),
 ('H','S4','▼ Sala SE#3 en sitio','2027-05-24',None,C_PROC,False,'ABB'),
 ('H','S5','▼ Válvulas de control','2027-05-31',None,C_PROC,False,''),
 ('H','S6','▼ Sistema PMS (ABB)','2027-06-02',None,C_PROC,False,''),
 ('H','S7','▼ Sala SE#4 en sitio','2027-06-23',None,C_PROC,False,'ABB'),
 ('H','S8','★ Sistema de Control PCS (Inauco)','2027-06-24',None,C_CRIT,True,'ancla control'),
 ('H','S10','★ Skids de combustible','2027-07-20',None,C_CRIT,True,'PISO DURO · última entrega'),

 ('IN','##','INSTRUMENTACIÓN (IN)  ·  15.752 puntas',None,None,None,False,''),
 ('IN','I1','Canalizaciones / conduit IN','2027-01-01','2027-03-31',C_IN,False,''),
 ('IN','I2','Tendido cables IN (inicio c/20% · 6 d/sem)','2027-02-15','2027-07-15',C_IN,True,'arranque progresivo'),
 ('IN','I3','Montaje de instrumentos (AESA)','2027-06-01','2027-11-15',C_IN,False,''),
 ('IN','I4','Conexionado IN — campo + Sala INS (2 turnos)','2027-07-01','2027-11-15',C_IN,True,'Sala INS doble turno'),
 ('IN','I5','Loop check IN (PCS/SIS)','2027-10-15','2028-01-31',C_IN,True,''),

 ('EL','##','ELECTRICIDAD (EL)  ·  5.178 puntas',None,None,None,False,''),
 ('EL','E1','Bandejas / canalizaciones EL','2027-02-01','2027-04-15',C_EL,False,''),
 ('EL','E2','Tendido cables EL (inicio c/20% · 6 d/sem)','2027-03-01','2027-07-31',C_EL,True,'arranque progresivo'),
 ('EL','E3','Armado salas SE#3 / SE#4 (campo)','2027-05-24','2027-07-15',C_EL,False,'módulos'),
 ('EL','E4','Montaje tableros / celdas / CCMs','2027-07-01','2027-08-31',C_EL,False,'37 tableros'),
 ('EL','E5','Montaje motores / equipos (skids 20-jul)','2027-08-01','2027-10-15',C_EL,False,'piso skids'),
 ('EL','E6','Conexionado EL — salas SE#3/SE#4 (2 turnos)','2027-07-15','2027-10-31',C_EL,True,'doble turno'),
 ('EL','E7','Pruebas Megger EL (HOLD POINT)','2027-10-15','2027-11-15',C_EL,True,''),

 ('T','##','ENERGIZACIÓN · PRECOM · COMISIONADO (solapamiento parcial)',None,None,None,False,''),
 ('T','P1','Energización progresiva MT→BT','2027-11-16','2027-12-15',C_TEST,True,'HOLD POINT'),
 ('T','P2','Precomisionado eléctrico funcional','2027-12-01','2028-01-31',C_TEST,True,''),
 ('T','P3','Precomisionado instrumentación (PCS/SIS)','2027-12-01','2028-02-15',C_TEST,True,''),
 ('T','P4','Comisionado integrado (solapa parcial con precom)','2028-01-15','2028-03-31',C_COM,True,'overlap por subsistema'),
 ('T','P5','SAT PMS (HOLD POINT)','2028-02-01','2028-03-31',C_COM,False,'ventana'),
 ('T','RBASE','◇ RFSU base (contractual)','2028-01-31',None,C_CRIT,False,'forzado'),
 ('T','RCONS','◇ RFSU conservador (escenario previo)','2028-05-31',None,C_CRIT,False,'sin palancas'),
 ('T','RFSU','★★★ RFSU OPTIMIZADO (estimado)','2028-03-31',None,C_CRIT,True,'≈ +2 meses vs base'),
]
# ---------------- PDF ----------------
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.units import mm
from reportlab.lib import colors
from reportlab.platypus import (BaseDocTemplate,PageTemplate,Frame,Paragraph,Spacer,Table,TableStyle,LongTable,Flowable,PageBreak)
from reportlab.lib.styles import getSampleStyleSheet,ParagraphStyle
from reportlab.lib.enums import TA_CENTER
PAGE=landscape(A4); st=getSampleStyleSheet()
H1=ParagraphStyle('H1',parent=st['Title'],fontSize=15,spaceAfter=2)
SUB=ParagraphStyle('SUB',parent=st['Normal'],fontSize=9,textColor=colors.HexColor('#555'),alignment=TA_CENTER)
H2=ParagraphStyle('H2',parent=st['Heading2'],fontSize=11,textColor=colors.HexColor('#1f3b63'),spaceBefore=6,spaceAfter=4)
CELL=ParagraphStyle('CELL',parent=st['Normal'],fontSize=7.5,leading=9)
MES=['Ene','Feb','Mar','Abr','May','Jun','Jul','Ago','Sep','Oct','Nov','Dic']
def months(a,b):
    o=[];y,m=a.year,a.month
    while (y,m)<=(b.year,b.month):
        o.append(date(y,m,1)); m+=1
        if m>12:m=1;y+=1
    return o
class Gantt(Flowable):
    def __init__(s,tasks,w,h,lw=88*mm): s.t=tasks;s.w=w;s.h=h;s.lw=lw;Flowable.__init__(s)
    def wrap(s,a,b): return (s.w,s.h)
    def x(s,dt): return s.lw+(s.w-s.lw)*((dt-T0).days)/((T1-T0).days)
    def draw(s):
        c=s.canv;W=s.w;H=s.h;lw=s.lw;head=10*mm;n=len(s.t);rh=(H-head)/n;ly=None
        for mdt in months(T0,T1):
            xx=s.x(mdt); c.setStrokeColor(colors.HexColor('#e8e8e8'));c.setLineWidth(0.3);c.line(xx,0,xx,H-head)
            nxt=date(mdt.year+(mdt.month//12),(mdt.month%12)+1,1)
            c.setFillColor(colors.HexColor('#666'));c.setFont('Helvetica',5);c.drawCentredString((xx+s.x(nxt))/2,H-head+1.5,MES[mdt.month-1])
            if mdt.year!=ly:
                c.setFillColor(colors.HexColor('#1f3b63'));c.setFont('Helvetica-Bold',7);c.drawString(xx+1,H-head+5.3,str(mdt.year));ly=mdt.year
                c.setStrokeColor(colors.HexColor('#b0b0b0'));c.setLineWidth(0.6);c.line(xx,0,xx,H-head+5)
        c.setStrokeColor(colors.HexColor('#bbb'));c.setLineWidth(0.5);c.rect(lw,0,W-lw,H-head,stroke=1,fill=0)
        for dt,col,lbl in [(RFSU_BASE,'#999999','RFSU base 31-ENE-28'),(RFSU_OPT,'#c00000','RFSU optim. 31-MAR-28'),(RFSU_CONS,'#bbbbbb','RFSU consv. 31-MAY-28')]:
            xx=s.x(dt);c.setStrokeColor(colors.HexColor(col));c.setLineWidth(1);c.setDash(2,2);c.line(xx,0,xx,H-head);c.setDash()
            c.saveState();c.setFillColor(colors.HexColor(col));c.setFont('Helvetica-Bold',5.2);c.translate(xx,H-head-1);c.rotate(90);c.drawString(2,1.2,lbl);c.restoreState()
        for i,t in enumerate(s.t):
            blk,tid,name,ini,fin,color,crit,note=t; y=H-head-(i+1)*rh
            if tid=='##':
                c.setFillColor(colors.HexColor('#1f3b63'));c.rect(0,y,W,rh,stroke=0,fill=1)
                c.setFillColor(colors.white);c.setFont('Helvetica-Bold',6.4);c.drawString(2,y+rh*0.28,name);continue
            if i%2==0: c.setFillColor(colors.HexColor('#f6f6f6'));c.rect(0,y,W,rh,stroke=0,fill=1)
            c.setFillColor(colors.black);c.setFont('Helvetica',5.7);c.drawString(3,y+rh*0.28,name[:60])
            c.setFillColor(colors.HexColor('#999'));c.setFont('Helvetica',4.7);c.drawRightString(lw-2,y+rh*0.28,note[:24])
            if fin is None and ini:
                xx=s.x(d(ini));sz=rh*0.30;c.setFillColor(colors.HexColor(color))
                c.saveState();c.translate(xx,y+rh/2);c.rotate(45);c.rect(-sz,-sz,2*sz,2*sz,stroke=0,fill=1);c.restoreState()
            elif ini and fin:
                x0=s.x(d(ini));x1=s.x(d(fin));bh=rh*0.5;yy=y+(rh-bh)/2
                c.setFillColor(colors.HexColor(color));c.roundRect(x0,yy,max(x1-x0,1),bh,1,stroke=0,fill=1)
                if crit: c.setStrokeColor(colors.HexColor('#c00000'));c.setLineWidth(0.8);c.roundRect(x0,yy,max(x1-x0,1),bh,1,stroke=1,fill=0)
def footer(cv,doc):
    cv.saveState();cv.setFont('Helvetica',7);cv.setFillColor(colors.HexColor('#888'))
    cv.drawString(12*mm,7*mm,"Constructibilidad CPF2 · Cronograma OPTIMIZADO EL+IN · Rev.0 · RFSU base 31-ENE → optim. 31-MAR-28 (consv. 31-MAY)")
    cv.drawRightString(285*mm,7*mm,"Pag. %d"%doc.page);cv.restoreState()
doc=BaseDocTemplate(OUT_PDF,pagesize=PAGE,leftMargin=12*mm,rightMargin=12*mm,topMargin=12*mm,bottomMargin=12*mm)
fr=Frame(doc.leftMargin,doc.bottomMargin,doc.width,doc.height,id='f'); doc.addPageTemplates([PageTemplate(id='m',frames=[fr],onPage=footer)])
E=[]
E.append(Paragraph("Cronograma OPTIMIZADO — Instrumentación y Electricidad",H1))
E.append(Paragraph("CPF2 La Calera II · palancas: cables progresivos (inicio c/20%) · 6 d/sem · conexionado salas y Sala INS a doble turno · "
 "solapamiento parcial precom-comisionado · <b>RFSU optimizado ≈ 31-MAR-2028</b> (base 31-ENE · conservador 31-MAY)",SUB))
E.append(Spacer(1,3*mm))
E.append(Gantt(T,doc.width,doc.height-34*mm))
leg=[["■ Suministros","■ Instrumentación","■ Electricidad","■ Energiz./Precom","■ Comisionado","▭ rojo=crítico · ◇ RFSU ref"]]
tl=Table(leg,colWidths=[doc.width/6]*6)
tl.setStyle(TableStyle([('FONTSIZE',(0,0),(-1,-1),7),('TEXTCOLOR',(0,0),(0,0),colors.HexColor(C_PROC)),('TEXTCOLOR',(1,0),(1,0),colors.HexColor(C_IN)),
  ('TEXTCOLOR',(2,0),(2,0),colors.HexColor(C_EL)),('TEXTCOLOR',(3,0),(3,0),colors.HexColor(C_TEST)),('TEXTCOLOR',(4,0),(4,0),colors.HexColor(C_COM)),('TEXTCOLOR',(5,0),(5,0),colors.HexColor(C_CRIT))]))
E.append(tl); E.append(PageBreak())
E.append(Paragraph("Detalle de tareas e hitos",H2))
BL={'H':'Suministro','IN':'Instrum.','EL':'Electric.','T':'Energ/Comis'}
data=[["Bloque","ID","Tarea / Hito","Inicio","Fin","Días","Nota"]]
for t in T:
    blk,tid,name,ini,fin,color,crit,note=t
    if tid=='##': data.append(['','',name,'','','','']);continue
    if fin is None: ini_s=d(ini).strftime('%d-%m-%y');fin_s='';dur='hito'
    else: ini_s=d(ini).strftime('%d-%m-%y');fin_s=d(fin).strftime('%d-%m-%y');dur=str((d(fin)-d(ini)).days+1)
    nm=(name+("  ◄ crítico" if crit and fin else "")).replace('&','&amp;')
    data.append([BL.get(blk,blk),tid,Paragraph(nm,CELL),ini_s,fin_s,dur,note])
tb=LongTable(data,colWidths=[18*mm,12*mm,110*mm,20*mm,20*mm,14*mm,79*mm],repeatRows=1)
sty=TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#1f3b63')),('TEXTCOLOR',(0,0),(-1,0),colors.white),('FONTNAME',(0,0),(-1,0),'Helvetica-Bold'),
  ('FONTSIZE',(0,0),(-1,-1),7.5),('GRID',(0,0),(-1,-1),0.3,colors.HexColor('#ccc')),('VALIGN',(0,0),(-1,-1),'MIDDLE'),('ALIGN',(3,0),(5,-1),'CENTER'),
  ('ROWBACKGROUNDS',(0,1),(-1,-1),[colors.white,colors.HexColor('#f4f4f4')])])
ri=1
for t in T:
    if t[1]=='##': sty.add('SPAN',(2,ri),(6,ri));sty.add('BACKGROUND',(0,ri),(-1,ri),colors.HexColor('#dce6f2'));sty.add('FONTNAME',(0,ri),(-1,ri),'Helvetica-Bold');ri+=1;continue
    if t[6]: sty.add('TEXTCOLOR',(2,ri),(2,ri),colors.HexColor('#c00000'))
    ri+=1
tb.setStyle(sty);E.append(tb);E.append(PageBreak())
E.append(Paragraph("Palancas aplicadas, comparación y conclusión",H2))
def kv(title,rows):
    dd=[[title,""]]+[[Paragraph("<b>%s</b>"%k,CELL),Paragraph(v,CELL)] for k,v in rows]
    t=Table(dd,colWidths=[55*mm,218*mm])
    t.setStyle(TableStyle([('SPAN',(0,0),(1,0)),('BACKGROUND',(0,0),(-1,0),colors.HexColor('#1f3b63')),('TEXTCOLOR',(0,0),(-1,0),colors.white),
      ('FONTNAME',(0,0),(-1,0),'Helvetica-Bold'),('FONTSIZE',(0,0),(-1,0),9),('GRID',(0,0),(-1,-1),0.3,colors.HexColor('#ccc')),
      ('VALIGN',(0,0),(-1,-1),'TOP'),('ROWBACKGROUNDS',(0,1),(-1,-1),[colors.white,colors.HexColor('#f4f4f4')])]))
    return t
E.append(kv("PALANCAS DE ACELERACIÓN APLICADAS",[
  ("Cables progresivos","La entrega de cables es progresiva; el tendido arranca sin esperar el total (inicio c/20%). Tendido IN desde feb-27, EL desde mar-27 (vs abr/may en el conservador)."),
  ("6 días/semana","Se trabaja sábados → ~+20% de avance semanal en todos los frentes (durations ×0,83)."),
  ("Doble turno en conexionado","Conexionado en salas eléctricas (SE#3/SE#4) y Sala INS a doble turno → conexionado EL y la porción en Sala INS del IN se comprimen ~40%."),
  ("Solapamiento precom-comisionado","El comisionado integrado arranca al completar el precom de los primeros subsistemas (no espera el 100%)."),
  ("Dimensionado por puntas","EL ~5.178 puntas · IN 15.752 puntas como base de los conexionados."),
]))
E.append(Spacer(1,3*mm))
E.append(kv("COMPARACIÓN Y CONCLUSIÓN",[
  ("RFSU optimizado","≈ 31-MAR-2028. Recupera ~2 meses respecto del escenario conservador (31-MAY-2028)."),
  ("Brecha residual vs base","~2 meses por encima de la base contractual 31-ENE-2028."),
  ("Por qué no llega al 31-ENE","Las entregas de equipo crítico son un PISO DURO que la mano de obra no acelera: PCS 24-jun, bombas/calentadores ~29-jun y SKIDS DE COMBUSTIBLE 20-jul-2027. Su montaje→conexionado→precom empuja el comisionado de planta completa a 2028."),
  ("Palancas adicionales para cerrar la brecha","(1) adelantar OC/entrega de skids y PCS; (2) RFSU por fases (arrancar generación/área lista antes); (3) más cuadrillas / 2 turnos también en tendido; (4) comisionado por subsistemas en paralelo."),
  ("Estado","Preliminar. Sensible a la confirmación de entregas firmes y a la productividad real de doble turno."),
]))
doc.build(E)
print("OK PDF ->",OUT_PDF)
# XLSX
import openpyxl
from openpyxl.styles import Font,PatternFill
wb=openpyxl.Workbook();ws=wb.active;ws.title="Cronograma_Optimizado"
ws.append(["Bloque","ID","Tarea / Hito","Inicio","Fin","Días","Crítico","Nota"])
for c in ws[1]: c.font=Font(bold=True,color="FFFFFF");c.fill=PatternFill("solid",fgColor="1F3B63")
for t in T:
    blk,tid,name,ini,fin,color,crit,note=t
    if tid=='##':
        ws.append([name,"","","","","","",""])
        for c in ws[ws.max_row]: c.font=Font(bold=True)
        continue
    dur=((d(fin)-d(ini)).days+1) if fin else ""
    ws.append([BL.get(blk,blk),tid,name,d(ini) if ini else "",d(fin) if fin else "",dur,"SÍ" if crit else "",note])
for col,w in zip("ABCDEFGH",[12,8,54,12,12,7,8,44]): ws.column_dimensions[col].width=w
wb.save(OUT_XLS); print("OK XLSX ->",OUT_XLS)
