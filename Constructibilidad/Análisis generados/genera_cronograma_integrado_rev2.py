#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Cronograma INTEGRADO CPF-2 — Rev2 (06-08-2026).
Provisiones: PCS (Inauco) + PSS/SIS (HIMA) + Salas Eléctricas y PMS (ABB Rev4 optimizado).
Novedades de esta revisión:
  - Inauco: se adopta el CRONOGRAMA OFICIAL (Inauco 06-08-2026, branch Inauco_Proyecto):
    inicio/Recepción OC 28-JUL-2026; FAT PCS 29-MAR→22-ABR; FAT SIS 29-MAR→26-ABR;
    Pruebas PCS↔PMS (Bs.As.) 28-ABR→04-MAY; Campo 18-MAY→05-OCT; MCE+CAO →07-DIC-2027.
  - HIMA: se respeta la fecha del proveedor -> inicio 03-JUL-2026 + 25 semanas =
    entrega/recepción de tableros 25-DIC-2026; +45 d de recableado interno (→08-FEB-2027);
    queda holgado, por lo que la FAT SIS pasa a estar gobernada por Inauco (29-MAR→26-ABR).
  - FAT como prueba conjunta: FIN FAT integral cierra con el más tardío = Pruebas PCS↔PMS 04-MAY.
Genera: Gantt (PDF A3), Flujo (PDF A3) y XLSX, con Rev2 y fecha.
"""
import os
from datetime import date
from reportlab.lib.pagesizes import A3, landscape
from reportlab.lib import colors
from reportlab.pdfgen import canvas

OUT="/home/user/OWN-Daro-1/Constructibilidad/Análisis generados"
REV="Rev2"; FECHA="2026-08-06"
D=lambda y,m,d: date(y,m,d)

C_ABB=colors.HexColor('#C62828'); C_ABBP=colors.HexColor('#EF6C00'); C_INA=colors.HexColor('#1565C0')
C_HIMA=colors.HexColor('#6A1B9A'); C_EPC=colors.HexColor('#2E7D32'); GREY=colors.HexColor('#555555')

TASKS=[
 # ---- ABB SALAS (Rev4) ----
 ("ABB — Salas Eléctricas (S#3 crítica / S#4)","Ingeniería IB→ID→IC (FP Constr. S#3 31-AGO)",D(2026,5,22),D(2026,8,31),C_ABB),
 (None,"Fabricación tableros MT/BT (Brasil M&B)",D(2026,9,1),D(2026,11,2),C_ABB),
 (None,"FAT tableros/celdas (Brasil)",D(2026,11,3),D(2026,11,23),C_ABB),
 (None,"Acopio de materiales (Argentina)",D(2026,12,7),D(2027,1,29),C_ABB),
 (None,"Montaje e integración Shelters (Mendoza)",D(2027,1,18),D(2027,3,26),C_ABB),
 (None,"FAT Salas Eléctricas (S#4 / S#3)",D(2027,3,8),D(2027,4,19),C_ABB),
 (None,"Despacho a Vaca Muerta (S#4 / S#3)",D(2027,4,2),D(2027,5,14),C_ABB),
 (None,"★ Entrega S#4 (08-ABR) / ★ Entrega S#3 (14-MAY-2027)",D(2027,5,14),None,C_ABB),
 # ---- ABB PMS ----
 ("ABB — PMS (Power Management System)","Ingeniería Básica (FP IB 21-AGO)",D(2026,6,8),D(2026,8,21),C_ABBP),
 (None,"Ingeniería Detalle + Procedimientos",D(2026,8,3),D(2026,10,23),C_ABBP),
 (None,"Procura (AC800M / switches / tableros)",D(2026,5,22),D(2026,11,5),C_ABBP),
 (None,"FAT de Maqueta (prueba temprana AESA)",D(2026,9,22),D(2026,9,25),C_ABBP),
 (None,"Armado tableros + Software 800XA",D(2026,9,11),D(2026,12,17),C_ABBP),
 (None,"FAT Tableros + Sistema PMS (witness AESA)",D(2027,1,25),D(2027,1,29),C_ABBP),
 (None,"Despacho + Montaje PMS en Shelter (Mendoza)",D(2027,2,4),D(2027,3,17),C_ABBP),
 (None,"★ FAT Integral Shelter PMS / DataBook (31-MAR-2027)",D(2027,3,31),None,C_ABBP),
 # ---- INAUCO (cronograma oficial 06-08-2026) ----
 ("Inauco — PCS / SCADA / Comunicaciones","◆ Recepción OC — inicio (28-JUL-2026)",D(2026,7,28),None,C_INA),
 (None,"Ingeniería (Rev A y 0)",D(2026,8,11),D(2026,11,4),C_INA),
 (None,"◆ Aprobación procedimientos FAT (19-NOV) · HW (26-NOV)",D(2026,11,26),None,C_INA),
 (None,"Construcción de tableros PCS (Remota 7A/7B)",D(2026,8,26),D(2027,2,16),C_INA),
 (None,"Configuración PLC / SCADA",D(2026,8,26),D(2027,3,23),C_INA),
 (None,"FAT PCS + Comunicaciones (29-MAR → 22-ABR)",D(2027,3,29),D(2027,4,22),C_INA),
 (None,"FAT SIS — conjunto con HIMA (29-MAR → 26-ABR)",D(2027,3,29),D(2027,4,26),C_INA),
 (None,"Pruebas PCS↔PMS ABB (Bs.As.) — dentro del período FAT",D(2027,4,28),D(2027,5,4),C_INA),
 (None,"★ FIN FAT integral (todos los sistemas — el más tardío)",D(2027,5,4),None,C_INA),
 (None,"Configuraciones Post-FAT",D(2027,4,27),D(2027,5,17),C_INA),
 (None,"SAT + Comisionado PCS — campo",D(2027,5,18),D(2027,8,3),C_INA),
 (None,"SAT + Comisionado SCADA SIS — campo",D(2027,5,18),D(2027,10,5),C_INA),
 (None,"Prueba MCE + CAO (as-built)",D(2027,10,6),D(2027,12,7),C_INA),
 # ---- HIMA (premisa del proveedor: inicio 03-JUL + 25 sem) ----
 ("HIMA — PSS / SIS (ESD · F&G)","Fabricación tablero SIS (inicio 03-JUL, 25 sem)",D(2026,7,3),D(2026,12,25),C_HIMA),
 (None,"★ Recepción tablero HIMA en Inauco (25-DIC-2026)",D(2026,12,25),None,C_HIMA),
 (None,"Recableado interno HIMA (45 d) — habilita FAT SIS",D(2026,12,25),D(2027,2,8),C_HIMA),
 (None,"FAT SIS (conjunto con PCS/Inauco)",D(2027,3,29),D(2027,4,26),C_HIMA),
 (None,"Soporte HIMA en campo (SAT/PEM SCADA SIS)",D(2027,5,18),D(2027,10,5),C_HIMA),
 # ---- INTEGRACIÓN EPC ----
 ("Integración EPC (AESA)","Precomisionado / montaje de campo",D(2027,5,3),D(2028,1,13),C_EPC),
 (None,"iFAT — Integración Shelters + PMS + INAUCO + HIMA",D(2027,6,25),D(2027,7,14),C_EPC),
 (None,"Comisionado",D(2027,7,16),D(2028,1,31),C_EPC),
 (None,"SAT PMS en sitio (hold point)",D(2027,10,1),D(2027,12,15),C_EPC),
 (None,"★ RFSU — Ready For Start Up (31-ENE-2028)",D(2028,1,31),None,C_EPC),
]
VLINES=[
 (D(2026,12,25),"Tablero HIMA",C_HIMA),
 (D(2027,2,8),"Recableado HIMA",C_HIMA),
 (D(2027,5,4),"FIN FAT integral",C_EPC),
 (D(2027,5,14),"Entrega S#3",C_ABB),
 (D(2027,7,14),"iFAT",C_EPC),
 (D(2028,1,31),"RFSU",C_EPC),
]

def month_iter(d0,d1):
    y,m=d0.year,d0.month; out=[]
    while (y<d1.year) or (y==d1.year and m<=d1.month):
        out.append((y,m)); m+=1
        if m>12: m=1;y+=1
    return out
MES=['','Ene','Feb','Mar','Abr','May','Jun','Jul','Ago','Sep','Oct','Nov','Dic']

def build_gantt():
    W,H=landscape(A3); path=os.path.join(OUT,f"Cronograma_Integrado_CPF2_Gantt_{REV}_{FECHA}.pdf")
    c=canvas.Canvas(path,pagesize=(W,H))
    d0,d1=D(2026,5,1),D(2028,2,29); months=month_iter(d0,d1)
    x0=345; x1=W-24; top=H-96; span=(d1-d0).days
    SHORT={"ABB — Salas Eléctricas (S#3 crítica / S#4)":"ABB · Salas","ABB — PMS (Power Management System)":"ABB · PMS",
           "Inauco — PCS / SCADA / Comunicaciones":"Inauco · PCS","HIMA — PSS / SIS (ESD · F&G)":"HIMA · SIS",
           "Integración EPC (AESA)":"Integr. EPC"}
    def X(dt): return x0+(x1-x0)*((dt-d0).days/span)
    c.setFont("Helvetica-Bold",15); c.setFillColor(colors.black)
    c.drawString(24,H-34,f"CRONOGRAMA INTEGRADO — CPF-2 La Calera II   ·   {REV} · {FECHA}")
    c.setFont("Helvetica",9.3); c.setFillColor(GREY)
    c.drawString(24,H-50,"Inauco (cronograma oficial 06-08-2026, inicio 28-JUL) · Tablero HIMA inicio 03-JUL + 25 sem → 25-DIC-2026 "
                         "+ recableado 45 d → FAT SIS · FIN FAT integral 04-MAY-2027 · RFSU 31-ENE-2028")
    yr_row=top+20; c.setFont("Helvetica-Bold",8); prev=None
    for (y,m) in months:
        mx=X(D(y,m,1)); mxn=X(D(y+(m//12),(m%12)+1,1))
        c.setStrokeColor(colors.HexColor('#DDDDDD')); c.setLineWidth(0.5); c.line(mx,top+8,mx,42)
        c.setFillColor(GREY); c.setFont("Helvetica",7); c.drawCentredString((mx+mxn)/2,top+2,MES[m])
        if y!=prev: c.setFillColor(colors.black); c.setFont("Helvetica-Bold",9); c.drawString(mx+2,yr_row,str(y)); prev=y
    for dt,lab,col in VLINES:
        vx=X(dt); c.setStrokeColor(col); c.setLineWidth(1); c.setDash(2,2); c.line(vx,42,vx,top+8); c.setDash()
        c.setFillColor(col); c.setFont("Helvetica-Bold",6.5)
        c.saveState(); c.translate(vx,top+30); c.rotate(90); c.drawString(0,2,lab); c.restoreState()
    rowh=(top-52)/len(TASKS); y=top
    from math import cos,sin,pi
    def star(cx,cy,r,col):
        p=c.beginPath()
        for i in range(10):
            ang=pi/2+i*pi/5; rr=r if i%2==0 else r*0.45
            (p.moveTo if i==0 else p.lineTo)(cx+rr*cos(ang),cy+rr*sin(ang))
        p.close(); c.setFillColor(col); c.drawPath(p,fill=1,stroke=0)
    for (track,label,s,e,col) in TASKS:
        y-=rowh
        if track:
            c.setFillColor(col); c.rect(6,y,10,rowh-3,fill=1,stroke=0)
            c.setFillColor(colors.black); c.setFont("Helvetica-Bold",7.2); c.drawString(20,y+rowh/2-3,SHORT.get(track,track)[:20])
        c.setFillColor(colors.HexColor('#333333')); c.setFont("Helvetica",6.6); c.drawString(126,y+rowh/2-3,label[:54])
        c.setStrokeColor(colors.HexColor('#EEEEEE')); c.setLineWidth(0.3); c.line(x0,y,x1,y)
        if e is None: star(X(s),y+rowh/2,4.2,col)
        else:
            bx=X(s); bw=max(X(e)-X(s),2.2); c.setFillColor(col); c.roundRect(bx,y+2.5,bw,rowh-6,2,fill=1,stroke=0)
    c.setStrokeColor(colors.HexColor('#BBBBBB')); c.setLineWidth(0.6); c.rect(x0,44,x1-x0,top-44,fill=0,stroke=1)
    c.setFont("Helvetica-Bold",7); ly=30; lx=24
    for lab,col in [("ABB Salas",C_ABB),("ABB PMS",C_ABBP),("Inauco PCS",C_INA),("HIMA PSS/SIS",C_HIMA),("Integración EPC",C_EPC),("★ Hito",colors.black)]:
        c.setFillColor(col); c.rect(lx,ly,10,7,fill=1,stroke=0); c.setFillColor(colors.black); c.drawString(lx+13,ly+1,lab); lx+=95
    c.setFont("Helvetica-Oblique",6.5); c.setFillColor(GREY)
    c.drawRightString(x1,30,"Fuentes: ABB Rev4 · Inauco cronograma 06-08-2026 · HIMA 03-JUL+25 sem. Fechas de referencia, sujetas a OC y freezing points.")
    c.showPage(); c.save(); return path

def build_flow():
    W,H=A3; path=os.path.join(OUT,f"Cronograma_Integrado_CPF2_Flujo_{REV}_{FECHA}.pdf")
    c=canvas.Canvas(path,pagesize=(W,H))
    c.setFont("Helvetica-Bold",15); c.drawString(24,H-36,f"FLUJO INTEGRADO DE PROVISIONES — CPF-2 La Calera II · {REV} · {FECHA}")
    c.setFont("Helvetica",9); c.setFillColor(GREY)
    c.drawString(24,H-52,"Convergencia de PCS (Inauco), PSS/SIS (HIMA), Salas y PMS (ABB) hacia iFAT → Comisionado → RFSU")
    def box(x,y,w,h,title,lines,col):
        c.setFillColor(col); c.roundRect(x,y,w,h,5,fill=1,stroke=0); c.setFillColor(colors.white)
        c.setFont("Helvetica-Bold",8.6); c.drawCentredString(x+w/2,y+h-13,title); c.setFont("Helvetica",6.9)
        yy=y+h-26
        for ln in lines: c.drawCentredString(x+w/2,yy,ln); yy-=9.2
    def arrow(x1,y1,x2,y2,col=GREY):
        from math import atan2,cos,sin
        c.setStrokeColor(col); c.setLineWidth(1.4); c.line(x1,y1,x2,y2); a=atan2(y2-y1,x2-x1)
        for da in (2.6,-2.6): c.line(x2,y2,x2-8*cos(a+da),y2-8*sin(a+da))
    colw=(W-48-3*12)/4; xs=[24+i*(colw+12) for i in range(4)]; ytop=H-90; bh=156
    cols=[
     ("ABB — Salas Eléctricas",C_ABB,["Ingeniería IB→ID→IC","FP Constr. S#3: 31-AGO-26","Fabricación tableros (Brasil)","FAT Brasil → Acopio Arg.","Montaje Shelters (Mendoza)","FAT Salas → Despacho","★ Entrega S#3: 14-MAY-27"]),
     ("ABB — PMS",C_ABBP,["Ing. Básica (FP 21-AGO)","Procura AC800M/switches","FAT Maqueta (22-SEP)","Armado + Software 800XA","FAT PMS witness (25-ENE)","Montaje Shelter Mendoza","★ DataBook PMS: 31-MAR-27"]),
     ("Inauco — PCS/SCADA",C_INA,["Recepción OC: 28-JUL-2026","Construcción tableros PCS","Config PLC/SCADA","FAT PCS+Comms 29-MAR→22-ABR","FAT SIS (conjunto) → 26-ABR","Pruebas PCS↔PMS → 04-MAY","Campo 18-MAY → 05-OCT"]),
     ("HIMA — PSS/SIS",C_HIMA,["Tablero SIS — inicio 03-JUL","25 sem → 25-DIC-2026","Recableado interno 45 d","(25-DIC → 08-FEB-2027)","HIMA listo, espera a Inauco","FAT SIS con Inauco (MAR-ABR)","Soporte campo → 05-OCT"]),
    ]
    for i,(t,col,lines) in enumerate(cols): box(xs[i],ytop-bh,colw,bh,t,lines,col)
    ifat_y=ytop-bh-70; ifat_w=W-160; ifat_x=(W-ifat_w)/2
    for i in range(4): arrow(xs[i]+colw/2,ytop-bh,(W)/2,ifat_y+52)
    box(ifat_x,ifat_y,ifat_w,52,"iFAT — INTEGRACIÓN Shelters + PMS + INAUCO + HIMA",["25-JUN → 14-JUL-2027   (tras FIN FAT integral 04-MAY)"],C_EPC)
    camp_y=ifat_y-64; arrow(W/2,ifat_y,W/2,camp_y+44)
    box(ifat_x,camp_y,ifat_w,44,"PRECOMISIONADO + COMISIONADO EN CAMPO",["Precom: 03-MAY-27 → 13-ENE-28   ·   Comisionado: 16-JUL-27 → 31-ENE-28   ·   SAT PMS sitio: OCT→DIC-27"],C_EPC)
    rfsu_y=camp_y-58; arrow(W/2,camp_y,W/2,rfsu_y+40)
    c.setFillColor(colors.HexColor('#1B5E20')); c.roundRect(ifat_x+ifat_w/2-150,rfsu_y,300,40,6,fill=1,stroke=0)
    c.setFillColor(colors.white); c.setFont("Helvetica-Bold",13); c.drawCentredString(W/2,rfsu_y+15,"★ RFSU — 31-ENE-2028")
    c.setFont("Helvetica-Oblique",7.5); c.setFillColor(GREY)
    c.drawString(24,rfsu_y-24,"Ruta que gobierna la integración: FAT SIS/FIN FAT integral (04-MAY) y Entrega S#3 (14-MAY) → iFAT (JUL) → Comisionado → RFSU.")
    c.drawString(24,rfsu_y-36,"Premisas: Inauco cronograma oficial 06-08-2026; HIMA inicio 03-JUL + 25 sem (tablero 25-DIC) + 45 d recableado; HIMA queda holgado y la FAT SIS pasa a gobernarla Inauco; PCS↔PMS dentro del período FAT (no crítico).")
    c.showPage(); c.save(); return path

def build_xlsx():
    import openpyxl
    from openpyxl.styles import Font,PatternFill,Alignment
    path=os.path.join(OUT,f"Cronograma_Integrado_CPF2_{REV}_{FECHA}.xlsx")
    wb=openpyxl.Workbook(); ws=wb.active; ws.title="Cronograma Integrado"
    ws.append([f"CRONOGRAMA INTEGRADO CPF-2 — {REV} · {FECHA}"]); ws.merge_cells("A1:E1")
    ws["A1"].font=Font(bold=True,size=12,color="FFFFFF"); ws["A1"].fill=PatternFill("solid",fgColor="1F3864")
    hdr=["Proveedor","Actividad / Hito","Inicio","Fin","Tipo"]; ws.append(hdr)
    for j in range(1,6): c=ws.cell(2,j); c.font=Font(bold=True,color="FFFFFF"); c.fill=PatternFill("solid",fgColor="37474F"); c.alignment=Alignment(horizontal="center")
    hexmap={C_ABB:"C62828",C_ABBP:"EF6C00",C_INA:"1565C0",C_HIMA:"6A1B9A",C_EPC:"2E7D32"}
    cur=""
    for (track,label,s,e,col) in TASKS:
        if track: cur=track
        ws.append([cur,label,s.strftime("%d-%m-%Y"),(e.strftime("%d-%m-%Y") if e else "—"),("Hito" if e is None else "Tarea")])
        r=ws.max_row; ws.cell(r,1).fill=PatternFill("solid",fgColor=hexmap.get(col,"999999")); ws.cell(r,1).font=Font(color="FFFFFF",bold=True,size=8)
        if e is None:
            for j in range(2,5): ws.cell(r,j).font=Font(bold=True,size=9)
    for j,w in zip("ABCDE",[34,54,12,12,8]): ws.column_dimensions[j].width=w
    ws.freeze_panes="A3"; wb.save(path); return path

g=build_gantt(); f=build_flow(); x=build_xlsx()
print("OK\n ",g,"\n ",f,"\n ",x)
