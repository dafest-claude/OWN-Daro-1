#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Rev. Alt 4 sobre la Rev. Alt 3. Aplica las sugerencias de la revisión final (1 a 4):
 (1) Corrige el typo de PALABRAS CLAVE: 'Eléctridad &Instrumentos' -> 'Electricidad & Instrumentos'.
 (2) Agrega nota en C.3: los soportes de JB (~80) y sus HH son indicativos (LM-K-0002 en revisión).
 (3) Desdobla el estándar de soporte de JB en dos (sobre estructura existente ~2 HH / pedestal
     anclado a piso ~4,5 HH) y lo pondera en C.3 (48 estructura + 32 pedestal = 240 HH).
 (4) Abre un renglón propio de 'Precomisionado y comisionado E&I' en C.4 (dotación).
No borra revisiones anteriores.
Base: Análisis generados/Pliego subcontrato E&I - Rev Alt 3 - 2026-07-07 (con Anexos).docx
Salida: Análisis generados/Pliego subcontrato E&I - Rev Alt 4 - 2026-07-07 (con Anexos).docx
"""
import os
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

GEN="/home/user/OWN-Daro-1/Constructibilidad/Análisis generados"
SRC=os.path.join(GEN,"Pliego subcontrato E&I - Rev Alt 3 - 2026-07-07 (con Anexos).docx")
OUT=os.path.join(GEN,"Pliego subcontrato E&I - Rev Alt 4 - 2026-07-07 (con Anexos).docx")
doc=Document(SRC)

def set_cell_text(cell,new):
    p=cell.paragraphs[0]
    if p.runs:
        p.runs[0].text=new
        for r in list(p.runs)[1:]: r._element.getparent().remove(r._element)
    else:
        p.add_run(new)

def insert_row_after(table,ref_idx,values,size=8.5):
    new=table.add_row()
    for j,v in enumerate(values):
        c=new.cells[j]; c.text=''
        r=c.paragraphs[0].add_run(str(v)); r.font.size=Pt(size); r.font.name='Calibri'
        if j>0 and len(str(v))<22: c.paragraphs[0].alignment=WD_ALIGN_PARAGRAPH.CENTER
    table.rows[ref_idx]._tr.addnext(new._tr)
    trPr=new._tr.get_or_add_trPr()
    if trPr.find(qn('w:cantSplit')) is None: trPr.append(OxmlElement('w:cantSplit'))
    return new

# ---------- (1) typo PALABRAS CLAVE ----------
for t in doc.tables:
    for row in t.rows:
        for c in row.cells:
            if 'Eléctridad' in c.text:
                for p in c.paragraphs:
                    for r in p.runs:
                        if 'Eléctridad' in r.text:
                            r.text=r.text.replace('Eléctridad &Instrumentos','Electricidad & Instrumentos')\
                                         .replace('Eléctridad','Electricidad')

# ---------- (3) C.1 (Tabla 17): desdoblar estándar de soporte de JB ----------
t17=doc.tables[17]
for i,r in enumerate(t17.rows):
    if 'soporte' in r.cells[1].text.lower() and 'JB' in r.cells[1].text:
        # convertir la fila existente en 'sobre estructura existente' (2 HH)
        set_cell_text(r.cells[1],"Montaje de soporte de JB sobre estructura existente")
        set_cell_text(r.cells[2],"2")
        # insertar debajo la fila de 'pedestal anclado a piso' (4,5 HH)
        insert_row_after(t17,i,
            ["IN","Montaje de pedestal de JB anclado a piso","4,5","HH/soporte","2","8"])
        break

# ---------- (3) C.3 (Tabla 19): ponderar 48 estructura + 32 pedestal = 240 HH ----------
t19=doc.tables[19]
for r in t19.rows:
    if 'soporte' in r.cells[0].text.lower():
        set_cell_text(r.cells[1],"~80 (48 estr. + 32 ped.)")
        set_cell_text(r.cells[2],"2 y 4,5 HH/soporte")
        set_cell_text(r.cells[3],"240")   # 48·2 + 32·4,5 = 240 (subtotal sin cambios)
        break

# ---------- (4) C.4 (Tabla 20): renglón propio de precom/comisionado E&I ----------
# Fila 'EL Montaje + pruebas + precomisionado' -> 'EL Montaje de equipos' (se extrae precom)
# Fila 'IN Loop check / pruebas'               -> 'E&I Precomisionado y comisionado' (pruebas EL + loop IN)
t20=doc.tables[20]
for r in t20.rows:
    txt=r.cells[1].text
    if 'Montaje + pruebas + precomisionado' in txt:
        set_cell_text(r.cells[1],"Montaje de equipos")
        set_cell_text(r.cells[2],"2.840")   # sin las pruebas/precom (3.677)
        set_cell_text(r.cells[4],"710")      # 2.840 / 4
        set_cell_text(r.cells[5],"~4")
    elif 'Loop check / pruebas' in txt:
        set_cell_text(r.cells[0],"E&I")
        set_cell_text(r.cells[1],"Precomisionado y comisionado (pruebas EL + loop check IN)")
        set_cell_text(r.cells[2],"5.447")    # 3.677 (EL) + 1.770 (IN)
        set_cell_text(r.cells[3],"5")
        set_cell_text(r.cells[4],"1.089")    # 5.447 / 5
        set_cell_text(r.cells[5],"~6")

# ---------- (2) nota en C.3 sobre indicatividad de soportes/JB ----------
def add_note_like(anchor_text,note_text):
    anchor=None
    for p in doc.paragraphs:
        if anchor_text in p.text: anchor=p; break
    if anchor is None: return
    new_p=doc.add_paragraph(); new_p.style=anchor.style
    r=new_p.add_run(note_text); r.italic=True; r.font.size=Pt(9); r.font.color.rgb=RGBColor(0x55,0x55,0x55)
    r.font.name='Calibri'
    anchor._p.addnext(new_p._p)   # justo después de la nota del TOTAL
add_note_like("TOTAL mano de obra estimada E&I",
 "Los soportes de cajas de conexión (JB) —~80, y las HH asociadas— son indicativos: el cómputo de JB "
 "se ajustará al cerrarse el LM-K-0002 (a la fecha incluye solo canalizaciones troncales). El desdoblado "
 "estructura/pedestal (48/32) es una referencia; se recalcula con las cantidades definitivas y los estándares de C.1.")

# ---------- tipografía Calibri en runs nuevos ----------
def force_cal(run):
    rpr=run._element.get_or_add_rPr(); rf=rpr.find(qn('w:rFonts'))
    if rf is None: rf=OxmlElement('w:rFonts'); rpr.insert(0,rf)
    for a in ('w:asciiTheme','w:hAnsiTheme','w:cstheme'):
        if rf.get(qn(a)) is not None: del rf.attrib[qn(a)]
    rf.set(qn('w:ascii'),'Calibri'); rf.set(qn('w:hAnsi'),'Calibri'); rf.set(qn('w:cs'),'Calibri')
for p in doc.paragraphs:
    for r in p.runs: force_cal(r)
for t in doc.tables:
    for row in t.rows:
        for c in row.cells:
            for p in c.paragraphs:
                for r in p.runs: force_cal(r)

doc.save(OUT)
print("OK ->",OUT,"| tablas:",len(doc.tables))
