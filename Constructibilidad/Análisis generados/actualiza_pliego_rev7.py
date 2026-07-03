#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Rev.7 (FINAL) sobre la Rev.6. Preserva todo el documento y:
- Completa la celda ANEXOS de la tabla 'RESUMEN DEL DOCUMENTO' (primera hoja) con los anexos generados.
- Agrega una fila 'DOCUMENTOS DE REFERENCIA' con las listas de cables, instrumentos y materiales usados.
- Actualiza la revisión (6 -> 7).
Base: Análisis generados/Pliego subcontrato E&I - Rev6 - 2026-07-02 (con Anexos).docx
Salida: Análisis generados/Pliego subcontrato E&I - Rev7 - 2026-07-02 (con Anexos).docx
"""
import os
from docx import Document
from docx.shared import Pt

BASE="/home/user/OWN-Daro-1/Constructibilidad/Análisis generados"
SRC=os.path.join(BASE,"Pliego subcontrato E&I - Rev6 - 2026-07-02 (con Anexos).docx")
OUT=os.path.join(BASE,"Pliego subcontrato E&I - Rev7 - 2026-07-02 (con Anexos).docx")
doc=Document(SRC)

# ---- 1) revisión 6 -> 7 ----
for p in doc.paragraphs:
    if 'Revisión 6' in p.text or 'Rev. 6' in p.text or 'Rev.6' in p.text:
        for r in p.runs:
            r.text=r.text.replace('Revisión 6','Revisión 7').replace('Rev. 6','Rev. 7').replace('Rev.6','Rev.7')

def set_cell(cell, lines, bold=False, size=9):
    """Reemplaza el contenido de la celda por 'lines' (lista de renglones)."""
    cell.text=''
    p0=cell.paragraphs[0]
    for i,ln in enumerate(lines):
        p=p0 if i==0 else cell.add_paragraph()
        r=p.add_run(ln); r.font.size=Pt(size); r.bold=bold

# ---- 2) ubicar la tabla 'RESUMEN' (tiene fila ANEXOS) ----
resumen=None
for t in doc.tables:
    labels=[t.rows[r].cells[0].text.strip().upper() for r in range(len(t.rows))]
    if 'ANEXOS' in labels and 'OBJETIVO' in labels:
        resumen=t; break
assert resumen is not None, "No se encontró la tabla RESUMEN con fila ANEXOS"

# ---- 3) completar celda ANEXOS ----
anexos=[
 "Anexo E — Electricidad (cables, conexionado, tableros, PAT, SPCDA, iluminación y personal).",
 "Anexo I — Instrumentación (cables, conexionado, instrumentos, tableros Sala INS, cajas de conexión, bandejas/conduits y personal).",
 "Anexo P — Precomisionado, Comisionado y Puesta en Marcha (instrumental y misceláneos).",
 "Anexo C — Planilla de cómputo y precios unitarios.",
]
for r in range(len(resumen.rows)):
    if resumen.rows[r].cells[0].text.strip().upper()=='ANEXOS':
        set_cell(resumen.rows[r].cells[1], anexos, size=9); break

# ---- 4) agregar fila DOCUMENTOS DE REFERENCIA ----
docs=[
 "Listas de instrumentos:",
 "   • ACAL-00102-LI-K-0001 Rev. A — Lista de Instrumentos (Procesos).",
 "   • ACAL-00670-LI-K-0001 Rev. B — Lista de Instrumentos (Generación).",
 "Listas de cables:",
 "   • ACAL-00102-LC-E-0001 — Lista de Cables Eléctricos (unifilar ACAL-00100-EE-E-0001).",
 "   • ACAL-00102-LC-K-0001 / ACAL-00670-LC-K-0001 — Lista de Cables de Instrumentación.",
 "Listados de materiales:",
 "   • ACAL-00102-LM-E-0017 Rev. 0 — Puesta a Tierra y SPCDA (underground).",
 "   • ACAL-102-LM-E-0020 Rev. 1 — Iluminación y Tomas.",
 "   • ACAL-00102-LM-K-0002 Rev. 0 — Bandejas y caños/conduits (montaje eléctrico de instrumentos).",
 "Análisis y documentos de soporte:",
 "   • Análisis Eléctrico CPF2 Campo Rev. 4 (cables, puntas, cargas).",
 "   • Análisis de Cables de Instrumentación Rev. 0 (formaciones, puntas).",
 "   • CPF2 — Arquitectura de Ampliación Rev. 3 (PCS, SIS, PMS, comunicaciones OT).",
]
row=resumen.add_row()
set_cell(row.cells[0], ["DOCUMENTOS DE REFERENCIA"], bold=True, size=9)
set_cell(row.cells[1], docs, size=8.5)

doc.save(OUT)
print("OK ->",OUT)
print("filas tabla resumen:",len(resumen.rows))
