#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
CRONOGRAMA EJECUTIVO — Provisiones críticas CPF-2 (Rev0, 2026-09-02).
Para seguimiento del cliente (PPSA): hitos ejecutivos por provisión, con corte en la
ENTREGA DE EQUIPOS (todo lo posterior al despacho/entrega queda fuera).
Fuentes (últimas revisiones cargadas):
  - ABB Salas SE#3/SE#4: 5155-00-0009-VE-CO-PG-001_A (21-08-2026)
  - ABB PMS:             5155-00-0007-VE-CO-PG-001_REV B (.mpp)
  - Inauco:              5155-00-0029-VE-CO-PG-001-A (Rev A, 24-08-2026)
  - HIMA marshalling:    103618-Schedule_16-August-26_draft
Genera Excel con: Resumen Ejecutivo · Cronograma Ejecutivo (hitos + seguimiento) ·
Timeline (Gantt de hitos) · Notas y criterios.
"""
import os
from datetime import date
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

GEN="/home/user/OWN-Daro-1/Constructibilidad/Análisis generados"
REV="Rev0"; FECHA="2026-09-02"; HOY=date(2026,9,2)
D=lambda y,m,d: date(y,m,d)

# (código, hito, inicio, fin, crítico, observación)
PROV=[
 ("ABB — Sala Eléctrica SE#3 (BT)","ABB","C62828",[
  ("H01","Inicio de contrato — Recepción de OC",D(2026,5,8),D(2026,5,8),0,"OC recibida en ABB"),
  ("H02","Kick-off Meeting (KOM) técnico",D(2026,5,15),D(2026,5,18),0,""),
  ("H03","Emisión de ingeniería básica",D(2026,5,20),D(2026,8,28),0,""),
  ("H04","Aprobación de ingeniería básica",D(2026,8,28),D(2026,9,11),0,"Aprobación AESA"),
  ("H05","Emisión de ingeniería de detalle",D(2026,5,22),D(2026,9,21),0,""),
  ("H06","Aprobación ing. de detalle — FREEZING POINT",D(2026,9,21),D(2026,10,5),1,"Cierra ingeniería del shelter"),
  ("H07","Inicio compra y acopio de materiales",D(2026,5,19),D(2026,5,19),0,""),
  ("H08","Acopio de materiales completo",D(2026,12,10),D(2026,12,10),1,"Materiales críticos en fábrica"),
  ("H09","Inicio de fabricación / armado",D(2026,11,2),D(2026,11,2),0,""),
  ("H10","Fin de fabricación de tableros",D(2027,2,26),D(2027,2,26),1,""),
  ("H11","Montaje e integración de equipos en sala",D(2027,2,19),D(2027,3,15),0,"Shelterista"),
  ("H12","Interconexión de equipos",D(2027,3,16),D(2027,3,23),0,""),
  ("H13","FAT Sala Eléctrica",D(2027,3,25),D(2027,3,30),1,"Witness AESA"),
  ("H14","Ensayos PMS en sala (contrato PMS)",D(2027,3,31),D(2027,4,6),0,"Interfase con contrato PMS"),
  ("H15","Despacho (entrega EXW)",D(2027,4,21),D(2027,4,23),1,""),
  ("H16","ENTREGA DE EQUIPOS EN SITIO",D(2027,6,8),D(2027,6,10),1,"HITO FINAL — cierre del alcance"),
 ]),
 ("ABB — Sala Eléctrica SE#4 (MT)","ABB","AD1457",[
  ("H01","Inicio de contrato — Recepción de OC",D(2026,5,8),D(2026,5,8),0,"OC recibida en ABB"),
  ("H02","Kick-off Meeting (KOM) técnico",D(2026,5,15),D(2026,5,18),0,""),
  ("H03","Emisión de ingeniería básica",D(2026,5,20),D(2026,8,28),0,""),
  ("H04","Aprobación de ingeniería básica",D(2026,8,28),D(2026,9,11),0,"Aprobación AESA"),
  ("H05","Emisión de ingeniería de detalle",D(2026,5,22),D(2026,9,21),0,""),
  ("H06","Aprobación ing. de detalle — FREEZING POINT",D(2026,9,21),D(2026,10,5),1,"Cierra ingeniería del shelter"),
  ("H07","Inicio compra y acopio de materiales",D(2026,5,19),D(2026,5,19),0,""),
  ("H08","Acopio de materiales completo",D(2026,12,10),D(2026,12,10),1,""),
  ("H09","Inicio de fabricación / armado",D(2026,10,5),D(2026,10,5),0,""),
  ("H10","Fin de fabricación de tableros",D(2027,1,21),D(2027,1,21),1,"Incluye celdas MT"),
  ("H11","Montaje e integración de equipos en sala",D(2027,1,21),D(2027,2,22),0,"Shelterista"),
  ("H12","Interconexión de equipos",D(2027,2,11),D(2027,2,23),0,""),
  ("H13","FAT Sala Eléctrica",D(2027,2,24),D(2027,3,1),1,"Witness AESA"),
  ("H14","Ensayos PMS en sala (contrato PMS)",D(2027,3,2),D(2027,3,9),0,"Interfase con contrato PMS"),
  ("H15","Despacho (entrega EXW)",D(2027,3,19),D(2027,3,25),1,""),
  ("H16","ENTREGA DE EQUIPOS EN SITIO",D(2027,6,8),D(2027,6,10),1,"HITO FINAL — cierre del alcance"),
 ]),
 ("ABB — PMS (Power Management System)","ABB","EF6C00",[
  ("H01","Inicio de contrato — Recepción / aceptación OC",D(2026,5,8),D(2026,5,11),0,""),
  ("H02","Kick-off Meeting (KOM)",D(2026,5,22),D(2026,5,22),0,""),
  ("H03","Emisión de ingeniería básica",D(2026,6,8),D(2026,7,31),0,"Hito cert. 10%"),
  ("H04","Aprobación de ingeniería básica",D(2026,7,6),D(2026,8,14),0,"Hito cert. 10%"),
  ("H05","Emisión de ingeniería Rev.0",D(2026,7,20),D(2026,8,28),0,""),
  ("H06","Aprobación ingeniería Rev.0 — FREEZING POINT",D(2026,8,3),D(2026,9,11),1,"Cierra ingeniería"),
  ("H07","Inicio compra de materiales",D(2026,5,22),D(2026,5,22),0,"Switches, S800, tableros"),
  ("H08","Acopio de materiales completo",D(2026,11,5),D(2026,11,5),1,"Hito cert. 25%"),
  ("H09","Inicio de armado de tableros",D(2026,9,11),D(2026,9,11),0,""),
  ("H10","Fin de armado de tableros",D(2026,12,17),D(2026,12,17),1,""),
  ("H11","Configuración / programación del sistema",D(2026,8,17),D(2026,12,11),1,"I/O, comms IEC61850/PROFI/MODBUS, lógicas"),
  ("H12","Pruebas de Maqueta Integrada PMS/CCM/PCS",D(2026,9,22),D(2026,9,25),0,"Prueba temprana de integración"),
  ("H13","Aprobación procedimientos FAT Rev.0",D(2026,10,12),D(2026,10,23),0,""),
  ("H14","Pruebas internas (pre-FAT)",D(2026,12,18),D(2027,1,7),0,""),
  ("H15","FAT Tableros + Sistema",D(2027,1,25),D(2027,1,29),1,"Hito cert. 20% · witness AESA"),
  ("H16","Ajustes post-FAT / liberación",D(2027,2,1),D(2027,2,3),0,""),
  ("H17","Despachos",D(2027,2,4),D(2027,2,17),1,""),
  ("H18","ENTREGA DE EQUIPAMIENTO EN SHELTER",D(2027,2,17),D(2027,2,17),1,"HITO FINAL — cert. 30%"),
 ]),
 ("Inauco — PCS / SCADA / Comunicaciones","Inauco","1565C0",[
  ("H01","Inicio de contrato — Recepción de OC",D(2026,7,28),D(2026,7,28),0,""),
  ("H02","Recepción de ing. Rev.0 p/ tableros (insumo AESA)",D(2026,8,10),D(2026,8,10),1,"Requisito 1 — insumo de AESA"),
  ("H03","Emisión de ingeniería (Rev. A y 0)",D(2026,8,11),D(2026,11,4),0,""),
  ("H04","Aprobación ing. p/ construcción — FREEZING POINT",D(2026,11,19),D(2026,11,19),1,"Requisito 3"),
  ("H05","Recepción de hardware (acopio)",D(2026,11,26),D(2026,11,26),1,"Requisito 6 — Rockwell/Stratix"),
  ("H06","Inicio de construcción de tableros",D(2026,11,27),D(2026,11,27),0,""),
  ("H07","Fin de construcción de tableros",D(2027,2,16),D(2027,2,16),1,""),
  ("H08","Configuración PLC — Etapa 1",D(2026,8,26),D(2026,9,22),0,""),
  ("H09","Recepción tableros HIMA en base Inauco (insumo)",D(2027,1,25),D(2027,1,25),1,"Requisito 2 — interfase HIMA"),
  ("H10","Configuración PLC — Etapa 2",D(2027,1,25),D(2027,3,23),0,""),
  ("H11","Configuración SCADA",D(2026,8,26),D(2027,3,23),1,""),
  ("H12","Aprobación de procedimientos FAT",D(2027,2,24),D(2027,2,24),0,"Requisito 4"),
  ("H13","FAT PCS + Comunicaciones",D(2027,3,29),D(2027,4,22),1,"Witness AESA"),
  ("H14","FAT SIS (conjunto con HIMA)",D(2027,3,29),D(2027,4,26),1,"Requiere presencia HIMA"),
  ("H15","Pruebas de integración PCS↔PMS (Bs.As.)",D(2027,4,28),D(2027,5,4),1,"Interfase con ABB PMS"),
  ("H16","Configuraciones post-FAT",D(2027,4,27),D(2027,5,17),0,""),
  ("H17","ENTREGA DE EQUIPOS EN SITIO",D(2027,5,17),D(2027,5,18),1,"HITO FINAL — habilita inicio de campo"),
 ]),
 ("HIMA — Tableros de marshalling (ESD/F&G/PSS)","HIMA","6A1B9A",[
  ("H01","Inicio de contrato / proyecto",D(2026,7,10),D(2026,7,10),0,""),
  ("H02","Inicio de ingeniería",D(2026,8,14),D(2026,8,14),0,""),
  ("H03","Ingeniería de detalle / especificación",D(2026,8,17),D(2026,9,11),0,""),
  ("H04","Fin de ingeniería / CAD — FREEZING POINT",D(2026,9,18),D(2026,9,18),1,"Cierra ingeniería"),
  ("H05","Fabricación de tableros de marshalling",D(2026,9,21),D(2026,11,6),1,""),
  ("H06","Hardware fabricado y probado (test de fábrica)",D(2026,11,6),D(2026,11,6),1,"FAT integrada a fabricación"),
  ("H07","Preparación de despacho (packing)",D(2026,11,9),D(2026,11,20),0,""),
  ("H08","Listo para despacho",D(2026,11,20),D(2026,11,20),1,""),
  ("H09","Despacho / shipping",D(2026,11,23),D(2026,12,14),1,""),
  ("H10","Documentación de entrega enviada",D(2026,12,4),D(2026,12,4),0,""),
  ("H11","Entrega DAP Ezeiza (Argentina)",D(2026,12,15),D(2026,12,15),1,""),
  ("H12","RECEPCIÓN DE TABLEROS EN BASE INAUCO",D(2027,1,25),D(2027,1,25),1,"HITO FINAL — nacionalización + transporte"),
 ]),
]
NOTAS=[
 "Alcance: cronograma EJECUTIVO por hitos, con corte en la ENTREGA DE EQUIPOS de cada provisión. Las tareas posteriores al despacho/entrega (montaje en planta, SAT, precomisionado, comisionado, PEM, CAO/DataBook) quedan FUERA de este cronograma.",
 "Fuentes (últimas revisiones cargadas): ABB Salas 5155-00-0009-VE-CO-PG-001_A (21-08-26) · ABB PMS 5155-00-0007-VE-CO-PG-001_REV B (.mpp) · Inauco 5155-00-0029-VE-CO-PG-001-A Rev.A (24-08-26) · HIMA 103618-Schedule 16-08-26 (draft).",
 "Estado (plan): se calcula contra la fecha de corte. 'Cumplido s/plan' = la fecha planificada ya venció y debe confirmarse el cumplimiento real con el proveedor; 'En curso' = el hito está en ejecución; 'Pendiente' = aún no inició.",
 "Las columnas Fecha real / % avance / Estado real / Observaciones quedan en blanco para la carga de seguimiento periódico (reporte a PPSA).",
 "Hitos marcados como CRÍTICOS: condicionan la fecha de entrega de la provisión o constituyen interfases con otras provisiones.",
 "Interfases clave entre provisiones: HIMA → recepción de tableros en Inauco (25-01-27) habilita la FAT SIS; ABB PMS → ensayos PMS en Salas y pruebas de integración PCS↔PMS con Inauco.",
 "Inauco: el cronograma no explicita un hito de despacho; la entrega de equipos se deriva del cierre de configuraciones post-FAT y el inicio de tareas de campo (18-05-27). A confirmar con el proveedor.",
 "Salas SE#3/SE#4: la entrega en sitio (08/10-06-27) corresponde al shelter completo; el despacho EXW desde el shelterista es previo (25-03-27 SE#4 / 23-04-27 SE#3).",
]

NAVY="1F3864"; HEAD="37474F"
thin=Side(style="thin",color="D9D9D9"); B=Border(thin,thin,thin,thin)
EST_FILL={"Cumplido s/plan":"2E7D32","En curso":"EF6C00","Pendiente":"90A4AE"}
def estado(ini,fin):
    if fin < HOY: return "Cumplido s/plan"
    if ini <= HOY <= fin: return "En curso"
    return "Pendiente"
def cell(ws,r,c,v,fill=None,color="000000",bold=False,size=9,center=False,wrap=False,fmt=None):
    cc=ws.cell(r,c,v); cc.font=Font(bold=bold,color=color,size=size)
    if fill: cc.fill=PatternFill("solid",fgColor=fill)
    cc.alignment=Alignment(horizontal="center" if center else "left",vertical="center",wrap_text=wrap)
    cc.border=B
    if fmt: cc.number_format=fmt
    return cc

# meses de la línea de tiempo
def month_list(a,b):
    out=[];y,m=a.year,a.month
    while (y<b.year) or (y==b.year and m<=b.month):
        out.append((y,m)); m+=1
        if m>12: m=1;y+=1
    return out
MESES=month_list(D(2026,5,1),D(2027,6,1))
M3=['','Ene','Feb','Mar','Abr','May','Jun','Jul','Ago','Sep','Oct','Nov','Dic']

wb=openpyxl.Workbook()

# ---------------- Hoja 1: Resumen Ejecutivo ----------------
ws=wb.active; ws.title="Resumen Ejecutivo"
ws.merge_cells("A1:K1"); cell(ws,1,1,f"CRONOGRAMA EJECUTIVO — PROVISIONES CRÍTICAS CPF-2  ·  {REV} · {FECHA}",NAVY,"FFFFFF",True,13,True); ws.row_dimensions[1].height=24
ws.merge_cells("A2:K2"); cell(ws,2,1,"Seguimiento para PPSA · Corte de alcance: hasta la ENTREGA DE EQUIPOS (excluye montaje en planta, SAT, precom/comisionado y PEM)",None,"555555",False,9)
ws.merge_cells("A3:K3"); cell(ws,3,1,f"Fecha de corte del estado: {HOY.strftime('%d/%m/%Y')}",None,"C62828",True,9)
hdr=["Proveedor","Provisión","Inicio contrato","Entrega de equipos","Hitos","Cumplidos s/plan","En curso","Pendientes","% avance plan","Próximo hito","Fecha próx. hito"]
for j,h in enumerate(hdr): cell(ws,5,1+j,h,HEAD,"FFFFFF",True,9,True,True)
ws.row_dimensions[5].height=30
r=6
for (name,ven,col,ms) in PROV:
    tot=len(ms); cum=sum(1 for m in ms if estado(m[2],m[3])=="Cumplido s/plan")
    enc=sum(1 for m in ms if estado(m[2],m[3])=="En curso"); pen=tot-cum-enc
    fin_ms=ms[-1]; prox=next((m for m in ms if m[3]>=HOY), None)
    cell(ws,r,1,ven,col,"FFFFFF",True,9,True)
    cell(ws,r,2,name,None,"000000",True,9,False,True)
    cell(ws,r,3,ms[0][2].strftime("%d/%m/%Y"),None,"000000",False,9,True)
    cell(ws,r,4,fin_ms[3].strftime("%d/%m/%Y"),"1F3864","FFFFFF",True,9,True)
    cell(ws,r,5,tot,None,"000000",False,9,True)
    cell(ws,r,6,cum,"2E7D32","FFFFFF",True,9,True)
    cell(ws,r,7,enc,"EF6C00","FFFFFF",True,9,True)
    cell(ws,r,8,pen,"90A4AE","FFFFFF",True,9,True)
    cell(ws,r,9,f"{100*cum/tot:.0f}%",None,"000000",True,10,True)
    cell(ws,r,10,(prox[1] if prox else "—"),None,"000000",False,8.5,False,True)
    cell(ws,r,11,(prox[3].strftime("%d/%m/%Y") if prox else "—"),None,"C62828",True,9,True)
    ws.row_dimensions[r].height=28; r+=1
for c,w in zip("ABCDEFGHIJK",[11,38,14,17,7,15,10,12,13,42,15]): ws.column_dimensions[c].width=w
r+=1
cell(ws,r,1,"Interfases clave entre provisiones",None,NAVY,True,11); r+=1
for t in ["HIMA → recepción de tableros en base Inauco (25/01/27): habilita la configuración PLC Etapa 2 y la FAT SIS.",
          "ABB PMS → entrega de equipamiento en shelter (17/02/27): habilita el montaje del PMS en las salas y los ensayos PMS.",
          "ABB PMS ↔ Inauco → pruebas de integración PCS↔PMS en Bs.As. (28/04 → 04/05/27).",
          "AESA → entrega de ingeniería Rev.0 a Inauco (10/08/26) y aprobaciones de ingeniería a ABB/HIMA: condicionan los freezing points."]:
    cell(ws,r,1,"• "+t,None,"000000",False,9,False,True); ws.merge_cells(start_row=r,start_column=1,end_row=r,end_column=11); ws.row_dimensions[r].height=16; r+=1

# ---------------- Hoja 2: Cronograma Ejecutivo ----------------
cw=wb.create_sheet("Cronograma Ejecutivo")
cw.merge_cells("A1:K1"); cell(cw,1,1,f"CRONOGRAMA EJECUTIVO POR HITOS  ·  {REV} · {FECHA}",NAVY,"FFFFFF",True,12,True); cw.row_dimensions[1].height=22
cw.merge_cells("A2:K2"); cell(cw,2,1,"Columnas de seguimiento (Fecha real / % avance / Estado real / Observaciones) para carga periódica.",None,"555555",False,9)
H2=["Nº","Hito ejecutivo","Inicio","Fin","Crítico","Estado (plan)","Fecha real","% avance","Estado real","Observaciones"]
r=4
for (name,ven,col,ms) in PROV:
    cw.merge_cells(start_row=r,start_column=1,end_row=r,end_column=10)
    cell(cw,r,1,f"▶  {name}   ·   {ven}   ·   Entrega de equipos: {ms[-1][3].strftime('%d/%m/%Y')}",col,"FFFFFF",True,10)
    cw.row_dimensions[r].height=18; r+=1
    for j,h in enumerate(H2): cell(cw,r,1+j,h,HEAD,"FFFFFF",True,8.5,True,True)
    cw.row_dimensions[r].height=26; r+=1
    for (cod,hito,ini,fin,crit,obs) in ms:
        est=estado(ini,fin)
        cell(cw,r,1,cod,None,"333333",True,8.5,True)
        cell(cw,r,2,hito,None,"000000",crit==1,9,False,True)
        cell(cw,r,3,ini.strftime("%d/%m/%Y"),None,"1F3864",False,8.5,True)
        cell(cw,r,4,fin.strftime("%d/%m/%Y"),None,"1F3864",True if crit else False,8.5,True)
        cell(cw,r,5,"SÍ" if crit else "",("C62828" if crit else None),"FFFFFF" if crit else "000000",True,8.5,True)
        cell(cw,r,6,est,EST_FILL[est],"FFFFFF",True,8.5,True)
        for k in (7,8,9): cell(cw,r,k,"",None,"000000",False,9,True)
        cell(cw,r,10,obs,None,"555555",False,8.5,False,True)
        r+=1
    r+=1
for c,w in zip("ABCDEFGHIJ",[6,46,12,12,8,16,12,10,14,40]): cw.column_dimensions[c].width=w
cw.freeze_panes="A4"

# ---------------- Hoja 3: Timeline ----------------
tw=wb.create_sheet("Timeline (Gantt hitos)")
tw.merge_cells(start_row=1,start_column=1,end_row=1,end_column=3+len(MESES))
cell(tw,1,1,f"TIMELINE DE HITOS EJECUTIVOS  ·  {REV} · {FECHA}  ·  corte: entrega de equipos",NAVY,"FFFFFF",True,12,True); tw.row_dimensions[1].height=22
r=3
cell(tw,r,1,"Provisión / Hito",HEAD,"FFFFFF",True,8.5,False,True); cell(tw,r,2,"Inicio",HEAD,"FFFFFF",True,8.5,True); cell(tw,r,3,"Fin",HEAD,"FFFFFF",True,8.5,True)
prevy=None
for j,(y,m) in enumerate(MESES):
    lbl=f"{M3[m]}\n{y if y!=prevy else ''}".strip(); prevy=y
    cell(tw,r,4+j,f"{M3[m]}-{str(y)[2:]}",HEAD,"FFFFFF",True,7.5,True,True)
    tw.column_dimensions[get_column_letter(4+j)].width=6.5
tw.row_dimensions[r].height=26; r+=1
def midx(d):
    for j,(y,m) in enumerate(MESES):
        if d.year==y and d.month==m: return j
    return None
for (name,ven,col,ms) in PROV:
    tw.merge_cells(start_row=r,start_column=1,end_row=r,end_column=3+len(MESES))
    cell(tw,r,1,f"▶  {name}",col,"FFFFFF",True,9.5); tw.row_dimensions[r].height=16; r+=1
    for (cod,hito,ini,fin,crit,obs) in ms:
        cell(tw,r,1,f"{cod} · {hito}",None,"000000",crit==1,8.3,False,True)
        cell(tw,r,2,ini.strftime("%d/%m/%y"),None,"1F3864",False,8,True)
        cell(tw,r,3,fin.strftime("%d/%m/%y"),None,"1F3864",crit==1,8,True)
        a,b=midx(ini),midx(fin)
        for j in range(len(MESES)):
            v=""; f=None
            if a is not None and b is not None and a<=j<=b:
                f = "C62828" if crit else col
                if a==b or j==b: v="◆" if crit else "■"
            cell(tw,r,4+j,v,f,"FFFFFF",True,8,True)
        r+=1
    r+=1
for c,w in zip("ABC",[52,10,10]): tw.column_dimensions[c].width=w
tw.freeze_panes="D4"

# ---------------- Hoja 4: Notas ----------------
nw=wb.create_sheet("Notas y criterios")
cell(nw,1,1,"NOTAS Y CRITERIOS",NAVY,"FFFFFF",True,12)
for i,n in enumerate(NOTAS):
    cell(nw,3+i,1,"• "+n,None,"000000",False,9,False,True); nw.row_dimensions[3+i].height=32
nw.column_dimensions['A'].width=160

out=os.path.join(GEN,f"Cronograma Ejecutivo Provisiones Criticas - {REV} - {FECHA}.xlsx")
wb.save(out)
print("OK ->",out)
for (name,ven,col,ms) in PROV:
    cum=sum(1 for m in ms if estado(m[2],m[3])=="Cumplido s/plan")
    print(f"  {name[:44]:44} hitos={len(ms):2} cumplidos={cum:2} entrega={ms[-1][3].strftime('%d/%m/%Y')}")
