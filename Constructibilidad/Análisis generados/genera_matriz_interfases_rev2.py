#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
MATRIZ DE INTERFASES / RESPONSABILIDADES — Provisiones críticas CPF-2 (Rev2, 2026-08-21).
Cambios respecto de Rev1:
  - Se separa la función Precom/Com de la de Calidad (ambas pertenecen al Depto. de Calidad
    Integral): AESA-Precom&Comisionado ejecuta precom/PEM; AESA-Calidad (QA/QC) revisa/aprueba/
    libera según estadios de fabricación y despacho, y aprueba/revisa la documentación de ensayos.
  - Se incorpora Activación de suministros (gestión de la OC/provisión — expediting).
Esquema RACI ampliado: E=Ejecuta · A=Aprueba · R=Revisa/comenta · P=Participa/Asiste · I=Informado.
Actores: Proveedor · Activación de suministros · AESA-Proyecto/Contrato · AESA-Ingeniería ·
         AESA-Construcciones · AESA-Precom&Comis. · AESA-Calidad(QA/QC) · Cliente-PPSA.
Genera Excel (matriz editable) + PDF (informe, A3 apaisado).
"""
import os, textwrap
GEN="/home/user/OWN-Daro-1/Constructibilidad/Análisis generados"
REV="Rev2"; FECHA="2026-08-21"

ACT_COLS=["PROVEEDOR","Activación de suministros","AESA — Proyecto / Contrato","AESA — Ingeniería",
          "AESA — Construcciones","AESA — Precom. y Comisionado","AESA — Calidad (QA/QC)","CLIENTE — PPSA"]
ACT_SHORT=["Proveedor","Activación suministros","AESA Proy/Contrato","AESA Ingeniería",
           "AESA Construcciones","AESA Precom/Com","AESA Calidad (QA/QC)","Cliente PPSA"]
RACI={'E':('1F3864','Ejecuta / responsable primario'),'A':('C62828','Aprueba / libera'),
      'R':('EF6C00','Revisa / comenta'),'P':('2E7D32','Participa / Asiste'),'I':('90A4AE','Informado')}

# codes = [Proveedor, Activación, Proyecto, Ingeniería, Construcciones, Precom&Com, Calidad, Cliente]
ACTIVITIES=[
 ("Contrato","Gestión de la OC / contrato y coordinación",           ['P','P','E','I','I','I','I','I']),
 ("Ingeniería","Emisión de documentación (IB / ID / IC)",            ['E','I','I','R','I','I','I','I']),
 ("Ingeniería","Revisión y aprobación de documentación",             ['P','I','I','A','R','I','R','I']),
 ("Ingeniería","Freezing / hold points · cierre de ingeniería",      ['E','I','P','A','R','I','R','I']),
 ("Procura / Fabricación","Procura y fabricación (seguimiento)",     ['E','E','R','I','I','I','R','I']),
 ("Procura / Fabricación","Certificación de hitos de pago / avance", ['P','P','E','I','I','I','I','I']),
 ("Pruebas de fábrica (FAT)","Procedimientos de FAT/pruebas (emisión y aprobación)",['E','I','I','R','I','P','A','I']),
 ("Pruebas de fábrica (FAT)","Ejecución de FAT / pruebas en fábrica (asistencia)",  ['E','P','P','P','I','P','A','P']),
 ("Pruebas de fábrica (FAT)","iFAT / pruebas de integración (donde aplica)",        ['E','I','P','P','P','P','R','P']),
 ("Despacho / Sitio","Despacho y entrega en sitio",                  ['E','E','R','I','P','I','A','I']),
 ("Despacho / Sitio","Montaje / integración en sitio",               ['P','I','R','R','E','I','R','I']),
 ("Precom. / Puesta en marcha","SAT / pruebas en sitio",             ['P','I','P','P','P','E','R','P']),
 ("Precom. / Puesta en marcha","Precomisionado",                     ['P','I','P','P','P','E','R','P']),
 ("Precom. / Puesta en marcha","Comisionado / Puesta en marcha (PEM)",['P','I','P','P','I','E','R','P']),
 ("Precom. / Puesta en marcha","Capacitación (operación / mantenimiento)",['E','I','P','I','I','P','I','P']),
 ("Cierre","Documentación final / DataBook / As-built",              ['E','I','I','A','P','P','A','I']),
]
ADJ_ROW=("Adjudicación","Proceso de adjudicación / cierre de contrato",['P','P','E','R','I','I','I','I'])

PROV=[
 ("ABB — Salas Eléctricas SE#3 y SE#4","ABB","Adjudicado — OC vigente + planificación","C62828",False),
 ("ABB — PMS (Power Management System)","ABB","Adjudicado — OC vigente + planificación","EF6C00",False),
 ("Inauco — PCS / SCADA / Comunicaciones","Inauco","Adjudicado — Obra 129-008","1565C0",False),
 ("HIMA — Tableros de marshalling (ESD / F&G / PSS)","HIMA","Adjudicado","6A1B9A",False),
 ("HIMA — Adecuación y configuración SIS","HIMA","EN ADJUDICACIÓN","00838F",True),
]
DATES={
 "ABB — Salas Eléctricas SE#3 y SE#4":{
   "Freezing / hold points · cierre de ingeniería":"FP Constr. S#3: 31-AGO-2026",
   "Procura y fabricación (seguimiento)":"Fabricación: 01-SEP → 02-NOV-2026",
   "Ejecución de FAT / pruebas en fábrica (asistencia)":"FAT Brasil 03-23-NOV-26 · FAT Salas 08-MAR→19-ABR-27",
   "iFAT / pruebas de integración (donde aplica)":"iFAT: 25-JUN → 14-JUL-2027",
   "Despacho y entrega en sitio":"Despacho 02-ABR→14-MAY · Entrega S#3: 14-MAY-2027",
   "Montaje / integración en sitio":"Montaje shelters (Mendoza): 18-ENE→26-MAR-2027",
   "Comisionado / Puesta en marcha (PEM)":"Comis. 16-JUL-27 → RFSU 31-ENE-2028"},
 "ABB — PMS (Power Management System)":{
   "Freezing / hold points · cierre de ingeniería":"FP Ing. Básica: 21-AGO-2026",
   "Procura y fabricación (seguimiento)":"Procura → 05-NOV · Armado+SW → 17-DIC-2026",
   "Ejecución de FAT / pruebas en fábrica (asistencia)":"FAT Maqueta 22-SEP-26 · FAT PMS 25-29-ENE-27",
   "iFAT / pruebas de integración (donde aplica)":"FAT Integral Shelter 31-MAR · iFAT 25-JUN→14-JUL-27",
   "Despacho y entrega en sitio":"Despacho + Montaje shelter: 04-FEB→17-MAR-2027",
   "Documentación final / DataBook / As-built":"DataBook PMS: 31-MAR-2027",
   "Comisionado / Puesta en marcha (PEM)":"SAT PMS sitio OCT→DIC-27 · RFSU 31-ENE-2028",
   "Capacitación (operación / mantenimiento)":"Capacitación PMS: FEB-MAR-2028"},
 "Inauco — PCS / SCADA / Comunicaciones":{
   "Gestión de la OC / contrato y coordinación":"Recepción OC: 28-JUL-2026",
   "Procura y fabricación (seguimiento)":"Construcción tableros: 26-AGO-26 → 16-FEB-27",
   "Ejecución de FAT / pruebas en fábrica (asistencia)":"FAT PCS 29-MAR→22-ABR · FAT SIS 29-MAR→26-ABR-27",
   "iFAT / pruebas de integración (donde aplica)":"Pruebas PCS↔PMS 28-ABR→04-MAY · iFAT JUL-27",
   "SAT / pruebas en sitio":"SAT/Comis campo: 18-MAY → 05-OCT-2027",
   "Comisionado / Puesta en marcha (PEM)":"Comisionado → 31-ENE-2028",
   "Documentación final / DataBook / As-built":"Prueba MCE + CAO: → 07-DIC-2027"},
 "HIMA — Tableros de marshalling (ESD / F&G / PSS)":{
   "Procura y fabricación (seguimiento)":"Fabricación tablero: 03-JUL → 25-DIC-2026",
   "Despacho y entrega en sitio":"Recepción en Inauco: 25-DIC-2026",
   "Montaje / integración en sitio":"Recableado interno (45 d): 25-DIC → 08-FEB-2027",
   "Ejecución de FAT / pruebas en fábrica (asistencia)":"FAT SIS (conjunto): 29-MAR → 26-ABR-2027",
   "Comisionado / Puesta en marcha (PEM)":"Soporte campo → 05-OCT-2027"},
 "HIMA — Adecuación y configuración SIS":{
   "Proceso de adjudicación / cierre de contrato":"A definir (en adjudicación)",
   "Ejecución de FAT / pruebas en fábrica (asistencia)":"FAT SIS (conjunto): 29-MAR → 26-ABR-2027",
   "Comisionado / Puesta en marcha (PEM)":"Asistencia PEM: 2027 (a confirmar)"},
}
NOTAS=[
 "Alcance: desde la adjudicación/OC en adelante. HIMA — Adecuación y configuración SIS está EN ADJUDICACIÓN; el resto está adjudicado.",
 "Activación de suministros: tiene a cargo la gestión de la OC/provisión (expediting) — seguimiento de fabricación, plazos y despacho.",
 "AESA — Precom. y Comisionado y AESA — Calidad (QA/QC) pertenecen al mismo Departamento de Calidad Integral, separados por función.",
 "AESA — Calidad (QA/QC): revisa / aprueba / libera según los estadios de fabricación y despacho, y aprueba/revisa la documentación de ensayos (protocolos, DataBook).",
 "AESA — Precom. y Comisionado: ejecuta el precomisionado y la puesta en marcha, y asiste a FAT/pruebas.",
 "AESA — Construcciones instala/monta en planta y recibe el material del despacho en sitio.",
 "AESA — Proyecto / Contrato es el responsable del contrato; AESA — Ingeniería revisa y aprueba la documentación técnica.",
 "Hitos/fechas: cronograma integrado Rev2 (06-08-2026). Fechas de referencia, sujetas a OC y freezing points. Documento en iteración.",
]
def dfor(prov,act): return DATES.get(prov,{}).get(act,"")

# ================= EXCEL =================
def build_xlsx():
    import openpyxl
    from openpyxl.styles import Font,PatternFill,Alignment,Border,Side
    thin=Side(style="thin",color="D9D9D9"); B=Border(thin,thin,thin,thin); NAVY="1F3864"; HEAD="37474F"
    def cell(ws,r,c,v,fill=None,color="000000",bold=False,size=9,center=False,wrap=False):
        cc=ws.cell(r,c,v); cc.font=Font(bold=bold,color=color,size=size)
        if fill: cc.fill=PatternFill("solid",fgColor=fill)
        cc.alignment=Alignment(horizontal="center" if center else "left",vertical="center",wrap_text=wrap); cc.border=B; return cc
    wb=openpyxl.Workbook(); ws=wb.active; ws.title="Matriz de Interfases"; NC=11
    ws.merge_cells(start_row=1,start_column=1,end_row=1,end_column=NC); cell(ws,1,1,f"MATRIZ DE INTERFASES / RESPONSABILIDADES — PROVISIONES CRÍTICAS CPF-2 · {REV} · {FECHA}",NAVY,"FFFFFF",True,12,True); ws.row_dimensions[1].height=22
    ws.merge_cells(start_row=2,start_column=1,end_row=2,end_column=NC); cell(ws,2,1,"CPF-2 La Calera II · RACI ampliado · Proveedor · Activación de suministros · AESA-Proyecto/Contrato · AESA-Ingeniería · AESA-Construcciones · AESA-Precom&Comis. · AESA-Calidad(QA/QC) · Cliente-PPSA",None,"555555",False,9)
    r=4; cell(ws,r,1,"Leyenda:",None,"000000",True,9)
    for i,(k,(col,desc)) in enumerate(RACI.items()): cell(ws,r,2+i,f"{k} = {desc}",col,"FFFFFF" if k!='I' else "000000",True,8,True,True)
    r=6
    def header(r):
        cell(ws,r,1,"Fase",HEAD,"FFFFFF",True,8.5,True); cell(ws,r,2,"Actividad / interfase",HEAD,"FFFFFF",True,8.5,True); cell(ws,r,3,"Hito / Fecha (cronograma Rev2)",HEAD,"FFFFFF",True,8,True,True)
        for j,a in enumerate(ACT_SHORT): cell(ws,r,4+j,a,HEAD,"FFFFFF",True,7.6,True,True)
        ws.row_dimensions[r].height=30
    for (name,ven,estado,col,adj) in PROV:
        ws.merge_cells(start_row=r,start_column=1,end_row=r,end_column=NC)
        cell(ws,r,1,f"▶  {name}   ·   Proveedor: {ven}   ·   Estado: {estado}",col,"FFFFFF",True,10); ws.row_dimensions[r].height=18; r+=1
        header(r); r+=1
        for fase,act,codes in (([ADJ_ROW]+ACTIVITIES) if adj else ACTIVITIES):
            cell(ws,r,1,fase,None,"333333",False,8); cell(ws,r,2,act,None,"000000",False,8.5,False,True); cell(ws,r,3,dfor(name,act),None,"1F3864",False,8,False,True)
            for j,code in enumerate(codes):
                fc,_=RACI[code]; cell(ws,r,4+j,code,fc,"FFFFFF" if code!='I' else "000000",True,9,True)
            r+=1
        r+=1
    for c,w in zip("ABCDEFGHIJK",[18,42,32,11,12,12,12,12,12,12,11]): ws.column_dimensions[c].width=w
    ws.freeze_panes="A7"
    ps=wb.create_sheet("Provisiones"); cell(ps,1,1,"PROVISIONES CRÍTICAS",NAVY,"FFFFFF",True,12); ps.merge_cells("A1:D1")
    for j,h in enumerate(["Provisión","Proveedor","Estado","Alcance de interfase"]): cell(ps,3,1+j,h,HEAD,"FFFFFF",True,9,True)
    alc={"ABB — Salas Eléctricas SE#3 y SE#4":"Ingeniería, fabricación, FAT, liberación/despacho, montaje shelters (Construcciones) y PEM.",
         "ABB — PMS (Power Management System)":"Ingeniería, procura, armado, FAT (incl. maqueta), montaje shelter, iFAT, PEM y capacitación.",
         "Inauco — PCS / SCADA / Comunicaciones":"Ingeniería, construcción tableros, FAT PCS/SIS, integración, SAT/comisionado en campo y CAO.",
         "HIMA — Tableros de marshalling (ESD / F&G / PSS)":"Ingeniería, fabricación tableros, FAT, liberación y entrega en Inauco.",
         "HIMA — Adecuación y configuración SIS":"Adjudicación, configuración/adecuación SIS, FAT SIS (conjunto), asistencia PEM."}
    r=4
    for (name,ven,estado,col,adj) in PROV:
        cell(ps,r,1,name,None,"000000",True,9,False,True); cell(ps,r,2,ven,None,"000000",False,9,True)
        cell(ps,r,3,estado,"C62828" if adj else "2E7D32","FFFFFF",True,9,True); cell(ps,r,4,alc[name],None,"000000",False,8.5,False,True); ps.row_dimensions[r].height=30; r+=1
    for c,w in zip("ABCD",[42,12,20,64]): ps.column_dimensions[c].width=w
    ns=wb.create_sheet("Notas y premisas"); cell(ns,1,1,"NOTAS Y PREMISAS",NAVY,"FFFFFF",True,12)
    for i,n in enumerate(NOTAS): cell(ns,3+i,1,f"• {n}",None,"000000",False,9,False,True); ns.row_dimensions[3+i].height=30
    ns.column_dimensions['A'].width=150
    path=os.path.join(GEN,f"Matriz de Interfases - Provisiones criticas - {REV} - {FECHA}.xlsx"); wb.save(path); return path

# ================= PDF (A3 apaisado) =================
def build_pdf():
    from reportlab.lib.pagesizes import A3, landscape
    from reportlab.lib import colors
    from reportlab.lib.units import mm
    from reportlab.pdfgen import canvas
    path=os.path.join(GEN,f"Matriz de Interfases - Provisiones criticas - {REV} - {FECHA}.pdf")
    W,H=landscape(A3); c=canvas.Canvas(path,pagesize=(W,H)); NAVY=colors.HexColor('#1F3864'); x0=10*mm; x1=W-10*mm
    def banner():
        c.setFillColor(NAVY); c.rect(0,H-13*mm,W,13*mm,fill=1,stroke=0); c.setFillColor(colors.white)
        c.setFont("Helvetica-Bold",13); c.drawString(x0,H-9*mm,"MATRIZ DE INTERFASES / RESPONSABILIDADES — PROVISIONES CRÍTICAS CPF-2")
        c.setFont("Helvetica",9); c.drawRightString(x1,H-9*mm,f"{REV} · {FECHA}")
    banner(); y=H-20*mm
    c.setFillColor(NAVY); c.setFont("Helvetica-Bold",11); c.drawString(x0,y,"1 · Provisiones y estado"); y-=6*mm
    colw=[95*mm,24*mm,52*mm,229*mm]
    def prow(vals,widths,y,fill=None,bold=False):
        cx=x0; tw=sum(widths)
        if fill: c.setFillColor(fill); c.rect(x0,y-5.4*mm,tw,5.4*mm,fill=1,stroke=0)
        for v,w in zip(vals,widths):
            c.setFillColor(colors.white if fill else colors.black); c.setFont("Helvetica-Bold" if bold else "Helvetica",8)
            for j,ln in enumerate(textwrap.wrap(str(v),int(w/1.55))[:2]): c.drawString(cx+1.3*mm,y-3.5*mm-j*3.1*mm,ln)
            cx+=w
        c.setStrokeColor(colors.HexColor('#DDDDDD')); c.setLineWidth(0.3); c.line(x0,y-5.4*mm,x0+tw,y-5.4*mm); return y-5.4*mm
    y=prow(["Provisión","Proveedor","Estado","Alcance de interfase"],colw,y,fill=colors.HexColor('#37474F'),bold=True)
    alc={"ABB — Salas Eléctricas SE#3 y SE#4":"Ingeniería, fabricación, FAT, liberación/despacho, montaje shelters (Construcciones) y PEM.",
         "ABB — PMS (Power Management System)":"Ingeniería, procura, armado, FAT (incl. maqueta), montaje shelter, iFAT, PEM y capacitación.",
         "Inauco — PCS / SCADA / Comunicaciones":"Ingeniería, construcción tableros, FAT PCS/SIS, integración, SAT/comisionado campo y CAO.",
         "HIMA — Tableros de marshalling (ESD / F&G / PSS)":"Ingeniería, fabricación tableros, FAT, liberación y entrega en Inauco.",
         "HIMA — Adecuación y configuración SIS":"Adjudicación, configuración/adecuación SIS, FAT SIS (conjunto), asistencia PEM."}
    for (name,ven,estado,col,adj) in PROV:
        c.setFillColor(colors.HexColor('#'+col)); c.rect(x0,y-5.4*mm,3*mm,5.4*mm,fill=1,stroke=0); y=prow([name,ven,estado,alc[name]],colw,y)
    y-=6*mm
    c.setFillColor(NAVY); c.setFont("Helvetica-Bold",11); c.drawString(x0,y,"2 · Leyenda RACI ampliado"); y-=6*mm
    lx=x0
    for k,(col,desc) in RACI.items():
        c.setFillColor(colors.HexColor('#'+col)); c.rect(lx,y-4.2*mm,5.2*mm,5.2*mm,fill=1,stroke=0); c.setFillColor(colors.black)
        c.setFont("Helvetica-Bold",8.5); c.drawString(lx+6.5*mm,y-2.6*mm,k); c.setFont("Helvetica",7.6); c.drawString(lx+10.5*mm,y-2.6*mm,desc); lx+=62*mm
    y-=10*mm
    c.setFillColor(NAVY); c.setFont("Helvetica-Bold",11); c.drawString(x0,y,"3 · Notas y premisas"); y-=5.5*mm
    c.setFillColor(colors.black); c.setFont("Helvetica",8)
    for n in NOTAS:
        for j,ln in enumerate(textwrap.wrap(n,235)): c.drawString(x0+(0 if j==0 else 3*mm),y,("• " if j==0 else "")+ln); y-=4*mm
        y-=0.8*mm
    # matrix
    aw=[19*mm,60*mm,54*mm]+[ (x1-x0-133*mm)/8 ]*8
    FSHORT={"Procura / Fabricación":"Procura/Fabric.","Pruebas de fábrica (FAT)":"Pruebas FAT","Precom. / Puesta en marcha":"Precom./PEM"}
    ASHORT={"Emisión de documentación (IB / ID / IC)":"Emisión de doc. (IB/ID/IC)","Revisión y aprobación de documentación":"Revisión y aprobación de doc.",
      "Freezing / hold points · cierre de ingeniería":"Freezing/hold points · cierre ing.","Certificación de hitos de pago / avance":"Certif. hitos de pago / avance",
      "Procedimientos de FAT/pruebas (emisión y aprobación)":"Procedimientos FAT (emisión/aprob.)","Ejecución de FAT / pruebas en fábrica (asistencia)":"Ejecución FAT / pruebas (asistencia)",
      "iFAT / pruebas de integración (donde aplica)":"iFAT / pruebas integración (si aplica)","Comisionado / Puesta en marcha (PEM)":"Comisionado / PEM",
      "Capacitación (operación / mantenimiento)":"Capacitación (op./mant.)","Documentación final / DataBook / As-built":"Doc. final / DataBook / As-built",
      "Gestión de la OC / contrato y coordinación":"Gestión de OC / contrato"}
    def mhead(y):
        cx=x0; c.setFillColor(colors.HexColor('#37474F')); c.rect(x0,y-9*mm,sum(aw),9*mm,fill=1,stroke=0); c.setFillColor(colors.white); c.setFont("Helvetica-Bold",6.8)
        labs=["Fase","Actividad / interfase","Hito / Fecha (cronograma)"]+ACT_SHORT
        for l,w in zip(labs,aw):
            for j,ln in enumerate(textwrap.wrap(l,int(w/1.45))[:2]): c.drawString(cx+1.1*mm,y-3.6*mm-j*3.1*mm,ln)
            cx+=w
        return y-9*mm
    def ensure(y,need):
        if y-need<12*mm: c.showPage(); banner(); return H-18*mm,True
        return y,False
    c.showPage(); banner(); y=H-18*mm
    for (name,ven,estado,col,adj) in PROV:
        y,_=ensure(y,26*mm)
        c.setFillColor(colors.HexColor('#'+col)); c.rect(x0,y-6*mm,sum(aw),6*mm,fill=1,stroke=0); c.setFillColor(colors.white); c.setFont("Helvetica-Bold",9)
        c.drawString(x0+2*mm,y-4.3*mm,f"{name}   ·   {ven}   ·   {estado}"); y-=6*mm; y=mhead(y)
        for ri,(fase,act,codes) in enumerate(([ADJ_ROW]+ACTIVITIES) if adj else ACTIVITIES):
            y,newp=ensure(y,6*mm)
            if newp: y=mhead(y)
            rowh=6.2*mm; cx=x0
            c.setFillColor(colors.HexColor('#F7F9FB') if ri%2 else colors.white); c.rect(x0,y-rowh,sum(aw),rowh,fill=1,stroke=0)
            c.setFillColor(colors.HexColor('#444444')); c.setFont("Helvetica",6.6); c.drawString(cx+1.2*mm,y-4*mm,FSHORT.get(fase,fase)[:22]); cx+=aw[0]
            c.setFillColor(colors.black); c.setFont("Helvetica",6.8); c.drawString(cx+1.2*mm,y-4*mm,ASHORT.get(act,act)[:44]); cx+=aw[1]
            c.setFillColor(colors.HexColor('#1F3864')); c.setFont("Helvetica",6.0)
            for j,ln in enumerate(textwrap.wrap(dfor(name,act),int(aw[2]/1.32))[:2]): c.drawString(cx+1.2*mm,y-3.2*mm-j*2.8*mm,ln)
            cx+=aw[2]
            for code in codes:
                fc,_=RACI[code]; c.setFillColor(colors.HexColor('#'+fc)); c.roundRect(cx+aw[3]/2-3.4*mm,y-rowh+1*mm,6.8*mm,rowh-2*mm,1,fill=1,stroke=0)
                c.setFillColor(colors.white if code!='I' else colors.black); c.setFont("Helvetica-Bold",7.2); c.drawCentredString(cx+aw[3]/2,y-4*mm,code); cx+=aw[3]
            y-=rowh
        y-=4*mm
    c.showPage(); c.save(); return path

p1=build_xlsx(); p2=build_pdf(); print("OK\n ",p1,"\n ",p2)
