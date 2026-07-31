#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Cronograma INTEGRADO CPF-2 La Calera II — Provisiones PCS (Inauco) + PSS/SIS (HIMA) + Salas
Eléctricas y PMS (ABB, Rev4 optimizado). Genera:
  - Gantt integrado (PDF, A3 apaisado)
  - Diagrama de flujo integrado (PDF, A3 vertical)
  - Planilla de tareas/hitos (XLSX)
Información nueva incorporada:
  - Cronogramas ABB de Salas y PMS: Rev4 optimizado (branch ABB).
  - Inauco (PCS): INICIO 01-AGO-2026 (antes 01-JUL); cadena corrida +31 días.
  - HIMA (PSS/SIS): tablero de seguridad con plazo de 25 semanas desde agosto -> 23-ENE-2027.
"""
import os
from datetime import date, timedelta
from reportlab.lib.pagesizes import A3, landscape
from reportlab.lib.units import mm
from reportlab.lib import colors
from reportlab.pdfgen import canvas

OUT="/home/user/OWN-Daro-1/Constructibilidad/Análisis generados"
os.makedirs(OUT,exist_ok=True)
D=lambda y,m,d: date(y,m,d)

# ---------------- paleta por proveedor ----------------
C_ABB   = colors.HexColor('#C62828')   # ABB Salas (rojo — cadena crítica)
C_ABBP  = colors.HexColor('#EF6C00')   # ABB PMS (naranja)
C_INA   = colors.HexColor('#1565C0')   # Inauco PCS (azul)
C_HIMA  = colors.HexColor('#6A1B9A')   # HIMA PSS/SIS (violeta)
C_EPC   = colors.HexColor('#2E7D32')   # Integración EPC (verde)
GREY    = colors.HexColor('#555555')

# ---------------- tareas: (track, etiqueta, inicio, fin, color) ----------------
# hito: fin=None  -> se dibuja como rombo/estrella
TASKS=[
 # ---- ABB — SALAS ELÉCTRICAS (Rev4) ----
 ("ABB — Salas Eléctricas (S#3 crítica / S#4)","Ingeniería IB→ID→IC (FP Constr. S#3 31-AGO)",D(2026,5,22),D(2026,8,31),C_ABB),
 (None,"Fabricación tableros MT/BT (Brasil M&B)",D(2026,9,1),D(2026,11,2),C_ABB),
 (None,"FAT tableros/celdas (Brasil)",D(2026,11,3),D(2026,11,23),C_ABB),
 (None,"Acopio de materiales (Argentina)",D(2026,12,7),D(2027,1,29),C_ABB),
 (None,"Montaje e integración Shelters (Mendoza)",D(2027,1,18),D(2027,3,26),C_ABB),
 (None,"FAT Salas Eléctricas (S#4 / S#3)",D(2027,3,8),D(2027,4,19),C_ABB),
 (None,"Despacho a Vaca Muerta (S#4 / S#3)",D(2027,4,2),D(2027,5,14),C_ABB),
 (None,"★ Entrega S#4 (08-ABR) / ★ Entrega S#3 (14-MAY-2027)",D(2027,5,14),None,C_ABB),
 # ---- ABB — PMS ----
 ("ABB — PMS (Power Management System)","Ingeniería Básica (FP IB 21-AGO)",D(2026,6,8),D(2026,8,21),C_ABBP),
 (None,"Ingeniería Detalle + Procedimientos",D(2026,8,3),D(2026,10,23),C_ABBP),
 (None,"Procura (AC800M / switches / tableros)",D(2026,5,22),D(2026,11,5),C_ABBP),
 (None,"FAT de Maqueta (prueba temprana AESA)",D(2026,9,22),D(2026,9,25),C_ABBP),
 (None,"Armado tableros + Software 800XA",D(2026,9,11),D(2026,12,17),C_ABBP),
 (None,"FAT Tableros + Sistema PMS (witness AESA)",D(2027,1,25),D(2027,1,29),C_ABBP),
 (None,"Despacho + Montaje PMS en Shelter (Mendoza)",D(2027,2,4),D(2027,3,17),C_ABBP),
 (None,"★ FAT Integral Shelter PMS / DataBook (31-MAR-2027)",D(2027,3,31),None,C_ABBP),
 # ---- INAUCO — PCS ----
 ("Inauco — PCS / SCADA / Comunicaciones","INICIO 01-AGO + Ingeniería detalle (Freezing 31-OCT)",D(2026,8,1),D(2026,10,31),C_INA),
 (None,"Procura HW Rockwell/Stratix (recepción 21-SEP)",D(2026,8,1),D(2026,9,21),C_INA),
 (None,"◆ Aprobación AESA Ingeniería Rev0 (10-DIC-2026)",D(2026,12,10),None,C_INA),
 (None,"Construcción tableros PCS (Remota 7A/7B)",D(2026,12,11),D(2027,2,26),C_INA),
 (None,"Configuración PLC / SCADA",D(2026,9,1),D(2027,2,26),C_INA),
 (None,"FAT PCS + Comunicaciones (Neuquén, 18 d)",D(2027,2,27),D(2027,3,22),C_INA),
 (None,"Integración PCS↔PMS ABB (Bs.As., 5 d)",D(2027,3,27),D(2027,4,2),C_INA),
 (None,"SAT + Comisionado PCS — campo (55 d)",D(2027,4,17),D(2027,7,2),C_INA),
 (None,"SAT + Comisionado SCADA SIS — campo (100 d)",D(2027,4,17),D(2027,9,3),C_INA),
 (None,"Prueba MCE + CAO (as-built)",D(2027,9,4),D(2027,10,22),C_INA),
 # ---- HIMA — PSS/SIS ----
 ("HIMA — PSS / SIS (ESD · F&G)","Fabricación tablero SIS (25 sem desde AGO)",D(2026,8,1),D(2027,1,23),C_HIMA),
 (None,"★ Recepción tablero HIMA en Inauco (23-ENE-2027)",D(2027,1,23),None,C_HIMA),
 (None,"FAT SIS (HIMA + Inauco, 20 d)",D(2027,2,27),D(2027,3,26),C_HIMA),
 (None,"Soporte HIMA en campo (SAT/PEM SCADA SIS)",D(2027,4,17),D(2027,9,3),C_HIMA),
 # ---- INTEGRACIÓN EPC (AESA) ----
 ("Integración EPC (AESA)","Precomisionado / montaje de campo",D(2027,5,3),D(2028,1,13),C_EPC),
 (None,"iFAT — Integración Shelters + PMS + INAUCO + HIMA",D(2027,6,25),D(2027,7,14),C_EPC),
 (None,"Comisionado",D(2027,7,16),D(2028,1,31),C_EPC),
 (None,"SAT PMS en sitio (hold point)",D(2027,10,1),D(2027,12,15),C_EPC),
 (None,"★ RFSU — Ready For Start Up (31-ENE-2028)",D(2028,1,31),None,C_EPC),
]

# hitos verticales de referencia (fecha, etiqueta, color)
VLINES=[
 (D(2027,1,23),"Tablero HIMA",C_HIMA),
 (D(2027,3,31),"DataBook PMS",C_ABBP),
 (D(2027,5,14),"Entrega S#3",C_ABB),
 (D(2027,7,14),"iFAT",C_EPC),
 (D(2028,1,31),"RFSU",C_EPC),
]

# ---------------- GANTT ----------------
def month_iter(d0,d1):
    y,m=d0.year,d0.month; out=[]
    while (y<d1.year) or (y==d1.year and m<=d1.month):
        out.append((y,m)); m+=1
        if m>12: m=1;y+=1
    return out
MES=['','Ene','Feb','Mar','Abr','May','Jun','Jul','Ago','Sep','Oct','Nov','Dic']

def build_gantt():
    W,H=landscape(A3)
    path=os.path.join(OUT,"Cronograma_Integrado_CPF2_Gantt.pdf")
    c=canvas.Canvas(path,pagesize=(W,H))
    d0,d1=D(2026,5,1),D(2028,2,29)
    months=month_iter(d0,d1)
    x0=345; x1=W-24; top=H-96
    SHORT={"ABB — Salas Eléctricas (S#3 crítica / S#4)":"ABB · Salas",
           "ABB — PMS (Power Management System)":"ABB · PMS",
           "Inauco — PCS / SCADA / Comunicaciones":"Inauco · PCS",
           "HIMA — PSS / SIS (ESD · F&G)":"HIMA · SIS",
           "Integración EPC (AESA)":"Integr. EPC"}
    span_days=(d1-d0).days
    def X(dt): return x0+(x1-x0)*((dt-d0).days/span_days)
    # título
    c.setFont("Helvetica-Bold",15); c.setFillColor(colors.black)
    c.drawString(24,H-34,"CRONOGRAMA INTEGRADO — CPF-2 La Calera II (Vaca Muerta)")
    c.setFont("Helvetica",9.5); c.setFillColor(GREY)
    c.drawString(24,H-50,"Provisiones: PCS (Inauco) · PSS/SIS (HIMA) · Salas Eléctricas y PMS (ABB, Rev4 optimizado)   |   "
                         "Inauco inicio 01-AGO-2026 · Tablero HIMA 25 sem (→23-ENE-2027) · RFSU 31-ENE-2028")
    # grilla meses + años
    yr_row=top+20
    c.setFont("Helvetica-Bold",8)
    prev_year=None
    for (y,m) in months:
        mx=X(D(y,m,1)); mxn=X(D(y+(m//12),(m%12)+1,1))
        c.setStrokeColor(colors.HexColor('#DDDDDD')); c.setLineWidth(0.5)
        c.line(mx,top+8,mx,42)
        c.setFillColor(GREY); c.setFont("Helvetica",7)
        c.drawCentredString((mx+mxn)/2,top+2,MES[m])
        if y!=prev_year:
            c.setFillColor(colors.black); c.setFont("Helvetica-Bold",9)
            c.drawString(mx+2,yr_row,str(y)); prev_year=y
    # líneas verticales de hitos
    for dt,lab,col in VLINES:
        vx=X(dt)
        c.setStrokeColor(col); c.setLineWidth(1); c.setDash(2,2)
        c.line(vx,42,vx,top+8); c.setDash()
        c.setFillColor(col); c.setFont("Helvetica-Bold",6.5)
        c.saveState(); c.translate(vx,top+30); c.rotate(90); c.drawString(0,2,lab); c.restoreState()
    # filas
    rowh=(top-52)/len(TASKS)
    y=top
    def star(cx,cy,r,col):
        from math import cos,sin,pi
        p=c.beginPath();
        for i in range(10):
            ang=pi/2+i*pi/5; rr=r if i%2==0 else r*0.45
            xx=cx+rr*cos(ang); yy=cy+rr*sin(ang)
            (p.moveTo if i==0 else p.lineTo)(xx,yy)
        p.close(); c.setFillColor(col); c.setStrokeColor(col); c.drawPath(p,fill=1,stroke=0)
    for (track,label,s,e,col) in TASKS:
        y-=rowh
        if track:
            c.setFillColor(col); c.rect(6,y,10,rowh-3,fill=1,stroke=0)
            c.setFillColor(colors.black); c.setFont("Helvetica-Bold",7.2)
            c.drawString(20,y+rowh/2-3,SHORT.get(track,track)[:20])
        # etiqueta de tarea (columna única alineada)
        c.setFillColor(colors.HexColor('#333333')); c.setFont("Helvetica",6.6)
        c.drawString(126,y+rowh/2-3,label[:54])
        # banda
        c.setStrokeColor(colors.HexColor('#EEEEEE')); c.setLineWidth(0.3); c.line(x0,y,x1,y)
        if e is None:  # hito
            star(X(s),y+rowh/2,4.2,col)
        else:
            bx=X(s); bw=max(X(e)-X(s),2.2)
            c.setFillColor(col); c.setStrokeColor(col)
            c.roundRect(bx,y+2.5,bw,rowh-6,2,fill=1,stroke=0)
    # marco timeline
    c.setStrokeColor(colors.HexColor('#BBBBBB')); c.setLineWidth(0.6)
    c.rect(x0,44,x1-x0,top-44,fill=0,stroke=1)
    # leyenda
    c.setFont("Helvetica-Bold",7); ly=30
    items=[("ABB Salas",C_ABB),("ABB PMS",C_ABBP),("Inauco PCS",C_INA),("HIMA PSS/SIS",C_HIMA),("Integración EPC",C_EPC),("★ Hito",colors.black)]
    lx=24
    for lab,col in items:
        c.setFillColor(col); c.rect(lx,ly,10,7,fill=1,stroke=0)
        c.setFillColor(colors.black); c.drawString(lx+13,ly+1,lab); lx+=95
    c.setFont("Helvetica-Oblique",6.5); c.setFillColor(GREY)
    c.drawRightString(x1,30,"Fuente: ABB Rev4 (27-JUL-2026) · Dashboard Inauco · plazo HIMA 25 sem. Fechas de referencia, sujetas a OC y freezing points.")
    c.showPage(); c.save(); return path

# ---------------- FLUJO ----------------
def build_flow():
    W,H=A3  # vertical
    path=os.path.join(OUT,"Cronograma_Integrado_CPF2_Flujo.pdf")
    c=canvas.Canvas(path,pagesize=(W,H))
    c.setFont("Helvetica-Bold",15)
    c.drawString(24,H-36,"FLUJO INTEGRADO DE PROVISIONES — CPF-2 La Calera II")
    c.setFont("Helvetica",9); c.setFillColor(GREY)
    c.drawString(24,H-52,"Convergencia de PCS (Inauco), PSS/SIS (HIMA), Salas y PMS (ABB) hacia iFAT → Comisionado → RFSU")
    def box(x,y,w,h,title,lines,col):
        c.setFillColor(col); c.roundRect(x,y,w,h,5,fill=1,stroke=0)
        c.setFillColor(colors.white); c.setFont("Helvetica-Bold",8.6)
        c.drawCentredString(x+w/2,y+h-13,title)
        c.setFont("Helvetica",6.9)
        yy=y+h-26
        for ln in lines:
            c.drawCentredString(x+w/2,yy,ln); yy-=9.2
    def arrow(x1,y1,x2,y2,col=GREY):
        c.setStrokeColor(col); c.setLineWidth(1.4); c.line(x1,y1,x2,y2)
        from math import atan2,cos,sin
        a=atan2(y2-y1,x2-x1)
        for da in (2.6,-2.6):
            c.line(x2,y2,x2-8*cos(a+da),y2-8*sin(a+da))
    colw=(W-48-3*12)/4
    xs=[24+i*(colw+12) for i in range(4)]
    ytop=H-90; bh=150
    cols=[
     ("ABB — Salas Eléctricas",C_ABB,["Ingeniería IB→ID→IC","FP Constr. S#3: 31-AGO-26","Fabricación tableros (Brasil)","FAT Brasil → Acopio Arg.","Montaje Shelters (Mendoza)","FAT Salas → Despacho","★ Entrega S#3: 14-MAY-27"]),
     ("ABB — PMS",C_ABBP,["Ing. Básica (FP 21-AGO)","Procura AC800M/switches","FAT Maqueta (22-SEP)","Armado + Software 800XA","FAT PMS witness (25-ENE)","Montaje Shelter Mendoza","★ DataBook PMS: 31-MAR-27"]),
     ("Inauco — PCS/SCADA",C_INA,["INICIO: 01-AGO-2026","Ing. detalle (Freezing 31-OCT)","Aprob. Rev0: 10-DIC-26","Construcción tableros PCS","Config PLC/SCADA","FAT PCS+Comms (27-FEB→22-MAR)","Integr. PCS↔PMS (Bs.As.)"]),
     ("HIMA — PSS/SIS",C_HIMA,["Tablero SIS (ESD·F&G·PSS)","Plazo 25 sem desde AGO","★ Recepción: 23-ENE-2027","FAT SIS con Inauco","(27-FEB → 26-MAR-2027)","Soporte SAT/PEM en campo","→ hasta 03-SEP-2027"]),
    ]
    for i,(t,col,lines) in enumerate(cols):
        box(xs[i],ytop-bh,colw,bh,t,lines,col)
    # convergencia -> iFAT
    ifat_y=ytop-bh-70; ifat_w=W-160; ifat_x=(W-ifat_w)/2
    for i in range(4):
        arrow(xs[i]+colw/2,ytop-bh,(W)/2,ifat_y+52)
    box(ifat_x,ifat_y,ifat_w,52,"iFAT — INTEGRACIÓN Shelters + PMS + INAUCO + HIMA",
        ["25-JUN → 14-JUL-2027   (Neuquén / staging)"],C_EPC)
    # -> campo
    camp_y=ifat_y-64
    arrow(W/2,ifat_y,W/2,camp_y+44)
    box(ifat_x,camp_y,ifat_w,44,"PRECOMISIONADO + COMISIONADO EN CAMPO",
        ["Precom: 03-MAY-27 → 13-ENE-28   ·   Comisionado: 16-JUL-27 → 31-ENE-28   ·   SAT PMS sitio: OCT→DIC-27"],C_EPC)
    # -> RFSU
    rfsu_y=camp_y-58
    arrow(W/2,camp_y,W/2,rfsu_y+40)
    c.setFillColor(colors.HexColor('#1B5E20')); c.roundRect(ifat_x+ifat_w/2-150,rfsu_y,300,40,6,fill=1,stroke=0)
    c.setFillColor(colors.white); c.setFont("Helvetica-Bold",13)
    c.drawCentredString(W/2,rfsu_y+15,"★ RFSU — 31-ENE-2028")
    # ruta crítica nota
    c.setFont("Helvetica-Oblique",7.5); c.setFillColor(GREY)
    c.drawString(24,rfsu_y-24,"Ruta que gobierna la integración: Tablero HIMA (23-ENE) → FAT SIS (MAR) y Entrega Sala S#3 (14-MAY) → iFAT (JUL) → Comisionado → RFSU.")
    c.drawString(24,rfsu_y-36,"Cambios respecto de la hipótesis previa: Inauco arranca 01-AGO (no 01-JUL) y el tablero HIMA corre 25 semanas desde agosto; los cronogramas ABB son la Rev4 optimizada.")
    c.showPage(); c.save(); return path

# ---------------- XLSX ----------------
def build_xlsx():
    import openpyxl
    from openpyxl.styles import Font,PatternFill,Alignment,Border,Side
    path=os.path.join(OUT,"Cronograma_Integrado_CPF2.xlsx")
    wb=openpyxl.Workbook(); ws=wb.active; ws.title="Cronograma Integrado"
    hdr=["Proveedor","Actividad / Hito","Inicio","Fin","Tipo"]
    ws.append(hdr)
    hf=Font(bold=True,color="FFFFFF"); hfill=PatternFill("solid",fgColor="37474F")
    for j,_ in enumerate(hdr,1):
        cc=ws.cell(1,j); cc.font=hf; cc.fill=hfill; cc.alignment=Alignment(horizontal="center")
    hexmap={C_ABB.hexval()[2:]:"C62828",C_ABBP.hexval()[2:]:"EF6C00",C_INA.hexval()[2:]:"1565C0",C_HIMA.hexval()[2:]:"6A1B9A",C_EPC.hexval()[2:]:"2E7D32"}
    cur=""
    for (track,label,s,e,col) in TASKS:
        if track: cur=track
        tipo="Hito" if e is None else "Tarea"
        ws.append([cur,label,s.strftime("%d-%m-%Y"),(e.strftime("%d-%m-%Y") if e else "—"),tipo])
        r=ws.max_row
        fillhex=hexmap.get(col.hexval()[2:],"999999")
        ws.cell(r,1).fill=PatternFill("solid",fgColor=fillhex); ws.cell(r,1).font=Font(color="FFFFFF",bold=True,size=8)
        if tipo=="Hito":
            for j in range(1,6): ws.cell(r,j).font=Font(bold=True,size=9,color=("FFFFFF" if j==1 else "000000"))
    widths=[34,52,12,12,8]
    for j,w in enumerate(widths,1): ws.column_dimensions[chr(64+j)].width=w
    ws.freeze_panes="A2"
    wb.save(path); return path

g=build_gantt(); f=build_flow(); x=build_xlsx()
print("OK\n ",g,"\n ",f,"\n ",x)
