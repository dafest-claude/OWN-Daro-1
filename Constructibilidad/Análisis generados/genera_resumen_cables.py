#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Resumen de PRIORIZACIÓN DE CABLES por especialidad, a partir de las requisiciones de ingeniería.
Objetivo: frente a largos plazos de entrega, separar por tipo/formación y priorizar la procura.

Premisas de priorización:
  ELECTRICIDAD  -> Cables de MT (todos) y BT de potencia con tendidos largos (mucho metraje).
  INSTRUMENTACIÓN -> Multipar de más de 8 pares; tendidos largos; fibras ópticas;
                     y cables resistentes al fuego (Fire Resistant, IEC 60331) de varios pares.

Entradas (Datos entrada):
  5155-00-0000-IG-EL-RI-006-0.xls   (Electricidad)
  5155-00-0000-IG-IN-RI-032-0.xls   (Instrumentación)
  5155-00-2000-IG-IN-RI-009-0.xls   (Instrumentación, área 2000)
Salidas (Análisis generados):
  Resumen Cables EL - Rev0 - 2026-08-04.xlsx
  Resumen Cables IN - Rev0 - 2026-08-04.xlsx
"""
import os, re, zipfile, shutil
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

ENT="/home/user/OWN-Daro-1/Constructibilidad/Datos entrada"
GEN="/home/user/OWN-Daro-1/Constructibilidad/Análisis generados"
FECHA="2026-08-04"; REV="Rev0"
SRC={"EL":["5155-00-0000-IG-EL-RI-006-0.xls"],
     "IN":["5155-00-0000-IG-IN-RI-032-0.xls","5155-00-2000-IG-IN-RI-009-0.xls"]}

def load(fn):
    tmp="/tmp/_"+os.path.basename(fn)+".xlsx"; shutil.copy(os.path.join(ENT,fn),tmp)
    ws=openpyxl.load_workbook(tmp,data_only=True).active
    out=[]
    for i,row in enumerate(ws.iter_rows(values_only=True),1):
        if i==1: continue
        pos,ident,qty,unit,mat=row[4],row[5],row[7],row[8],row[12]
        if ident is None or qty is None: continue
        out.append(dict(req=fn.split('-IG-')[1][:6], pos=pos, ident=str(ident),
                        m=int(round(float(qty))), unit=unit or "m", desc=str(mat or "")))
    return out

def num(rx,txt,default=0):
    m=re.search(rx,txt,re.I); return int(m.group(1)) if m else default

def attrs(d):
    t=d["desc"]; tl=t.lower()
    pairs=num(r'(\d+)\s*PAIRS?',t); triads=num(r'(\d+)\s*TRIADS?',t)
    conductors=num(r'(\d+)\s*conductors?',t); fibers=num(r'(\d+)\s*fibers?',t)
    fiber = 'fiber optic' in tl
    resistant = ('fire resistant' in tl) or ('iec 60331' in tl) or ('60331' in t)
    retardant = ('flame retardant' in tl) or ('fire retardant' in tl) or ('60332' in t)
    sec=""
    ms=re.search(r'Section:\s*([0-9x/,\.]+mm2)',t,re.I)
    if ms: sec=ms.group(1).replace(' ','')
    awg=""
    ma=re.search(r'(\d+)\s*AWG',t)
    if ma: awg=ma.group(1)+"AWG"
    # categoría
    if 'medium voltage' in tl: cat="MT — Potencia"
    elif 'low voltage power' in tl: cat="BT — Potencia"
    elif 'low voltage control' in tl or 'control cable' in tl: cat="BT — Control"
    elif 'fiber optic' in tl: cat="INS — Fibra óptica"
    elif 'communication' in tl: cat="INS — Comunicación"
    elif 'digital' in tl: cat="INS — Digital"
    elif 'analog' in tl: cat="INS — Analógico"
    else: cat="Otro"
    # formación / detalle
    if fiber: form=f"{fibers} fibras (SM G-652)"
    elif pairs: form=f"{pairs} pares"+(f" · {awg}" if awg else "")
    elif triads: form=f"{triads} ternas"+(f" · {awg}" if awg else "")
    elif conductors: form=f"{conductors} cond."+(f" · {sec}" if sec else "")
    elif sec: form=sec
    else: form=(awg or "—")
    fire = "Fire Resistant (IEC 60331)" if resistant else ("Fire Retardant (IEC 60332)" if retardant else "")
    return dict(pairs=pairs,triads=triads,conductors=conductors,fibers=fibers,fiber=fiber,
                resistant=resistant,sec=sec,awg=awg,cat=cat,form=form,fire=fire)

def prio_EL(d,a):
    m=d["m"]
    if a["cat"]=="MT — Potencia": return "ALTA","MT — provisión crítica / largo plazo"
    if a["cat"]=="BT — Potencia":
        if m>=3000: return "ALTA",f"BT potencia — tendido largo ({m:,} m)".replace(',','.')
        if m>=1500: return "MEDIA",f"BT potencia ({m:,} m)".replace(',','.')
        return "BAJA","BT potencia — tramo corto"
    if m>=5000: return "MEDIA",f"Tendido muy largo ({m:,} m)".replace(',','.')
    return "BAJA","No prioritario s/premisa EL"

def prio_IN(d,a):
    m=d["m"]; mot=[]
    hi=False
    if a["pairs"]>8: hi=True; mot.append(f"multipar {a['pairs']} pares (>8)")
    if a["fiber"]: hi=True; mot.append(f"fibra óptica ({a['fibers']} f)")
    if a["resistant"] and (a["pairs"]>=2 or a["triads"]>=2): hi=True; mot.append("resistente al fuego (varios pares)")
    if m>=5000: hi=True; mot.append(f"tendido largo ({m:,} m)".replace(',','.'))
    if hi: return "ALTA"," · ".join(mot)
    med=[]
    if a["pairs"]==8 or (a["triads"]>=8): med.append(f"{a['pairs'] or a['triads']} {'pares' if a['pairs'] else 'ternas'}")
    if 2000<=m<5000: med.append(f"{m:,} m".replace(',','.'))
    if a["resistant"]: med.append("resistente al fuego")
    if med: return "MEDIA"," · ".join(med)
    return "BAJA","No prioritario s/premisa IN"

# ---------- estilos ----------
PR_FILL={"ALTA":"C62828","MEDIA":"EF6C00","BAJA":"90A4AE"}
CAT_FILL={"MT — Potencia":"7B1FA2","BT — Potencia":"C62828","BT — Control":"8D6E63",
 "INS — Fibra óptica":"00838F","INS — Analógico":"1565C0","INS — Digital":"5E35B1",
 "INS — Comunicación":"455A64","Otro":"607080"}
thin=Side(style="thin",color="D9D9D9"); BORD=Border(left=thin,right=thin,top=thin,bottom=thin)

def build(spec):
    items=[]
    for fn in SRC[spec]:
        for d in load(fn):
            a=attrs(d)
            p,mot=(prio_EL if spec=="EL" else prio_IN)(d,a)
            items.append({**d,**a,"prio":p,"motivo":mot})
    order={"ALTA":0,"MEDIA":1,"BAJA":2}
    items.sort(key=lambda r:(order[r["prio"]], -r["m"]))
    wb=openpyxl.Workbook()
    ws=wb.active; ws.title="Priorización"
    title=f"RESUMEN DE PRIORIZACIÓN DE CABLES — {'ELECTRICIDAD' if spec=='EL' else 'INSTRUMENTACIÓN'}"
    ws.append([title]); ws.merge_cells("A1:L1")
    ws["A1"].font=Font(bold=True,size=13,color="FFFFFF"); ws["A1"].fill=PatternFill("solid",fgColor="1F3864")
    ws["A1"].alignment=Alignment(horizontal="left",vertical="center"); ws.row_dimensions[1].height=24
    sub=("CPF-2 La Calera II · "+("Req. 5155-00-0000-IG-EL-RI-006-0" if spec=="EL"
         else "Req. 5155-00-0000-IG-IN-RI-032-0 + 5155-00-2000-IG-IN-RI-009-0")
         +f" · {REV} · {FECHA}")
    prem=("Premisa EL: priorizar MT (todos) y BT de potencia con tendidos largos."
          if spec=="EL" else
          "Premisa IN: priorizar multipar >8 pares, tendidos largos, fibras ópticas y "
          "resistentes al fuego (IEC 60331) de varios pares.")
    ws.append([sub]); ws.merge_cells("A2:L2"); ws["A2"].font=Font(size=9,italic=True,color="555555")
    ws.append([prem]); ws.merge_cells("A3:L3"); ws["A3"].font=Font(size=9,bold=True,color="C62828")
    ws.append([])
    hdr=["Prioridad","Req.","Pos","Código","Categoría","Formación / Detalle","Metros",
         "Fuego / Atributos","Motivo prioridad","Plazo (sem)","Proveedor","Estado"]
    ws.append(hdr); hr=ws.max_row
    for j in range(1,len(hdr)+1):
        c=ws.cell(hr,j); c.font=Font(bold=True,color="FFFFFF",size=9)
        c.fill=PatternFill("solid",fgColor="37474F"); c.alignment=Alignment(horizontal="center",wrap_text=True)
        c.border=BORD
    for r in items:
        ws.append([r["prio"],r["req"],r["pos"],r["ident"],r["cat"],r["form"],r["m"],
                   r["fire"],r["motivo"],"","",""])
        rr=ws.max_row
        ws.cell(rr,1).fill=PatternFill("solid",fgColor=PR_FILL[r["prio"]]); ws.cell(rr,1).font=Font(bold=True,color="FFFFFF",size=9)
        ws.cell(rr,1).alignment=Alignment(horizontal="center")
        ws.cell(rr,5).fill=PatternFill("solid",fgColor=CAT_FILL.get(r["cat"],"607080")); ws.cell(rr,5).font=Font(color="FFFFFF",size=8.5)
        ws.cell(rr,7).number_format='#,##0'; ws.cell(rr,7).alignment=Alignment(horizontal="right")
        for j in range(1,len(hdr)+1):
            ws.cell(rr,j).border=BORD
            if ws.cell(rr,j).font.size is None: ws.cell(rr,j).font=Font(size=8.5)
    widths=[9,7,5,20,17,26,9,26,34,9,14,10]
    for j,w in enumerate(widths,1): ws.column_dimensions[chr(64+j)].width=w
    ws.freeze_panes="A6"

    # ---- hoja Resumen ----
    rs=wb.create_sheet("Resumen")
    rs.append(["RESUMEN — "+("ELECTRICIDAD" if spec=="EL" else "INSTRUMENTACIÓN")]); rs.merge_cells("A1:D1")
    rs["A1"].font=Font(bold=True,size=12,color="FFFFFF"); rs["A1"].fill=PatternFill("solid",fgColor="1F3864")
    rs.append([f"{REV} · {FECHA}"]); rs["A2"].font=Font(italic=True,size=9,color="555555"); rs.append([])
    def block(titulo,key):
        rs.append([titulo]); rs.cell(rs.max_row,1).font=Font(bold=True,size=10)
        rs.append(["Grupo","Ítems","Metros"]); h=rs.max_row
        for j in range(1,4): rs.cell(h,j).font=Font(bold=True,color="FFFFFF",size=9); rs.cell(h,j).fill=PatternFill("solid",fgColor="37474F")
        agg={}
        for r in items:
            k=r[key]; agg.setdefault(k,[0,0]); agg[k][0]+=1; agg[k][1]+=r["m"]
        keyorder=(["ALTA","MEDIA","BAJA"] if key=="prio" else sorted(agg,key=lambda k:-agg[k][1]))
        ti=tm=0
        for k in keyorder:
            if k not in agg: continue
            n,mm=agg[k]; ti+=n; tm+=mm
            rs.append([k,n,mm]); rr=rs.max_row; rs.cell(rr,3).number_format='#,##0'
            if key=="prio": rs.cell(rr,1).fill=PatternFill("solid",fgColor=PR_FILL[k]); rs.cell(rr,1).font=Font(bold=True,color="FFFFFF",size=9)
        rs.append(["TOTAL",ti,tm]); rr=rs.max_row
        for j in range(1,4): rs.cell(rr,j).font=Font(bold=True)
        rs.cell(rr,3).number_format='#,##0'; rs.append([])
    block("Por PRIORIDAD","prio")
    block("Por CATEGORÍA / TIPO","cat")
    for j,w in enumerate([22,9,12],1): rs.column_dimensions[chr(64+j)].width=w

    out=os.path.join(GEN,f"Resumen Cables {spec} - {REV} - {FECHA}.xlsx")
    wb.save(out)
    # resumen consola
    from collections import Counter,defaultdict
    pc=Counter(r["prio"] for r in items); mm=defaultdict(int)
    for r in items: mm[r["prio"]]+=r["m"]
    print(f"[{spec}] {len(items)} ítems -> ALTA {pc['ALTA']} ({mm['ALTA']:,} m) · MEDIA {pc['MEDIA']} · BAJA {pc['BAJA']}  -> {os.path.basename(out)}".replace(',','.'))
    return out,items

for spec in ("EL","IN"):
    build(spec)
