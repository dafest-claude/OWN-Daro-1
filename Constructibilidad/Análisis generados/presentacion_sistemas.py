#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Presentacion - Sistemas de Electricidad e Instrumentacion (alcance por proveedor)
CPF2 La Calera II. Mismo esquema que Alcance_Especialidades_CPF2.
Cubre: PMS (ABB) · PCS/SCADA/SIS y PCS remotas (Inauco) · tablero marshalling
ESD/F&G/PSS (HIMA) · config/readecuacion ESD/F&G/PSS (PPSA, a definir).
Complementa con fechas hito + diagrama de flujo de provisiones + Planning de hitos.
Salida: Sistemas_EL_IN_CPF2.pptx
"""
import os
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

BASE=os.path.dirname(os.path.abspath(__file__))
OUT=os.path.join(BASE,"Sistemas_EL_IN_CPF2.pptx")
FLUJO=os.path.join(BASE,"Flujo_Provisiones_CPF2.pdf")
PLAN=os.path.join(BASE,"Planning_PMS_Inauco_CPF2.pdf")

NAVY=RGBColor(0x1F,0x3B,0x63); IN=RGBColor(0x2E,0x75,0xB6); EL=RGBColor(0xC5,0x5A,0x11)
HIMA=RGBColor(0x70,0x30,0xA0); DEF=RGBColor(0x80,0x80,0x80); GREEN=RGBColor(0x54,0x82,0x35)
GOLD=RGBColor(0xBF,0x90,0x00); RED=RGBColor(0xC0,0x00,0x00); GREY=RGBColor(0x59,0x59,0x59)
LGREY=RGBColor(0xF2,0xF2,0xF2); WHITE=RGBColor(0xFF,0xFF,0xFF); LBLUE=RGBColor(0xDC,0xE6,0xF2); DK=RGBColor(0x22,0x22,0x22)

prs=Presentation(); prs.slide_width=Inches(13.333); prs.slide_height=Inches(7.5)
SW,SH=prs.slide_width,prs.slide_height; BLANK=prs.slide_layouts[6]
def slide(): return prs.slides.add_slide(BLANK)
def rect(s,x,y,w,h,color,line=None):
    sp=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,x,y,w,h); sp.fill.solid(); sp.fill.fore_color.rgb=color
    if line is None: sp.line.fill.background()
    else: sp.line.color.rgb=line; sp.line.width=Pt(0.75)
    sp.shadow.inherit=False; return sp
def txt(s,x,y,w,h,text,size=18,color=NAVY,bold=False,align=PP_ALIGN.LEFT,anchor=MSO_ANCHOR.TOP):
    tb=s.shapes.add_textbox(x,y,w,h); tf=tb.text_frame; tf.word_wrap=True; tf.vertical_anchor=anchor
    tf.margin_left=Pt(2); tf.margin_right=Pt(2)
    for i,line in enumerate(text.split('\n')):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph(); p.alignment=align
        r=p.add_run(); r.text=line; f=r.font; f.size=Pt(size); f.bold=bold; f.color.rgb=color; f.name='Calibri'
    return tb
def header(s,title,accent,kick="SISTEMAS EL · IN · ALCANCE POR PROVEEDOR · CPF2 LA CALERA II"):
    rect(s,0,0,SW,Inches(1.12),NAVY); rect(s,0,Inches(1.12),SW,Inches(0.06),accent)
    txt(s,Inches(0.45),Inches(0.12),Inches(11),Inches(0.3),kick,11,RGBColor(0x9D,0xB8,0xD8),True)
    txt(s,Inches(0.45),Inches(0.40),Inches(12.4),Inches(0.7),title,25,WHITE,True)
    txt(s,Inches(10.9),Inches(0.12),Inches(2.2),Inches(0.3),"RFSU 20-DIC-2027",11,GOLD,True,PP_ALIGN.RIGHT)
def footer(s,n):
    txt(s,Inches(0.45),Inches(7.08),Inches(9),Inches(0.3),"Constructibilidad CPF2 · Sistemas EL · IN · Rev. 0 · 13-06-2026",9,GREY)
    txt(s,Inches(12.2),Inches(7.08),Inches(0.9),Inches(0.3),str(n),9,GREY,align=PP_ALIGN.RIGHT)
def table(s,x,y,w,rows,colw,hfill=NAVY,fs=11.5,hfs=11.5,rowh=Inches(0.34),accent_rows=None):
    nr=len(rows); nc=len(rows[0]); gf=s.shapes.add_table(nr,nc,x,y,w,rowh*nr); tb=gf.table
    tb.first_row=False; tb.horz_banding=False; tot=sum(colw)
    for j,cw in enumerate(colw): tb.columns[j].width=Emu(int(w*cw/tot))
    for i,row in enumerate(rows):
        tb.rows[i].height=rowh
        for j,val in enumerate(row):
            cell=tb.cell(i,j); cell.margin_left=Pt(5); cell.margin_right=Pt(4); cell.margin_top=Pt(1); cell.margin_bottom=Pt(1)
            cell.vertical_anchor=MSO_ANCHOR.MIDDLE; p=cell.text_frame.paragraphs[0]; cell.text_frame.word_wrap=True
            r=p.add_run(); r.text=str(val); f=r.font; f.name='Calibri'
            if i==0:
                cell.fill.solid(); cell.fill.fore_color.rgb=hfill; f.size=Pt(hfs); f.bold=True; f.color.rgb=WHITE
            else:
                cell.fill.solid()
                cell.fill.fore_color.rgb=(LBLUE if accent_rows and i in accent_rows else (WHITE if i%2 else LGREY))
                f.size=Pt(fs); f.color.rgb=DK
                if accent_rows and i in accent_rows and j==0: f.bold=True
            if j>0 and len(str(val))<16: p.alignment=PP_ALIGN.CENTER
    return tb
def card(s,x,y,w,h,big,small,accent):
    rect(s,x,y,w,h,WHITE,line=RGBColor(0xCC,0xCC,0xCC)); rect(s,x,y,w,Inches(0.10),accent)
    txt(s,x,y+Inches(0.16),w,Inches(0.58),big,24,accent,True,PP_ALIGN.CENTER,MSO_ANCHOR.MIDDLE)
    txt(s,x,y+h-Inches(0.66),w,Inches(0.6),small,11,GREY,False,PP_ALIGN.CENTER)
def bullets(s,x,y,w,h,items,size=14,gap=6):
    tb=s.shapes.add_textbox(x,y,w,h); tf=tb.text_frame; tf.word_wrap=True
    for i,(lvl,t,c) in enumerate(items):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph(); p.space_after=Pt(gap); p.level=lvl
        r=p.add_run(); r.text=("• " if lvl==0 else "– ")+t
        r.font.size=Pt(size-(1 if lvl else 0)); r.font.color.rgb=c or DK; r.font.name='Calibri'; r.font.bold=(lvl==0 and c is not None)
    return tb

# render PDFs a PNG
import fitz; fitz.TOOLS.mupdf_display_errors(False)
def render(pdf,dpi=170):
    png=os.path.join(BASE,"_tmp_"+os.path.basename(pdf)+".png")
    p=fitz.open(pdf)[0]; p.get_pixmap(dpi=dpi).save(png); return png,p.rect.width,p.rect.height
TMP=[]

# ===== 1 PORTADA =====
s=slide(); rect(s,0,0,SW,SH,NAVY); rect(s,0,Inches(4.4),SW,Inches(0.07),GOLD)
txt(s,Inches(0.7),Inches(1.4),Inches(12),Inches(0.4),"CPF2 LA CALERA II · VACA MUERTA",16,RGBColor(0x9D,0xB8,0xD8),True)
txt(s,Inches(0.7),Inches(1.95),Inches(12),Inches(1.6),"Sistemas de Electricidad e Instrumentación\nAlcance por proveedor",36,WHITE,True)
txt(s,Inches(0.7),Inches(4.55),Inches(12),Inches(0.6),"PMS (ABB) · PCS/SCADA/SIS (Inauco) · marshalling ESD/F&G/PSS (HIMA) · config ESD/F&G/PSS (PPSA)",15,LBLUE)
txt(s,Inches(0.7),Inches(5.2),Inches(12),Inches(0.5),"Fechas hito · diagrama de flujo de provisiones · Planning de hitos · RFSU 20-DIC-2027",15,GOLD,True)
txt(s,Inches(0.7),Inches(6.7),Inches(9),Inches(0.4),"Rev. 0 · 13-06-2026 · Primera revisión para depuración",11,RGBColor(0x9D,0xB8,0xD8))
footer(s,1)

# ===== 2 MAPA DE SISTEMAS Y PROVEEDORES =====
s=slide(); header(s,"Mapa de sistemas y proveedores",NAVY)
rows=[["Sistema","Descripción","Provee","Hito clave"],
 ["PMS","Power Management System — gestión de energía, deslastre/DAG (800XA)","ABB","FAT PMS 29-ENE-27"],
 ["PCS","Control de procesos (DCS) + PCS remotas 7A/7B — Sala INS","Inauco","FAT PCS 19-FEB-27"],
 ["SCADA","Supervisión y adquisición de datos","Inauco","Matriz C&E AGO-27"],
 ["ESD / SIS","Parada de emergencia / sistema instrumentado de seguridad","HIMA (HW) + Inauco","FAT SIS 23-FEB-27"],
 ["F&G","Fuego y gas — detectores AESA + hardware de sistema","AESA + HIMA","Matriz C&E"],
 ["PSS","Sistema de seguridad/parada de procesos","HIMA","con SIS"],
 ["Tablero marshalling","Hardware de ESD / F&G / PSS","HIMA","listo 23-DIC-26"],
 ["Config / readecuación ESD·F&G·PSS","Modificación de provisión y reconfiguración","PPSA [A DEFINIR]","completo al iFAT"]]
table(s,Inches(0.45),Inches(1.45),Inches(12.45),rows,[2.6,5.0,2.2,2.6],fs=11,hfs=11.5,rowh=Inches(0.6),accent_rows=[7,8])
footer(s,2)

# ===== 3 PMS ABB =====
s=slide(); header(s,"PMS — Power Management System (ABB)",EL)
card(s,Inches(0.45),Inches(1.5),Inches(2.9),Inches(1.4),"PMS-001","SE#3",EL)
card(s,Inches(3.55),Inches(1.5),Inches(2.9),Inches(1.4),"PMS-101","SE#4",EL)
card(s,Inches(6.65),Inches(1.5),Inches(2.9),Inches(1.4),"800XA","plataforma ABB",EL)
card(s,Inches(9.75),Inches(1.5),Inches(2.9),Inches(1.4),"OC","4508945953",EL)
bullets(s,Inches(0.45),Inches(3.25),Inches(7.3),Inches(3.4),[
 (0,"Gestión de energía: deslastre de cargas, DAG, control y monitoreo de SS.EE.",EL),
 (0,"Comunicaciones IEC 61850 / PROFINET / MODBUS TCP-IP (interfaz con PCS).",None),
 (0,"Ampliación de la red de control PMS SE#3 ↔ SE#4 (fibra óptica).",None),
 (0,"FAT Maqueta temprana solicitada por el CLIENTE (mitigación de riesgo).",RED),
],size=13)
rows=[["Hito PMS","Fecha"],["KOM","22-MAY-26"],["FAT Maqueta (cliente)","22-25-SEP-26"],
 ["FAT PMS (tableros+sistema)","29-ENE-27"],["FAT Integral Shelter","31-MAR-27"],
 ["Entrega a sitio (con salas)","08-ABR / 14-MAY-27"],["SAT PMS en sitio (HOLD POINT)","06-OCT→05-DIC-27"]]
table(s,Inches(8.0),Inches(3.25),Inches(4.9),rows,[3.4,1.6],hfill=EL,fs=10.5,hfs=10.5,rowh=Inches(0.40))
footer(s,3)

# ===== 4 INAUCO =====
s=slide(); header(s,"PCS · SCADA · SIS — Inauco",IN)
bullets(s,Inches(0.45),Inches(1.5),Inches(7.3),Inches(4.5),[
 (0,"PCS — Sistema de Control de Procesos (DCS) en la Sala INS (Sala 7).",IN),
 (1,"PCS remotas 7A + 7B (tableros).",None),
 (0,"SCADA — supervisión, adquisición y visualización.",IN),
 (0,"SIS — sistema instrumentado de seguridad (hardware HIMA integrado).",IN),
 (0,"Comunicaciones: integración PCS ↔ PMS (Modbus TCP/IP).",None),
 (0,"OC pendiente: el planning se ancla en arranque JUL-2026.",GOLD),
 (0,"Alcance Inauco hasta la ENTREGA; trabajos de campo (SAT/comisionado) los ejecuta AESA.",RED),
],size=13)
rows=[["Hito Inauco","Fecha"],["INICIO (OC + Ing. Rev0 + HW)","01-JUL-26"],["Recepción HW (HIMA + HW)","21-AGO-26"],
 ["FAT PCS + Comunicaciones","27-ENE→19-FEB-27"],["FAT SIS","27-ENE→23-FEB-27"],
 ["Integración PCS↔PMS (Bs.As.)","24-FEB→02-MAR-27"],["Config post-FAT + entrega","FEB→MAR-27"]]
table(s,Inches(8.0),Inches(1.5),Inches(4.9),rows,[3.4,1.8],hfill=IN,fs=10.5,hfs=10.5,rowh=Inches(0.45))
footer(s,4)

# ===== 5 HIMA + PPSA =====
s=slide(); header(s,"ESD · F&G · PSS — HIMA (hardware) + PPSA (config)",HIMA)
rect(s,Inches(0.45),Inches(1.45),Inches(6.05),Inches(0.45),HIMA)
txt(s,Inches(0.55),Inches(1.48),Inches(5.9),Inches(0.4),"HIMA — Tablero marshalling (hardware)",13,WHITE,True)
bullets(s,Inches(0.5),Inches(2.0),Inches(6.0),Inches(3.0),[
 (0,"Provee el hardware de ESD, F&G y PSS (tablero marshalling).",HIMA),
 (0,"Lead de provisión: 25 semanas desde julio-2026.",None),
 (0,"Listo ~23-DIC-2026 → disponible para FAT SIS e iFAT.",None),
 (0,"Se integra con el SIS de Inauco.",None),
],size=12.5)
rect(s,Inches(6.85),Inches(1.45),Inches(6.05),Inches(0.45),DEF)
txt(s,Inches(6.95),Inches(1.48),Inches(5.9),Inches(0.4),"PPSA — Config / readecuación  [A DEFINIR]",13,WHITE,True)
bullets(s,Inches(6.9),Inches(2.0),Inches(6.0),Inches(3.0),[
 (0,"Modificación de la provisión y reconfiguración de ESD, F&G y PSS.",DEF),
 (0,"Alcance y fechas A DEFINIR.",RED),
 (0,"Arranca AGO-2026.",None),
 (0,"Debe confluir y estar COMPLETO (FAT incluido) al momento del iFAT (25-JUN-27).",None),
],size=12.5)
txt(s,Inches(0.45),Inches(5.3),Inches(12.4),Inches(1.2),
    "ESD/F&G/PSS combina hardware HIMA (marshalling) + integración SIS (Inauco) + detectores F&G (AESA) + "
    "configuración/readecuación (PPSA). Todos deben estar operativos para el iFAT.",12.5,GREY)
footer(s,5)

# ===== 6 FECHAS HITO =====
s=slide(); header(s,"Fechas hito clave",GOLD)
rows=[["Hito","Fecha","Proveedor"],
 ["KOM PMS","22-MAY-26","ABB"],["INICIO Inauco (OC pendiente)","01-JUL-26","Inauco"],
 ["Tablero marshalling HIMA listo","23-DIC-26","HIMA"],
 ["FAT PMS","29-ENE-27","ABB"],["FAT PCS / FAT SIS","19-FEB / 23-FEB-27","Inauco"],
 ["Integración PCS↔PMS (Bs.As.)","24-FEB→02-MAR-27","ABB + Inauco"],
 ["FAT Integral PMS","31-MAR-27","ABB"],["Entrega salas SE#4 / SE#3","08-ABR / 14-MAY-27","ABB"],
 ["FAT ESD/F&G/PSS completo (PPSA)","24-JUN-27","PPSA [A DEFINIR]"],
 ["★ iFAT (confluencia total + PCS↔PMS en sitio)","25-JUN→14-JUL-27","Todos"],
 ["SAT PMS en sitio (HOLD POINT)","06-OCT→05-DIC-27","ABB"],
 ["★★★ RFSU","20-DIC-27","—"]]
table(s,Inches(0.45),Inches(1.4),Inches(12.45),rows,[5.4,3.0,3.0],fs=10.5,hfs=11,rowh=Inches(0.42),accent_rows=[10,12])
footer(s,6)

# ===== 7 FLUJO DE PROVISIONES =====
png,iw,ih=render(FLUJO); TMP.append(png)
s=slide(); header(s,"Planning · Diagrama de flujo de provisiones",NAVY)
w=Inches(12.6); h=w*ih/iw
if h>Inches(5.4): h=Inches(5.4); w=h*iw/ih
s.shapes.add_picture(png,int(SW/2-w/2),Inches(1.4),width=int(w),height=int(h))
txt(s,Inches(0.45),Inches(6.95),Inches(12.4),Inches(0.3),"Carriles por proveedor (ABB·Inauco·HIMA·PPSA·AESA): cuándo ocurre cada actividad y dónde confluyen → iFAT → RFSU.",10.5,GREY)
footer(s,7)

# ===== 8 PLANNING HITOS (Gantt) =====
png2,iw2,ih2=render(PLAN); TMP.append(png2)
s=slide(); header(s,"Planning · Hitos importantes (Gantt PMS + Inauco)",NAVY)
w=Inches(12.6); h=w*ih2/iw2
if h>Inches(5.4): h=Inches(5.4); w=h*iw2/ih2
s.shapes.add_picture(png2,int(SW/2-w/2),Inches(1.4),width=int(w),height=int(h))
txt(s,Inches(0.45),Inches(6.95),Inches(12.4),Inches(0.3),"Gantt integrado de provisiones con FATs, integración, iFAT y convergencia a comisionado. Detalle en Planning_PMS_Inauco_CPF2.",10.5,GREY)
footer(s,8)

prs.save(OUT)
for t in TMP:
    try: os.remove(t)
    except OSError: pass
print("OK ->",OUT,"| slides:",len(prs.slides._sldIdLst))
