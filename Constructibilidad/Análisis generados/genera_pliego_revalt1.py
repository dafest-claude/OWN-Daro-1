#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Rev. Alt 1 (variante) sobre Rev.9. Reorienta el pliego a NECESIDAD DE MANO DE OBRA:
transforma el Anexo C en un marco de HH / dotación + cuadros a completar por el oferente
(comparación entre proveedores) + disponibilidad de personal. Mantiene el formato corporativo
y los anexos técnicos (E/I) como base. No borra revisiones anteriores.
Base: Rev.9. Salida: 'Pliego subcontrato E&I - Rev Alt 1 - 2026-07-05 (con Anexos).docx'
"""
import os
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

BASE="/home/user/OWN-Daro-1/Constructibilidad/Análisis generados"
SRC=os.path.join(BASE,"Pliego subcontrato E&I - Rev9 - 2026-07-05 (con Anexos).docx")
OUT=os.path.join(BASE,"Pliego subcontrato E&I - Rev Alt 1 - 2026-07-05 (con Anexos).docx")
doc=Document(SRC)
TBL='Tabla Cuadro'
try: doc.styles[TBL]
except KeyError: TBL='Table Grid'

# revisión -> Alt 1
for p in doc.paragraphs:
    if 'Revisión 9' in p.text:
        for r in p.runs: r.text=r.text.replace('Revisión 9','Revisión Alt 1')

def edit_para(start,new):
    for p in doc.paragraphs:
        if p.text.strip().startswith(start):
            for r in list(p.runs)[1:]: r._element.getparent().remove(r._element)
            if p.runs: p.runs[0].text=new
            else: p.add_run(new)
            return p
def remove_para(start):
    for p in doc.paragraphs:
        if p.text.strip().startswith(start):
            p._element.getparent().remove(p._element); return True
    return False
def remove_table_hdr(h0):
    for t in list(doc.tables):
        if t.rows and t.rows[0].cells[0].text.strip()==h0:
            t._element.getparent().remove(t._element); return
def add_h(txt,lvl=2):
    p=doc.add_heading(txt,level=lvl); p.paragraph_format.keep_with_next=True; return p
def add_p(txt,note=False):
    p=doc.add_paragraph(); r=p.add_run(txt)
    if note: r.italic=True; r.font.size=Pt(9); r.font.color.rgb=RGBColor(0x55,0x55,0x55)
    return p
def add_table(headers,rows,bold_total=False):
    t=doc.add_table(rows=1,cols=len(headers)); t.style=TBL
    for j,h in enumerate(headers):
        c=t.rows[0].cells[j]; c.text=''; r=c.paragraphs[0].add_run(h); r.bold=True; r.font.size=Pt(8.5)
    for ri,row in enumerate(rows):
        cs=t.add_row().cells
        for j,v in enumerate(row):
            cs[j].text=''; rr=cs[j].paragraphs[0].add_run(str(v)); rr.font.size=Pt(8.5)
            if bold_total and ri==len(rows)-1: rr.bold=True
            if j>0 and len(str(v))<18: cs[j].paragraphs[0].alignment=WD_ALIGN_PARAGRAPH.CENTER
    for row in t.rows:
        trPr=row._tr.get_or_add_trPr()
        if trPr.find(qn('w:cantSplit')) is None: trPr.append(OxmlElement('w:cantSplit'))
    return t

# ---- Anexo C: reorientar a mano de obra ----
edit_para('PERSONAL REQUERIDO Y PRECIO MES-HOMBRE','NECESIDAD DE MANO DE OBRA Y COTIZACIÓN')
edit_para('El presente pliego tiene por objeto la contratación de personal especializado',
    "El presente pliego tiene por objeto la contratación de personal especializado de Electricidad e "
    "Instrumentación. La cotización se estructura sobre la NECESIDAD DE MANO DE OBRA: horas-hombre (HH) y "
    "dotación por tarea/especialidad (estimación de referencia de AESA, derivada de las cantidades de los "
    "Anexos E e I) y cuadros a completar por el Oferente para facilitar la comparación entre proveedores.")
remove_para('C.1  Personal requerido'); remove_para('C.1 Personal requerido')
remove_para('El Oferente deberá indicar el precio mes-hombre')
remove_table_hdr('Especialidad / perfil requerido')

# C.1 — Estimación de HH por tarea
add_h("C.1  Estimación de mano de obra por tarea (referencia AESA)")
add_p("Estimación de referencia derivada de las cantidades de los Anexos E e I y de rendimientos O&G. "
      "Supuestos: cuadrilla EL ~3 personas, cuadrilla IN ~2 personas, jornada 8 h.")
add_table(["Especialidad","Tarea","Cantidad de referencia","Días-cuadrilla","HH estimadas"],
 [["Electricidad","Tendido de cables (todos los tipos)","81.231 m","1.026","24.620"],
  ["Electricidad","Conexionado y terminaciones","5.178 puntas","284","6.820"],
  ["Electricidad","Montaje de equipos (motores/celdas; salas vienen montadas)","79 mot. + 6 celdas","89","2.130"],
  ["Electricidad","Pruebas, energización y precomisionado","513 circuitos","153","3.677"],
  ["Instrumentación","Tendido de cables","65.315 m","452","7.226"],
  ["Instrumentación","Conexionado de puntas","15.752 puntas","—","1.974"],
  ["Instrumentación","Montaje de instrumentos (AESA)","1.163 instr.","—","5.815"],
  ["Instrumentación","Montaje de bandejas / conduits","2.838 U","—","1.780"],
  ["Instrumentación","Loop check / prueba de señales","885 lazos","—","1.770"],
  ["TOTAL","Mano de obra estimada E&I","—","—","≈ 55.812 HH"]],bold_total=True)
add_p("Equivalente ≈ 317 persona-mes (1 persona-mes ≈ 176 HH). Valores de referencia, sujetos a los "
      "rendimientos reales y al programa de obra.",note=True)

# C.2 — Dotación y HH-mes por frente
add_h("C.2  Dotación y dedicación mensual (HH-mes) por frente")
add_table(["Frente / especialidad","HH estim.","Duración estim. (meses)","HH-mes","Dotación media (pers.)"],
 [["Electricidad — tendido + conexionado","31.440","6","5.240","~30"],
  ["Electricidad — montaje + pruebas + precom","5.807","4","1.452","~8"],
  ["Instrumentación — tendido + conexionado","9.200","7","1.314","~7,5"],
  ["Instrumentación — montaje instrumentos + bandejas","7.595","6","1.266","~7"],
  ["Instrumentación — loop check / pruebas","1.770","3","590","~3,5"],
  ["PICO SIMULTÁNEO (referencia)","—","—","—","~40-50"]],bold_total=True)
add_p("HH-mes = HH / duración; dotación media = HH-mes / 176. El pico simultáneo se da en la etapa de tendido "
      "(FEB-MAR). Referencia para dimensionar la oferta.",note=True)

# C.3 — Cuadro de cotización (a completar por el Oferente)
add_h("C.3  Cuadro de cotización por perfil — a completar por el Oferente")
PERF=["Supervisor / Capataz de Electricidad","Oficial electricista — montaje y tendido",
 "Oficial electricista — conexionado y terminaciones","Técnico MT certificado (terminaciones y energización)",
 "Técnico de pruebas eléctricas / precomisionado","Supervisor / Capataz de Instrumentación",
 "Oficial instrumentista — montaje y tendido","Técnico instrumentista — conexionado y loop check",
 "Técnico de calibración / banco de pruebas","Ayudantes","Ingeniería / técnico de precom y comisionado",
 "QA/QC — control de calidad E&I","Higiene y Seguridad"]
add_table(["Especialidad / perfil","Precio hora-hombre (USD)","Precio mes-hombre (USD)"],
 [[p,"",""] for p in PERF])
add_p("El Oferente completa el precio por hora-hombre y/o mes-hombre por perfil. La dedicación total (HH-mes) "
      "surge de las estimaciones C.1/C.2 y del programa; permite comparar ofertas de forma homogénea.",note=True)

# C.4 — Disponibilidad de personal del Oferente
add_h("C.4  Disponibilidad de personal — a completar por el Oferente")
add_table(["Especialidad / perfil","Cantidad disponible","Fecha de disponibilidad","Certificaciones / observaciones"],
 [[p,"","",""] for p in PERF])
add_p("El Oferente indica la cantidad de personal por perfil que puede ofrecer y su fecha de disponibilidad, "
      "para evaluar la capacidad de movilización frente al programa (pico MAR-2027).",note=True)

# ---- actualizar cuadro ANEXOS ----
def set_cell(cell,lines,size=9):
    cell.text=''; p0=cell.paragraphs[0]
    for i,ln in enumerate(lines):
        p=p0 if i==0 else cell.add_paragraph(); r=p.add_run(ln); r.font.size=Pt(size)
for t in doc.tables:
    labs=[t.rows[r].cells[0].text.strip().upper() for r in range(len(t.rows))]
    if 'ANEXOS' in labs and 'OBJETIVO' in labs:
        for r in range(len(t.rows)):
            if t.rows[r].cells[0].text.strip().upper()=='ANEXOS':
                set_cell(t.rows[r].cells[1],[
                 "Anexo E — Electricidad (cables, conexionado, tableros, PAT, SPCDA, iluminación y personal).",
                 "Anexo I — Instrumentación (cables, conexionado, instrumentos, tableros Sala INS, cajas de conexión, bandejas/conduits y personal).",
                 "Anexo P — Precomisionado, Comisionado y Puesta en Marcha.",
                 "Anexo C — Necesidad de mano de obra y cotización (HH-mes, dotación y cuadros a completar por el Oferente).",
                 "Anexo R — Matriz de Responsabilidades AESA / Subcontratista.",
                ])
        break

# ---- tipografía Calibri en runs nuevos ----
def fc(run):
    rpr=run._element.get_or_add_rPr(); rf=rpr.find(qn('w:rFonts'))
    if rf is None: rf=OxmlElement('w:rFonts'); rpr.insert(0,rf)
    for a in ('w:asciiTheme','w:hAnsiTheme','w:cstheme'):
        if rf.get(qn(a)) is not None: del rf.attrib[qn(a)]
    rf.set(qn('w:ascii'),'Calibri'); rf.set(qn('w:hAnsi'),'Calibri'); rf.set(qn('w:cs'),'Calibri')
for p in doc.paragraphs:
    for r in p.runs: fc(r)
for t in doc.tables:
    for row in t.rows:
        for c in row.cells:
            for p in c.paragraphs:
                for r in p.runs: fc(r)

doc.save(OUT); print("OK ->",OUT,"| tablas:",len(doc.tables))
