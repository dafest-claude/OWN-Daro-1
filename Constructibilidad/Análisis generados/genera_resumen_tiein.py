#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Análisis del LISTADO DE TIE-IN ELÉCTRICOS (ACAL-00102-LG-E-0004 Rev.0, 29-06-2026).
Genera: informe ejecutivo (PDF), dashboard (XLSX) y cuadro de conteo/seguimiento (HTML).
Discrimina: cantidad total, de sistemas vs no sistemas, por nivel de tensión, y —lo más
importante— cuáles deben ejecutarse durante un PARO de planta.
Rev0 · fecha en el nombre para seguimiento.
"""
import os, json
import fitz
from collections import Counter, defaultdict

fitz.TOOLS.mupdf_display_errors(False)
ENT="/home/user/OWN-Daro-1/Constructibilidad/Datos entrada"
GEN="/home/user/OWN-Daro-1/Constructibilidad/Análisis generados"
SRC=os.path.join(ENT,"ACAL-00102-LG-E-0004-0.pdf")
REV="Rev0"; FECHA="2026-08-18"; DOC="ACAL-00102-LG-E-0004 Rev.0 (29-06-2026)"

# ---------------- 1) extracción ----------------
def extract():
    d=fitz.open(SRC); rows=[]
    for pi in range(d.page_count):
        tabs=d[pi].find_tables()
        if not tabs.tables: continue
        ext=tabs.tables[0].extract(); hdr=None
        for ri,r in enumerate(ext):
            up=[(c or '').strip().upper() for c in r]
            if 'ITEM' in up and 'ETAPA' in up: hdr=ri; header=[(c or '').strip() for c in r]; break
        if hdr is None: continue
        idx={h.upper():j for j,h in enumerate(header) if h}
        gc=lambda r,n:(r[idx[n]] or '').replace('\n',' ').strip() if n in idx else ''
        carry={}
        for r in ext[hdr+1:]:
            if not gc(r,'ITEM'): continue
            rec={k:gc(r,k) for k in ['ITEM','AREA','UBICACION','EQUIPO','TAREA','FUNCIÓN','ETAPA','CABLE','RESPONSABLE']}
            for k in ['AREA','UBICACION','EQUIPO','TAREA','ETAPA','RESPONSABLE']:
                if rec[k]: carry[k]=rec[k]
                else: rec[k]=carry.get(k,'')
            rows.append(rec)
    return rows

def norm_etapa(e):
    e=e.upper().replace(' ','')
    if e=='PARO': return 'PARO'
    if 'POSPARO' in e and 'PREPARO' in e: return 'PREPARO/POSPARO'
    if e=='POSPARO': return 'POSPARO'
    if 'PREPARO' in e: return 'PREPARO/POSPARO'
    return e or '—'

def classify(r):
    fu=r['FUNCIÓN'].lower(); ta=r['TAREA'].lower(); eq=r['EQUIPO'].lower()
    if 'puesta a tierra' in fu or 'pat' in ta or 'malla de pat' in eq: return 'PAT','No sistema'
    if 'potencia mt' in fu: return 'MT','No sistema'
    if fu.startswith('potencia bt') or fu=='potencia bt' or 'iluminación' in fu or 'reposición de reserva' in fu or 'resistencia calefactora' in fu:
        return 'BT','No sistema'
    return 'Señales/Comunic.','Sistema'

rows=extract()
for r in rows:
    r['etapa']=norm_etapa(r['ETAPA']); r['nivel'],r['sistema']=classify(r); r['paro']=r['etapa']=='PARO'
TOTAL=len(rows)
NIVELES=['MT','BT','Señales/Comunic.','PAT']; ETAPAS=['PARO','PREPARO/POSPARO','POSPARO']
cnt_niv={n:sum(1 for r in rows if r['nivel']==n) for n in NIVELES}
cnt_sis={'Sistema':sum(1 for r in rows if r['sistema']=='Sistema'),'No sistema':sum(1 for r in rows if r['sistema']=='No sistema')}
cnt_eta={e:sum(1 for r in rows if r['etapa']==e) for e in ETAPAS}
N_PARO=cnt_eta['PARO']
mx_niv={n:{e:sum(1 for r in rows if r['nivel']==n and r['etapa']==e) for e in ETAPAS} for n in NIVELES}
mx_sis={s:{e:sum(1 for r in rows if r['sistema']==s and r['etapa']==e) for e in ETAPAS} for s in ['Sistema','No sistema']}
cnt_area={a:sum(1 for r in rows if r['AREA']==a) for a in sorted({r['AREA'] for r in rows})}
pct=lambda n:f"{100*n/TOTAL:.0f}%"

# ---------------- 2) PDF informe ejecutivo ----------------
def build_pdf():
    from reportlab.lib.pagesizes import A4
    from reportlab.lib import colors
    from reportlab.lib.units import mm
    from reportlab.pdfgen import canvas
    path=os.path.join(GEN,f"Tie-in Electricos - Informe Ejecutivo - {REV} - {FECHA}.pdf")
    W,H=A4; c=canvas.Canvas(path,pagesize=A4)
    NAVY=colors.HexColor('#1F3864'); RED=colors.HexColor('#C62828'); GREY=colors.HexColor('#555555')
    x0=18*mm; x1=W-18*mm; y=H-20*mm
    c.setFillColor(NAVY); c.rect(0,H-16*mm,W,16*mm,fill=1,stroke=0)
    c.setFillColor(colors.white); c.setFont("Helvetica-Bold",14); c.drawString(x0,H-11*mm,"TIE-IN ELÉCTRICOS — INFORME EJECUTIVO")
    c.setFont("Helvetica",8.5); c.drawRightString(x1,H-11*mm,f"{REV} · {FECHA}")
    y=H-22*mm; c.setFillColor(GREY); c.setFont("Helvetica",8.5)
    c.drawString(x0,y,f"CPF-2 La Calera II · Fuente: {DOC} · Alcance AESA (Electricidad)")
    # KPIs
    y-=9*mm; kpis=[("Total tie-in",str(TOTAL),NAVY),("Requieren PARO",f"{N_PARO}  ({pct(N_PARO)})",RED),
                   ("De sistemas",str(cnt_sis['Sistema']),colors.HexColor('#1565C0')),("No sistemas",str(cnt_sis['No sistema']),colors.HexColor('#2E7D32'))]
    bw=(x1-x0-3*4*mm)/4
    for i,(lab,val,col) in enumerate(kpis):
        bx=x0+i*(bw+4*mm); c.setFillColor(colors.HexColor('#F2F4F7')); c.roundRect(bx,y-16*mm,bw,16*mm,3,fill=1,stroke=0)
        c.setFillColor(col); c.rect(bx,y-16*mm,2.2*mm,16*mm,fill=1,stroke=0)
        c.setFillColor(GREY); c.setFont("Helvetica",7.5); c.drawString(bx+4*mm,y-5*mm,lab)
        c.setFillColor(col); c.setFont("Helvetica-Bold",15); c.drawString(bx+4*mm,y-13*mm,val)
    y-=22*mm
    def tabla(titulo,headers,data,widths,y,hl_row=None):
        c.setFillColor(NAVY); c.setFont("Helvetica-Bold",10); c.drawString(x0,y,titulo); y-=5*mm
        tw=sum(widths)
        c.setFillColor(colors.HexColor('#37474F')); c.rect(x0,y-5*mm,tw,5*mm,fill=1,stroke=0)
        cx=x0; c.setFillColor(colors.white); c.setFont("Helvetica-Bold",7.8)
        for h,w in zip(headers,widths): c.drawString(cx+1.5*mm,y-3.6*mm,h); cx+=w
        y-=5*mm
        for ri,row in enumerate(data):
            bg=colors.HexColor('#FBE9E7') if (hl_row and hl_row(row)) else (colors.HexColor('#F7F9FB') if ri%2 else colors.white)
            c.setFillColor(bg); c.rect(x0,y-5*mm,tw,5*mm,fill=1,stroke=0)
            cx=x0; c.setFont("Helvetica",7.8); c.setFillColor(colors.black)
            for j,(v,w) in enumerate(zip(row,widths)):
                bold = j==0
                c.setFont("Helvetica-Bold" if (bold or (hl_row and hl_row(row))) else "Helvetica",7.8)
                if j==0: c.drawString(cx+1.5*mm,y-3.6*mm,str(v))
                else: c.drawRightString(cx+w-1.5*mm,y-3.6*mm,str(v))
                cx+=w
            y-=5*mm
        c.setStrokeColor(colors.HexColor('#CCCCCC')); c.setLineWidth(0.4); c.rect(x0,y,tw,(len(data)+1)*5*mm,fill=0,stroke=1)
        return y-6*mm
    Wts=[62*mm,26*mm,26*mm]
    y=tabla("1 · Por nivel de tensión / tipo",["Nivel / tipo","Cantidad","%"],
            [[n,cnt_niv[n],pct(cnt_niv[n])] for n in NIVELES],Wts,y)
    y=tabla("2 · De sistemas vs. no sistemas",["Clase","Cantidad","%"],
            [[s,cnt_sis[s],pct(cnt_sis[s])] for s in ['Sistema','No sistema']],Wts,y)
    y=tabla("3 · Por etapa de ejecución",["Etapa","Cantidad","%"],
            [[e,cnt_eta[e],pct(cnt_eta[e])] for e in ETAPAS],Wts,y,hl_row=lambda r:r[0]=='PARO')
    W4=[50*mm,22*mm,34*mm,22*mm]
    y=tabla("4 · ¿Requiere PARO?  ·  cruce por nivel de tensión",["Nivel / tipo","PARO","PREPARO/POSPARO","POSPARO"],
            [[n,mx_niv[n]['PARO'],mx_niv[n]['PREPARO/POSPARO'],mx_niv[n]['POSPARO']] for n in NIVELES],W4,y)
    y=tabla("5 · ¿Requiere PARO?  ·  cruce sistemas / no sistemas",["Clase","PARO","PREPARO/POSPARO","POSPARO"],
            [[s,mx_sis[s]['PARO'],mx_sis[s]['PREPARO/POSPARO'],mx_sis[s]['POSPARO']] for s in ['Sistema','No sistema']],W4,y)
    # conclusiones
    c.setFillColor(NAVY); c.setFont("Helvetica-Bold",10); c.drawString(x0,y,"Conclusiones ejecutivas"); y-=5.5*mm
    concl=[
     f"• {N_PARO} de {TOTAL} tie-in ({pct(N_PARO)}) deben ejecutarse DURANTE el paro de planta — es el condicionante principal de la ventana de parada.",
     f"• De los que requieren paro, {mx_sis['Sistema']['PARO']} son de SISTEMAS (interconexiones de generación GE10/GE11 a TMC/TTLM, PMS y comunicaciones).",
     f"• La potencia BT es mayormente PREPARO/POSPARO ({mx_niv['BT']['PREPARO/POSPARO']} de {cnt_niv['BT']}): puede resolverse fuera de la ventana de paro.",
     f"• Los {cnt_niv['MT']} tie-in de MT (13,2 kV) son todos de paro; la PAT es casi toda preparo/posparo.",
     f"• Concentración por área: "+", ".join(f"{a}={n}" for a,n in cnt_area.items())+".",
    ]
    c.setFont("Helvetica",8.2); c.setFillColor(colors.black)
    for ln in concl:
        c.drawString(x0,y,ln[:150]); y-=4.6*mm
    c.setFont("Helvetica-Oblique",7); c.setFillColor(GREY)
    c.drawString(x0,14*mm,"Clasificación: 'Sistema' = control/comunicaciones/señales/PMS; 'No sistema' = potencia MT/BT y PAT. Nivel de señales/comunicaciones = corriente débil.")
    c.drawString(x0,11*mm,"Etapa PARO = ejecución durante parada de planta. Valores de referencia; verificar contra revisión vigente del listado.")
    c.showPage(); c.save(); return path

# ---------------- 3) XLSX dashboard ----------------
def build_xlsx():
    import openpyxl
    from openpyxl.styles import Font,PatternFill,Alignment,Border,Side
    path=os.path.join(GEN,f"Dashboard Tie-in Electricos - {REV} - {FECHA}.xlsx")
    wb=openpyxl.Workbook()
    thin=Side(style="thin",color="D9D9D9"); B=Border(thin,thin,thin,thin)
    NAVY="1F3864"; HEAD="37474F"; RED="C62828"
    # --- Dashboard ---
    ws=wb.active; ws.title="Dashboard"
    ws["A1"]=f"DASHBOARD TIE-IN ELÉCTRICOS · {REV} · {FECHA}"; ws.merge_cells("A1:F1")
    ws["A1"].font=Font(bold=True,size=13,color="FFFFFF"); ws["A1"].fill=PatternFill("solid",fgColor=NAVY); ws.row_dimensions[1].height=22
    ws["A2"]=f"Fuente: {DOC} · CPF-2 La Calera II · Alcance AESA (Electricidad)"; ws.merge_cells("A2:F2"); ws["A2"].font=Font(italic=True,size=9,color="555555")
    kp=[("Total tie-in",TOTAL,NAVY),("Requieren PARO",N_PARO,RED),("De sistemas",cnt_sis['Sistema'],"1565C0"),("No sistemas",cnt_sis['No sistema'],"2E7D32")]
    for i,(lab,val,col) in enumerate(kp):
        col_l=chr(65+i)
        ws[f"{col_l}4"]=lab; ws[f"{col_l}4"].font=Font(size=9,color="FFFFFF",bold=True); ws[f"{col_l}4"].fill=PatternFill("solid",fgColor=col); ws[f"{col_l}4"].alignment=Alignment(horizontal="center")
        ws[f"{col_l}5"]=val; ws[f"{col_l}5"].font=Font(size=20,bold=True,color=col); ws[f"{col_l}5"].alignment=Alignment(horizontal="center")
    r=7
    def block(title,headers,data,hl=None):
        nonlocal r
        ws.cell(r,1,title).font=Font(bold=True,size=11,color=NAVY); r+=1
        for j,h in enumerate(headers): c=ws.cell(r,1+j,h); c.font=Font(bold=True,color="FFFFFF",size=9); c.fill=PatternFill("solid",fgColor=HEAD); c.alignment=Alignment(horizontal="center" if j else "left"); c.border=B
        r+=1
        for row in data:
            for j,v in enumerate(row):
                c=ws.cell(r,1+j,v); c.border=B; c.alignment=Alignment(horizontal="center" if j else "left")
                if j==0: c.font=Font(bold=True,size=9)
                if hl and hl(row): c.fill=PatternFill("solid",fgColor="FBE9E7"); c.font=Font(bold=True,size=9,color=RED if j else "000000")
            r+=1
        r+=1
    block("Por nivel de tensión / tipo",["Nivel / tipo","Cantidad","%"],[[n,cnt_niv[n],pct(cnt_niv[n])] for n in NIVELES])
    block("Sistemas vs no sistemas",["Clase","Cantidad","%"],[[s,cnt_sis[s],pct(cnt_sis[s])] for s in ['Sistema','No sistema']])
    block("Por etapa de ejecución",["Etapa","Cantidad","%"],[[e,cnt_eta[e],pct(cnt_eta[e])] for e in ETAPAS],hl=lambda row:row[0]=='PARO')
    block("¿Requiere PARO? · por nivel",["Nivel / tipo","PARO","PREPARO/POSPARO","POSPARO","Total"],
          [[n,mx_niv[n]['PARO'],mx_niv[n]['PREPARO/POSPARO'],mx_niv[n]['POSPARO'],cnt_niv[n]] for n in NIVELES])
    block("¿Requiere PARO? · sistemas / no sistemas",["Clase","PARO","PREPARO/POSPARO","POSPARO"],
          [[s,mx_sis[s]['PARO'],mx_sis[s]['PREPARO/POSPARO'],mx_sis[s]['POSPARO']] for s in ['Sistema','No sistema']])
    block("Por área",["Área","Cantidad"],[[a,n] for a,n in cnt_area.items()])
    for col,w in zip("ABCDE",[34,16,18,12,10]): ws.column_dimensions[col].width=w
    # --- Listado ---
    ls=wb.create_sheet("Listado")
    HED=["Ítem","Área","Ubicación","Equipo","Tarea","Función","Cable","Nivel","Sistema","Etapa","¿PARO?","Estado","Fecha ejec.","Observaciones"]
    ls.append(HED)
    for j,_ in enumerate(HED,1): c=ls.cell(1,j); c.font=Font(bold=True,color="FFFFFF",size=9); c.fill=PatternFill("solid",fgColor=HEAD); c.alignment=Alignment(horizontal="center",wrap_text=True); c.border=B
    NIVFILL={'MT':'7B1FA2','BT':'C62828','Señales/Comunic.':'1565C0','PAT':'8D6E63'}
    for r_ in rows:
        ls.append([r_['ITEM'],r_['AREA'],r_['UBICACION'],r_['EQUIPO'],r_['TAREA'],r_['FUNCIÓN'],r_['CABLE'],
                   r_['nivel'],r_['sistema'],r_['etapa'],'SÍ' if r_['paro'] else 'no','','',''])
        rr=ls.max_row
        ls.cell(rr,8).fill=PatternFill("solid",fgColor=NIVFILL.get(r_['nivel'],'607080')); ls.cell(rr,8).font=Font(color="FFFFFF",size=8)
        if r_['paro']:
            ls.cell(rr,11).fill=PatternFill("solid",fgColor=RED); ls.cell(rr,11).font=Font(bold=True,color="FFFFFF")
        for j in range(1,len(HED)+1): ls.cell(rr,j).border=B;
        if ls.cell(rr,1).font.size is None: pass
    for col,w in zip("ABCDEFGHIJKLMN",[7,7,16,15,26,30,16,15,11,16,8,12,12,26]): ls.column_dimensions[col].width=w
    ls.freeze_panes="A2"; ls.auto_filter.ref=f"A1:N{ls.max_row}"
    wb.save(path); return path

# ---------------- 4) HTML cuadro de conteo / seguimiento ----------------
def build_html():
    path=os.path.join(GEN,f"Dashboard Tie-in Electricos - {REV} - {FECHA}.html")
    data=[{"item":r['ITEM'],"area":r['AREA'],"ub":r['UBICACION'],"eq":r['EQUIPO'],"tarea":r['TAREA'],
           "func":r['FUNCIÓN'],"cable":r['CABLE'],"niv":r['nivel'],"sis":r['sistema'],"eta":r['etapa'],"paro":r['paro']} for r in rows]
    payload=json.dumps(data,ensure_ascii=False)
    html=HTML_TMPL.replace("__REV__",REV).replace("__FECHA__",FECHA).replace("__DOC__",DOC)\
        .replace("__TOTAL__",str(TOTAL)).replace("__NPARO__",str(N_PARO)).replace("__PCTPARO__",pct(N_PARO))\
        .replace("__NSIS__",str(cnt_sis['Sistema'])).replace("__NNOSIS__",str(cnt_sis['No sistema']))\
        .replace("__DATA__",payload)
    open(path,"w",encoding="utf-8").write(html); return path

HTML_TMPL=r"""<!doctype html><html lang="es"><head><meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1"><title>Tie-in Eléctricos · Cuadro de conteo</title>
<style>
:root{--bg:#f4f6f8;--card:#fff;--ink:#1f2933;--muted:#607080;--line:#e2e8f0;--red:#C62828;--blue:#1565C0;--green:#2E7D32;--navy:#1F3864;--pur:#7B1FA2;--brown:#8D6E63;--shadow:0 1px 3px rgba(0,0,0,.08)}
@media(prefers-color-scheme:dark){:root{--bg:#0f141a;--card:#1a222c;--ink:#e6edf3;--muted:#93a1b0;--line:#2a343f;--shadow:0 1px 3px rgba(0,0,0,.5)}}
*{box-sizing:border-box}body{margin:0;font-family:-apple-system,Segoe UI,Roboto,Helvetica,Arial,sans-serif;background:var(--bg);color:var(--ink);font-size:14px}
header{background:var(--navy);color:#fff;padding:16px 22px}h1{margin:0;font-size:19px}.sub{opacity:.85;font-size:12px;margin-top:3px}
.wrap{max-width:1300px;margin:0 auto;padding:16px 22px 60px}
.kpis{display:flex;gap:12px;flex-wrap:wrap;margin:14px 0}.kpi{flex:1;min-width:150px;background:var(--card);border:1px solid var(--line);border-left:4px solid var(--navy);border-radius:10px;padding:12px 14px;box-shadow:var(--shadow)}
.kpi.red{border-left-color:var(--red)}.kpi.blue{border-left-color:var(--blue)}.kpi.green{border-left-color:var(--green)}
.kpi .l{font-size:11px;color:var(--muted);text-transform:uppercase;letter-spacing:.04em}.kpi .v{font-size:24px;font-weight:800;margin-top:2px}
.grid{display:grid;grid-template-columns:repeat(auto-fit,minmax(280px,1fr));gap:14px;margin:10px 0 4px}
.card{background:var(--card);border:1px solid var(--line);border-radius:10px;padding:14px;box-shadow:var(--shadow)}
.card h3{margin:0 0 10px;font-size:13px}
.brow{display:flex;align-items:center;gap:8px;margin:5px 0;font-size:12.5px}.brow .nm{width:120px;flex-shrink:0}.brow .barwrap{flex:1;background:var(--line);border-radius:5px;height:16px;overflow:hidden}.brow .bar{height:100%;border-radius:5px}.brow .n{width:56px;text-align:right;font-weight:700}
.toolbar{display:flex;gap:8px;flex-wrap:wrap;align-items:center;margin:16px 0 8px}
.chip{border:1.5px solid var(--line);background:var(--card);border-radius:20px;padding:5px 12px;font-size:12px;cursor:pointer;font-weight:600;color:var(--ink)}.chip.on{background:var(--navy);color:#fff;border-color:var(--navy)}
table{width:100%;border-collapse:collapse;background:var(--card);border-radius:10px;overflow:hidden;box-shadow:var(--shadow);font-size:12px}
th,td{padding:6px 8px;border-bottom:1px solid var(--line);text-align:left}th{background:rgba(127,127,127,.08);position:sticky;top:0;font-size:10.5px;text-transform:uppercase;color:var(--muted);cursor:pointer}
.tag{color:#fff;border-radius:4px;padding:1px 6px;font-size:10px;font-weight:700;white-space:nowrap}
.paro{background:var(--red);color:#fff;font-weight:800;text-align:center}
input[type=text]{padding:6px 9px;border:1px solid var(--line);border-radius:7px;background:var(--card);color:var(--ink);font-size:12px;min-width:180px}
.count{font-size:12px;color:var(--muted);margin:6px 0}
</style></head><body>
<header><h1>Tie-in Eléctricos — Cuadro de conteo y seguimiento</h1>
<div class="sub">CPF-2 La Calera II · Fuente __DOC__ · __REV__ · __FECHA__</div></header>
<div class="wrap">
 <div class="kpis">
  <div class="kpi"><div class="l">Total tie-in</div><div class="v">__TOTAL__</div></div>
  <div class="kpi red"><div class="l">Requieren PARO</div><div class="v">__NPARO__ <span style="font-size:14px">(__PCTPARO__)</span></div></div>
  <div class="kpi blue"><div class="l">De sistemas</div><div class="v">__NSIS__</div></div>
  <div class="kpi green"><div class="l">No sistemas</div><div class="v">__NNOSIS__</div></div>
 </div>
 <div class="grid">
  <div class="card"><h3>Por nivel de tensión / tipo</h3><div id="c-niv"></div></div>
  <div class="card"><h3>Por etapa (¿durante paro?)</h3><div id="c-eta"></div></div>
  <div class="card"><h3>Sistemas vs no sistemas</h3><div id="c-sis"></div></div>
  <div class="card"><h3>¿PARO? · por nivel</h3><div id="c-mx"></div></div>
 </div>
 <div class="toolbar">
  <b style="font-size:12px">Filtrar:</b>
  <span class="chip on" data-f="all">Todos</span>
  <span class="chip" data-f="paro">Solo PARO</span>
  <span class="chip" data-f="Sistema">Sistemas</span>
  <span class="chip" data-f="No sistema">No sistemas</span>
  <span class="chip" data-f="MT">MT</span><span class="chip" data-f="BT">BT</span>
  <span class="chip" data-f="Señales/Comunic.">Señales/Comunic.</span><span class="chip" data-f="PAT">PAT</span>
  <input type="text" id="q" placeholder="buscar equipo, función, cable…">
 </div>
 <div class="count" id="count"></div>
 <div style="max-height:560px;overflow:auto;border-radius:10px">
 <table id="tbl"><thead><tr>
  <th>Ítem</th><th>Área</th><th>Ubicación</th><th>Equipo</th><th>Función</th><th>Cable</th><th>Nivel</th><th>Sistema</th><th>Etapa</th><th>¿PARO?</th>
 </tr></thead><tbody></tbody></table></div>
 <div style="font-size:11px;color:var(--muted);margin-top:10px">Etapa <b>PARO</b> = ejecución durante parada de planta. Clasificación: Sistema = control/comunicaciones/señales/PMS; No sistema = potencia MT/BT y PAT. Valores de referencia, sujetos a la revisión vigente del listado.</div>
</div>
<script>
const DATA=__DATA__;
const NIVC={'MT':'var(--pur)','BT':'var(--red)','Señales/Comunic.':'var(--blue)','PAT':'var(--brown)'};
const ETAC={'PARO':'var(--red)','PREPARO/POSPARO':'var(--green)','POSPARO':'#EF6C00'};
function bars(el,pairs,colf){const mx=Math.max(...pairs.map(p=>p[1]),1);
 el.innerHTML=pairs.map(([k,n])=>`<div class="brow"><span class="nm">${k}</span><span class="barwrap"><span class="bar" style="width:${100*n/mx}%;background:${colf(k)}"></span></span><span class="n">${n}</span></div>`).join('');}
const by=(f)=>{const m={};DATA.forEach(d=>{const k=f(d);m[k]=(m[k]||0)+1});return m;}
const niv=by(d=>d.niv), eta=by(d=>d.eta), sis=by(d=>d.sis);
bars(document.getElementById('c-niv'),['MT','BT','Señales/Comunic.','PAT'].map(k=>[k,niv[k]||0]),k=>NIVC[k]);
bars(document.getElementById('c-eta'),['PARO','PREPARO/POSPARO','POSPARO'].map(k=>[k,eta[k]||0]),k=>ETAC[k]);
bars(document.getElementById('c-sis'),['Sistema','No sistema'].map(k=>[k,sis[k]||0]),k=>k==='Sistema'?'var(--blue)':'var(--green)');
const mx=['MT','BT','Señales/Comunic.','PAT'].map(n=>[n,DATA.filter(d=>d.niv===n&&d.paro).length]);
bars(document.getElementById('c-mx'),mx,k=>NIVC[k]);
let flt='all',q='';
function render(){
 const tb=document.querySelector('#tbl tbody');
 const rows=DATA.filter(d=>{
   if(flt==='paro'&&!d.paro)return false;
   if(['Sistema','No sistema'].includes(flt)&&d.sis!==flt)return false;
   if(['MT','BT','Señales/Comunic.','PAT'].includes(flt)&&d.niv!==flt)return false;
   if(q){const s=(d.eq+' '+d.func+' '+d.cable+' '+d.ub+' '+d.tarea).toLowerCase();if(!s.includes(q))return false;}
   return true;});
 document.getElementById('count').textContent=`${rows.length} de ${DATA.length} tie-in`;
 tb.innerHTML=rows.map(d=>`<tr>
  <td><b>${d.item}</b></td><td>${d.area}</td><td>${d.ub}</td><td>${d.eq}</td><td>${d.func}</td><td>${d.cable}</td>
  <td><span class="tag" style="background:${NIVC[d.niv]}">${d.niv}</span></td>
  <td>${d.sis}</td><td>${d.eta}</td><td class="${d.paro?'paro':''}">${d.paro?'SÍ':'no'}</td></tr>`).join('');
}
document.querySelectorAll('.chip').forEach(ch=>ch.onclick=()=>{document.querySelectorAll('.chip').forEach(x=>x.classList.remove('on'));ch.classList.add('on');flt=ch.dataset.f;render();});
document.getElementById('q').addEventListener('input',e=>{q=e.target.value.toLowerCase().trim();render();});
render();
</script></body></html>"""

p1=build_pdf(); p2=build_xlsx(); p3=build_html()
print("TOTAL",TOTAL,"| PARO",N_PARO,"| Sistema",cnt_sis['Sistema'],"| niveles",cnt_niv)
print("OK\n ",p1,"\n ",p2,"\n ",p3)
