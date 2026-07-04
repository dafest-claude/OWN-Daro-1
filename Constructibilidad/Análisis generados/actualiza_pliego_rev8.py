#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Rev.8 sobre la Rev.7. Cambios:
- Uniforma tipografía a Calibri 11 (quita ambigüedad 'Calibri' vs 'Calibri (cuerpo)') e interlineado.
- Anexo C: quita la planilla de precios unitarios/total y la reemplaza por una tabla de PERSONAL
  requerido por especialidad con PRECIO MES-HOMBRE (la estimación del trabajo total se toma de las
  cantidades de referencia de los Anexos E e I).
- Agrega Anexo R — Matriz de Responsabilidades AESA / Subcontratista (adaptada de otro proyecto).
- Revisión 7 -> 8, fecha 04-07-2026.
Base: Análisis generados/Pliego subcontrato E&I - Rev7 - 2026-07-02 (con Anexos).docx
Salida: Análisis generados/Pliego subcontrato E&I - Rev8 - 2026-07-04 (con Anexos).docx
"""
import os
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

BASE="/home/user/OWN-Daro-1/Constructibilidad/Análisis generados"
SRC=os.path.join(BASE,"Pliego subcontrato E&I - Rev7 - 2026-07-02 (con Anexos).docx")
OUT=os.path.join(BASE,"Pliego subcontrato E&I - Rev8 - 2026-07-04 (con Anexos).docx")
NAVY=RGBColor(0x1F,0x3B,0x63)
doc=Document(SRC)
TBL='Tabla Cuadro'
try: doc.styles[TBL]
except KeyError: TBL='Table Grid'

# ---------- 1) revisión / fecha ----------
for p in doc.paragraphs:
    if 'Revisión 7' in p.text:
        for r in p.runs:
            r.text=r.text.replace('Revisión 7','Revisión 8').replace('02-07-2026','04-07-2026')

# ---------- 2) helpers ----------
def edit_para(match_start, new_text):
    for p in doc.paragraphs:
        if p.text.strip().startswith(match_start):
            # dejar un solo run con el nuevo texto (preserva estilo del párrafo)
            for r in list(p.runs)[1:]: r._element.getparent().remove(r._element)
            if p.runs: p.runs[0].text=new_text
            else: p.add_run(new_text)
            return p
def remove_para(match_start):
    for p in doc.paragraphs:
        if p.text.strip().startswith(match_start):
            p._element.getparent().remove(p._element); return True
    return False
def remove_tables_with_header(h0):
    for t in list(doc.tables):
        if t.rows and t.rows[0].cells[0].text.strip()==h0:
            t._element.getparent().remove(t._element)
def add_h(txt,lvl=2): return doc.add_heading(txt,level=lvl)
def add_p(txt,note=False):
    p=doc.add_paragraph(); r=p.add_run(txt)
    if note: r.italic=True; r.font.size=Pt(9); r.font.color.rgb=RGBColor(0x55,0x55,0x55)
    return p
def add_table(headers,rows):
    t=doc.add_table(rows=1,cols=len(headers)); t.style=TBL
    for j,h in enumerate(headers):
        c=t.rows[0].cells[j]; c.text=''; r=c.paragraphs[0].add_run(h); r.bold=True; r.font.size=Pt(8.5)
    for row in rows:
        cs=t.add_row().cells
        for j,v in enumerate(row):
            cs[j].text=''; rr=cs[j].paragraphs[0].add_run(str(v)); rr.font.size=Pt(8.5)
            if j>0 and len(str(v))<16: cs[j].paragraphs[0].alignment=WD_ALIGN_PARAGRAPH.CENTER
    return t
def cover(title,subtitle):
    p=doc.add_paragraph(); p.add_run().add_break(WD_BREAK.PAGE)
    for _ in range(7): doc.add_paragraph()
    hp=doc.add_heading(title,level=1); hp.alignment=WD_ALIGN_PARAGRAPH.CENTER
    for r in hp.runs: r.font.size=Pt(28); r.font.color.rgb=NAVY
    ps=doc.add_paragraph(); ps.alignment=WD_ALIGN_PARAGRAPH.CENTER
    rs=ps.add_run(subtitle); rs.font.size=Pt(13); rs.font.color.rgb=RGBColor(0x55,0x55,0x55)
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

# ---------- 3) Anexo C: sacar precios, poner personal / mes-hombre ----------
edit_para('PLANILLA DE CÓMPUTO','PERSONAL REQUERIDO Y PRECIO MES-HOMBRE')
edit_para('El Oferente deberá completar los precios unitarios',
    "El presente pliego tiene por objeto la contratación de personal especializado en las especialidades de "
    "Electricidad e Instrumentación. Por tal motivo, la cotización NO se realiza por precio unitario y total de "
    "obra, sino por el valor MES-HOMBRE de cada especialidad/perfil requerido. La estimación del trabajo total "
    "se obtiene a partir de las CANTIDADES DE REFERENCIA y sus especificidades detalladas en los Anexos E "
    "(Electricidad) e I (Instrumentación).")
remove_para('C.1  Electricidad'); remove_para('C.1 Electricidad')
remove_para('C.2  Instrumentación'); remove_para('C.2 Instrumentación')
remove_para('Nota: la planilla es indicativa')
remove_tables_with_header('Ítem de obra')
# nuevo contenido (se agrega al final = justo tras el intro del Anexo C)
add_h("C.1  Personal requerido por especialidad — precio mes-hombre",2)
add_table(["Especialidad / perfil requerido","Unidad","Precio mes-hombre (USD)"],
 [["Supervisor / Capataz de Electricidad","mes-hombre",""],
  ["Oficial electricista — montaje y tendido","mes-hombre",""],
  ["Oficial electricista — conexionado y terminaciones","mes-hombre",""],
  ["Técnico MT certificado — terminaciones y energización","mes-hombre",""],
  ["Técnico de pruebas eléctricas (megger / precomisionado)","mes-hombre",""],
  ["Supervisor / Capataz de Instrumentación","mes-hombre",""],
  ["Oficial instrumentista — montaje y tendido","mes-hombre",""],
  ["Técnico instrumentista — conexionado y loop check","mes-hombre",""],
  ["Técnico de calibración / banco de pruebas","mes-hombre",""],
  ["Ayudantes","mes-hombre",""],
  ["Ingeniería / técnico de precomisionado y comisionado","mes-hombre",""],
  ["QA/QC — control de calidad E&I","mes-hombre",""],
  ["Higiene y Seguridad","mes-hombre",""]])
add_p("El Oferente deberá indicar el precio mes-hombre por cada especialidad. La dedicación (cantidad de "
      "meses-hombre por perfil) se estima a partir del programa de trabajo y de las cantidades de referencia "
      "de los Anexos E e I (metros de cable, puntas, instrumentos, bandejas, PAT/SPCDA, iluminación, etc.).",note=True)

# ---------- 4) Anexo R: Matriz de Responsabilidades ----------
cover("ANEXO R","MATRIZ DE RESPONSABILIDADES AESA / SUBCONTRATISTA")
add_p("Matriz de responsabilidades entre AESA (Comitente/EPC) y el Subcontratista para la provisión de personal "
      "especializado de Electricidad e Instrumentación (montaje, instalación, conexionado, precomisionado, "
      "comisionado y puesta en marcha). Adaptada al alcance del presente pliego. La marca X indica el responsable "
      "de cada concepto; las observaciones precisan el criterio de reparto.")
MATRIZ=[
 ["1","Gerenciamiento global del proyecto","X","",""],
 ["2","Ingeniería de detalle y planos 'For Construction'","X","",""],
 ["3","Provisión de materiales principales (cables, bandejas, instrumentos, tableros, PAT/SPCDA, iluminación)","X","","Según listas de materiales"],
 ["4","Consumibles y misceláneos de montaje (terminales menores, ferrules, precintos, tornillería)","","X","Salvo indicación en contrario"],
 ["5","Herramientas y equipos de montaje","","X",""],
 ["6","Instrumental de medición y prueba (megger, multímetro, calibrador, comunicador HART, banco)","","X","Calibrado y certificado"],
 ["7","Mano de obra calificada (montaje, tendido, conexionado, precom, comisionado)","","X","CS, seguros, ART s/normativa"],
 ["8","Supervisión técnica de las tareas","","X",""],
 ["9","Movilización y desmovilización de personal y equipos","","X",""],
 ["10","Andamios y plataformas para trabajo en altura","X","","Según disponibilidad de obra"],
 ["11","Izaje / grúas para montaje de equipos pesados (tableros, motores, celdas)","X","",""],
 ["12","Energía, aire comprimido y servicios en obrador y frentes","X","",""],
 ["13","Permisos de trabajo (PTW) y habilitación de tareas","X","","Emisión AESA; cumplimiento Subcontratista"],
 ["14","EPP para todo el personal","","X","S/normativa vigente"],
 ["15","Vehículos del subcontratista y choferes","","X",""],
 ["16","Cumplimiento de Seguridad, Higiene y Medio Ambiente","","X","Supervisión AESA / Cliente"],
 ["17","Plan de calidad, ITP, protocolos y registros","","X","Aprobación AESA"],
 ["18","Inspección y liberación de puntos de espera (hold points)","X","","Ejecución y autocontrol Subcontratista"],
 ["19","Documentación conforme a obra (as-built) de sus tareas","","X",""],
 ["20","Partes diarios, reportes de avance y certificación","","X",""],
 ["21","Gestión administrativa de documentación laboral/legal","","X",""],
 ["22","Alojamiento, comida y traslado del personal (campamento)","X","","Campamento de obra"],
 ["23","Servicio médico en obra y emergencias","X","",""],
 ["24","Gestión de residuos de sus tareas","","X","Dispone donde AESA indique"],
 ["25","Instrumental y personal para precom/comisionado y pruebas","","X","Coordinación con AESA / vendors"],
 ["26","Coordinación con vendors (ABB, Inauco, HIMA) para FAT/SAT/iFAT","X","",""],
 ["27","Seguros del personal y equipos (ART, RC)","","X",""],
]
add_table(["Nº","Concepto","AESA","Subcon.","Observaciones"],MATRIZ)
add_p("Matriz preliminar; el reparto definitivo de responsabilidades se ajustará en la negociación del "
      "subcontrato y en el Kick-off Meeting.",note=True)

# ---------- 5) actualizar cuadro ANEXOS del RESUMEN ----------
def set_cell(cell,lines,bold=False,size=9):
    cell.text=''; p0=cell.paragraphs[0]
    for i,ln in enumerate(lines):
        p=p0 if i==0 else cell.add_paragraph()
        r=p.add_run(ln); r.font.size=Pt(size); r.bold=bold
for t in doc.tables:
    labels=[t.rows[r].cells[0].text.strip().upper() for r in range(len(t.rows))]
    if 'ANEXOS' in labels and 'OBJETIVO' in labels:
        for r in range(len(t.rows)):
            if t.rows[r].cells[0].text.strip().upper()=='ANEXOS':
                set_cell(t.rows[r].cells[1],[
                 "Anexo E — Electricidad (cables, conexionado, tableros, PAT, SPCDA, iluminación y personal).",
                 "Anexo I — Instrumentación (cables, conexionado, instrumentos, tableros Sala INS, cajas de conexión, bandejas/conduits y personal).",
                 "Anexo P — Precomisionado, Comisionado y Puesta en Marcha (instrumental y misceláneos).",
                 "Anexo C — Personal requerido y precio mes-hombre.",
                 "Anexo R — Matriz de Responsabilidades AESA / Subcontratista.",
                ],size=9)
        break

# ---------- 6) uniformar tipografía e interlineado ----------
def force_calibri(run):
    rpr=run._element.get_or_add_rPr()
    rf=rpr.find(qn('w:rFonts'))
    if rf is None:
        rf=OxmlElement('w:rFonts'); rpr.insert(0,rf)
    for a in ('w:asciiTheme','w:hAnsiTheme','w:cstheme','w:eastAsiaTheme'):
        if rf.get(qn(a)) is not None: del rf.attrib[qn(a)]
    rf.set(qn('w:ascii'),'Calibri'); rf.set(qn('w:hAnsi'),'Calibri'); rf.set(qn('w:cs'),'Calibri')
# estilo Normal base
nst=doc.styles['Normal']; nst.font.name='Calibri'; nst.font.size=Pt(11)
nst.paragraph_format.line_spacing=1.15; nst.paragraph_format.space_after=Pt(6)
# todos los runs del cuerpo
for p in doc.paragraphs:
    for r in p.runs: force_calibri(r)
    if p.style.name=='Normal':
        p.paragraph_format.line_spacing=1.15; p.paragraph_format.space_after=Pt(6)
# runs en tablas
for t in doc.tables:
    for row in t.rows:
        for c in row.cells:
            for p in c.paragraphs:
                for r in p.runs: force_calibri(r)

doc.save(OUT)
print("OK ->",OUT)
print("tablas:",len(doc.tables))
