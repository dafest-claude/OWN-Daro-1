#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
MATRIZ DE INTERFASES / RESPONSABILIDADES — Provisiones críticas CPF-2 (Rev0, 2026-08-21).
Provisiones: ABB (Salas Eléctricas SE#3/SE#4 y PMS), Inauco (PCS/SCADA), HIMA (tableros de
marshalling) — todas ADJUDICADAS con OC/planificación vigente — y HIMA (adecuación y
configuración SIS) que está EN ADJUDICACIÓN.
Esquema RACI ampliado: E=Ejecuta · A=Aprueba · R=Revisa/comenta · P=Participa/Asiste · I=Informado.
Actores: Proveedor · AESA-Proyecto/Contrato · AESA-Ingeniería · AESA-Precom&Comisionado · Cliente-PPSA.
Genera: Excel (matriz editable) + PDF (informe).
"""
import os, textwrap
GEN="/home/user/OWN-Daro-1/Constructibilidad/Análisis generados"
REV="Rev0"; FECHA="2026-08-21"

ACT_COLS=["PROVEEDOR","AESA — Proyecto / Contrato","AESA — Ingeniería","AESA — Precom. y Comisionado","CLIENTE — PPSA"]
ACT_SHORT=["Proveedor","AESA\nProyecto/\nContrato","AESA\nIngeniería","AESA\nPrecom. y\nComis.","Cliente\nPPSA"]
RACI={'E':('1F3864','Ejecuta / responsable primario'),'A':('C62828','Aprueba'),
      'R':('EF6C00','Revisa / comenta'),'P':('2E7D32','Participa / Asiste'),'I':('90A4AE','Informado')}

ACTIVITIES=[
 ("Contrato","Gestión de la OC / contrato y coordinación",           ['P','E','I','I','I']),
 ("Ingeniería","Emisión de documentación (IB / ID / IC)",            ['E','I','R','I','R']),
 ("Ingeniería","Revisión y aprobación de documentación",             ['P','I','A','I','A']),
 ("Ingeniería","Freezing / hold points · cierre de ingeniería",      ['E','P','A','I','I']),
 ("Procura / Fabricación","Procura y fabricación (seguimiento)",     ['E','R','I','I','I']),
 ("Procura / Fabricación","Certificación de hitos de pago / avance", ['P','E','I','I','A']),
 ("Pruebas de fábrica (FAT)","Procedimientos de FAT/pruebas (emisión y aprobación)",['E','I','A','P','A']),
 ("Pruebas de fábrica (FAT)","Ejecución de FAT / pruebas en fábrica (asistencia)",  ['E','P','P','P','P']),
 ("Pruebas de fábrica (FAT)","iFAT / pruebas de integración (donde aplica)",        ['E','P','P','P','P']),
 ("Despacho / Sitio","Despacho y entrega en sitio",                  ['E','R','I','I','I']),
 ("Despacho / Sitio","Montaje / integración en sitio",               ['P','E','R','I','I']),
 ("Precom. / Puesta en marcha","SAT / pruebas en sitio",             ['P','P','P','E','P']),
 ("Precom. / Puesta en marcha","Precomisionado",                     ['P','P','P','E','I']),
 ("Precom. / Puesta en marcha","Comisionado / Puesta en marcha (PEM)",['P','P','P','E','A']),
 ("Precom. / Puesta en marcha","Capacitación (operación / mantenimiento)",['E','P','I','P','P']),
 ("Cierre","Documentación final / DataBook / As-built",              ['E','I','A','P','A']),
]
ADJ_ROW=("Adjudicación","Proceso de adjudicación / cierre de contrato",['P','E','R','I','A'])

PROV=[
 ("ABB — Salas Eléctricas SE#3 y SE#4","ABB","Adjudicado — OC vigente + planificación","C62828",False),
 ("ABB — PMS (Power Management System)","ABB","Adjudicado — OC vigente + planificación","EF6C00",False),
 ("Inauco — PCS / SCADA / Comunicaciones","Inauco","Adjudicado — Obra 129-008","1565C0",False),
 ("HIMA — Tableros de marshalling (ESD / F&G / PSS)","HIMA","Adjudicado","6A1B9A",False),
 ("HIMA — Adecuación y configuración SIS","HIMA","EN ADJUDICACIÓN","00838F",True),
]
NOTAS=[
 "Alcance de la matriz: desde la adjudicación/OC en adelante. Para las provisiones adjudicadas (ABB Salas y PMS, Inauco, HIMA marshalling) las etapas comercial/adjudicación ya están cerradas; la matriz aplica a ingeniería, fabricación, FAT, sitio y puesta en marcha.",
 "HIMA — Adecuación y configuración SIS está EN ADJUDICACIÓN: se incorpora la etapa de adjudicación (Proyecto ejecuta, Ingeniería evalúa técnica, Cliente aprueba).",
 "AESA — Proyecto / Contrato es el responsable del contrato y la coordinación con cada proveedor.",
 "AESA — Ingeniería revisa y aprueba la documentación del proveedor y asiste técnicamente a FAT/pruebas.",
 "AESA — Precom. y Comisionado y AESA — Ingeniería (junto con Proyecto) asisten a FAT/pruebas y ejecutan/acompañan la puesta en marcha.",
 "Documento de arranque para iteración: los códigos por actividad son una propuesta base y se ajustan por provisión.",
]

# ================= EXCEL =================
def build_xlsx():
    import openpyxl
    from openpyxl.styles import Font,PatternFill,Alignment,Border,Side
    thin=Side(style="thin",color="D9D9D9"); B=Border(thin,thin,thin,thin)
    NAVY="1F3864"; HEAD="37474F"
    def cell(ws,r,c,v,fill=None,color="000000",bold=False,size=9,center=False,wrap=False):
        cc=ws.cell(r,c,v); cc.font=Font(bold=bold,color=color,size=size)
        if fill: cc.fill=PatternFill("solid",fgColor=fill)
        cc.alignment=Alignment(horizontal="center" if center else "left",vertical="center",wrap_text=wrap); cc.border=B; return cc
    wb=openpyxl.Workbook()
    ws=wb.active; ws.title="Matriz de Interfases"
    ws.merge_cells("A1:G1"); cell(ws,1,1,f"MATRIZ DE INTERFASES / RESPONSABILIDADES — PROVISIONES CRÍTICAS CPF-2 · {REV} · {FECHA}",NAVY,"FFFFFF",True,12,True); ws.row_dimensions[1].height=22
    ws.merge_cells("A2:G2"); cell(ws,2,1,"CPF-2 La Calera II · Esquema RACI ampliado · Actores: Proveedor · AESA-Proyecto/Contrato · AESA-Ingeniería · AESA-Precom&Comis. · Cliente-PPSA",None,"555555",False,9)
    # leyenda
    r=4; cell(ws,r,1,"Leyenda:",None,"000000",True,9)
    for i,(k,(col,desc)) in enumerate(RACI.items()):
        cell(ws,r,2+i,f"{k} = {desc}",col,"FFFFFF" if k!='I' else "000000",True,8,True,True)
    r=6
    def header(r):
        cell(ws,r,1,"Fase",HEAD,"FFFFFF",True,8.5,True); cell(ws,r,2,"Actividad / interfase",HEAD,"FFFFFF",True,8.5,True)
        for j,a in enumerate(ACT_SHORT): cell(ws,r,3+j,a.replace('\n',' '),HEAD,"FFFFFF",True,8,True,True)
        ws.row_dimensions[r].height=26
    for (name,ven,estado,col,adj) in PROV:
        ws.merge_cells(start_row=r,start_column=1,end_row=r,end_column=7)
        est_col = "C62828" if adj else col
        cell(ws,r,1,f"▶  {name}   ·   Proveedor: {ven}   ·   Estado: {estado}",col,"FFFFFF",True,10)
        ws.row_dimensions[r].height=18; r+=1
        header(r); r+=1
        acts = ([ADJ_ROW]+ACTIVITIES) if adj else ACTIVITIES
        for fase,act,codes in acts:
            cell(ws,r,1,fase,None,"333333",False,8); cell(ws,r,2,act,None,"000000",False,8.5)
            for j,code in enumerate(codes):
                fc,_=RACI[code]; cell(ws,r,3+j,code,fc,"FFFFFF" if code!='I' else "000000",True,9,True)
            r+=1
        r+=1
    for c,w in zip("ABCDEFG",[20,46,12,12,12,13,11]): ws.column_dimensions[c].width=w
    ws.freeze_panes="A7"
    # Provisiones
    ps=wb.create_sheet("Provisiones")
    cell(ps,1,1,"PROVISIONES CRÍTICAS",NAVY,"FFFFFF",True,12); ps.merge_cells("A1:D1")
    for j,h in enumerate(["Provisión","Proveedor","Estado","Alcance de interfase"]): cell(ps,3,1+j,h,HEAD,"FFFFFF",True,9,True)
    alc={"ABB — Salas Eléctricas SE#3 y SE#4":"Ingeniería, fabricación, FAT, despacho, montaje shelters y PEM.",
         "ABB — PMS (Power Management System)":"Ingeniería, procura, armado, FAT (incl. maqueta), montaje shelter, iFAT, PEM y capacitación.",
         "Inauco — PCS / SCADA / Comunicaciones":"Ingeniería, construcción tableros, FAT PCS/SIS, integración, SAT/comisionado en campo y CAO.",
         "HIMA — Tableros de marshalling (ESD / F&G / PSS)":"Ingeniería, fabricación tableros, FAT y entrega en Inauco.",
         "HIMA — Adecuación y configuración SIS":"Adjudicación, configuración/adecuación SIS, FAT SIS (conjunto), asistencia PEM."}
    r=4
    for (name,ven,estado,col,adj) in PROV:
        cell(ps,r,1,name,None,"000000",True,9,False,True); cell(ps,r,2,ven,None,"000000",False,9,True)
        cell(ps,r,3,estado,"C62828" if adj else "2E7D32","FFFFFF",True,9,True); cell(ps,r,4,alc[name],None,"000000",False,8.5,False,True)
        ps.row_dimensions[r].height=30; r+=1
    for c,w in zip("ABCD",[42,12,20,60]): ps.column_dimensions[c].width=w
    # Notas
    ns=wb.create_sheet("Notas y premisas")
    cell(ns,1,1,"NOTAS Y PREMISAS",NAVY,"FFFFFF",True,12); ns.merge_cells("A1:A1")
    for i,n in enumerate(NOTAS): cell(ns,3+i,1,f"• {n}",None,"000000",False,9,False,True); ns.row_dimensions[3+i].height=30
    ns.column_dimensions['A'].width=130
    path=os.path.join(GEN,f"Matriz de Interfases - Provisiones criticas - {REV} - {FECHA}.xlsx"); wb.save(path); return path

# ================= PDF =================
def build_pdf():
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib import colors
    from reportlab.lib.units import mm
    from reportlab.pdfgen import canvas
    path=os.path.join(GEN,f"Matriz de Interfases - Provisiones criticas - {REV} - {FECHA}.pdf")
    W,H=landscape(A4); c=canvas.Canvas(path,pagesize=(W,H))
    NAVY=colors.HexColor('#1F3864'); GREY=colors.HexColor('#555555'); x0=12*mm; x1=W-12*mm
    def banner():
        c.setFillColor(NAVY); c.rect(0,H-13*mm,W,13*mm,fill=1,stroke=0); c.setFillColor(colors.white)
        c.setFont("Helvetica-Bold",12); c.drawString(x0,H-9*mm,"MATRIZ DE INTERFASES / RESPONSABILIDADES — PROVISIONES CRÍTICAS CPF-2")
        c.setFont("Helvetica",8.5); c.drawRightString(x1,H-9*mm,f"{REV} · {FECHA}")
    # ---- page 1: provisiones + leyenda + notas ----
    banner(); y=H-20*mm
    c.setFillColor(NAVY); c.setFont("Helvetica-Bold",10); c.drawString(x0,y,"1 · Provisiones y estado de adjudicación"); y-=6*mm
    colw=[86*mm,22*mm,50*mm,110*mm]; hdr=["Provisión","Proveedor","Estado","Alcance de interfase"]
    def row(vals,widths,y,fill=None,bold=False,hcol=None):
        cx=x0; tw=sum(widths)
        if fill: c.setFillColor(fill); c.rect(x0,y-5.4*mm,tw,5.4*mm,fill=1,stroke=0)
        for i,(v,w) in enumerate(zip(vals,widths)):
            c.setFillColor(hcol[i] if hcol else (colors.white if fill else colors.black))
            c.setFont("Helvetica-Bold" if bold else "Helvetica",7.6)
            for j,ln in enumerate(textwrap.wrap(str(v),int(w/1.55))[:2]):
                c.drawString(cx+1.3*mm,y-3.4*mm-j*3.1*mm,ln)
            cx+=w
        c.setStrokeColor(colors.HexColor('#DDDDDD')); c.setLineWidth(0.3); c.line(x0,y-5.4*mm,x0+tw,y-5.4*mm)
        return y-5.4*mm
    y=row(hdr,colw,y,fill=colors.HexColor('#37474F'),bold=True)
    alc={"ABB — Salas Eléctricas SE#3 y SE#4":"Ingeniería, fabricación, FAT, despacho, montaje shelters y PEM.",
         "ABB — PMS (Power Management System)":"Ingeniería, procura, armado, FAT (incl. maqueta), montaje shelter, iFAT, PEM y capacitación.",
         "Inauco — PCS / SCADA / Comunicaciones":"Ingeniería, construcción tableros, FAT PCS/SIS, integración, SAT/comisionado campo y CAO.",
         "HIMA — Tableros de marshalling (ESD / F&G / PSS)":"Ingeniería, fabricación tableros, FAT y entrega en Inauco.",
         "HIMA — Adecuación y configuración SIS":"Adjudicación, configuración/adecuación SIS, FAT SIS (conjunto), asistencia PEM."}
    for (name,ven,estado,col,adj) in PROV:
        c.setFillColor(colors.HexColor('#'+col)); c.rect(x0,y-5.4*mm,3*mm,5.4*mm,fill=1,stroke=0)
        y=row([name,ven,estado,alc[name]],colw,y)
    y-=6*mm
    c.setFillColor(NAVY); c.setFont("Helvetica-Bold",10); c.drawString(x0,y,"2 · Leyenda RACI ampliado"); y-=6*mm
    lx=x0
    for k,(col,desc) in RACI.items():
        c.setFillColor(colors.HexColor('#'+col)); c.rect(lx,y-4*mm,5*mm,5*mm,fill=1,stroke=0)
        c.setFillColor(colors.black); c.setFont("Helvetica-Bold",8); c.drawString(lx+6*mm,y-2.5*mm,f"{k}");
        c.setFont("Helvetica",7.3); c.drawString(lx+10*mm,y-2.5*mm,desc); lx+=54*mm
    y-=10*mm
    c.setFillColor(NAVY); c.setFont("Helvetica-Bold",10); c.drawString(x0,y,"3 · Notas y premisas"); y-=5.5*mm
    c.setFillColor(colors.black); c.setFont("Helvetica",7.6)
    for n in NOTAS:
        for j,ln in enumerate(textwrap.wrap(n,155)):
            c.drawString(x0+(0 if j==0 else 3*mm),y,("• " if j==0 else "")+ln); y-=3.8*mm
        y-=0.8*mm
    # ---- matrix pages ----
    aw=[26*mm,74*mm]+[ (x1-x0-100*mm)/5 ]*5
    def matrix_header(y):
        cx=x0; c.setFillColor(colors.HexColor('#37474F')); c.rect(x0,y-8*mm,sum(aw),8*mm,fill=1,stroke=0)
        c.setFillColor(colors.white); c.setFont("Helvetica-Bold",7)
        labs=["Fase","Actividad / interfase"]+["Proveedor","AESA Proy/Contr.","AESA Ingeniería","AESA Precom/Com","Cliente PPSA"]
        for i,(l,w) in enumerate(zip(labs,aw)):
            for j,ln in enumerate(textwrap.wrap(l,int(w/1.5))[:2]): c.drawString(cx+1.2*mm,y-3.4*mm-j*3*mm,ln)
            cx+=w
        return y-8*mm
    def ensure(y,need):
        if y-need < 14*mm:
            c.showPage(); banner(); return H-20*mm, True
        return y, False
    c.showPage(); banner(); y=H-20*mm
    for (name,ven,estado,col,adj) in PROV:
        y,_=ensure(y,26*mm)
        c.setFillColor(colors.HexColor('#'+col)); c.rect(x0,y-6*mm,sum(aw),6*mm,fill=1,stroke=0)
        c.setFillColor(colors.white); c.setFont("Helvetica-Bold",8.5)
        c.drawString(x0+2*mm,y-4.2*mm,f"{name}   ·   {ven}   ·   {estado}"); y-=6*mm
        y=matrix_header(y)
        acts=([ADJ_ROW]+ACTIVITIES) if adj else ACTIVITIES
        for ri,(fase,act,codes) in enumerate(acts):
            y,newp=ensure(y,5.4*mm)
            if newp: y=matrix_header(y)
            cx=x0; rowh=5.0*mm
            c.setFillColor(colors.HexColor('#F7F9FB') if ri%2 else colors.white); c.rect(x0,y-rowh,sum(aw),rowh,fill=1,stroke=0)
            FSHORT={"Procura / Fabricación":"Procura/Fabric.","Pruebas de fábrica (FAT)":"Pruebas FAT","Precom. / Puesta en marcha":"Precom./PEM"}
            c.setFillColor(colors.HexColor('#444444')); c.setFont("Helvetica",6.4); c.drawString(cx+1.2*mm,y-3.4*mm,FSHORT.get(fase,fase)[:22]); cx+=aw[0]
            c.setFillColor(colors.black); c.setFont("Helvetica",6.8)
            c.drawString(cx+1.2*mm,y-3.4*mm,act[:64]); cx+=aw[1]
            for code in codes:
                fc,_=RACI[code]; c.setFillColor(colors.HexColor('#'+fc)); c.roundRect(cx+aw[2]/2-3.2*mm,y-rowh+0.9*mm,6.4*mm,rowh-1.8*mm,1,fill=1,stroke=0)
                c.setFillColor(colors.white if code!='I' else colors.black); c.setFont("Helvetica-Bold",7); c.drawCentredString(cx+aw[2]/2,y-3.4*mm,code); cx+=aw[2]
            y-=rowh
        y-=4*mm
    c.showPage(); c.save(); return path

p1=build_xlsx(); p2=build_pdf()
print("OK\n ",p1,"\n ",p2)
