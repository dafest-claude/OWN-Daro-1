#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Rev. Alt 6 sobre la Rev. Alt 5 corregida por el usuario (en 'Datos entrada').
El usuario cambió los ESTÁNDARES de la Tabla C.1 (marcados en amarillo) y pasó las puntas
de min/punta a HH/punta. Este script:
 - Recalcula TODAS las HH de C.2 (EL) y C.3 (IN) con los nuevos estándares y las cantidades
   de referencia (tablas de cables/terminaciones/instrumentos), y actualiza subtotales y TOTAL.
 - Unifica las puntas en HH/punta (corrige la unidad 'mHH/punta' -> 'HH/punta' y el texto de C.3).
 - Recalcula la dotación (C.4) y el pico simultáneo.
 - Agrega una condición en el Anexo C: los estándares son una estimación propia de AESA; el Oferente
   presenta su oferta de HH; toda mejora del estándar se valora en la adjudicación; el alcance puede
   adjudicarse de manera total o parcial.
 - Quita las marcas amarillas (eran anotaciones del usuario).
No borra revisiones anteriores.
Base:  Datos entrada/Pliego subcontrato E&I - Rev Alt 5 - 2026-07-07 (con Anexos).docx
Salida: Análisis generados/Pliego subcontrato E&I - Rev Alt 6 - 2026-07-07 (con Anexos).docx
"""
import os
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_COLOR_INDEX
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

ENT="/home/user/OWN-Daro-1/Constructibilidad/Datos entrada"
GEN="/home/user/OWN-Daro-1/Constructibilidad/Análisis generados"
SRC=os.path.join(ENT,"Pliego subcontrato E&I - Rev Alt 5 - 2026-07-07 (con Anexos).docx")
OUT=os.path.join(GEN,"Pliego subcontrato E&I - Rev Alt 6 - 2026-07-07 (con Anexos).docx")
doc=Document(SRC)

def set_cell_text(cell,new):
    p=cell.paragraphs[0]
    if p.runs:
        p.runs[0].text=new
        for r in list(p.runs)[1:]: r._element.getparent().remove(r._element)
    else:
        p.add_run(new)

# ---------- C.1: unificar unidad de punta de armadura 'mHH/punta' -> 'HH/punta' ----------
t17=doc.tables[17]
for r in t17.rows:
    if 'armadura' in r.cells[1].text.lower():
        set_cell_text(r.cells[3],"HH/punta")

# ---------- C.2 (EL): recomputo con nuevos estándares ----------
t18=doc.tables[18]
set_cell_text(t18.rows[1].cells[3],"36.400")   # tendido (rend. nuevos C.1)
set_cell_text(t18.rows[2].cells[2],"6,4 / 2,88 / 1,04 HH·ext · 4,48 HH/cable")
set_cell_text(t18.rows[2].cells[3],"2.101")    # terminaciones/conexionado
# rows[3] montaje 2.840 y rows[4] pruebas 3.677 sin cambio de estándar
set_cell_text(t18.rows[5].cells[3],"≈ 45.018 HH")   # subtotal EL

# ---------- C.3 (IN): recomputo con nuevos estándares ----------
t19=doc.tables[19]
set_cell_text(t19.rows[1].cells[2],"50 / 80 m/día·cuadrilla"); set_cell_text(t19.rows[1].cells[3],"16.469")
set_cell_text(t19.rows[2].cells[2],"0,18 / 0,13 / 0,26 HH/punta"); set_cell_text(t19.rows[2].cells[3],"2.896")
set_cell_text(t19.rows[3].cells[2],"6,40 HH/instrumento"); set_cell_text(t19.rows[3].cells[3],"7.443")
set_cell_text(t19.rows[4].cells[2],"4,80 y 1,44 HH/unidad"); set_cell_text(t19.rows[4].cells[3],"7.554")
set_cell_text(t19.rows[5].cells[2],"8,0 HH/lazo"); set_cell_text(t19.rows[5].cells[3],"7.080")
set_cell_text(t19.rows[6].cells[2],"8 HH/caja + 6,4 y 8,0 HH/soporte"); set_cell_text(t19.rows[6].cells[3],"1.203")
set_cell_text(t19.rows[7].cells[3],"≈ 42.645 HH")   # subtotal IN

# ---------- C.4 (dotación): recomputo ----------
t20=doc.tables[20]
# (HH, HH-mes, dotación)
c4=[(1,"38.501","6.417","~36"),
    (3,"19.365","2.766","~16"),
    (4,"16.200","2.700","~15"),
    (5,"10.757","2.151","~12")]
for idx,hh,hhmes,dot in c4:
    set_cell_text(t20.rows[idx].cells[2],hh)
    set_cell_text(t20.rows[idx].cells[4],hhmes)
    set_cell_text(t20.rows[idx].cells[5],dot)
set_cell_text(t20.rows[6].cells[5],"~60-70")   # pico simultáneo

# ---------- TOTAL E&I ----------
for p in doc.paragraphs:
    if 'TOTAL mano de obra estimada' in p.text:
        for r in p.runs:
            r.text=r.text.replace('60.746','87.663').replace('345 persona-mes','498 persona-mes')

# ---------- Condición de la oferta (Anexo C, tras la Tabla C.1) ----------
note=doc.add_paragraph()
try: note.style=doc.styles['Normal']
except KeyError: pass
lead=note.add_run("Condición de la oferta (estándares y HH): "); lead.bold=True
note.add_run(
 "los estándares de rendimiento de la Tabla C.1 constituyen una estimación propia de AESA, de "
 "carácter orientativo. El Oferente deberá presentar su propia oferta de horas-hombre (HH) para la "
 "ejecución de los trabajos. Toda mejora sobre los estándares aquí planteados será valorada "
 "positivamente al momento de la adjudicación. Asimismo, AESA podrá adjudicar el alcance de manera "
 "total o parcial.")
t17._tbl.addnext(note._p)   # queda entre la Tabla C.1 y el título C.2

# ---------- quitar marcas amarillas (anotaciones del usuario) ----------
def clear_hl(paras):
    for p in paras:
        for r in p.runs:
            if r.font.highlight_color is not None and r.font.highlight_color!=WD_COLOR_INDEX.AUTO:
                r.font.highlight_color=None
clear_hl(doc.paragraphs)
for t in doc.tables:
    for row in t.rows:
        for c in row.cells:
            clear_hl(c.paragraphs)

# ---------- tipografía Calibri en runs nuevos ----------
def force_cal(run):
    rpr=run._element.get_or_add_rPr(); rf=rpr.find(qn('w:rFonts'))
    if rf is None: rf=OxmlElement('w:rFonts'); rpr.insert(0,rf)
    for a in ('w:asciiTheme','w:hAnsiTheme','w:cstheme'):
        if rf.get(qn(a)) is not None: del rf.attrib[qn(a)]
    rf.set(qn('w:ascii'),'Calibri'); rf.set(qn('w:hAnsi'),'Calibri'); rf.set(qn('w:cs'),'Calibri')
for p in doc.paragraphs:
    for r in p.runs: force_cal(r)
for t in doc.tables:
    for row in t.rows:
        for c in row.cells:
            for p in c.paragraphs:
                for r in p.runs: force_cal(r)

doc.save(OUT)
print("OK ->",OUT,"| tablas:",len(doc.tables))
