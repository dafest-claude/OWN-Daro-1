"""
Genera planilla Excel con datos del clima de Buenos Aires
Período: 01/05/2026 - 21/05/2026
Fuente: Open-Meteo Archive API (open-meteo.com)
"""
import requests
import openpyxl
from openpyxl.styles import (
    Font, PatternFill, Alignment, Border, Side, numbers
)
from openpyxl.utils import get_column_letter
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.chart.series import DataPoint
from datetime import datetime, date

# ── Descripción de códigos WMO de clima ──────────────────────────────────────
WMO_CODES = {
    0: "Cielo despejado", 1: "Mayormente despejado", 2: "Parcialmente nublado",
    3: "Nublado", 45: "Niebla", 48: "Niebla con escarcha",
    51: "Llovizna leve", 53: "Llovizna moderada", 55: "Llovizna intensa",
    61: "Lluvia leve", 63: "Lluvia moderada", 65: "Lluvia intensa",
    71: "Nevada leve", 73: "Nevada moderada", 75: "Nevada intensa",
    80: "Chubascos leves", 81: "Chubascos moderados", 82: "Chubascos violentos",
    95: "Tormenta eléctrica", 96: "Tormenta con granizo leve",
    99: "Tormenta con granizo intenso",
}


def obtener_datos_clima():
    url = "https://archive-api.open-meteo.com/v1/archive"
    params = {
        "latitude": -34.6037,
        "longitude": -58.3816,
        "start_date": "2026-05-01",
        "end_date": "2026-05-21",
        "daily": (
            "temperature_2m_max,temperature_2m_min,temperature_2m_mean,"
            "precipitation_sum,windspeed_10m_max,windgusts_10m_max,"
            "weathercode,sunshine_duration"
        ),
        "timezone": "America/Argentina/Buenos_Aires",
        "wind_speed_unit": "kmh",
    }
    r = requests.get(url, params=params, timeout=30)
    r.raise_for_status()
    return r.json()["daily"]


def crear_estilo_borde():
    thin = Side(style="thin", color="BFBFBF")
    return Border(left=thin, right=thin, top=thin, bottom=thin)


def crear_excel(datos):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Clima Buenos Aires"

    # ── Paleta de colores ────────────────────────────────────────────────────
    AZUL_OSCURO  = "1F3864"
    AZUL_MEDIO   = "2E75B6"
    AZUL_CLARO   = "BDD7EE"
    CELESTE_FILA = "DEEAF1"
    BLANCO       = "FFFFFF"
    GRIS_CLARO   = "F2F2F2"
    VERDE        = "70AD47"
    NARANJA      = "ED7D31"

    borde = crear_estilo_borde()

    # ── Título principal ─────────────────────────────────────────────────────
    ws.merge_cells("A1:J1")
    titulo = ws["A1"]
    titulo.value = "DATOS METEOROLÓGICOS – BUENOS AIRES, ARGENTINA"
    titulo.font = Font(name="Calibri", bold=True, size=16, color=BLANCO)
    titulo.fill = PatternFill("solid", fgColor=AZUL_OSCURO)
    titulo.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 30

    # ── Subtítulo ────────────────────────────────────────────────────────────
    ws.merge_cells("A2:J2")
    sub = ws["A2"]
    sub.value = "Período: 01/05/2026 al 21/05/2026  |  Fuente: Open-Meteo Archive API"
    sub.font = Font(name="Calibri", italic=True, size=10, color=BLANCO)
    sub.fill = PatternFill("solid", fgColor=AZUL_MEDIO)
    sub.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[2].height = 18

    # ── Encabezados ──────────────────────────────────────────────────────────
    encabezados = [
        ("Fecha",            "A"),
        ("Día",              "B"),
        ("Condición",        "C"),
        ("T° Máx (°C)",      "D"),
        ("T° Mín (°C)",      "E"),
        ("T° Media (°C)",    "F"),
        ("Precipitación (mm)","G"),
        ("Viento Máx (km/h)","H"),
        ("Ráfaga Máx (km/h)","I"),
        ("Sol (hs)",         "J"),
    ]

    ws.row_dimensions[3].height = 40
    for titulo_col, col_letra in encabezados:
        celda = ws[f"{col_letra}3"]
        celda.value = titulo_col
        celda.font = Font(name="Calibri", bold=True, size=10, color=BLANCO)
        celda.fill = PatternFill("solid", fgColor=AZUL_MEDIO)
        celda.alignment = Alignment(
            horizontal="center", vertical="center", wrap_text=True
        )
        celda.border = borde

    # ── Días de la semana en español ─────────────────────────────────────────
    DIAS_ES = {
        "Monday": "Lunes", "Tuesday": "Martes", "Wednesday": "Miércoles",
        "Thursday": "Jueves", "Friday": "Viernes",
        "Saturday": "Sábado", "Sunday": "Domingo",
    }

    fechas = datos["time"]
    for i, fecha_str in enumerate(fechas):
        fila = i + 4
        es_par = i % 2 == 0
        color_fondo = CELESTE_FILA if es_par else BLANCO

        fecha_dt = datetime.strptime(fecha_str, "%Y-%m-%d").date()
        dia_ingles = fecha_dt.strftime("%A")
        dia_es = DIAS_ES[dia_ingles]

        t_max  = datos["temperature_2m_max"][i]
        t_min  = datos["temperature_2m_min"][i]
        t_med  = datos["temperature_2m_mean"][i]
        precip = datos["precipitation_sum"][i]
        viento = datos["windspeed_10m_max"][i]
        rafaga = datos["windgusts_10m_max"][i]
        wcode  = datos["weathercode"][i]
        sol_s  = datos["sunshine_duration"][i]
        sol_h  = round(sol_s / 3600, 1) if sol_s is not None else None
        condicion = WMO_CODES.get(wcode, f"Código {wcode}")

        valores = [
            fecha_dt, dia_es, condicion,
            t_max, t_min, t_med,
            precip, viento, rafaga, sol_h,
        ]

        for j, (col_letra, val) in enumerate(
            zip("ABCDEFGHIJ", valores)
        ):
            celda = ws[f"{col_letra}{fila}"]
            celda.value = val
            celda.fill = PatternFill("solid", fgColor=color_fondo)
            celda.border = borde
            celda.font = Font(name="Calibri", size=10)

            # Alineación por columna
            if col_letra == "A":
                celda.number_format = "DD/MM/YYYY"
                celda.alignment = Alignment(horizontal="center")
            elif col_letra in ("B", "C"):
                celda.alignment = Alignment(horizontal="left")
            else:
                celda.alignment = Alignment(horizontal="center")

            # Colores condicionales temperatura
            if col_letra == "D" and t_max is not None:
                if t_max >= 22:
                    celda.fill = PatternFill("solid", fgColor="FFE699")
                elif t_max <= 12:
                    celda.fill = PatternFill("solid", fgColor="BDD7EE")
            if col_letra == "G" and precip is not None and precip > 0:
                celda.fill = PatternFill("solid", fgColor="DDEBF7")
                celda.font = Font(name="Calibri", size=10, color="1F78B4", bold=True)

    # ── Fila de totales/promedios ─────────────────────────────────────────────
    fila_total = len(fechas) + 4
    ws.row_dimensions[fila_total].height = 18

    ws.merge_cells(f"A{fila_total}:C{fila_total}")
    celda_res = ws[f"A{fila_total}"]
    celda_res.value = "RESUMEN DEL PERÍODO"
    celda_res.font = Font(name="Calibri", bold=True, size=10, color=BLANCO)
    celda_res.fill = PatternFill("solid", fgColor=AZUL_OSCURO)
    celda_res.alignment = Alignment(horizontal="center")
    celda_res.border = borde

    n = len(fechas)
    resumen = {
        "D": f"=AVERAGE(D4:D{fila_total-1})",
        "E": f"=AVERAGE(E4:E{fila_total-1})",
        "F": f"=AVERAGE(F4:F{fila_total-1})",
        "G": f"=SUM(G4:G{fila_total-1})",
        "H": f"=AVERAGE(H4:H{fila_total-1})",
        "I": f"=MAX(I4:I{fila_total-1})",
        "J": f"=AVERAGE(J4:J{fila_total-1})",
    }
    etiquetas_res = {
        "D": "Prom", "E": "Prom", "F": "Prom",
        "G": "Total", "H": "Prom", "I": "Máx", "J": "Prom",
    }
    for col_l, formula in resumen.items():
        celda = ws[f"{col_l}{fila_total}"]
        celda.value = formula
        celda.font = Font(name="Calibri", bold=True, size=10, color=BLANCO)
        celda.fill = PatternFill("solid", fgColor=AZUL_OSCURO)
        celda.alignment = Alignment(horizontal="center")
        celda.border = borde
        celda.number_format = "0.0"

    # ── Anchos de columna ────────────────────────────────────────────────────
    anchos = {
        "A": 13, "B": 12, "C": 24, "D": 13, "E": 13,
        "F": 14, "G": 18, "H": 17, "I": 17, "J": 10,
    }
    for col, ancho in anchos.items():
        ws.column_dimensions[col].width = ancho

    # ── Congelar paneles ─────────────────────────────────────────────────────
    ws.freeze_panes = "A4"

    # ── Segunda hoja: gráfico de temperaturas ────────────────────────────────
    ws2 = wb.create_sheet("Gráficos")
    ws2["A1"] = "Gráficos de temperatura y precipitación"
    ws2["A1"].font = Font(name="Calibri", bold=True, size=13, color=AZUL_OSCURO)

    # Gráfico de líneas: temperaturas
    chart_temp = LineChart()
    chart_temp.title = "Temperaturas Diarias – Buenos Aires (Mayo 2026)"
    chart_temp.style = 10
    chart_temp.y_axis.title = "Temperatura (°C)"
    chart_temp.x_axis.title = "Fecha"
    chart_temp.height = 12
    chart_temp.width = 22

    datos_max = Reference(
        wb["Clima Buenos Aires"], min_col=4, max_col=4,
        min_row=3, max_row=3 + n
    )
    datos_min = Reference(
        wb["Clima Buenos Aires"], min_col=5, max_col=5,
        min_row=3, max_row=3 + n
    )
    datos_med = Reference(
        wb["Clima Buenos Aires"], min_col=6, max_col=6,
        min_row=3, max_row=3 + n
    )
    fechas_ref = Reference(
        wb["Clima Buenos Aires"], min_col=1,
        min_row=4, max_row=3 + n
    )
    chart_temp.add_data(datos_max, titles_from_data=True)
    chart_temp.add_data(datos_min, titles_from_data=True)
    chart_temp.add_data(datos_med, titles_from_data=True)
    chart_temp.set_categories(fechas_ref)
    chart_temp.series[0].graphicalProperties.line.solidFill = NARANJA
    chart_temp.series[1].graphicalProperties.line.solidFill = AZUL_MEDIO
    chart_temp.series[2].graphicalProperties.line.solidFill = VERDE
    ws2.add_chart(chart_temp, "A3")

    # Gráfico de barras: precipitación
    chart_prec = BarChart()
    chart_prec.title = "Precipitación Diaria – Buenos Aires (Mayo 2026)"
    chart_prec.style = 10
    chart_prec.y_axis.title = "mm"
    chart_prec.x_axis.title = "Fecha"
    chart_prec.height = 12
    chart_prec.width = 22

    datos_prec = Reference(
        wb["Clima Buenos Aires"], min_col=7, max_col=7,
        min_row=3, max_row=3 + n
    )
    chart_prec.add_data(datos_prec, titles_from_data=True)
    chart_prec.set_categories(fechas_ref)
    chart_prec.series[0].graphicalProperties.solidFill = AZUL_MEDIO
    ws2.add_chart(chart_prec, "A23")

    # ── Guardar ──────────────────────────────────────────────────────────────
    nombre_archivo = "clima_buenos_aires_mayo2026.xlsx"
    wb.save(nombre_archivo)
    print(f"Archivo generado: {nombre_archivo}")
    return nombre_archivo


if __name__ == "__main__":
    print("Obteniendo datos del clima...")
    datos = obtener_datos_clima()
    print(f"Datos recibidos para {len(datos['time'])} días.")
    archivo = crear_excel(datos)
    print("¡Listo!")
