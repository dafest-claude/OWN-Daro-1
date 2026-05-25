"""
Análisis comparativo: Cronograma Global CPF-2 (Primavera P6 XER) vs Gantt ABB (Excel)
Verifica compatibilidad de fechas entre el cronograma comercial del proyecto y las
actividades de ABB (Salas #3, #4 y PMS).
"""

import os
import io
from datetime import datetime, date
from collections import defaultdict
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
PLANNING_DIR = os.path.join(BASE_DIR, "02_Planning")

# ── Estilos ───────────────────────────────────────────────────────────────────
def thin_border():
    s = Side(style="thin", color="CCCCCC")
    return Border(left=s, right=s, top=s, bottom=s)

def hdr(ws, row, cols, bg="CC0000", fg="FFFFFF", sz=10):
    for c in range(1, cols + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = PatternFill("solid", fgColor=bg)
        cell.font = Font(bold=True, color=fg, size=sz)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border()

def data_cell(ws, row, col, value, bg=None, bold=False, align="left"):
    cell = ws.cell(row=row, column=col, value=value)
    cell.border = thin_border()
    cell.alignment = Alignment(vertical="center", horizontal=align, wrap_text=True)
    if bg:
        cell.fill = PatternFill("solid", fgColor=bg)
    if bold:
        cell.font = Font(bold=True)
    return cell

# ── Parser XER ────────────────────────────────────────────────────────────────
def parse_xer(filepath):
    """Lee un archivo Primavera P6 XER y devuelve dict de tablas."""
    tables = {}
    current_table = None
    current_fields = []

    encodings = ["latin-1", "utf-8", "cp1252"]
    content = None
    for enc in encodings:
        try:
            with open(filepath, encoding=enc, errors="replace") as f:
                content = f.read()
            break
        except Exception:
            continue

    if not content:
        raise ValueError("No se pudo leer el archivo XER")

    for line in content.splitlines():
        line = line.rstrip("\r")
        if not line:
            continue
        if line.startswith("%T\t"):
            current_table = line[3:].strip()
            tables[current_table] = []
            current_fields = []
        elif line.startswith("%F\t"):
            current_fields = line[3:].split("\t")
        elif line.startswith("%R\t") and current_table and current_fields:
            vals = line[3:].split("\t")
            record = {}
            for i, field in enumerate(current_fields):
                record[field] = vals[i] if i < len(vals) else ""
            tables[current_table].append(record)

    return tables


def parse_date(s):
    """Convierte string de fecha P6 a date."""
    if not s:
        return None
    for fmt in ["%Y-%m-%d %H:%M", "%Y-%m-%d"]:
        try:
            return datetime.strptime(s[:16] if len(s) > 10 else s, fmt).date()
        except ValueError:
            continue
    return None


# ── Extraer datos del XER ─────────────────────────────────────────────────────
def extraer_proyecto(tables):
    project = {}
    if "PROJECT" in tables and tables["PROJECT"]:
        p = tables["PROJECT"][0]
        project = {
            "nombre":     p.get("proj_short_name", ""),
            "inicio":     parse_date(p.get("plan_start_date", "")),
            "fin":        parse_date(p.get("scd_end_date", "") or p.get("anticip_end_date", "")),
            "data_date":  parse_date(p.get("last_recalc_date", "")),
        }
    return project


def extraer_wbs(tables):
    wbs_map = {}
    if "PROJWBS" in tables:
        for w in tables["PROJWBS"]:
            wbs_map[w.get("wbs_id", "")] = {
                "nombre": w.get("wbs_name", ""),
                "codigo": w.get("wbs_short_name", ""),
                "parent": w.get("parent_wbs_id", ""),
            }
    return wbs_map


def extraer_tareas(tables, wbs_map):
    tareas = []
    if "TASK" not in tables:
        return tareas

    for t in tables["TASK"]:
        nombre = t.get("task_name", "")
        tipo   = t.get("task_type", "")
        inicio = parse_date(t.get("target_start_date") or t.get("act_start_date") or t.get("early_start_date", ""))
        fin    = parse_date(t.get("target_end_date")   or t.get("act_end_date")   or t.get("early_end_date", ""))
        avance = t.get("phys_complete_pct", t.get("target_complete_pct", "0"))
        wbs_id = t.get("wbs_id", "")
        wbs_info = wbs_map.get(wbs_id, {})

        tareas.append({
            "id":        t.get("task_id", ""),
            "codigo":    t.get("task_code", ""),
            "nombre":    nombre,
            "tipo":      tipo,
            "inicio":    inicio,
            "fin":       fin,
            "avance":    avance,
            "wbs_id":    wbs_id,
            "wbs_nombre": wbs_info.get("nombre", ""),
            "wbs_codigo": wbs_info.get("codigo", ""),
            "duracion":  ((fin - inicio).days if inicio and fin else None),
        })
    return tareas


# ── Datos de referencia: Gantt ABB ────────────────────────────────────────────
# Extraídos del archivo CPF2 Gantt EPC Integrado.xlsx
ABB_HITOS = [
    # (Actividad, Fecha_Plan, Tipo)
    ("Inicio Proyecto ABB (S#3,4 y PMS)",              date(2026,  5,  8), "Hito"),
    ("Freezing Point IB PMS (S#3 y S#4)",              date(2026,  6, 19), "Freezing Point"),
    ("Freezing Point IB (S#4)",                        date(2026,  6, 19), "Freezing Point"),
    ("Freezing Point IB (S#3)",                        date(2026,  7,  3), "Freezing Point"),
    ("FAT PMS",                                        date(2026, 10, 23), "FAT"),
    ("Ensayos Celdas y Tableros S#4",                  date(2026, 11, 23), "Ensayo"),
    ("Ensayos Celdas y Tableros S#3",                  date(2026, 11, 23), "Ensayo"),
    ("FAT Sala Eléctrica #4",                          date(2027,  3, 16), "FAT"),
    ("FAT Sala Eléctrica #3",                          date(2027,  4, 19), "FAT"),
    ("Despacho Sala #4",                               date(2027,  4,  8), "Despacho"),
    ("Despacho Sala #3",                               date(2027,  5, 14), "Despacho"),
    ("iFAT Integrado EPC",                             date(2027,  7, 14), "iFAT"),
    ("Precomisionado (inicio)",                        date(2027,  5,  3), "Comisionado"),
    ("Comisionado / RFSU",                             date(2028,  1, 31), "RFSU"),
]

ABB_ACTIVIDADES = [
    # (Actividad, Inicio, Fin)
    ("Emisión IB / APR IB S#4",          date(2026,  5, 22), date(2026,  6, 26)),
    ("Emisión ID / APR ID S#4",          date(2026,  6, 26), date(2026,  8, 14)),
    ("APR Constructiva S#4",             date(2026,  8, 10), date(2026,  8, 24)),
    ("Emisión IB / APR IB S#3",          date(2026,  5, 22), date(2026,  6, 26)),
    ("Emisión ID / APR ID S#3",          date(2026,  6, 26), date(2026,  7, 31)),
    ("APR Constructiva S#3",             date(2026,  8, 17), date(2026,  8, 31)),
    ("APR IB PMS (S#3 y S#4)",           date(2026,  6,  8), date(2026,  6, 19)),
    ("APR ID PMS (S#3 y S#4)",           date(2026,  7, 20), date(2026,  7, 31)),
    ("Ensayos PMS (S#3 y S#4)",          date(2026, 10, 19), date(2026, 10, 23)),
]


# ── Clasificar tareas P6 relevantes ──────────────────────────────────────────
KEYWORDS_ABB      = ["sala", "eléctric", "electric", "abb", "pms", "power manag", "shelter", "celda", "tablero", "unigear", "mt ", "bt ", "baja tensión", "media tensión", "despacho", "fat sala", "fat pms"]
KEYWORDS_RFSU     = ["rfsu", "ready for start", "start up", "comisionado", "commissioning", "puesta en marcha", "arranque"]
KEYWORDS_IFAT     = ["ifat", "fat integrado", "integrated fat"]
KEYWORDS_CIVIL    = ["civil", "montaje shelter", "obra"]

def clasificar(nombre):
    n = nombre.lower()
    if any(k in n for k in KEYWORDS_ABB):
        return "ABB"
    if any(k in n for k in KEYWORDS_RFSU):
        return "RFSU/Comisionado"
    if any(k in n for k in KEYWORDS_IFAT):
        return "iFAT"
    if any(k in n for k in KEYWORDS_CIVIL):
        return "Civil/Montaje"
    return "Proyecto General"


def semaforo(delta_dias, es_hito=False):
    """Verde / Amarillo / Rojo según desvío."""
    limite = 7 if es_hito else 14
    if delta_dias is None:
        return "⬜ Sin datos", "F2F2F2"
    if abs(delta_dias) <= limite:
        return "🟢 Compatible", "C6EFCE"
    elif abs(delta_dias) <= 30:
        return "🟡 Desvío leve", "FFEB9C"
    else:
        return f"🔴 Desvío {'adelanto' if delta_dias < 0 else 'atraso'} {abs(delta_dias)}d", "FFC7CE"


# ── Buscar fecha en P6 por palabra clave ──────────────────────────────────────
def buscar_fecha_p6(tareas_p6, *keywords):
    """Busca la primera tarea que contenga todas las keywords."""
    for kw_set in keywords:
        if isinstance(kw_set, str):
            kw_set = [kw_set]
        kw_set = [k.lower() for k in kw_set]
        for t in tareas_p6:
            n = t["nombre"].lower()
            if all(k in n for k in kw_set):
                return t
    return None


# ── Generar reporte Excel ─────────────────────────────────────────────────────
def generar_reporte(proyecto_p6, tareas_p6, wbs_map):
    wb = Workbook()

    # ═══════════════════════════════════════════════════════════════════
    # 1. RESUMEN EJECUTIVO
    # ═══════════════════════════════════════════════════════════════════
    ws1 = wb.active
    ws1.title = "Resumen Ejecutivo"

    ws1.merge_cells("A1:G1")
    ws1["A1"] = "ANÁLISIS COMPARATIVO DE CRONOGRAMAS — CPF-2 VACA MUERTA"
    ws1["A1"].fill = PatternFill("solid", fgColor="CC0000")
    ws1["A1"].font = Font(bold=True, color="FFFFFF", size=14)
    ws1["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws1.row_dimensions[1].height = 38

    ws1["A2"] = f"Generado: {datetime.today().strftime('%d/%m/%Y %H:%M')}   |   Fuente P6: Cronograma CPF2 Rev.0 18-03-26 LB Comercial   |   Fuente ABB: CPF2 Gantt EPC Integrado"
    ws1["A2"].font = Font(italic=True, size=9, color="888888")
    ws1.merge_cells("A2:G2")

    # Datos del proyecto P6
    row = 4
    ws1.merge_cells(f"A{row}:G{row}")
    ws1.cell(row=row, column=1, value="DATOS DEL PROYECTO (Cronograma Global P6)")
    hdr(ws1, row, 7)
    ws1.row_dimensions[row].height = 22

    campos_proy = [
        ("Proyecto P6:",        proyecto_p6.get("nombre", "N/D")),
        ("Data Date (Rev.0):",  str(proyecto_p6.get("data_date", "N/D"))),
        ("Inicio planificado:", str(proyecto_p6.get("inicio", "N/D"))),
        ("Fin planificado:",    str(proyecto_p6.get("fin", "N/D"))),
        ("Total actividades P6:", str(len(tareas_p6))),
        ("RFSU objetivo ABB:",  "31/01/2028"),
    ]
    for i, (label, val) in enumerate(campos_proy, row + 1):
        ws1.cell(row=i, column=1, value=label).font = Font(bold=True)
        ws1.cell(row=i, column=2, value=val)
        ws1.merge_cells(f"B{i}:G{i}")
        ws1.row_dimensions[i].height = 18

    # Estadísticas P6
    row = row + len(campos_proy) + 2
    ws1.merge_cells(f"A{row}:G{row}")
    ws1.cell(row=row, column=1, value="ESTADÍSTICAS DEL CRONOGRAMA P6")
    hdr(ws1, row, 7)
    ws1.row_dimensions[row].height = 22

    por_tipo = defaultdict(int)
    for t in tareas_p6:
        por_tipo[t["tipo"]] += 1

    wbs_top = defaultdict(int)
    for t in tareas_p6:
        wbs_top[t["wbs_nombre"]] += 1

    est_headers = ["Métrica", "Valor", "", "Tipo de Tarea", "Cantidad", "", "WBS Principal (Top 5)", "N° Tareas"]
    for col, h in enumerate(est_headers, 1):
        ws1.cell(row=row + 1, column=col, value=h)
    hdr(ws1, row + 1, 8, bg="880000", sz=9)

    metricas = [
        ("Total tareas", len(tareas_p6)),
        ("Con fechas completas", sum(1 for t in tareas_p6 if t["inicio"] and t["fin"])),
        ("Hitos (milestones)", sum(1 for t in tareas_p6 if t["tipo"] in ("TT_Mile", "TT_FinMile"))),
        ("Tareas con avance > 0%", sum(1 for t in tareas_p6 if t["avance"] not in ("0", "0.0", "", "0.00"))),
    ]
    tipos_list = sorted(por_tipo.items(), key=lambda x: -x[1])[:6]
    wbs_top5   = sorted(wbs_top.items(), key=lambda x: -x[1])[:5]

    for i in range(max(len(metricas), len(tipos_list), len(wbs_top5))):
        r = row + 2 + i
        if i < len(metricas):
            data_cell(ws1, r, 1, metricas[i][0], bold=True)
            data_cell(ws1, r, 2, metricas[i][1], align="center")
        if i < len(tipos_list):
            data_cell(ws1, r, 4, tipos_list[i][0])
            data_cell(ws1, r, 5, tipos_list[i][1], align="center")
        if i < len(wbs_top5):
            data_cell(ws1, r, 7, wbs_top5[i][0])
            data_cell(ws1, r, 8, wbs_top5[i][1], align="center")
        ws1.row_dimensions[r].height = 18

    ws1.column_dimensions["A"].width = 28
    ws1.column_dimensions["B"].width = 22
    ws1.column_dimensions["C"].width = 4
    ws1.column_dimensions["D"].width = 20
    ws1.column_dimensions["E"].width = 12
    ws1.column_dimensions["F"].width = 4
    ws1.column_dimensions["G"].width = 35
    ws1.column_dimensions["H"].width = 12


    # ═══════════════════════════════════════════════════════════════════
    # 2. COMPARATIVA HITOS ABB
    # ═══════════════════════════════════════════════════════════════════
    ws2 = wb.create_sheet("Comparativa Hitos ABB")
    ws2.merge_cells("A1:H1")
    ws2["A1"] = "COMPARATIVA DE HITOS ABB: Gantt ABB vs Cronograma Global P6"
    ws2["A1"].fill = PatternFill("solid", fgColor="CC0000")
    ws2["A1"].font = Font(bold=True, color="FFFFFF", size=13)
    ws2["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws2.row_dimensions[1].height = 32

    cols_h = ["Hito / Actividad ABB", "Tipo", "Fecha ABB (Gantt)", "Fecha P6 (Global)",
              "Desvío (días)", "Estado", "Actividad P6 encontrada", "Observación"]
    for col, h in enumerate(cols_h, 1):
        ws2.cell(row=2, column=col, value=h)
    hdr(ws2, 2, len(cols_h))
    ws2.row_dimensions[2].height = 28

    # Mapeo manual de búsqueda: hito ABB → keywords en P6
    busquedas = {
        "Inicio Proyecto ABB (S#3,4 y PMS)": [["kom", "abb"], ["inicio", "abb"], ["recepcion", "oc"]],
        "Freezing Point IB PMS (S#3 y S#4)": [["freezing", "pms", "ib"], ["freeze", "pms"]],
        "Freezing Point IB (S#4)":           [["freezing", "s4", "ib"], ["freezing", "sala 4", "ib"], ["freeze", "sala 4"]],
        "Freezing Point IB (S#3)":           [["freezing", "s3", "ib"], ["freezing", "sala 3", "ib"]],
        "FAT PMS":                            [["fat", "pms"], ["fat", "power manag"]],
        "Ensayos Celdas y Tableros S#4":      [["fat", "sala", "4"], ["fat", "eléctrica", "4"], ["fat", "electrica", "4"]],
        "Ensayos Celdas y Tableros S#3":      [["fat", "sala", "3"], ["fat", "eléctrica", "3"], ["fat", "electrica", "3"]],
        "FAT Sala Eléctrica #4":              [["fat", "sala", "#4"], ["fat", "sala 4"], ["fat", "electrica", "4"]],
        "FAT Sala Eléctrica #3":              [["fat", "sala", "#3"], ["fat", "sala 3"], ["fat", "electrica", "3"]],
        "Despacho Sala #4":                   [["despacho", "4"], ["ship", "sala 4"], ["entrega", "sala 4"]],
        "Despacho Sala #3":                   [["despacho", "3"], ["ship", "sala 3"], ["entrega", "sala 3"]],
        "iFAT Integrado EPC":                 [["ifat", "integrado"], ["integrated fat"], ["ifat epc"]],
        "Precomisionado (inicio)":            [["precomision"], ["pre-comision"], ["precommission"]],
        "Comisionado / RFSU":                 [["rfsu"], ["start up", "cpf"], ["ready for start"], ["comisionado", "fin"]],
    }

    resumen_semaforo = {"🟢 Compatible": 0, "🟡 Desvío leve": 0, "🔴": 0, "⬜ Sin datos": 0}

    for row_i, (nombre, fecha_abb, tipo) in enumerate(ABB_HITOS, 3):
        kw_list = busquedas.get(nombre, [[nombre.lower()[:15]]])
        encontrada = None
        for kws in kw_list:
            encontrada = buscar_fecha_p6(tareas_p6, kws)
            if encontrada:
                break

        fecha_p6 = encontrada["fin"] if encontrada and encontrada["fin"] else (
                   encontrada["inicio"] if encontrada else None)
        delta = (fecha_p6 - fecha_abb).days if fecha_p6 and fecha_abb else None
        estado, color = semaforo(delta, es_hito=True)

        # conteo semaforo
        if "🟢" in estado:   resumen_semaforo["🟢 Compatible"] += 1
        elif "🟡" in estado: resumen_semaforo["🟡 Desvío leve"] += 1
        elif "🔴" in estado: resumen_semaforo["🔴"] += 1
        else:                resumen_semaforo["⬜ Sin datos"] += 1

        obs = ""
        if delta is not None:
            if delta > 0:   obs = f"P6 llega {delta}d después del Gantt ABB"
            elif delta < 0: obs = f"P6 anticipa {abs(delta)}d respecto al Gantt ABB"
            else:           obs = "Fechas exactamente iguales"
        else:
            obs = "No encontrado en P6 — verificar manualmente"

        data_cell(ws2, row_i, 1, nombre, bold=True)
        data_cell(ws2, row_i, 2, tipo)
        data_cell(ws2, row_i, 3, fecha_abb.strftime("%d/%m/%Y") if fecha_abb else "—", align="center")
        data_cell(ws2, row_i, 4, fecha_p6.strftime("%d/%m/%Y") if fecha_p6 else "—", align="center")
        data_cell(ws2, row_i, 5, delta if delta is not None else "—", align="center")
        data_cell(ws2, row_i, 6, estado, bg=color)
        data_cell(ws2, row_i, 7, encontrada["nombre"] if encontrada else "⚠ No encontrado")
        data_cell(ws2, row_i, 8, obs)
        ws2.row_dimensions[row_i].height = 22

    # Totales
    total_row = 3 + len(ABB_HITOS) + 1
    ws2.merge_cells(f"A{total_row}:E{total_row}")
    ws2.cell(row=total_row, column=1, value="RESUMEN SEMÁFORO").font = Font(bold=True)
    for col, (k, v) in enumerate(resumen_semaforo.items(), 1):
        ws2.cell(row=total_row + 1, column=col, value=k)
        ws2.cell(row=total_row + 2, column=col, value=v)
        hdr(ws2, total_row + 1, 1, bg="333333")

    ws2.column_dimensions["A"].width = 40
    ws2.column_dimensions["B"].width = 16
    ws2.column_dimensions["C"].width = 16
    ws2.column_dimensions["D"].width = 16
    ws2.column_dimensions["E"].width = 14
    ws2.column_dimensions["F"].width = 22
    ws2.column_dimensions["G"].width = 45
    ws2.column_dimensions["H"].width = 45
    ws2.freeze_panes = ws2["A3"]


    # ═══════════════════════════════════════════════════════════════════
    # 3. TAREAS ABB EN P6 (todas las relacionadas a ABB/Sala)
    # ═══════════════════════════════════════════════════════════════════
    ws3 = wb.create_sheet("Tareas ABB en P6")
    ws3.merge_cells("A1:G1")
    ws3["A1"] = "TAREAS DEL CRONOGRAMA P6 RELACIONADAS CON ABB / SALAS ELÉCTRICAS"
    ws3["A1"].fill = PatternFill("solid", fgColor="CC0000")
    ws3["A1"].font = Font(bold=True, color="FFFFFF", size=12)
    ws3["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws3.row_dimensions[1].height = 28

    cols3 = ["Código", "Nombre de Tarea", "WBS", "Inicio P6", "Fin P6", "Duración (d)", "% Avance"]
    for col, h in enumerate(cols3, 1):
        ws3.cell(row=2, column=col, value=h)
    hdr(ws3, 2, len(cols3))

    tareas_abb = [t for t in tareas_p6 if clasificar(t["nombre"]) == "ABB"]
    for row_i, t in enumerate(sorted(tareas_abb, key=lambda x: x["inicio"] or date(2099,1,1)), 3):
        data_cell(ws3, row_i, 1, t["codigo"])
        data_cell(ws3, row_i, 2, t["nombre"], bold=True)
        data_cell(ws3, row_i, 3, t["wbs_nombre"])
        data_cell(ws3, row_i, 4, t["inicio"].strftime("%d/%m/%Y") if t["inicio"] else "—", align="center")
        data_cell(ws3, row_i, 5, t["fin"].strftime("%d/%m/%Y") if t["fin"] else "—", align="center")
        data_cell(ws3, row_i, 6, t["duracion"], align="center")
        data_cell(ws3, row_i, 7, t["avance"], align="center")
        ws3.row_dimensions[row_i].height = 20

    ws3.column_dimensions["A"].width = 16
    ws3.column_dimensions["B"].width = 50
    ws3.column_dimensions["C"].width = 30
    ws3.column_dimensions["D"].width = 14
    ws3.column_dimensions["E"].width = 14
    ws3.column_dimensions["F"].width = 14
    ws3.column_dimensions["G"].width = 12
    ws3.freeze_panes = ws3["A3"]


    # ═══════════════════════════════════════════════════════════════════
    # 4. HITOS RFSU / COMISIONADO EN P6
    # ═══════════════════════════════════════════════════════════════════
    ws4 = wb.create_sheet("RFSU y Comisionado P6")
    ws4.merge_cells("A1:G1")
    ws4["A1"] = "HITOS DE RFSU, COMISIONADO E iFAT EN EL CRONOGRAMA GLOBAL P6"
    ws4["A1"].fill = PatternFill("solid", fgColor="CC0000")
    ws4["A1"].font = Font(bold=True, color="FFFFFF", size=12)
    ws4["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws4.row_dimensions[1].height = 28

    cols4 = ["Código", "Nombre de Tarea", "WBS", "Inicio P6", "Fin P6", "Duración (d)", "Categoría"]
    for col, h in enumerate(cols4, 1):
        ws4.cell(row=2, column=col, value=h)
    hdr(ws4, 2, len(cols4))

    tareas_crit = [t for t in tareas_p6 if clasificar(t["nombre"]) in ("RFSU/Comisionado", "iFAT")]
    for row_i, t in enumerate(sorted(tareas_crit, key=lambda x: x["inicio"] or date(2099,1,1)), 3):
        cat = clasificar(t["nombre"])
        bg = "E2EFDA" if cat == "RFSU/Comisionado" else "DDEEFF"
        data_cell(ws4, row_i, 1, t["codigo"])
        data_cell(ws4, row_i, 2, t["nombre"], bold=True)
        data_cell(ws4, row_i, 3, t["wbs_nombre"])
        data_cell(ws4, row_i, 4, t["inicio"].strftime("%d/%m/%Y") if t["inicio"] else "—", align="center")
        data_cell(ws4, row_i, 5, t["fin"].strftime("%d/%m/%Y") if t["fin"] else "—", align="center")
        data_cell(ws4, row_i, 6, t["duracion"], align="center")
        data_cell(ws4, row_i, 7, cat, bg=bg)
        ws4.row_dimensions[row_i].height = 20

    for col, w in zip("ABCDEFG", [16, 50, 30, 14, 14, 14, 20]):
        ws4.column_dimensions[col].width = w
    ws4.freeze_panes = ws4["A3"]


    # ═══════════════════════════════════════════════════════════════════
    # 5. CRONOGRAMA GLOBAL P6 COMPLETO
    # ═══════════════════════════════════════════════════════════════════
    ws5 = wb.create_sheet("Cronograma P6 Completo")
    ws5.merge_cells("A1:H1")
    ws5["A1"] = "CRONOGRAMA GLOBAL CPF-2 — TODAS LAS TAREAS (P6 Rev.0)"
    ws5["A1"].fill = PatternFill("solid", fgColor="333333")
    ws5["A1"].font = Font(bold=True, color="FFFFFF", size=12)
    ws5["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws5.row_dimensions[1].height = 28

    cols5 = ["Código", "Nombre de Tarea", "WBS", "Inicio", "Fin", "Dur (d)", "% Av.", "Categoría"]
    for col, h in enumerate(cols5, 1):
        ws5.cell(row=2, column=col, value=h)
    hdr(ws5, 2, len(cols5), bg="333333")

    color_cat = {
        "ABB":              "FFE0E0",
        "RFSU/Comisionado": "E2EFDA",
        "iFAT":             "DDEEFF",
        "Civil/Montaje":    "FFF2CC",
        "Proyecto General": "F5F5F5",
    }
    for row_i, t in enumerate(sorted(tareas_p6, key=lambda x: x["inicio"] or date(2099,1,1)), 3):
        cat = clasificar(t["nombre"])
        bg = color_cat.get(cat, "F5F5F5")
        data_cell(ws5, row_i, 1, t["codigo"])
        data_cell(ws5, row_i, 2, t["nombre"])
        data_cell(ws5, row_i, 3, t["wbs_nombre"])
        data_cell(ws5, row_i, 4, t["inicio"].strftime("%d/%m/%Y") if t["inicio"] else "—", align="center")
        data_cell(ws5, row_i, 5, t["fin"].strftime("%d/%m/%Y") if t["fin"] else "—", align="center")
        data_cell(ws5, row_i, 6, t["duracion"], align="center")
        data_cell(ws5, row_i, 7, t["avance"], align="center")
        data_cell(ws5, row_i, 8, cat, bg=bg)
        ws5.row_dimensions[row_i].height = 16

    for col, w in zip("ABCDEFGH", [16, 50, 30, 14, 14, 10, 8, 20]):
        ws5.column_dimensions[col].width = w
    ws5.freeze_panes = ws5["A3"]

    # Guardar
    out = os.path.join(PLANNING_DIR, "Analisis_Cronogramas_CPF2.xlsx")
    wb.save(out)
    return out, tareas_abb, tareas_crit, resumen_semaforo


# ── MAIN ──────────────────────────────────────────────────────────────────────
if __name__ == "__main__":
    print("\n=== Análisis Comparativo de Cronogramas CPF-2 ===\n")

    # Buscar archivo XER
    xer_path = None
    for f in os.listdir(PLANNING_DIR):
        if f.endswith(".xer"):
            xer_path = os.path.join(PLANNING_DIR, f)
            break
    if not xer_path:
        print("❌ No se encontró archivo .xer en 02_Planning/")
        exit(1)

    print(f"📂 Leyendo: {os.path.basename(xer_path)}")
    tables = parse_xer(xer_path)
    print(f"   Tablas encontradas: {list(tables.keys())[:10]}...")

    proyecto_p6 = extraer_proyecto(tables)
    wbs_map     = extraer_wbs(tables)
    tareas_p6   = extraer_tareas(tables, wbs_map)

    print(f"   Proyecto: {proyecto_p6.get('nombre','N/D')}")
    print(f"   Data Date: {proyecto_p6.get('data_date','N/D')}")
    print(f"   Tareas extraídas: {len(tareas_p6)}")
    print(f"   WBS nodes: {len(wbs_map)}")

    tareas_con_fecha = [t for t in tareas_p6 if t["inicio"] and t["fin"]]
    if tareas_con_fecha:
        inicio_min = min(t["inicio"] for t in tareas_con_fecha)
        fin_max    = max(t["fin"]    for t in tareas_con_fecha)
        print(f"   Rango del cronograma: {inicio_min} → {fin_max}")

    print("\n📊 Generando reporte comparativo...")
    out, tareas_abb, tareas_crit, semaforo_res = generar_reporte(proyecto_p6, tareas_p6, wbs_map)

    print(f"\n✅ Reporte generado: {out}")
    print(f"\n── Tareas ABB encontradas en P6: {len(tareas_abb)}")
    print(f"── Hitos RFSU/Comisionado en P6: {len(tareas_crit)}")
    print(f"\n── Semáforo de compatibilidad:")
    for k, v in semaforo_res.items():
        print(f"   {k}: {v}")

    print("\n── Muestra de tareas ABB en P6:")
    for t in sorted(tareas_abb, key=lambda x: x["inicio"] or date(2099,1,1))[:8]:
        print(f"   [{t['codigo']}] {t['nombre'][:55]:<55} {str(t['inicio'])} → {str(t['fin'])}")
