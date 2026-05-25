"""
Gantt por cadena crítica — Provisión ABB (CPF-2 Vaca Muerta)
Salas Eléctricas #3, #4 y PMS — RFSU 31/01/2028
"""

import os
from datetime import date, timedelta
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

BASE_DIR = os.path.dirname(os.path.abspath(__file__))

# ── Paleta ────────────────────────────────────────────────────────────────────
C = {
    "red_abb":     "CC0000",
    "red_dark":    "880000",
    "critico":     "CC0000",   # cadena crítica (S#3)
    "no_critico":  "5B9BD5",   # actividades no críticas (S#4, PMS)
    "fat":         "FF6600",   # FAT / Ensayos
    "despacho":    "7030A0",   # Despacho
    "integracion": "375623",   # iFAT / Comisionado
    "hito":        "FF0000",   # ◆ Freezing Points y hitos
    "hito_final":  "1F4E79",   # RFSU
    "header_mes":  "1F4E79",
    "header_bg":   "2E3A4E",
    "sub_header":  "404040",
    "fase_s3":     "FFE0E0",
    "fase_s4":     "E0EEFF",
    "fase_pms":    "FFFDE0",
    "fase_int":    "E0FFE5",
    "row_alt":     "F7F7F7",
    "row_white":   "FFFFFF",
    "borde":       "BBBBBB",
}

def bd():
    s = Side(style="thin", color=C["borde"])
    return Border(left=s, right=s, top=s, bottom=s)

def fill(c): return PatternFill("solid", fgColor=c)

# ── Actividades ABB con cadena crítica marcada ────────────────────────────────
# Campos: (grupo, actividad, inicio, fin, tipo, critica, nota)
# tipo: "actividad" | "hito" | "fat" | "despacho" | "integracion"
# critica: True = cadena crítica (rojo), False = paralela (azul)

ACTIVIDADES = [
    # ─── PMS ─────────────────────────────────────────────────────────────────
    ("PMS", "APR Ingeniería Básica PMS (S#3 y S#4)",
     date(2026,  6,  8), date(2026,  6, 19), "actividad", False,
     "Aprobación IB Tablero Control PMS"),

    ("PMS", "◆ FP Ing. Básica PMS",
     date(2026,  6, 19), date(2026,  6, 19), "hito", False,
     "Freezing Point IB PMS — S#3 y S#4"),

    ("PMS", "APR Ingeniería de Detalle PMS (S#3 y S#4)",
     date(2026,  7, 20), date(2026,  7, 31), "actividad", False,
     "Aprobación ID Tablero Control PMS"),

    ("PMS", "◆ FP Ing. Detalle PMS",
     date(2026,  7, 31), date(2026,  7, 31), "hito", False,
     "Freezing Point ID PMS — S#3 y S#4"),

    ("PMS", "Ensayos + FAT PMS (S#3 y S#4)",
     date(2026, 10, 19), date(2026, 10, 23), "fat", False,
     "FAT PMS — 4 días. HOLGURA 222 días vs demanda P6"),

    # ─── SALA #4 (no crítica) ─────────────────────────────────────────────────
    ("Sala #4", "Emisión IB / APR IB Sala #4",
     date(2026,  5, 22), date(2026,  6, 26), "actividad", False,
     "35 días. Aprobación AESA: 12/06/2026"),

    ("Sala #4", "◆ FP Ing. Básica S#4",
     date(2026,  6, 26), date(2026,  6, 26), "hito", False,
     "Freezing Point post IB Sala #4"),

    ("Sala #4", "Emisión ID / APR ID Sala #4",
     date(2026,  6, 26), date(2026,  8, 14), "actividad", False,
     "49 días. Aprobación AESA: 31/07/2026"),

    ("Sala #4", "◆ FP Ing. Detalle S#4",
     date(2026,  8, 14), date(2026,  8, 14), "hito", False,
     "Freezing Point post ID Sala #4"),

    ("Sala #4", "APR Ing. Constructiva Sala #4",
     date(2026,  8, 10), date(2026,  8, 24), "actividad", False,
     "14 días"),

    ("Sala #4", "◆ FP Ing. Constructiva S#4",
     date(2026,  8, 24), date(2026,  8, 24), "hito", False,
     "Cierre ingeniería S#4 — 24/08/2026"),

    ("Sala #4", "Ensayos Celdas y Tableros Sala #4",
     date(2026, 11,  3), date(2026, 11, 23), "fat", False,
     "20 días. Fabricación + pruebas de componentes"),

    ("Sala #4", "FAT Sala Eléctrica #4 (Shelter completo)",
     date(2027,  3,  8), date(2027,  3, 16), "fat", False,
     "8 días. FAT final S#4"),

    ("Sala #4", "▶ Despacho Sala #4",
     date(2027,  4,  2), date(2027,  4,  8), "despacho", False,
     "Despacho a Vaca Muerta. P6 da margen hasta 23/06/2027"),

    # ─── SALA #3 — CADENA CRÍTICA ─────────────────────────────────────────────
    ("Sala #3  ★", "Emisión IB / APR IB Sala #3",
     date(2026,  5, 22), date(2026,  6, 26), "actividad", True,
     "35 días. Paralela a S#4"),

    ("Sala #3  ★", "◆ FP Ing. Básica S#3",
     date(2026,  7,  3), date(2026,  7,  3), "hito", True,
     "Freezing Point post IB Sala #3 — 1 semana después que S#4"),

    ("Sala #3  ★", "Emisión ID / APR ID Sala #3",
     date(2026,  6, 26), date(2026,  7, 31), "actividad", True,
     "35 días. Aprobación AESA: 17/07/2026"),

    ("Sala #3  ★", "◆ FP Ing. Detalle S#3",
     date(2026,  7, 31), date(2026,  7, 31), "hito", True,
     "Freezing Point post ID Sala #3"),

    ("Sala #3  ★", "APR Ing. Constructiva Sala #3",
     date(2026,  8, 17), date(2026,  8, 31), "actividad", True,
     "14 días — 1 semana después que S#4"),

    ("Sala #3  ★", "◆ FP Ing. Constructiva S#3  ← ÚLTIMO FP",
     date(2026,  8, 31), date(2026,  8, 31), "hito", True,
     "ÚLTIMO FREEZING POINT ABB — 31/08/2026. Cierra toda la ingeniería"),

    ("Sala #3  ★", "Ensayos Celdas y Tableros Sala #3",
     date(2026, 11,  3), date(2026, 11, 23), "fat", True,
     "20 días. Paralela a S#4"),

    ("Sala #3  ★", "FAT Sala Eléctrica #3  ← FAT MÁS TARDÍO",
     date(2027,  4,  2), date(2027,  4, 19), "fat", True,
     "17 días. ÚLTIMO FAT ABB — más tardío de los dos"),

    ("Sala #3  ★", "▶ Despacho Sala #3  ← DESPACHO MÁS TARDÍO",
     date(2027,  5, 10), date(2027,  5, 14), "despacho", True,
     "ÚLTIMO DESPACHO ABB. Campo lo necesita el 28/06/2027 (+45d holgura)"),

    # ─── INTEGRACIÓN Y CIERRE ─────────────────────────────────────────────────
    ("Integración", "iFAT Integrado EPC (Shelters + PMS + INAUCO + HIMA)",
     date(2027,  6, 25), date(2027,  7, 14), "integracion", True,
     "19 días. P6 hito TA1124 = 28/06/2027. Todos los vendors"),

    ("Integración", "Precomisionado",
     date(2027,  5,  3), date(2028,  1, 13), "integracion", False,
     "255 días. CDW1000 — inicio idéntico en P6 y Gantt ABB"),

    ("Integración", "Comisionado",
     date(2027,  7, 16), date(2028,  1, 31), "integracion", True,
     "199 días. CDW1010 — inicio idéntico en P6 y Gantt ABB"),

    ("Integración", "★ RFSU — Ready For Start Up  31/01/2028",
     date(2028,  1, 31), date(2028,  1, 31), "hito", True,
     "A10290 — IDÉNTICO en P6 y Gantt ABB ✅"),
]

# ── Generación del Gantt ──────────────────────────────────────────────────────

INICIO_GANTT = date(2026,  5,  1)
FIN_GANTT    = date(2028,  2, 28)
SEMANAS      = []
d = INICIO_GANTT
while d <= FIN_GANTT:
    SEMANAS.append(d)
    d += timedelta(weeks=1)

N_INFO_COLS = 4   # Fase | Actividad | Inicio | Fin
COL_OFFSET  = N_INFO_COLS + 1


def semana_idx(fecha):
    for i, s in enumerate(SEMANAS):
        if s + timedelta(days=6) >= fecha >= s:
            return i
    return None


def rango_semanas(inicio, fin):
    """Devuelve (col_inicio, col_fin) en índices de SEMANAS."""
    if inicio == fin:  # hito puntual
        idx = semana_idx(inicio)
        return (idx, idx) if idx is not None else (None, None)
    ci = None
    cf = None
    for i, s in enumerate(SEMANAS):
        fin_semana = s + timedelta(days=6)
        if ci is None and fin_semana >= inicio:
            ci = i
        if s <= fin:
            cf = i
    return ci, cf


def color_barra(tipo, critica):
    if tipo == "hito":
        return C["hito"] if critica else C["hito"]
    if tipo == "fat":
        return C["fat"]
    if tipo == "despacho":
        return C["despacho"]
    if tipo == "integracion":
        return C["integracion"]
    return C["critico"] if critica else C["no_critico"]


def generar():
    wb = Workbook()
    ws = wb.active
    ws.title = "Gantt Cadena Crítica"

    # ── Configuración de columnas ──────────────────────────────────────────────
    ws.column_dimensions["A"].width = 14   # Fase
    ws.column_dimensions["B"].width = 42   # Actividad
    ws.column_dimensions["C"].width = 11   # Inicio
    ws.column_dimensions["D"].width = 11   # Fin

    for i, _ in enumerate(SEMANAS):
        col = get_column_letter(COL_OFFSET + i)
        ws.column_dimensions[col].width = 2.3

    # ── Fila 1: Título ─────────────────────────────────────────────────────────
    last_col = get_column_letter(COL_OFFSET + len(SEMANAS) - 1)
    ws.merge_cells(f"A1:{last_col}1")
    c = ws.cell(row=1, column=1,
                value="GANTT POR CADENA CRÍTICA — PROVISIÓN ABB  |  CPF-2 VACA MUERTA  |  RFSU 31/01/2028")
    c.fill = fill(C["red_abb"])
    c.font = Font(bold=True, color="FFFFFF", size=13)
    c.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 32

    # ── Fila 2: Leyenda ────────────────────────────────────────────────────────
    leyenda = [
        (C["critico"],    "Cadena Crítica (S#3)"),
        (C["no_critico"], "Actividades Paralelas (S#4 / PMS)"),
        (C["fat"],        "FAT / Ensayos"),
        (C["despacho"],   "Despacho"),
        (C["integracion"],"iFAT / Comisionado / RFSU"),
        (C["hito"],       "◆ Freezing Point / Hito"),
    ]
    ws.merge_cells(f"A2:{last_col}2")
    leg_txt = "   ".join(f"■ {txt}" for _, txt in leyenda)
    c2 = ws.cell(row=2, column=1, value=leg_txt)
    c2.font = Font(size=8, color="444444")
    c2.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[2].height = 14

    # ── Fila 3: Encabezado años ────────────────────────────────────────────────
    row_anio = 3
    anio_grupos = {}
    for i, s in enumerate(SEMANAS):
        anio = s.year
        if anio not in anio_grupos:
            anio_grupos[anio] = {"start": i, "end": i}
        else:
            anio_grupos[anio]["end"] = i

    for anio, rng in anio_grupos.items():
        c_ini = get_column_letter(COL_OFFSET + rng["start"])
        c_fin = get_column_letter(COL_OFFSET + rng["end"])
        try:
            ws.merge_cells(f"{c_ini}{row_anio}:{c_fin}{row_anio}")
        except Exception:
            pass
        c = ws.cell(row=row_anio, column=COL_OFFSET + rng["start"], value=str(anio))
        c.fill = fill(C["header_mes"])
        c.font = Font(bold=True, color="FFFFFF", size=9)
        c.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[row_anio].height = 16

    # ── Fila 4: Meses ─────────────────────────────────────────────────────────
    row_mes = 4
    MESES_ES = ["", "Ene", "Feb", "Mar", "Abr", "May", "Jun",
                "Jul", "Ago", "Sep", "Oct", "Nov", "Dic"]
    mes_grupos = {}
    for i, s in enumerate(SEMANAS):
        key = (s.year, s.month)
        if key not in mes_grupos:
            mes_grupos[key] = {"start": i, "end": i}
        else:
            mes_grupos[key]["end"] = i

    for (anio, mes), rng in mes_grupos.items():
        c_ini = get_column_letter(COL_OFFSET + rng["start"])
        c_fin = get_column_letter(COL_OFFSET + rng["end"])
        try:
            ws.merge_cells(f"{c_ini}{row_mes}:{c_fin}{row_mes}")
        except Exception:
            pass
        label = f"{MESES_ES[mes]}"
        c = ws.cell(row=row_mes, column=COL_OFFSET + rng["start"], value=label)
        bg = C["sub_header"] if mes % 2 == 0 else "555555"
        c.fill = fill(bg)
        c.font = Font(bold=True, color="FFFFFF", size=7)
        c.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[row_mes].height = 14

    # ── Fila 5: Encabezados de info ────────────────────────────────────────────
    row_hdr = 5
    for col, txt in enumerate(["Fase", "Actividad / Hito", "Inicio", "Fin"], 1):
        c = ws.cell(row=row_hdr, column=col, value=txt)
        c.fill = fill(C["header_bg"])
        c.font = Font(bold=True, color="FFFFFF", size=9)
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = bd()
    ws.row_dimensions[row_hdr].height = 18

    # ── Líneas de mes (marcadores verticales) ─────────────────────────────────
    # Los aplicamos a cada fila de datos al rellenar

    # ── Filas de datos ─────────────────────────────────────────────────────────
    grupo_anterior = None
    row_data = 6
    colores_grupo = {
        "PMS":          C["fase_pms"],
        "Sala #4":      C["fase_s4"],
        "Sala #3  ★":   C["fase_s3"],
        "Integración":  C["fase_int"],
    }

    for (grupo, nombre, ini, fin, tipo, critica, nota) in ACTIVIDADES:

        # Separador de grupo
        if grupo != grupo_anterior:
            ws.merge_cells(f"A{row_data}:{last_col}{row_data}")
            lbl = f"  ▶  {grupo.upper()}"
            if grupo == "Sala #3  ★":
                lbl += "   ← CADENA CRÍTICA"
            c = ws.cell(row=row_data, column=1, value=lbl)
            bg_sep = C["red_dark"] if "★" in grupo else C["sub_header"]
            c.fill = fill(bg_sep)
            c.font = Font(bold=True, color="FFFFFF", size=9)
            c.alignment = Alignment(horizontal="left", vertical="center")
            ws.row_dimensions[row_data].height = 18
            row_data += 1
            grupo_anterior = grupo

        bg_info = colores_grupo.get(grupo, "F5F5F5")
        es_hito = (tipo == "hito")

        # Columnas de info
        ws.cell(row=row_data, column=1, value=grupo).fill = fill(bg_info)
        ws.cell(row=row_data, column=1).font  = Font(size=8, color="444444")
        ws.cell(row=row_data, column=1).border = bd()
        ws.cell(row=row_data, column=1).alignment = Alignment(horizontal="center", vertical="center")

        ws.cell(row=row_data, column=2, value=nombre).fill = fill(bg_info)
        ws.cell(row=row_data, column=2).font = Font(
            bold=es_hito or critica,
            size=9,
            color=C["red_dark"] if (critica and not es_hito) else "000000"
        )
        ws.cell(row=row_data, column=2).border = bd()
        ws.cell(row=row_data, column=2).alignment = Alignment(vertical="center", wrap_text=False)

        ws.cell(row=row_data, column=3, value=ini.strftime("%d/%m/%Y")).fill = fill(bg_info)
        ws.cell(row=row_data, column=3).font = Font(size=8)
        ws.cell(row=row_data, column=3).border = bd()
        ws.cell(row=row_data, column=3).alignment = Alignment(horizontal="center", vertical="center")

        ws.cell(row=row_data, column=4, value=fin.strftime("%d/%m/%Y")).fill = fill(bg_info)
        ws.cell(row=row_data, column=4).font = Font(size=8)
        ws.cell(row=row_data, column=4).border = bd()
        ws.cell(row=row_data, column=4).alignment = Alignment(horizontal="center", vertical="center")

        # ── Barras del Gantt ───────────────────────────────────────────────────
        ci, cf = rango_semanas(ini, fin)
        color_bar = color_barra(tipo, critica)

        for i in range(len(SEMANAS)):
            col = COL_OFFSET + i
            semana_actual = SEMANAS[i]
            c = ws.cell(row=row_data, column=col)

            # Borde de mes (línea izquierda más oscura)
            izq_style = "medium" if semana_actual.day <= 7 else "thin"
            izq_color  = "888888" if semana_actual.day <= 7 else C["borde"]
            c.border = Border(
                left=Side(style=izq_style, color=izq_color),
                right=Side(style="thin", color=C["borde"]),
                top=Side(style="thin", color=C["borde"]),
                bottom=Side(style="thin", color=C["borde"]),
            )

            if ci is not None and cf is not None and ci <= i <= cf:
                if tipo == "hito" and ci == cf:
                    # Hito puntual — celda con símbolo y color de fondo
                    c.fill = fill(color_bar)
                    c.value = "◆"
                    c.font = Font(bold=True, color="FFFFFF", size=7)
                    c.alignment = Alignment(horizontal="center", vertical="center")
                elif tipo == "hito" and ci == i:
                    c.fill = fill(color_bar)
                    c.value = "◆"
                    c.font = Font(bold=True, color="FFFFFF", size=7)
                    c.alignment = Alignment(horizontal="center", vertical="center")
                else:
                    c.fill = fill(color_bar)
                    # Para el RFSU usar símbolo especial
                    if "RFSU" in nombre and ci == i:
                        c.value = "★"
                        c.font = Font(bold=True, color="FFFFFF", size=8)
                        c.alignment = Alignment(horizontal="center", vertical="center")
                        c.fill = fill(C["hito_final"])
            else:
                c.fill = fill("FAFAFA" if i % 2 == 0 else "F3F3F3")

        # Altura de fila
        ws.row_dimensions[row_data].height = 16 if not es_hito else 14
        row_data += 1

    # ── Agregar nota de duración en días ──────────────────────────────────────
    row_data += 1
    ws.merge_cells(f"A{row_data}:{last_col}{row_data}")
    nota_gen = (
        "★ CADENA CRÍTICA: Sala #3 determina el despacho más tardío (14/05/2027). "
        "Holgura vs P6: Despacho S#3 +45d | FAT S#3 +35d | PMS +222d | RFSU = 31/01/2028 (idéntico en P6)"
    )
    c = ws.cell(row=row_data, column=1, value=nota_gen)
    c.fill = fill("1F4E79")
    c.font = Font(bold=True, color="FFFFFF", size=9)
    c.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    ws.row_dimensions[row_data].height = 20

    # ── Freeze y ajustes finales ───────────────────────────────────────────────
    ws.freeze_panes = "E6"

    # ── Hoja 2: Tabla de hitos críticos con duraciones ────────────────────────
    ws2 = wb.create_sheet("Tabla Cadena Crítica")

    ws2.merge_cells("A1:G1")
    c = ws2.cell(row=1, column=1, value="TABLA DE ACTIVIDADES — CADENA CRÍTICA ABB (CPF-2)")
    c.fill = fill(C["red_abb"]); c.font = Font(bold=True, color="FFFFFF", size=12)
    c.alignment = Alignment(horizontal="center", vertical="center")
    ws2.row_dimensions[1].height = 30

    hdrs2 = ["Fase", "Actividad", "Inicio", "Fin", "Duración (d)", "Tipo", "Cadena Crítica"]
    for col, h in enumerate(hdrs2, 1):
        c = ws2.cell(row=2, column=col, value=h)
        c.fill = fill(C["red_dark"]); c.font = Font(bold=True, color="FFFFFF", size=9)
        c.alignment = Alignment(horizontal="center", vertical="center"); c.border = bd()
    ws2.row_dimensions[2].height = 22

    for r, (grupo, nombre, ini, fin, tipo, critica, nota) in enumerate(ACTIVIDADES, 3):
        dur = (fin - ini).days if ini != fin else 0
        bg = C["fase_s3"] if critica else ("FFFFFF" if r % 2 == 0 else "F5F5F5")
        vals = [grupo, nombre, ini.strftime("%d/%m/%Y"), fin.strftime("%d/%m/%Y"),
                dur if dur > 0 else "Hito", tipo.capitalize(), "★ CRÍTICA" if critica else "Paralela"]
        for col, v in enumerate(vals, 1):
            c = ws2.cell(row=r, column=col, value=v)
            c.fill = fill(bg); c.font = Font(bold=critica, size=9,
                                              color=C["red_dark"] if critica else "000000")
            c.border = bd()
            c.alignment = Alignment(vertical="center", horizontal="center" if col > 2 else "left")
        ws2.cell(row=r, column=8, value=nota).border = bd()
        ws2.cell(row=r, column=8).font = Font(size=8, color="666666")
        ws2.cell(row=r, column=8).alignment = Alignment(wrap_text=True, vertical="center")
        ws2.row_dimensions[r].height = 18

    for col, w in zip("ABCDEFGH", [14, 48, 12, 12, 13, 14, 14, 50]):
        ws2.column_dimensions[get_column_letter(ord(col)-64)].width = w
    ws2.freeze_panes = ws2["A3"]

    out = os.path.join(BASE_DIR, "Gantt_Cadena_Critica_ABB.xlsx")
    wb.save(out)
    return out


if __name__ == "__main__":
    print("\n=== Generando Gantt Cadena Crítica ABB ===\n")
    print(f"  Período: {INICIO_GANTT} → {FIN_GANTT}")
    print(f"  Semanas en el Gantt: {len(SEMANAS)}")
    print(f"  Actividades: {len(ACTIVIDADES)}")
    out = generar()
    print(f"\n✅ Gantt generado: {out}")
    print("\nCadena crítica (S#3 — fechas más tardías):")
    for (g, n, i, f, t, crit, _) in ACTIVIDADES:
        if crit:
            dur = (f - i).days
            print(f"  [{i.strftime('%d/%m/%y')} → {f.strftime('%d/%m/%y')}]  {n[:60]}")
