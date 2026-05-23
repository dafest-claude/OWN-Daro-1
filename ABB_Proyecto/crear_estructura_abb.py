"""
Script de inicialización del sistema de seguimiento de proyecto ABB.
Crea todos los archivos Excel estructurados y el Dashboard maestro.
"""

import os
from datetime import datetime, date
from openpyxl import Workbook
from openpyxl.styles import (
    PatternFill, Font, Alignment, Border, Side, GradientFill
)
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

# ── Colores corporativos ABB ──────────────────────────────────────────────────
COLOR_ABB_RED    = "FF0000"
COLOR_ABB_DARK   = "1A1A1A"
COLOR_HEADER_BG  = "CC0000"
COLOR_HEADER_FG  = "FFFFFF"
COLOR_SUBHEADER  = "F5F5F5"
COLOR_ACCENT     = "E8E8E8"
COLOR_OK         = "C6EFCE"
COLOR_WARN       = "FFEB9C"
COLOR_ERROR      = "FFC7CE"
COLOR_BORDER     = "CCCCCC"

BASE_DIR = os.path.dirname(os.path.abspath(__file__))


# ── Helpers de estilo ─────────────────────────────────────────────────────────

def header_fill(color=COLOR_HEADER_BG):
    return PatternFill("solid", fgColor=color)

def thin_border():
    side = Side(style="thin", color=COLOR_BORDER)
    return Border(left=side, right=side, top=side, bottom=side)

def style_header_row(ws, row, cols, bg=COLOR_HEADER_BG, fg=COLOR_HEADER_FG, bold=True, size=10):
    for col in range(1, cols + 1):
        cell = ws.cell(row=row, column=col)
        cell.fill = PatternFill("solid", fgColor=bg)
        cell.font = Font(bold=bold, color=fg, size=size)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border()

def style_title(ws, cell_ref, text, bg=COLOR_HEADER_BG, fg=COLOR_HEADER_FG):
    cell = ws[cell_ref]
    cell.value = text
    cell.fill = PatternFill("solid", fgColor=bg)
    cell.font = Font(bold=True, color=fg, size=14)
    cell.alignment = Alignment(horizontal="center", vertical="center")

def set_col_widths(ws, widths: dict):
    for col_letter, width in widths.items():
        ws.column_dimensions[col_letter].width = width

def freeze(ws, cell="A2"):
    ws.freeze_panes = ws[cell]


# ── 1. OFERTAS ────────────────────────────────────────────────────────────────

def crear_ofertas():
    wb = Workbook()
    ws = wb.active
    ws.title = "Ofertas"
    ws.row_dimensions[1].height = 14
    ws.row_dimensions[2].height = 30

    ws.merge_cells("A1:K1")
    style_title(ws, "A1", "SEGUIMIENTO DE OFERTAS — PROYECTO ABB")

    headers = [
        "N° Oferta", "Descripción", "Proveedor / Origen",
        "Fecha Emisión", "Fecha Vencimiento", "Moneda",
        "Importe", "Estado", "Responsable", "Observaciones", "Archivos Adjuntos"
    ]
    for col, h in enumerate(headers, 1):
        ws.cell(row=2, column=col, value=h)
    style_header_row(ws, 2, len(headers))

    sample = [
        ["OFE-001", "Descripción de la oferta", "Proveedor ejemplo",
         date.today(), date.today(), "USD", 0, "En evaluación", "", "", ""],
    ]
    for r, row in enumerate(sample, 3):
        for c, val in enumerate(row, 1):
            cell = ws.cell(row=r, column=c, value=val)
            cell.border = thin_border()
            if c == 8:
                cell.fill = PatternFill("solid", fgColor=COLOR_WARN)
            cell.alignment = Alignment(vertical="center")

    set_col_widths(ws, {
        "A": 12, "B": 35, "C": 25, "D": 14, "E": 16,
        "F": 10, "G": 14, "H": 16, "I": 18, "J": 35, "K": 25
    })
    freeze(ws, "A3")

    ws2 = wb.create_sheet("Leyenda Estados")
    estados = [
        ("En evaluación", COLOR_WARN),
        ("Aprobada",      COLOR_OK),
        ("Rechazada",     COLOR_ERROR),
        ("Vencida",       "D9D9D9"),
        ("En negociación","FFD966"),
    ]
    ws2["A1"] = "Estado"
    ws2["B1"] = "Descripción"
    style_header_row(ws2, 1, 2)
    for i, (estado, color) in enumerate(estados, 2):
        ws2.cell(row=i, column=1, value=estado).fill = PatternFill("solid", fgColor=color)
        ws2.cell(row=i, column=2, value="")
    ws2.column_dimensions["A"].width = 20
    ws2.column_dimensions["B"].width = 40

    path = os.path.join(BASE_DIR, "01_Ofertas", "Ofertas.xlsx")
    wb.save(path)
    print(f"  ✓ {path}")


# ── 2. PLANNING ───────────────────────────────────────────────────────────────

def crear_planning():
    wb = Workbook()

    # ── Pestaña: Cronograma
    ws = wb.active
    ws.title = "Cronograma"
    ws.merge_cells("A1:N1")
    style_title(ws, "A1", "PLANNING DEL PROYECTO — PROYECTO ABB")

    headers = [
        "ID", "Fase", "Actividad / Entregable", "Responsable",
        "Fecha Inicio Plan", "Fecha Fin Plan",
        "Fecha Inicio Real", "Fecha Fin Real",
        "Duración (días)", "% Avance", "Estado",
        "Predecesores", "Hitos Clave", "Observaciones"
    ]
    for col, h in enumerate(headers, 1):
        ws.cell(row=2, column=col, value=h)
    style_header_row(ws, 2, len(headers))

    fases = [
        ("1", "Inicio",          "Kick-Off Meeting (KOM)",             "", "", "", "", "", "", 0),
        ("2", "Ingeniería",      "Revisión de especificaciones",        "", "", "", "", "", "", 0),
        ("3", "Ingeniería",      "Emisión de planos / documentación",   "", "", "", "", "", "", 0),
        ("4", "Procurement",     "Órdenes de compra",                   "", "", "", "", "", "", 0),
        ("5", "Procurement",     "Seguimiento de entregas",             "", "", "", "", "", "", 0),
        ("6", "Ejecución",       "Instalación / Montaje",               "", "", "", "", "", "", 0),
        ("7", "Commissioning",   "Pruebas y puesta en marcha",          "", "", "", "", "", "", 0),
        ("8", "Cierre",          "Documentación As-Built",              "", "", "", "", "", "", 0),
        ("9", "Cierre",          "Acta de cierre y aceptación",         "", "", "", "", "", "", 0),
    ]
    for r, (id_, fase, act, resp, fi_p, ff_p, fi_r, ff_r, dur, av) in enumerate(fases, 3):
        vals = [id_, fase, act, resp, fi_p, ff_p, fi_r, ff_r, dur, av, "Pendiente", "", "", ""]
        for c, v in enumerate(vals, 1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.border = thin_border()
            cell.alignment = Alignment(vertical="center")
            if c == 11:
                cell.fill = PatternFill("solid", fgColor=COLOR_WARN)

    set_col_widths(ws, {
        "A": 6, "B": 16, "C": 35, "D": 18,
        "E": 16, "F": 16, "G": 16, "H": 16,
        "I": 14, "J": 10, "K": 14, "L": 14, "M": 18, "N": 35
    })
    freeze(ws, "A3")

    # ── Pestaña: Hitos
    wh = wb.create_sheet("Hitos")
    wh.merge_cells("A1:F1")
    style_title(wh, "A1", "HITOS DEL PROYECTO")
    hitos_headers = ["ID Hito", "Descripción", "Fecha Planificada", "Fecha Real", "Estado", "Observaciones"]
    for col, h in enumerate(hitos_headers, 1):
        wh.cell(row=2, column=col, value=h)
    style_header_row(wh, 2, len(hitos_headers))
    set_col_widths(wh, {"A": 10, "B": 40, "C": 18, "D": 18, "E": 16, "F": 35})
    freeze(wh, "A3")

    path = os.path.join(BASE_DIR, "02_Planning", "Planning.xlsx")
    wb.save(path)
    print(f"  ✓ {path}")


# ── 3. MINUTAS DE REUNIONES ───────────────────────────────────────────────────

def crear_minutas():
    wb = Workbook()
    ws = wb.active
    ws.title = "Índice Minutas"

    ws.merge_cells("A1:I1")
    style_title(ws, "A1", "ÍNDICE DE MINUTAS DE REUNIÓN — PROYECTO ABB")

    headers = [
        "N° Minuta", "Tipo", "Fecha", "Tema Principal",
        "Participantes", "Moderador", "N° Acuerdos", "Archivo", "Estado"
    ]
    for col, h in enumerate(headers, 1):
        ws.cell(row=2, column=col, value=h)
    style_header_row(ws, 2, len(headers))

    set_col_widths(ws, {
        "A": 12, "B": 18, "C": 14, "D": 40,
        "E": 35, "F": 20, "G": 14, "H": 35, "I": 14
    })
    freeze(ws, "A3")

    # ── Plantilla de minuta individual
    wm = wb.create_sheet("Plantilla Minuta")
    campos = [
        ("A1", "MINUTA DE REUNIÓN — PROYECTO ABB", "A1:F1"),
        ("A3", "N° Minuta:",       None),
        ("A4", "Tipo de Reunión:", None),
        ("A5", "Fecha:",           None),
        ("A6", "Hora Inicio:",     None),
        ("A7", "Hora Fin:",        None),
        ("A8", "Lugar / Plataforma:", None),
        ("A9", "Convocado por:",   None),
        ("A10", "Moderador:",      None),
    ]
    wm.merge_cells("A1:F1")
    style_title(wm, "A1", "MINUTA DE REUNIÓN — PROYECTO ABB")

    meta = [
        ("N° Minuta:",         "B3"),
        ("Tipo de Reunión:",   "B4"),
        ("Fecha:",             "B5"),
        ("Hora Inicio:",       "B6"),
        ("Hora Fin:",          "B7"),
        ("Lugar / Plataforma:","B8"),
        ("Convocado por:",     "B9"),
        ("Moderador:",         "B10"),
    ]
    for i, (label, _) in enumerate(meta, 3):
        wm.cell(row=i, column=1, value=label).font = Font(bold=True)

    wm["A12"] = "PARTICIPANTES"
    wm.merge_cells("A12:F12")
    style_header_row(wm, 12, 6, bg="333333")
    part_h = ["Nombre", "Empresa / Área", "Cargo", "Email", "Presente", "Firma"]
    for col, h in enumerate(part_h, 1):
        wm.cell(row=13, column=col, value=h)
    style_header_row(wm, 13, 6, bg="555555", size=9)
    for r in range(14, 20):
        for c in range(1, 7):
            wm.cell(row=r, column=c).border = thin_border()

    wm["A21"] = "TEMAS TRATADOS"
    wm.merge_cells("A21:F21")
    style_header_row(wm, 21, 6, bg="333333")
    temas_h = ["#", "Tema", "Presentado por", "Descripción / Resolución", "Decisión", "Notas"]
    for col, h in enumerate(temas_h, 1):
        wm.cell(row=22, column=col, value=h)
    style_header_row(wm, 22, 6, bg="555555", size=9)
    for r in range(23, 30):
        for c in range(1, 7):
            wm.cell(row=r, column=c).border = thin_border()

    wm["A31"] = "ACUERDOS Y ACTION ITEMS"
    wm.merge_cells("A31:F31")
    style_header_row(wm, 31, 6, bg=COLOR_ABB_RED)
    ai_h = ["#", "Acuerdo / Acción", "Responsable", "Fecha Límite", "Estado", "Observaciones"]
    for col, h in enumerate(ai_h, 1):
        wm.cell(row=32, column=col, value=h)
    style_header_row(wm, 32, 6, bg="880000", size=9)
    for r in range(33, 40):
        for c in range(1, 7):
            cell = wm.cell(row=r, column=c)
            cell.border = thin_border()

    set_col_widths(wm, {"A": 6, "B": 35, "C": 22, "D": 40, "E": 16, "F": 30})

    path = os.path.join(BASE_DIR, "03_Minutas", "Minutas_Reuniones.xlsx")
    wb.save(path)
    print(f"  ✓ {path}")


# ── 4. MINUTA KOM ─────────────────────────────────────────────────────────────

def crear_minuta_kom():
    wb = Workbook()
    ws = wb.active
    ws.title = "KOM"

    def sec_title(ws, row, text, merge_end="F"):
        ws.merge_cells(f"A{row}:{merge_end}{row}")
        cell = ws.cell(row=row, column=1, value=text)
        cell.fill = PatternFill("solid", fgColor=COLOR_ABB_RED)
        cell.font = Font(bold=True, color="FFFFFF", size=11)
        cell.alignment = Alignment(horizontal="left", vertical="center")
        ws.row_dimensions[row].height = 22

    def field_row(ws, row, label, val_col="B", label_col="A"):
        ws.cell(row=row, column=1, value=label).font = Font(bold=True, size=10)
        ws.cell(row=row, column=2).border = thin_border()
        ws.row_dimensions[row].height = 18

    ws.merge_cells("A1:F1")
    style_title(ws, "A1", "MINUTA DE KICK-OFF MEETING (KOM) — PROYECTO ABB")
    ws.row_dimensions[1].height = 35

    # Datos generales
    row = 3
    sec_title(ws, row, "1. DATOS GENERALES DEL PROYECTO")
    campos_gen = [
        "Nombre del Proyecto:", "N° de Proyecto / Contrato:", "Cliente:",
        "Gerente de Proyecto (ABB):", "Gerente de Proyecto (Cliente):",
        "Fecha de KOM:", "Lugar / Plataforma:", "Duración:",
    ]
    for i, campo in enumerate(campos_gen, row + 1):
        field_row(ws, i, campo)
        ws.cell(row=i, column=3).border = thin_border()
        ws.merge_cells(f"B{i}:F{i}")

    row = row + len(campos_gen) + 2
    sec_title(ws, row, "2. PARTICIPANTES")
    ws.row_dimensions[row].height = 22
    part_h = ["Nombre", "Empresa", "Cargo", "Email", "Teléfono", "Rol en el Proyecto"]
    for col, h in enumerate(part_h, 1):
        ws.cell(row=row + 1, column=col, value=h)
    style_header_row(ws, row + 1, 6, bg="880000", size=9)
    for r in range(row + 2, row + 14):
        for c in range(1, 7):
            ws.cell(row=r, column=c).border = thin_border()
        ws.row_dimensions[r].height = 18

    row = row + 15
    sec_title(ws, row, "3. ALCANCE DEL PROYECTO")
    scope_items = [
        "Descripción general del alcance:",
        "Entregables principales:",
        "Límites del alcance (lo que NO incluye):",
        "Supuestos y exclusiones:",
    ]
    for i, item in enumerate(scope_items, row + 1):
        ws.cell(row=i, column=1, value=item).font = Font(bold=True, size=10)
        ws.merge_cells(f"B{i}:F{i}")
        ws.cell(row=i, column=2).border = thin_border()
        ws.row_dimensions[i].height = 35

    row = row + len(scope_items) + 2
    sec_title(ws, row, "4. CRONOGRAMA Y HITOS PRINCIPALES")
    hito_h = ["Hito", "Descripción", "Fecha Planificada", "Responsable", "Criterio de Aceptación", "Notas"]
    for col, h in enumerate(hito_h, 1):
        ws.cell(row=row + 1, column=col, value=h)
    style_header_row(ws, row + 1, 6, bg="880000", size=9)
    for r in range(row + 2, row + 8):
        for c in range(1, 7):
            ws.cell(row=r, column=c).border = thin_border()
        ws.row_dimensions[r].height = 20

    row = row + 10
    sec_title(ws, row, "5. ORGANIGRAMA Y RESPONSABILIDADES")
    org_h = ["Rol", "Nombre", "Empresa", "Responsabilidades Clave", "Autorización de Cambios", "Contacto de Emergencia"]
    for col, h in enumerate(org_h, 1):
        ws.cell(row=row + 1, column=col, value=h)
    style_header_row(ws, row + 1, 6, bg="880000", size=9)
    for r in range(row + 2, row + 8):
        for c in range(1, 7):
            ws.cell(row=r, column=c).border = thin_border()
        ws.row_dimensions[r].height = 20

    row = row + 10
    sec_title(ws, row, "6. PROCEDIMIENTOS DE COMUNICACIÓN")
    comm_items = [
        "Reuniones de seguimiento (frecuencia / plataforma):",
        "Canales oficiales de comunicación:",
        "Gestión documental (sistema / repositorio):",
        "Distribución de reportes:",
        "Idioma de trabajo:",
    ]
    for i, item in enumerate(comm_items, row + 1):
        ws.cell(row=i, column=1, value=item).font = Font(bold=True, size=10)
        ws.merge_cells(f"B{i}:F{i}")
        ws.cell(row=i, column=2).border = thin_border()
        ws.row_dimensions[i].height = 25

    row = row + len(comm_items) + 2
    sec_title(ws, row, "7. GESTIÓN DE RIESGOS IDENTIFICADOS EN KOM")
    riesgo_h = ["#", "Descripción del Riesgo", "Probabilidad", "Impacto", "Plan de Mitigación", "Responsable"]
    for col, h in enumerate(riesgo_h, 1):
        ws.cell(row=row + 1, column=col, value=h)
    style_header_row(ws, row + 1, 6, bg="880000", size=9)
    for r in range(row + 2, row + 7):
        for c in range(1, 7):
            ws.cell(row=r, column=c).border = thin_border()
        ws.row_dimensions[r].height = 20

    row = row + 9
    sec_title(ws, row, "8. ACUERDOS Y ACTION ITEMS DEL KOM")
    ai_h = ["#", "Acuerdo / Acción", "Responsable", "Fecha Límite", "Estado", "Observaciones"]
    for col, h in enumerate(ai_h, 1):
        ws.cell(row=row + 1, column=col, value=h)
    style_header_row(ws, row + 1, 6, bg=COLOR_ABB_RED, size=9)
    for r in range(row + 2, row + 10):
        for c in range(1, 7):
            ws.cell(row=r, column=c).border = thin_border()
        ws.row_dimensions[r].height = 22

    row = row + 12
    sec_title(ws, row, "9. FIRMAS DE CONFORMIDAD")
    firmas_h = ["Nombre", "Empresa", "Cargo", "Firma", "Fecha", "Aclaración"]
    for col, h in enumerate(firmas_h, 1):
        ws.cell(row=row + 1, column=col, value=h)
    style_header_row(ws, row + 1, 6, bg="333333", size=9)
    for r in range(row + 2, row + 6):
        for c in range(1, 7):
            ws.cell(row=r, column=c).border = thin_border()
        ws.row_dimensions[r].height = 30

    set_col_widths(ws, {
        "A": 32, "B": 22, "C": 22, "D": 22, "E": 22, "F": 25
    })

    path = os.path.join(BASE_DIR, "03_Minutas", "Minuta_KOM.xlsx")
    wb.save(path)
    print(f"  ✓ {path}")


# ── 5. DASHBOARD MAESTRO ──────────────────────────────────────────────────────

def crear_dashboard():
    wb = Workbook()
    ws = wb.active
    ws.title = "Dashboard"

    ws.merge_cells("A1:H1")
    style_title(ws, "A1", "DASHBOARD MAESTRO — PROYECTO ABB")
    ws.row_dimensions[1].height = 40

    ws["A2"] = f"Última actualización: {datetime.today().strftime('%d/%m/%Y %H:%M')}"
    ws["A2"].font = Font(italic=True, size=9, color="888888")
    ws.row_dimensions[2].height = 16

    # ── Bloque: Datos del Proyecto
    row = 4
    ws.merge_cells(f"A{row}:H{row}")
    ws.cell(row=row, column=1, value="DATOS DEL PROYECTO")
    style_header_row(ws, row, 8)
    ws.row_dimensions[row].height = 24

    project_fields = [
        ("Nombre del Proyecto:", ""),
        ("N° de Contrato / Orden:", ""),
        ("Cliente:", ""),
        ("Gerente de Proyecto:", ""),
        ("Fecha de Inicio:", ""),
        ("Fecha de Fin Planificada:", ""),
        ("Estado General:", "En ejecución"),
        ("% Avance Global:", "0%"),
    ]
    for i, (label, val) in enumerate(project_fields, row + 1):
        ws.cell(row=i, column=1, value=label).font = Font(bold=True, size=10)
        ws.cell(row=i, column=2, value=val).border = thin_border()
        ws.merge_cells(f"B{i}:D{i}")
        ws.row_dimensions[i].height = 20

    # ── Bloque: Resumen Financiero
    row = row + len(project_fields) + 2
    ws.merge_cells(f"A{row}:H{row}")
    ws.cell(row=row, column=1, value="RESUMEN FINANCIERO")
    style_header_row(ws, row, 8)
    ws.row_dimensions[row].height = 24

    fin_headers = ["Concepto", "Presupuesto", "Comprometido", "Ejecutado", "Saldo", "% Ejecutado", "Desvío", "Notas"]
    for col, h in enumerate(fin_headers, 1):
        ws.cell(row=row + 1, column=col, value=h)
    style_header_row(ws, row + 1, 8, bg="880000", size=9)
    conceptos = ["Ingeniería", "Procurement", "Construcción / Montaje", "Commissioning", "Gestión", "Contingencia", "TOTAL"]
    for i, conc in enumerate(conceptos, row + 2):
        ws.cell(row=i, column=1, value=conc).font = Font(bold=(conc == "TOTAL"))
        for c in range(1, 9):
            cell = ws.cell(row=i, column=c)
            cell.border = thin_border()
            if conc == "TOTAL":
                cell.fill = PatternFill("solid", fgColor=COLOR_ACCENT)
        ws.row_dimensions[i].height = 18

    # ── Bloque: Estado Entregables
    row = row + len(conceptos) + 4
    ws.merge_cells(f"A{row}:H{row}")
    ws.cell(row=row, column=1, value="ESTADO DE ENTREGABLES CLAVE")
    style_header_row(ws, row, 8)
    ws.row_dimensions[row].height = 24

    ent_headers = ["Entregable", "Responsable", "Fecha Plan", "Fecha Real", "% Avance", "Estado", "Bloqueantes", "Notas"]
    for col, h in enumerate(ent_headers, 1):
        ws.cell(row=row + 1, column=col, value=h)
    style_header_row(ws, row + 1, 8, bg="880000", size=9)
    for r in range(row + 2, row + 10):
        for c in range(1, 9):
            ws.cell(row=r, column=c).border = thin_border()
        ws.row_dimensions[r].height = 20

    # ── Bloque: Action Items Abiertos
    row = row + 12
    ws.merge_cells(f"A{row}:H{row}")
    ws.cell(row=row, column=1, value="ACTION ITEMS ABIERTOS")
    style_header_row(ws, row, 8, bg="CC0000")
    ws.row_dimensions[row].height = 24

    ai_headers = ["#", "Descripción", "Responsable", "Fecha Límite", "Días Restantes", "Estado", "Origen (Minuta)", "Prioridad"]
    for col, h in enumerate(ai_headers, 1):
        ws.cell(row=row + 1, column=col, value=h)
    style_header_row(ws, row + 1, 8, bg="880000", size=9)
    for r in range(row + 2, row + 10):
        for c in range(1, 9):
            ws.cell(row=r, column=c).border = thin_border()
        ws.row_dimensions[r].height = 20

    # ── Bloque: Próximas Reuniones
    row = row + 12
    ws.merge_cells(f"A{row}:H{row}")
    ws.cell(row=row, column=1, value="PRÓXIMAS REUNIONES / HITOS")
    style_header_row(ws, row, 8)
    ws.row_dimensions[row].height = 24

    prox_headers = ["Fecha", "Tipo", "Tema", "Convocados", "Plataforma", "Preparación Requerida", "Responsable", "Estado"]
    for col, h in enumerate(prox_headers, 1):
        ws.cell(row=row + 1, column=col, value=h)
    style_header_row(ws, row + 1, 8, bg="880000", size=9)
    for r in range(row + 2, row + 7):
        for c in range(1, 9):
            ws.cell(row=r, column=c).border = thin_border()
        ws.row_dimensions[r].height = 20

    set_col_widths(ws, {
        "A": 28, "B": 18, "C": 16, "D": 16,
        "E": 14, "F": 16, "G": 30, "H": 25
    })
    freeze(ws, "A3")

    # ── Pestaña: Log de Cambios
    wl = wb.create_sheet("Log de Cambios")
    wl.merge_cells("A1:F1")
    style_title(wl, "A1", "LOG DE CAMBIOS DEL PROYECTO")
    log_headers = ["Fecha", "Versión", "Descripción del Cambio", "Solicitado por", "Aprobado por", "Impacto"]
    for col, h in enumerate(log_headers, 1):
        wl.cell(row=2, column=col, value=h)
    style_header_row(wl, 2, 6)
    set_col_widths(wl, {"A": 14, "B": 10, "C": 45, "D": 20, "E": 20, "F": 30})
    freeze(wl, "A3")

    path = os.path.join(BASE_DIR, "Dashboard_ABB.xlsx")
    wb.save(path)
    print(f"  ✓ {path}")


# ── MAIN ──────────────────────────────────────────────────────────────────────

if __name__ == "__main__":
    print("\n=== Creando estructura del proyecto ABB ===\n")
    crear_ofertas()
    crear_planning()
    crear_minutas()
    crear_minuta_kom()
    crear_dashboard()
    print("\n✅ Estructura creada exitosamente.\n")
    print("Archivos generados:")
    print("  ABB_Proyecto/01_Ofertas/Ofertas.xlsx")
    print("  ABB_Proyecto/02_Planning/Planning.xlsx")
    print("  ABB_Proyecto/03_Minutas/Minutas_Reuniones.xlsx")
    print("  ABB_Proyecto/03_Minutas/Minuta_KOM.xlsx")
    print("  ABB_Proyecto/Dashboard_ABB.xlsx")
    print("\nCarpetas disponibles para cargar documentos:")
    print("  04_Documentos/   → PDFs, planos, especificaciones")
    print("  05_Tecnica/      → Documentación técnica")
