"""
Genera Analisis_KOM_y_Compromisos.xlsx con toda la información extraída
de las minutas KOM (18-MAY y 22-MAY) y la presentación KOM (PPTX).
Proyecto: CPF-2 / La Calera II - ABB
"""
import os
from datetime import date, datetime
from openpyxl import Workbook
from openpyxl.styles import (
    Font, PatternFill, Alignment, Border, Side, numbers
)
from openpyxl.utils import get_column_letter

PLANNING_DIR = os.path.dirname(os.path.abspath(__file__))
OUTPUT = os.path.join(PLANNING_DIR, "Analisis_KOM_y_Compromisos.xlsx")

# ── paleta ───────────────────────────────────────────────────────────────────
C_DARK   = "1F3864"   # azul marino ABB
C_MED    = "2E75B6"   # azul medio
C_LIGHT  = "BDD7EE"   # azul claro
C_ORANGE = "C55A11"   # naranja advertencia
C_GREEN  = "375623"   # verde OK
C_LG     = "E2EFDA"   # verde claro
C_RED    = "C00000"   # rojo riesgo
C_LR     = "FFCCCC"   # rojo claro
C_YELLOW = "FFFF00"   # amarillo
C_LY     = "FFFFCC"   # amarillo claro
C_GRAY   = "595959"   # gris texto
C_LGRAY  = "F2F2F2"   # gris claro fondo

def hdr(ws, col, row, text, bg=C_DARK, fg="FFFFFF", bold=True, size=11, wrap=False, center=True):
    c = ws.cell(row=row, column=col, value=text)
    c.font = Font(name="Calibri", bold=bold, color=fg, size=size)
    c.fill = PatternFill("solid", fgColor=bg)
    c.alignment = Alignment(horizontal="center" if center else "left",
                             vertical="center", wrap_text=wrap)
    return c

def cell(ws, col, row, value, bg=None, bold=False, size=10, wrap=True,
         left=False, color=None, italic=False, number_fmt=None):
    c = ws.cell(row=row, column=col, value=value)
    c.font = Font(name="Calibri", bold=bold, size=size,
                  color=color or "000000", italic=italic)
    if bg:
        c.fill = PatternFill("solid", fgColor=bg)
    c.alignment = Alignment(horizontal="left" if left else "center",
                             vertical="center", wrap_text=wrap)
    if number_fmt:
        c.number_format = number_fmt
    return c

def thin_border(ws, r1, c1, r2, c2):
    thin = Side(style="thin", color="AAAAAA")
    bd   = Border(left=thin, right=thin, top=thin, bottom=thin)
    for r in range(r1, r2 + 1):
        for c in range(c1, c2 + 1):
            ws.cell(r, c).border = bd

def merge_hdr(ws, r, c1, c2, text, bg=C_DARK, fg="FFFFFF", size=12):
    ws.merge_cells(start_row=r, start_column=c1, end_row=r, end_column=c2)
    hdr(ws, c1, r, text, bg=bg, fg=fg, size=size)

def row_bg(ws, r, c1, c2, bg):
    for c in range(c1, c2 + 1):
        ws.cell(r, c).fill = PatternFill("solid", fgColor=bg)

# ─────────────────────────────────────────────────────────────────────────────
def sheet_resumen(wb):
    ws = wb.create_sheet("1_Resumen_Ejecutivo")
    ws.sheet_view.showGridLines = False

    # anchos
    ws.column_dimensions["A"].width = 3
    ws.column_dimensions["B"].width = 28
    ws.column_dimensions["C"].width = 45
    ws.column_dimensions["D"].width = 20
    ws.column_dimensions["E"].width = 20
    ws.column_dimensions["F"].width = 3

    # ── título ───────────────────────────────────────────────────────────────
    ws.merge_cells("B1:E1")
    ws.row_dimensions[1].height = 14
    ws.merge_cells("B2:E2")
    ws.row_dimensions[2].height = 30
    hdr(ws, 2, 2, "PROYECTO CPF-2 / LA CALERA II — ABB", bg=C_DARK, size=16)
    ws.merge_cells("B3:E3")
    ws.row_dimensions[3].height = 22
    hdr(ws, 2, 3, "ANÁLISIS KOM Y COMPROMISOS — Mayo 2026", bg=C_MED, size=13)
    ws.row_dimensions[4].height = 8

    # ── datos generales ──────────────────────────────────────────────────────
    def blk(r, label, value, bg_v=C_LGRAY):
        ws.row_dimensions[r].height = 20
        hdr(ws, 2, r, label, bg=C_MED, fg="FFFFFF", size=10, center=False)
        ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=5)
        cell(ws, 3, r, value, bg=bg_v, left=True, size=10, bold=False)

    blk(5,  "Cliente",            "Pluspetrol / AESA")
    blk(6,  "Proyecto",           "CPF-2 — La Calera II (Vaca Muerta)")
    blk(7,  "Referencia ABB",     "R-2631036 (Salas) | E-2611027 (PMS)")
    blk(8,  "PM ABB (Salas)",     "Rodrigo Mack — rodrigo.mack@ar.abb.com")
    blk(9,  "PM ABB (PMS)",       "Pablo Kalis — pablo.kalis@ar.abb.com")
    blk(10, "PM AESA",            "Pablo Spinetta — pablo.spinetta@aesa.com.ar")
    blk(11, "Activador AESA",     "Julián Correa — julian.correa@aesa.com.ar")
    blk(12, "Fecha KOM Salas",    "18-MAY-2026")
    blk(13, "Fecha KOM PMS",      "22-MAY-2026")
    ws.row_dimensions[14].height = 8

    # ── OCs ──────────────────────────────────────────────────────────────────
    ws.merge_cells("B15:E15")
    ws.row_dimensions[15].height = 22
    hdr(ws, 2, 15, "ÓRDENES DE COMPRA", bg=C_ORANGE, fg="FFFFFF", size=12)

    for c, t in enumerate(["OC", "Descripción", "Monto USD", "Fecha Entrega"], 2):
        ws.row_dimensions[16].height = 18
        hdr(ws, c, 16, t, bg=C_MED, size=10)

    ocs = [
        ("OC 4508944973", "S4-A SALA ELÉCTRICA MT (11 meses)", "$3,686,311.55", "08-ABR-2027"),
        ("OC 4508944971", "S3-A SALA ELÉCTRICA BT (12 meses)", "$8,127,765.50", "14-MAY-2027"),
        ("OC 4508945953", "PMS-A Equipamientos (12 meses)",    "Por definir",   "14-MAY-2027"),
    ]
    total_row_colors = [C_LGRAY, "FFFFFF", C_LGRAY]
    for i, (oc, desc, monto, fecha) in enumerate(ocs):
        r = 17 + i
        ws.row_dimensions[r].height = 20
        bg = total_row_colors[i]
        cell(ws, 2, r, oc, bg=bg, bold=True, size=10)
        cell(ws, 3, r, desc, bg=bg, left=True, size=10)
        cell(ws, 4, r, monto, bg=bg, size=10, bold=True, color=C_DARK)
        cell(ws, 5, r, fecha, bg=bg, size=10)

    ws.merge_cells("B20:C20")
    ws.row_dimensions[20].height = 20
    hdr(ws, 2, 20, "TOTAL VISIBLE (sin PMS)", bg=C_DARK, fg="FFFFFF", size=10)
    ws.merge_cells("D20:E20")
    cell(ws, 4, 20, "$11,814,077.05", bg=C_DARK, bold=True, size=11, color="FFFFFF")
    ws.row_dimensions[21].height = 8

    # ── penalidades y garantías ───────────────────────────────────────────────
    ws.merge_cells("B22:E22")
    ws.row_dimensions[22].height = 22
    hdr(ws, 2, 22, "PENALIDADES Y GARANTÍAS", bg=C_RED, fg="FFFFFF", size=12)

    pens = [
        ("Penalidad por atraso",   "1% por semana, tope máximo 10% del contrato", C_LR),
        ("Fiel Cumplimiento",      "10% del Precio del Contrato — Póliza de Seguro de Caución", C_LGRAY),
        ("Fondo de Reparo",        "10% sin IVA — Póliza en 5 días hábiles post-aceptación", C_LGRAY),
        ("Avance de Fabricación",  "Póliza por cada hito certificable — necesaria antes de inspección", C_LGRAY),
        ("Garantía Extendida",     "Hasta 31-ENE-2029", C_LG),
    ]
    for i, (lbl, val, bg) in enumerate(pens):
        r = 23 + i
        ws.row_dimensions[r].height = 20
        hdr(ws, 2, r, lbl, bg=C_GRAY, fg="FFFFFF", size=10, center=False)
        ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=5)
        cell(ws, 3, r, val, bg=bg, left=True, size=10)

    ws.row_dimensions[28].height = 8

    # ── canal comunicación ────────────────────────────────────────────────────
    ws.merge_cells("B29:E29")
    ws.row_dimensions[29].height = 22
    hdr(ws, 2, 29, "CANAL DE COMUNICACIÓN", bg=C_MED, size=12)

    comms = [
        ("Canal oficial",             "Email — asunto obligatorio con Nº OC + Nº Orden Interna ABB"),
        ("Ref. Salas en asunto",      "C5155-OC 4508944971/44973 + R-2631036"),
        ("Ref. PMS en asunto",        "C5155-4508945953 + E-2611027"),
        ("Portal documentos",         "Sharefile AESA con transmittal"),
        ("Primer Informe",            "05-JUN-2026 (formato ABB)"),
        ("Reuniones seguimiento",     "Semanales"),
        ("Reuniones técnicas",        "Quincenales con especialistas"),
        ("Plazo revisión documentos", "10 días hábiles (ambas partes)"),
    ]
    for i, (lbl, val) in enumerate(comms):
        r = 30 + i
        ws.row_dimensions[r].height = 20
        bg = C_LGRAY if i % 2 == 0 else "FFFFFF"
        hdr(ws, 2, r, lbl, bg=C_LIGHT, fg=C_DARK, size=10, center=False)
        ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=5)
        cell(ws, 3, r, val, bg=bg, left=True, size=10)

    thin_border(ws, 2, 2, 37, 5)
    return ws


# ─────────────────────────────────────────────────────────────────────────────
def sheet_action_items(wb):
    ws = wb.create_sheet("2_Action_Items")
    ws.sheet_view.showGridLines = False

    cols = {"A": 3, "B": 5, "C": 42, "D": 22, "E": 16, "F": 14, "G": 3}
    for col, w in cols.items():
        ws.column_dimensions[col].width = w

    ws.merge_cells("B1:F1"); ws.row_dimensions[1].height = 8
    ws.merge_cells("B2:F2"); ws.row_dimensions[2].height = 28
    hdr(ws, 2, 2, "ACTION ITEMS — KOM 18-MAY y 22-MAY-2026", bg=C_DARK, size=14)
    ws.merge_cells("B3:F3"); ws.row_dimensions[3].height = 18
    hdr(ws, 2, 3, "Proyecto CPF-2 / La Calera II — ABB | Ref: R-2631036", bg=C_MED, size=11)
    ws.row_dimensions[4].height = 8

    hdrs = ["#", "Acción / Compromiso", "Responsable", "Fecha Límite", "Estado"]
    for i, h in enumerate(hdrs):
        ws.row_dimensions[5].height = 22
        hdr(ws, i + 2, 5, h, bg=C_DARK, size=10)

    STATE_COLORS = {
        "Completado":   (C_LG,   C_GREEN),
        "Pendiente":    (C_LY,   C_ORANGE),
        "En proceso":   (C_LIGHT, C_MED),
        "Por confirmar":(C_LGRAY, C_GRAY),
    }

    items = [
        # #   Accion                                                         Responsable               Fecha           Estado
        (1,  "Aceptación del Contrato por ABB",                             "Christian Bombace (ABB)", "12-MAY-26",    "Completado"),
        (2,  "Enviar Matriz de Comunicación con especialidades (Salas)",    "AESA (Julián Correa)",    "SEM 18-MAY",   "Pendiente"),
        (3,  "Completar Matriz de Comunicación con pares (Salas)",          "ABB (Rodrigo Mack)",      "SEM 18-MAY",   "Pendiente"),
        (4,  "Confirmación de usuario ABB para acceso a portal ShareFile",  "AESA",                    "18-MAY-26",    "Completado"),
        (5,  "Verificación acceso interno portal ShareFile",                "ABB",                     "Por definir",  "En proceso"),
        (6,  "Envío Listado de Documentos Preliminar (Salas)",              "ABB (Gustavo Pantolini)", "12-MAY-26",    "Completado"),
        (7,  "Análisis de LD y devolución con comentarios (Salas)",        "AESA",                    "Por definir",  "Pendiente"),
        (8,  "Envío documentación técnica ductos de barra (ubicación/conexión trafo)", "AESA",         "Por definir",  "Pendiente"),
        (9,  "Definición unifilares tableros CCMs (posible cambio de cargas)","AESA",                  "Por definir",  "Pendiente"),
        (10, "Reunión Técnica de Ingeniería",                               "ABB + AESA",              "21-MAY-26",    "Completado"),
        (11, "Documentos Sistema Detección Incendio (agregar a LD)",        "ABB",                     "Por definir",  "Pendiente"),
        (12, "Propuesta elementos de izaje (Perchas certificadas)",         "ABB",                     "Por definir",  "Pendiente"),
        (13, "Cronograma Preliminar Salas (enviado 12-MAY)",                "ABB",                     "12-MAY-26",    "Completado"),
        (14, "Cronograma Final Salas con Hold Points y hitos certificación","ABB",                     "SEM 01-JUN-26","Pendiente"),
        (15, "Primer Informe de Avance",                                    "ABB",                     "05-JUN-26",    "Pendiente"),
        (16, "Confirmación cambio CCM7: 3x22kW → 2x45kW+VFD",             "AESA",                    "05-JUN-26",    "Pendiente"),
        (17, "Enviar Matriz de Comunicación con especialidades (PMS)",      "AESA (Julián Correa)",    "SEM 26-MAY",   "Pendiente"),
        (18, "Completar Matriz de Comunicación con pares (PMS)",            "ABB (Pablo Kalis)",       "SEM 26-MAY",   "Pendiente"),
        (19, "Alta de usuarios en ShareFile (PMS)",                         "AESA",                    "18-MAY-26",    "Completado"),
        (20, "Habilitación acceso interno ShareFile (PMS)",                 "ABB",                     "Por definir",  "En proceso"),
        (21, "Envío LD Preliminar PMS",                                     "ABB",                     "20-MAY-26",    "Completado"),
        (22, "Análisis LD PMS y documentos críticos/prioritarios",          "AESA",                    "SEM 26-MAY",   "Pendiente"),
        (23, "Plan de Inspección y Ensayos Preliminar (PMS)",               "ABB",                     "Por definir",  "Pendiente"),
        (24, "Análisis prueba integrada PMS en Shelter (previo despacho)",  "ABB",                     "Por definir",  "Pendiente"),
        (25, "Cronograma Preliminar PMS",                                   "ABB",                     "SEM 26-MAY",   "Pendiente"),
        (26, "Cronograma Final PMS",                                        "ABB",                     "SEM 08-JUN-26","Pendiente"),
        (27, "Pólizas de Caución y Seguros (según hitos de certificación)", "ABB",                     "Según hitos",  "Pendiente"),
        (28, "Formato Consultas Técnicas",                                  "AESA",                     "Por definir",  "Pendiente"),
        (29, "Propuesta FAT sobre equipos locales y exterior (PMS)",        "ABB",                     "Por definir",  "Pendiente"),
    ]

    for i, (num, accion, resp, fecha, estado) in enumerate(items):
        r = 6 + i
        ws.row_dimensions[r].height = 22
        bg_row, fc = STATE_COLORS.get(estado, (C_LGRAY, C_GRAY))
        cell(ws, 2, r, num,    bg=bg_row, bold=True,  size=10, color=C_DARK)
        cell(ws, 3, r, accion, bg=bg_row, left=True,  size=10, wrap=True)
        cell(ws, 4, r, resp,   bg=bg_row, left=True,  size=10)
        cell(ws, 5, r, fecha,  bg=bg_row, bold=True,  size=10, color=fc)
        hdr(ws, 6, r, estado,  bg=bg_row, fg=fc, bold=True, size=10)

    thin_border(ws, 2, 2, 6 + len(items) - 1, 6)

    # leyenda
    lr = 6 + len(items) + 2
    ws.row_dimensions[lr].height = 18
    hdr(ws, 2, lr, "LEYENDA", bg=C_DARK, fg="FFFFFF", size=10)
    for i, (estado, (bg, fc)) in enumerate(STATE_COLORS.items()):
        hdr(ws, 3 + i, lr, estado, bg=bg, fg=fc, size=10, bold=True)

    return ws


# ─────────────────────────────────────────────────────────────────────────────
def sheet_pagos(wb):
    ws = wb.create_sheet("3_Hitos_Pagos")
    ws.sheet_view.showGridLines = False

    cols = {"A": 3, "B": 5, "C": 8, "D": 38, "E": 28, "F": 22, "G": 18, "H": 3}
    for col, w in cols.items():
        ws.column_dimensions[col].width = w

    ws.merge_cells("B1:G1"); ws.row_dimensions[1].height = 8
    ws.merge_cells("B2:G2"); ws.row_dimensions[2].height = 28
    hdr(ws, 2, 2, "HITOS DE CERTIFICACIÓN Y ESQUEMA DE PAGOS", bg=C_DARK, size=14)
    ws.merge_cells("B3:G3"); ws.row_dimensions[3].height = 18
    hdr(ws, 2, 3, "Aplica a las 3 OC — CPF-2 / La Calera II — ABB", bg=C_MED, size=11)
    ws.row_dimensions[4].height = 8

    # cabeceras
    hdrs5 = ["#", "%", "Hito", "Condición de Certificación", "Póliza / Garantía", "Observaciones"]
    for i, h in enumerate(hdrs5):
        ws.row_dimensions[5].height = 22
        hdr(ws, i + 2, 5, h, bg=C_DARK, size=10)

    hitos = [
        (1, "10%", "Envío Ingeniería Básica",
         "Emisión documentos principales definidos en KOM\n(planos principales, layout, etc.)",
         "Póliza Caución Avance de Fabricación\n(debe estar aprobada antes de certificar)",
         "Acordar previamente los documentos a emitir"),
        (2, "10%", "Aprobación Ingeniería Básica",
         "Aprobación de AESA según cronograma",
         "—",
         "Según cronograma acordado"),
        (3, "20%", "Finalización de FAT",
         "FAT de integración de tableros en fábrica\n(Brasil para tableros de M&B)",
         "Póliza Caución Avance de Fabricación",
         "PMS: Se analizará prueba integrada en Shelter\n(PMS + Gabinetes, previo despacho)"),
        (4, "25%", "Acopio Materiales Principales",
         "Arribo a Argentina de:\n• Celdas MT\n• Tableros BT Arco Interno\n• Ductos de Barras\n• Variadores MT y BT\n• Tablero PMS\n• UPS + Cargadores + Baterías\n• Tableros Auxiliares",
         "Póliza Caución Acopio de Materiales",
         "PMS: Se analiza acopio materiales de 3ros\n(a montar en Shelter)\nRectificar división Hito #3 y #4 con Dpto Compras AESA"),
        (5, "30%", "Contra Entrega",
         "Shelter listo para transportar hacia la obra\n(liberación por Calidad del COMPRADOR)",
         "—",
         "Entrega sobre camión en planta Shelterista (Mendoza)\nRetira AESA"),
        (6, "5%",  "Entrega Databook",
         "Databook completo y aprobado por el COMPRADOR",
         "—",
         "Según cronograma"),
    ]

    hito_colors = [C_LIGHT, C_LY, C_LIGHT, C_LY, C_LIGHT, C_LY]

    for i, (num, pct, hito, cond, poliza, obs) in enumerate(hitos):
        r = 6 + i
        ws.row_dimensions[r].height = 60
        bg = hito_colors[i]
        cell(ws, 2, r, num,    bg=bg, bold=True, size=11, color=C_DARK)
        cell(ws, 3, r, pct,    bg=bg, bold=True, size=13, color=C_ORANGE)
        cell(ws, 4, r, hito,   bg=bg, bold=True, size=11, left=True)
        cell(ws, 5, r, cond,   bg=bg, size=10, left=True, wrap=True)
        cell(ws, 6, r, poliza, bg=bg, size=10, left=True, wrap=True, color=C_RED)
        cell(ws, 7, r, obs,    bg=bg, size=9, left=True, wrap=True, italic=True)

    thin_border(ws, 2, 2, 6 + len(hitos) - 1, 7)

    # totales referencia
    tr = 6 + len(hitos) + 1
    ws.row_dimensions[tr].height = 22
    ws.merge_cells(start_row=tr, start_column=2, end_row=tr, end_column=4)
    hdr(ws, 2, tr, "Monto referencia S3 (OC 4508944971)", bg=C_DARK, fg="FFFFFF", size=10)
    cell(ws, 5, tr, "$8,127,765.50", bg=C_DARK, bold=True, size=11, color="FFFFFF")
    ws.merge_cells(start_row=tr, start_column=5, end_row=tr, end_column=7)

    tr2 = tr + 1
    ws.row_dimensions[tr2].height = 22
    ws.merge_cells(start_row=tr2, start_column=2, end_row=tr2, end_column=4)
    hdr(ws, 2, tr2, "Monto referencia S4 (OC 4508944973)", bg=C_GRAY, fg="FFFFFF", size=10)
    cell(ws, 5, tr2, "$3,686,311.55", bg=C_GRAY, bold=True, size=11, color="FFFFFF")
    ws.merge_cells(start_row=tr2, start_column=5, end_row=tr2, end_column=7)

    return ws


# ─────────────────────────────────────────────────────────────────────────────
def sheet_tecnica(wb):
    ws = wb.create_sheet("4_Datos_Tecnicos")
    ws.sheet_view.showGridLines = False

    cols = {"A": 3, "B": 32, "C": 32, "D": 32, "E": 3}
    for col, w in cols.items():
        ws.column_dimensions[col].width = w

    ws.merge_cells("B1:D1"); ws.row_dimensions[1].height = 8
    ws.merge_cells("B2:D2"); ws.row_dimensions[2].height = 28
    hdr(ws, 2, 2, "DATOS TÉCNICOS POR SALA — CPF-2 / LA CALERA II", bg=C_DARK, size=14)
    ws.row_dimensions[3].height = 8

    # cabeceras
    ws.merge_cells("B4:B4"); ws.row_dimensions[4].height = 22
    hdr(ws, 2, 4, "PARÁMETRO / ITEM", bg=C_DARK, size=10)
    hdr(ws, 3, 4, "SALA ELÉCTRICA #3 (BT)", bg=C_MED, size=11)
    hdr(ws, 4, 4, "SALA ELÉCTRICA #4 (MT)", bg=C_ORANGE, size=11)

    rows_data = [
        ("OC",                        "OC 4508944971", "OC 4508944973"),
        ("Monto OC",                  "USD 8,127,765.50", "USD 3,686,311.55"),
        ("Plazo OC",                  "12 meses",      "11 meses"),
        ("Fecha Entrega",             "14-MAY-2027",   "08-ABR-2027"),
        ("Tipo",                      "Baja Tensión (BT)", "Media Tensión (MT)"),
        ("Dimensiones Sala Tableros", "37 x 6 m",      "28 x 5 m"),
        ("Sala de Baterías",          "10,5 x 3 m",    "—"),
        ("Tableros Principales",      "CCM 4000A 65kA — Cantidad: 37", "ZS1 13,2kV y 6,6kV — 630A 25kA — 25 columnas"),
        ("Ductos de Barras",          "Cantidad: 6",   "—"),
        ("Variadores BT (LV VFD)",    "Cantidad: 3 (Provisión MO)", "—"),
        ("Variadores MT (MV VFD)",    "—",             "Cantidad: 3"),
        ("PMS",                       "Incluido (OC 4508945953)", "Incluido (OC 4508945953)"),
        ("HVAC",                      "Con ductos + Presurización + Det. Incendio + Extracción sala batería", "Presurización + Detección Incendio"),
        ("CCTV",                      "Provisión equipos + Montaje y cableado hasta rack", "Provisión equipos + Montaje y cableado hasta rack"),
        ("Racks",                     "CCTV + CA + IT", "CCTV + CA + IT"),
        ("UPS / Baterías",            "KMT (local)", "—"),
        ("Tableros Auxiliares",       "Mehcco (Buenos Aires — 3ros)", "Mehcco (Buenos Aires — 3ros)"),
        ("Interconexiones",           "En Shelter",    "En Shelter"),
        ("Repuestos",                 "PEM + 2 años de operación", "PEM + 2 años de operación"),
        ("Garantía Extendida",        "Hasta 31-ENE-2029", "Hasta 31-ENE-2029"),
        ("Entrega",                   "Sobre camión en planta Shelterista — Retira AESA", "Sobre camión en planta Shelterista — Retira AESA"),
    ]

    for i, (param, s3, s4) in enumerate(rows_data):
        r = 5 + i
        bg = C_LGRAY if i % 2 == 0 else "FFFFFF"
        ws.row_dimensions[r].height = 32
        hdr(ws, 2, r, param, bg=C_LIGHT, fg=C_DARK, size=10, center=False)
        cell(ws, 3, r, s3, bg=bg, left=True, size=10, wrap=True)
        cell(ws, 4, r, s4, bg=bg, left=True, size=10, wrap=True)

    thin_border(ws, 2, 2, 5 + len(rows_data) - 1, 4)

    # ── fabricación ────────────────────────────────────────────────────────
    fr = 5 + len(rows_data) + 2
    ws.merge_cells(start_row=fr, start_column=2, end_row=fr, end_column=4)
    ws.row_dimensions[fr].height = 22
    hdr(ws, 2, fr, "LOCACIONES DE FABRICACIÓN Y ENSAYOS", bg=C_DARK, size=12)

    fab_hdrs = ["Componente / Actividad", "Ubicación", "Observaciones"]
    for i, h in enumerate(fab_hdrs):
        ws.row_dimensions[fr + 1].height = 20
        hdr(ws, i + 2, fr + 1, h, bg=C_MED, size=10)

    fab = [
        ("Tableros CCM BT (M&B)",       "Brasil",             "ABB Brasil — Ensayos en Brasil"),
        ("Tableros Auxiliares",          "Buenos Aires",       "3ros — Mehcco"),
        ("Ductos de Barras",             "Turquía",            "Suministro ABB"),
        ("Variadores MT (MV VFD)",       "China",              "Suministro ABB"),
        ("Variadores BT (LV VFD)",       "Finlandia",          "Suministro MO"),
        ("Shelters / Salas",             "Mendoza",            "Bottino (Shelterista)"),
        ("UPS + Baterías",               "Local (Buenos Aires)","KMT"),
        ("FAT Maqueta PMS",              "Buenos Aires",       "Set-Oct 2026 — fábrica ABB BsAs"),
        ("FAT Tableros PMS",             "Buenos Aires",       "Enero 2027 — fábrica ABB BsAs"),
        ("Administración del Contrato",  "Buenos Aires",       "ABB Argentina"),
    ]

    for i, (comp, ubic, obs) in enumerate(fab):
        r = fr + 2 + i
        bg = C_LGRAY if i % 2 == 0 else "FFFFFF"
        ws.row_dimensions[r].height = 20
        cell(ws, 2, r, comp, bg=bg, left=True, size=10, bold=True)
        cell(ws, 3, r, ubic, bg=bg, left=True, size=10)
        cell(ws, 4, r, obs,  bg=bg, left=True, size=10, italic=True)

    thin_border(ws, 2, fr, fr + 2 + len(fab) - 1, 4)
    return ws


# ─────────────────────────────────────────────────────────────────────────────
def sheet_fechas(wb):
    ws = wb.create_sheet("5_Fechas_Criticas")
    ws.sheet_view.showGridLines = False

    cols = {"A": 3, "B": 22, "C": 45, "D": 22, "E": 22, "F": 3}
    for col, w in cols.items():
        ws.column_dimensions[col].width = w

    ws.merge_cells("B1:E1"); ws.row_dimensions[1].height = 8
    ws.merge_cells("B2:E2"); ws.row_dimensions[2].height = 28
    hdr(ws, 2, 2, "FECHAS CRÍTICAS DEL PROYECTO — CPF-2 / LA CALERA II", bg=C_DARK, size=14)
    ws.merge_cells("B3:E3"); ws.row_dimensions[3].height = 18
    hdr(ws, 2, 3, "ABB — Referencia R-2631036 / E-2611027", bg=C_MED, size=11)
    ws.row_dimensions[4].height = 8

    hdrs5 = ["Fecha", "Hito / Evento", "Responsable", "Estado"]
    for i, h in enumerate(hdrs5):
        ws.row_dimensions[5].height = 22
        hdr(ws, i + 2, 5, h, bg=C_DARK, size=10)

    STATE_COLORS = {
        "Completado":   (C_LG,    C_GREEN),
        "Pendiente":    (C_LY,    C_ORANGE),
        "En proceso":   (C_LIGHT, C_MED),
        "CRÍTICO":      (C_LR,    C_RED),
    }

    fechas = [
        ("12-MAY-2026",    "Aceptación del Contrato",                        "Christian Bombace (ABB)", "Completado"),
        ("12-MAY-2026",    "Envío LD Preliminar y Cronograma Preliminar (Salas)","ABB",                 "Completado"),
        ("18-MAY-2026",    "KOM Salas (18-MAY) — Lanzamiento oficial",        "ABB + AESA",             "Completado"),
        ("20-MAY-2026",    "Envío LD Preliminar PMS",                         "ABB",                    "Completado"),
        ("21-MAY-2026",    "Reunión Técnica de Ingeniería",                   "ABB + AESA",             "Completado"),
        ("22-MAY-2026",    "KOM PMS (22-MAY) — Lanzamiento oficial",          "ABB + AESA",             "Completado"),
        ("SEM 26-MAY-26",  "Cronograma Preliminar PMS",                       "ABB",                    "Pendiente"),
        ("SEM 26-MAY-26",  "Documentos críticos/prioritarios PMS",            "AESA",                   "Pendiente"),
        ("SEM 26-MAY-26",  "Matriz Comunicación PMS",                         "AESA + ABB",             "Pendiente"),
        ("01-JUN-2026",    "Cronograma Final Salas (con Hold Points)",         "ABB",                    "CRÍTICO"),
        ("05-JUN-2026",    "Primer Informe de Avance",                        "ABB",                    "CRÍTICO"),
        ("05-JUN-2026",    "Confirmación cambio CCM7 (3x22kW → 2x45kW+VFD)", "AESA",                   "CRÍTICO"),
        ("08-JUN-2026",    "Cronograma Final PMS (actualizado)",               "ABB",                    "Pendiente"),
        ("08-ABR-2027",    "ENTREGA SALA 4 (MT) — OC 4508944973",             "ABB",                    "CRÍTICO"),
        ("14-MAY-2027",    "ENTREGA SALA 3 (BT) — OC 4508944971",             "ABB",                    "CRÍTICO"),
        ("14-MAY-2027",    "ENTREGA PMS — OC 4508945953",                     "ABB",                    "CRÍTICO"),
        ("31-ENE-2029",    "Vencimiento Garantía Extendida",                  "ABB",                    "Pendiente"),
    ]

    for i, (fecha, evento, resp, estado) in enumerate(fechas):
        r = 6 + i
        ws.row_dimensions[r].height = 22
        bg, fc = STATE_COLORS.get(estado, (C_LGRAY, C_GRAY))
        cell(ws, 2, r, fecha,  bg=bg, bold=True,  size=10, color=fc)
        cell(ws, 3, r, evento, bg=bg, left=True,  size=10, bold=(estado == "CRÍTICO"),
             color=(C_RED if estado == "CRÍTICO" else "000000"))
        cell(ws, 4, r, resp,   bg=bg, left=True,  size=10)
        hdr(ws, 5, r, estado,  bg=bg, fg=fc, bold=True, size=10)

    thin_border(ws, 2, 2, 6 + len(fechas) - 1, 5)

    # leyenda
    lr = 6 + len(fechas) + 2
    ws.row_dimensions[lr].height = 18
    hdr(ws, 2, lr, "LEYENDA", bg=C_DARK, fg="FFFFFF", size=10)
    for i, (estado, (bg, fc)) in enumerate(STATE_COLORS.items()):
        hdr(ws, 3 + i, lr, estado, bg=bg, fg=fc, size=10, bold=True)

    return ws


# ─────────────────────────────────────────────────────────────────────────────
def sheet_equipo(wb):
    ws = wb.create_sheet("6_Equipo_Proyecto")
    ws.sheet_view.showGridLines = False

    cols = {"A": 3, "B": 28, "C": 30, "D": 38, "E": 3}
    for col, w in cols.items():
        ws.column_dimensions[col].width = w

    ws.merge_cells("B1:D1"); ws.row_dimensions[1].height = 8
    ws.merge_cells("B2:D2"); ws.row_dimensions[2].height = 28
    hdr(ws, 2, 2, "EQUIPO DEL PROYECTO — ABB / AESA", bg=C_DARK, size=14)
    ws.row_dimensions[3].height = 8

    def bloque(title, rows_list, start_row, bg_hdr, bg_rows):
        ws.merge_cells(start_row=start_row, start_column=2, end_row=start_row, end_column=4)
        ws.row_dimensions[start_row].height = 22
        hdr(ws, 2, start_row, title, bg=bg_hdr, size=12)
        ws.row_dimensions[start_row + 1].height = 18
        for i, h in enumerate(["Nombre", "Rol", "Email / Contacto"]):
            hdr(ws, i + 2, start_row + 1, h, bg=C_MED, size=10)
        for j, (nombre, rol, email) in enumerate(rows_list):
            r = start_row + 2 + j
            bg = bg_rows if j % 2 == 0 else "FFFFFF"
            ws.row_dimensions[r].height = 20
            cell(ws, 2, r, nombre, bg=bg, left=True, size=10, bold=True)
            cell(ws, 3, r, rol,    bg=bg, left=True, size=10)
            cell(ws, 4, r, email,  bg=bg, left=True, size=10, italic=True)
        thin_border(ws, start_row, 2, start_row + 2 + len(rows_list) - 1, 4)
        return start_row + 2 + len(rows_list) + 1

    abb_elds = [
        ("Rodrigo Mack",       "Project Manager (Salas)",    "rodrigo.mack@ar.abb.com"),
        ("Gustavo Pantolini",  "Líder de Ingeniería (Salas)","gustavo.pantolini@ar.abb.com"),
        ("Juan Palma",         "Ingeniero de Proyectos",     "juan.palma@ar.abb.com"),
        ("Christian Bombace",  "Ejecutivo de Cuentas",       "christian.bombace@ar.abb.com"),
        ("Paulo Souza",        "Ing. Proyectos (Brasil)",    "paulo.souza@br.abb.com"),
        ("Vinicius Rossi",     "Ing. Proyectos (Brasil)",    "vinicius.rossi@br.abb.com"),
        ("Robson Rodrigues",   "PM Brasil",                  "robson.rodrigues@br.abb.com"),
        ("William Cabral",     "PM Brasil",                  "William.Cabral@br.abb.com"),
        ("Daniela Breglia",    "Order Handler VFDs (MO)",    "daniela.Breglia@ar.abb.com"),
    ]

    abb_pms = [
        ("Pablo Kalis",        "Project Manager (PMS)",      "pablo.kalis@ar.abb.com"),
        ("Mariano Montironi",  "Coordinador Ingeniería PMS", "mariano.montironi@ar.abb.com"),
        ("Dario Alvarez",      "Líder de Ingeniería PMS",    "dario.alvarez@ar.abb.com"),
        ("Matias Seguí",       "Adm. de Documentos (PMS)",   "matias.segui@ar.abb.com"),
        ("Sergio Giacomantone","Resp. Ofertas PMS",          "—"),
    ]

    aesa = [
        ("Pablo Spinetta",    "Gerente de Proyecto",         "pablo.spinetta@aesa.com.ar"),
        ("Dario Stirparo",    "Ing. de Proyecto / Lider Ing","dario.stirparo@aesa.com.ar"),
        ("Marcelo Zappa",     "Responsable de Suministros",  "marcelo.h.zappa@aesa.com.ar"),
        ("Julián Correa",     "Activador (contacto principal)","julian.correa@aesa.com.ar"),
        ("Micaela Velasquez", "Gestión Documental",          "micaela.a.velasquez@aesa.com.ar"),
        ("Carlos Kovach",     "Jefe Ing. de Proyecto",       "—"),
        ("Andres Zizzutti",   "Líder de Ingeniería",         "—"),
        ("Ezequiel Fattori",  "Jefe de Proyecto",            "—"),
        ("Pedro Cardoso",     "Jefe de Compras Proyecto",    "—"),
        ("Cesar Sánchez",     "Jefe Inspección y Control",   "—"),
        ("Gustavo Pinto",     "Coord. Activación",           "—"),
        ("Tulio De La Torre", "Esp. Ing. Eléctrico",        "—"),
    ]

    next_row = bloque("ABB — ELDS (Salas Eléctricas)", abb_elds, 4,  C_DARK, C_LGRAY)
    next_row = bloque("ABB — PMS",                     abb_pms,  next_row, C_ORANGE, C_LY)
    next_row = bloque("AESA (Cliente)",                 aesa,     next_row, C_GREEN, C_LG)

    return ws


# ─────────────────────────────────────────────────────────────────────────────
def main():
    wb = Workbook()
    wb.remove(wb.active)   # quitar hoja por defecto

    sheet_resumen(wb)
    sheet_action_items(wb)
    sheet_pagos(wb)
    sheet_tecnica(wb)
    sheet_fechas(wb)
    sheet_equipo(wb)

    wb.save(OUTPUT)
    print(f"Generado: {OUTPUT}")

if __name__ == "__main__":
    main()
