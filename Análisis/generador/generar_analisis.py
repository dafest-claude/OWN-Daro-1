#!/usr/bin/env python3
"""
Generador de análisis de contratos — Legal PP Project LC
=========================================================

Toma un archivo JSON con el análisis estructurado de un contrato y genera
tres documentos de resultado en la carpeta de salida:

    - PDF   (.pdf)   -> informe legible para compartir
    - Word  (.docx)  -> documento editable
    - Excel (.xlsx)  -> datos tabulados por secciones

Uso:
    python generar_analisis.py <archivo.json> [carpeta_salida] [--formatos pdf,docx,xlsx]

Ejemplos:
    python generar_analisis.py ejemplo_contrato.json
    python generar_analisis.py mi_contrato.json ../   --formatos pdf,docx
    python generar_analisis.py mi_contrato.json salida --formatos xlsx

Si no se indica carpeta de salida, se usa la carpeta 'Análisis/'.
El esquema JSON esperado está documentado en ejemplo_contrato.json.
"""

import argparse
import json
import os
import re
import sys
import unicodedata
from datetime import datetime

# --- Word ---
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

# --- Excel ---
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# --- PDF ---
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import cm
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, ListFlowable, ListItem,
)

# Paleta de colores del proyecto
AZUL = "#1F3864"
AZUL_CLARO = "#2E5496"
GRIS = "#F2F2F2"


# --------------------------------------------------------------------------- #
# Utilidades
# --------------------------------------------------------------------------- #
def cargar_datos(ruta_json):
    with open(ruta_json, "r", encoding="utf-8") as f:
        return json.load(f)


def slugify(texto):
    """Convierte un nombre en un identificador seguro para archivos."""
    texto = unicodedata.normalize("NFKD", texto).encode("ascii", "ignore").decode("ascii")
    texto = re.sub(r"[^\w\s-]", "", texto).strip().lower()
    return re.sub(r"[-\s]+", "_", texto) or "analisis_contrato"


def g(d, clave, defecto=""):
    v = d.get(clave, defecto)
    return v if v is not None else defecto


# --------------------------------------------------------------------------- #
# Generador WORD (.docx)
# --------------------------------------------------------------------------- #
def generar_docx(datos, ruta_salida):
    doc = Document()

    # Estilo base
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    # Título
    t = doc.add_heading(level=0)
    run = t.add_run(f"Análisis de Contrato — {g(datos, 'nombre_contrato', 'Sin nombre')}")
    run.font.color.rgb = RGBColor(0x1F, 0x38, 0x64)

    # Metadatos
    meta = doc.add_paragraph()
    meta.add_run("Proyecto: ").bold = True
    meta.add_run("Legal PP Project LC\n")
    meta.add_run("Archivo fuente: ").bold = True
    meta.add_run(f"{g(datos, 'archivo_fuente')}\n")
    meta.add_run("Fecha de análisis: ").bold = True
    meta.add_run(f"{g(datos, 'fecha_analisis')}\n")
    meta.add_run("Analista: ").bold = True
    meta.add_run(f"{g(datos, 'analista')}")

    def encabezado(n, txt):
        h = doc.add_heading(f"{n}. {txt}", level=1)
        for r in h.runs:
            r.font.color.rgb = RGBColor(0x2E, 0x54, 0x96)

    def tabla(headers, filas):
        if not filas:
            doc.add_paragraph("Sin datos.")
            return
        tab = doc.add_table(rows=1, cols=len(headers))
        tab.style = "Light Grid Accent 1"
        for i, htxt in enumerate(headers):
            c = tab.rows[0].cells[i]
            c.text = htxt
            for p in c.paragraphs:
                for r in p.runs:
                    r.font.bold = True
        for fila in filas:
            cells = tab.add_row().cells
            for i, val in enumerate(fila):
                cells[i].text = str(val)

    # 1. Datos generales
    encabezado(1, "Datos generales del contrato")
    dg = g(datos, "datos_generales", {})
    tabla(["Campo", "Detalle"], [[k, v] for k, v in dg.items()])

    # 2. Resumen ejecutivo
    encabezado(2, "Resumen ejecutivo")
    doc.add_paragraph(g(datos, "resumen_ejecutivo", "—"))

    # 3. Obligaciones
    encabezado(3, "Obligaciones de las partes")
    for ob in g(datos, "obligaciones", []):
        p = doc.add_paragraph()
        p.add_run(g(ob, "parte", "Parte")).bold = True
        for it in g(ob, "items", []):
            doc.add_paragraph(str(it), style="List Bullet")

    # 4. Cláusulas clave
    encabezado(4, "Cláusulas clave")
    tabla(
        ["Cláusula", "Ubicación", "Resumen", "Observación"],
        [[g(c, "clausula"), g(c, "ubicacion"), g(c, "resumen"), g(c, "observacion")]
         for c in g(datos, "clausulas_clave", [])],
    )

    # 5. Fechas críticas
    encabezado(5, "Fechas y plazos críticos")
    tabla(
        ["Evento", "Fecha", "Acción requerida", "Responsable"],
        [[g(c, "evento"), g(c, "fecha"), g(c, "accion"), g(c, "responsable")]
         for c in g(datos, "fechas_criticas", [])],
    )

    # 6. Riesgos
    encabezado(6, "Riesgos identificados")
    tabla(
        ["Riesgo", "Nivel", "Cláusula", "Recomendación"],
        [[g(c, "riesgo"), g(c, "nivel"), g(c, "clausula"), g(c, "recomendacion")]
         for c in g(datos, "riesgos", [])],
    )

    # 7. Banderas rojas
    encabezado(7, "Banderas rojas / puntos de atención")
    for b in g(datos, "banderas_rojas", []):
        doc.add_paragraph(str(b), style="List Bullet")

    # 8. Cláusulas recomendadas
    encabezado(8, "Cláusulas ausentes o recomendadas")
    for c in g(datos, "clausulas_recomendadas", []):
        marca = "[x]" if g(c, "presente", False) else "[ ]"
        doc.add_paragraph(f"{marca} {g(c, 'clausula')}", style="List Bullet")

    # 9. Conclusiones
    encabezado(9, "Conclusiones y recomendaciones")
    doc.add_paragraph(g(datos, "conclusiones", "—"))

    # Pie
    doc.add_paragraph()
    pie = doc.add_paragraph(
        f"Documento generado automáticamente — {datetime.now():%Y-%m-%d %H:%M}"
    )
    pie.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for r in pie.runs:
        r.font.size = Pt(8)
        r.font.color.rgb = RGBColor(0x80, 0x80, 0x80)

    doc.save(ruta_salida)
    return ruta_salida


# --------------------------------------------------------------------------- #
# Generador EXCEL (.xlsx)
# --------------------------------------------------------------------------- #
def generar_xlsx(datos, ruta_salida):
    wb = Workbook()

    header_fill = PatternFill("solid", fgColor="1F3864")
    header_font = Font(bold=True, color="FFFFFF")
    title_font = Font(bold=True, size=14, color="1F3864")
    thin = Side(style="thin", color="CCCCCC")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    wrap = Alignment(wrap_text=True, vertical="top")

    def estilo_hoja(ws, headers, filas, titulo=None):
        fila_ini = 1
        if titulo:
            ws.cell(row=1, column=1, value=titulo).font = title_font
            fila_ini = 3
        for i, h in enumerate(headers, start=1):
            c = ws.cell(row=fila_ini, column=i, value=h)
            c.fill = header_fill
            c.font = header_font
            c.border = border
            c.alignment = wrap
        for r, fila in enumerate(filas, start=fila_ini + 1):
            for i, val in enumerate(fila, start=1):
                c = ws.cell(row=r, column=i, value=val)
                c.border = border
                c.alignment = wrap
        # Ancho de columnas
        for i, h in enumerate(headers, start=1):
            maxlen = len(str(h))
            for fila in filas:
                if i - 1 < len(fila):
                    maxlen = max(maxlen, len(str(fila[i - 1])))
            ws.column_dimensions[get_column_letter(i)].width = min(max(maxlen + 2, 12), 60)

    # Hoja 1: Resumen / datos generales
    ws = wb.active
    ws.title = "Datos generales"
    dg = g(datos, "datos_generales", {})
    estilo_hoja(
        ws,
        ["Campo", "Detalle"],
        [[k, v] for k, v in dg.items()],
        titulo=f"Análisis: {g(datos, 'nombre_contrato', 'Sin nombre')}",
    )
    # Metadatos al pie
    base = len(dg) + 5
    for i, (k, v) in enumerate([
        ("Proyecto", "Legal PP Project LC"),
        ("Archivo fuente", g(datos, "archivo_fuente")),
        ("Fecha de análisis", g(datos, "fecha_analisis")),
        ("Analista", g(datos, "analista")),
        ("Resumen ejecutivo", g(datos, "resumen_ejecutivo")),
    ]):
        ws.cell(row=base + i, column=1, value=k).font = Font(bold=True)
        ws.cell(row=base + i, column=2, value=v).alignment = wrap

    # Hoja 2: Cláusulas clave
    estilo_hoja(
        wb.create_sheet("Cláusulas clave"),
        ["Cláusula", "Ubicación", "Resumen", "Observación"],
        [[g(c, "clausula"), g(c, "ubicacion"), g(c, "resumen"), g(c, "observacion")]
         for c in g(datos, "clausulas_clave", [])],
    )

    # Hoja 3: Fechas críticas
    estilo_hoja(
        wb.create_sheet("Fechas críticas"),
        ["Evento", "Fecha", "Acción requerida", "Responsable"],
        [[g(c, "evento"), g(c, "fecha"), g(c, "accion"), g(c, "responsable")]
         for c in g(datos, "fechas_criticas", [])],
    )

    # Hoja 4: Riesgos
    estilo_hoja(
        wb.create_sheet("Riesgos"),
        ["Riesgo", "Nivel", "Cláusula", "Recomendación"],
        [[g(c, "riesgo"), g(c, "nivel"), g(c, "clausula"), g(c, "recomendacion")]
         for c in g(datos, "riesgos", [])],
    )

    # Hoja 5: Obligaciones
    filas_ob = []
    for ob in g(datos, "obligaciones", []):
        for it in g(ob, "items", []):
            filas_ob.append([g(ob, "parte"), it])
    estilo_hoja(wb.create_sheet("Obligaciones"), ["Parte", "Obligación"], filas_ob)

    # Hoja 6: Checklist y banderas
    ws6 = wb.create_sheet("Checklist y banderas")
    filas_chk = [[g(c, "clausula"), "Sí" if g(c, "presente", False) else "No"]
                 for c in g(datos, "clausulas_recomendadas", [])]
    estilo_hoja(ws6, ["Cláusula recomendada", "¿Presente?"], filas_chk,
                titulo="Cláusulas recomendadas")
    fila_b = len(filas_chk) + 6
    ws6.cell(row=fila_b, column=1, value="Banderas rojas").font = title_font
    for i, b in enumerate(g(datos, "banderas_rojas", []), start=fila_b + 1):
        ws6.cell(row=i, column=1, value=f"⚠ {b}").alignment = wrap

    wb.save(ruta_salida)
    return ruta_salida


# --------------------------------------------------------------------------- #
# Generador PDF (.pdf)
# --------------------------------------------------------------------------- #
def generar_pdf(datos, ruta_salida):
    doc = SimpleDocTemplate(
        ruta_salida, pagesize=A4,
        leftMargin=2 * cm, rightMargin=2 * cm, topMargin=2 * cm, bottomMargin=2 * cm,
    )
    styles = getSampleStyleSheet()
    h0 = ParagraphStyle("h0", parent=styles["Title"], textColor=colors.HexColor(AZUL), fontSize=18)
    h1 = ParagraphStyle("h1", parent=styles["Heading1"], textColor=colors.HexColor(AZUL_CLARO), fontSize=13, spaceBefore=14)
    normal = styles["BodyText"]
    small = ParagraphStyle("small", parent=normal, fontSize=8, textColor=colors.grey, alignment=1)

    elems = []

    def tabla(headers, filas, anchos=None):
        data = [[Paragraph(f"<b>{h}</b>", normal) for h in headers]]
        for fila in filas:
            data.append([Paragraph(str(v).replace("\n", "<br/>"), normal) for v in fila])
        t = Table(data, colWidths=anchos, repeatRows=1)
        t.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(AZUL)),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#CCCCCC")),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor(GRIS)]),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("FONTSIZE", (0, 0), (-1, -1), 9),
            ("LEFTPADDING", (0, 0), (-1, -1), 5),
            ("RIGHTPADDING", (0, 0), (-1, -1), 5),
            ("TOPPADDING", (0, 0), (-1, -1), 4),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ]))
        elems.append(t)
        elems.append(Spacer(1, 8))

    def lista(items):
        if not items:
            elems.append(Paragraph("Sin datos.", normal))
            return
        elems.append(ListFlowable(
            [ListItem(Paragraph(str(i), normal)) for i in items],
            bulletType="bullet", start="•",
        ))
        elems.append(Spacer(1, 6))

    ancho = doc.width

    # Título y metadatos
    elems.append(Paragraph(f"Análisis de Contrato — {g(datos, 'nombre_contrato', 'Sin nombre')}", h0))
    elems.append(Spacer(1, 6))
    meta = (
        f"<b>Proyecto:</b> Legal PP Project LC<br/>"
        f"<b>Archivo fuente:</b> {g(datos, 'archivo_fuente')}<br/>"
        f"<b>Fecha de análisis:</b> {g(datos, 'fecha_analisis')}<br/>"
        f"<b>Analista:</b> {g(datos, 'analista')}"
    )
    elems.append(Paragraph(meta, normal))
    elems.append(Spacer(1, 10))

    # 1. Datos generales
    elems.append(Paragraph("1. Datos generales del contrato", h1))
    dg = g(datos, "datos_generales", {})
    tabla(["Campo", "Detalle"], [[k, v] for k, v in dg.items()], anchos=[ancho * 0.35, ancho * 0.65])

    # 2. Resumen
    elems.append(Paragraph("2. Resumen ejecutivo", h1))
    elems.append(Paragraph(g(datos, "resumen_ejecutivo", "—"), normal))

    # 3. Obligaciones
    elems.append(Paragraph("3. Obligaciones de las partes", h1))
    for ob in g(datos, "obligaciones", []):
        elems.append(Paragraph(f"<b>{g(ob, 'parte', 'Parte')}</b>", normal))
        lista(g(ob, "items", []))

    # 4. Cláusulas clave
    elems.append(Paragraph("4. Cláusulas clave", h1))
    tabla(
        ["Cláusula", "Ubicación", "Resumen", "Observación"],
        [[g(c, "clausula"), g(c, "ubicacion"), g(c, "resumen"), g(c, "observacion")]
         for c in g(datos, "clausulas_clave", [])],
        anchos=[ancho * 0.22, ancho * 0.15, ancho * 0.40, ancho * 0.23],
    )

    # 5. Fechas críticas
    elems.append(Paragraph("5. Fechas y plazos críticos", h1))
    tabla(
        ["Evento", "Fecha", "Acción", "Responsable"],
        [[g(c, "evento"), g(c, "fecha"), g(c, "accion"), g(c, "responsable")]
         for c in g(datos, "fechas_criticas", [])],
        anchos=[ancho * 0.28, ancho * 0.15, ancho * 0.35, ancho * 0.22],
    )

    # 6. Riesgos
    elems.append(Paragraph("6. Riesgos identificados", h1))
    tabla(
        ["Riesgo", "Nivel", "Cláusula", "Recomendación"],
        [[g(c, "riesgo"), g(c, "nivel"), g(c, "clausula"), g(c, "recomendacion")]
         for c in g(datos, "riesgos", [])],
        anchos=[ancho * 0.30, ancho * 0.12, ancho * 0.20, ancho * 0.38],
    )

    # 7. Banderas rojas
    elems.append(Paragraph("7. Banderas rojas / puntos de atención", h1))
    lista([f"⚠ {b}" for b in g(datos, "banderas_rojas", [])])

    # 8. Cláusulas recomendadas
    elems.append(Paragraph("8. Cláusulas ausentes o recomendadas", h1))
    lista([f"{'☑' if g(c, 'presente', False) else '☐'} {g(c, 'clausula')}"
           for c in g(datos, "clausulas_recomendadas", [])])

    # 9. Conclusiones
    elems.append(Paragraph("9. Conclusiones y recomendaciones", h1))
    elems.append(Paragraph(g(datos, "conclusiones", "—"), normal))

    elems.append(Spacer(1, 16))
    elems.append(Paragraph(f"Documento generado automáticamente — {datetime.now():%Y-%m-%d %H:%M}", small))

    doc.build(elems)
    return ruta_salida


# --------------------------------------------------------------------------- #
# Main
# --------------------------------------------------------------------------- #
def main():
    parser = argparse.ArgumentParser(description="Genera análisis de contrato en PDF, Word y Excel.")
    parser.add_argument("json", help="Archivo JSON con el análisis estructurado.")
    parser.add_argument("salida", nargs="?", default=None, help="Carpeta de salida (por defecto: Análisis/).")
    parser.add_argument("--formatos", default="pdf,docx,xlsx",
                        help="Formatos a generar, separados por coma (pdf,docx,xlsx).")
    args = parser.parse_args()

    if not os.path.isfile(args.json):
        print(f"ERROR: no se encontró el archivo JSON: {args.json}", file=sys.stderr)
        sys.exit(1)

    datos = cargar_datos(args.json)

    # Carpeta de salida: por defecto la carpeta 'Análisis' (padre de 'generador')
    if args.salida:
        carpeta = args.salida
    else:
        carpeta = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    os.makedirs(carpeta, exist_ok=True)

    base = slugify(g(datos, "nombre_contrato", "analisis_contrato"))
    nombre = f"analisis_{base}"
    formatos = [f.strip().lower() for f in args.formatos.split(",") if f.strip()]

    generados = []
    if "docx" in formatos:
        ruta = os.path.join(carpeta, f"{nombre}.docx")
        generar_docx(datos, ruta)
        generados.append(ruta)
    if "xlsx" in formatos:
        ruta = os.path.join(carpeta, f"{nombre}.xlsx")
        generar_xlsx(datos, ruta)
        generados.append(ruta)
    if "pdf" in formatos:
        ruta = os.path.join(carpeta, f"{nombre}.pdf")
        generar_pdf(datos, ruta)
        generados.append(ruta)

    print("Documentos generados:")
    for r in generados:
        print(f"  - {r}")


if __name__ == "__main__":
    main()
