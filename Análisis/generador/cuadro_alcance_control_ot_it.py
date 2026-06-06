#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Cuadro de Alcance y Responsabilidades — Sistemas de Control e IT/OT
EPC CPF2 La Calera (actores: AESA y PP=Pluspetrol)

Detalla, por cada ítem de alcance de PCS, SCADA, PSS, ESD, F&G y OT/IT:
descripción, responsable de PROCURA, responsable de INSTALACIÓN/MONTAJE y
las páginas del contrato (P = principal "ANEXO A - EPC CPF 2"; Ad = "Adenda Dic25").

Uso:  python cuadro_alcance_control_ot_it.py [carpeta_salida]
"""
import os, sys
from datetime import datetime
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import cm
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle

AZUL="#1F3864"; AZUL2="#2E5496"; GRIS="#F2F2F2"
VERDE="#C6EFCE"; AMAR="#FFEB9C"; ROJO="#FFC7CE"

META={
 "titulo":"Cuadro de Alcance y Responsabilidades — Sistemas de Control (PCS/SCADA/PSS/ESD/F&G) e IT/OT",
 "subtitulo":"EPC CPF2 — La Calera · Actores: AESA (Contratista) y PP (Pluspetrol)",
 "proyecto":"Legal PP Project LC",
 "fuentes":"P = ANEXO A - EPC CPF 2 (682 págs) · Ad = CPF 2 - AESA - Adenda Dic25 (221 págs)",
 "fecha":"2026-06-06","analista":"Daro",
}

COLS=["#","Sistema","Alcance (ítem)","Breve descripción","Procura","Instalación / Montaje","Páginas"]

# Filas: [#, sistema, ítem, descripción, procura, instalación, páginas]
ROWS=[
 # ---------- PCS ----------
 ["1","PCS","Gabinetes de sistema PCS (CPF2-PCS-SC-07A/B, -08)","Tableros del sistema de control de proceso para las nuevas salas remotas INS-007/008. Marcas: Rockwell/Allen Bradley (ControlLogix, Stratix).","AESA","AESA","P-144, P-408"],
 ["2","PCS","Gabinetes de marshalling PCS (CPF2-PCS-MC-07A/B, -08A)","Marshalling para conexionado de campo, reserva mín. 20% capacidad / 25% ocupación; cables FTA.","AESA","AESA","P-144/145, P-408"],
 ["3","PCS","Tableros de comunicaciones (equipos activos) en salas remotas","Switches y equipos activos de la red de control en cada sala remota nueva.","AESA","AESA","P-144"],
 ["4","PCS","Programación de lógicas de control PCS","Nuevas lógicas (tableros nuevos y señales cableadas a remotas existentes). Subcontrato INAUCO.","AESA","AESA","P-144/145, P-303"],
 ["5","PCS","Instrumentos de campo off-skid","Provisión y montaje/contraste de transmisores, válvulas de control, caudalímetros, etc.","AESA","AESA","Ad-67 (01.C.07.03)"],
 ["6","PCS","Instrumentos sobre equipos/skids provistos por PP","Vienen montados y cableados a cajas de conexión sobre su propio skid (excepto P&ID de proveedores, p.ej. Propak).","PP","PP","P-408"],
 # ---------- SCADA proceso ----------
 ["7","SCADA proceso","Actualización e integración del SCADA de CPF","Nuevas pantallas, adecuación de existentes, base OPC, vínculos y documentación. Wonderware/OMI (AVEVA), integrador INAUCO.","AESA","AESA","P-145, P-155"],
 ["8","SCADA proceso","Integración a nivel lógico / configuración al sistema existente","La solución de AESA es completa SALVO la integración a nivel lógico/configuración al sistema en operación.","PP","PP","P-297"],
 ["9","SCADA proceso","Pruebas de integración (SIT) PCS↔SCADA anticipadas en taller","Para el PCS, dentro del subcontrato de AESA. Se anticipa el SIT respecto a IEC-62381.","AESA","AESA","P-303"],
 ["10","SCADA proceso","Ampliación de SCADA CPF1/CPF2","Incluida en montaje/programación/integración (sin precio separado en la planilla).","AESA","AESA","Ad-68 (01.C.07.04.03), P-282"],
 # ---------- SCADA eléctrico ----------
 ["11","SCADA eléctrico (PMS)","Ampliación del SCADA Eléctrico / PMS (ABB)","Ampliación de HW con tableros ABB, nuevas pantallas y modificación de existentes; desconexión de cargas.","AESA","AESA","P-144, P-281, Ad-67 (01.C.06.05.18)"],
 ["12","SCADA eléctrico (PMS)","Red redundante IEC 61850","Criterio de red redundante para SCADA eléctrico confirmado.","AESA","AESA","P-304"],
 # ---------- PSS ----------
 ["13","PSS","Gabinetes de seguridad PSS — hardware (CPF-PSS-SC-07/08)","Hardware del sistema de seguridad de proceso (HIMA/HIMAX). PP provee; AESA monta y cablea.","PP","AESA","P-144, P-408"],
 ["14","PSS","Ampliación de remota PSS","Módulos de I/O (según listado de señales + reserva), connector boards, cables FTA y tablero de marshalling completo (sala remota 6).","AESA","AESA","P-145, P-303"],
 ["15","PSS","Configuración del sistema PSS","Excluida del alcance del Contratista.","PP","PP","P-303"],
 # ---------- ESD ----------
 ["16","ESD","Gabinetes de seguridad ESD — hardware (CPF-ESD-SC-07/08)","Hardware del sistema de parada de emergencia (HIMA/HIMAX). PP provee; AESA monta y cablea.","PP","AESA","P-144, P-408"],
 ["17","ESD","Gabinetes de marshalling ESD (CPF2-ESD-MC-07A/08A)","Marshalling del ESD (ingeniería, provisión y montaje).","AESA","AESA","P-144/145, P-408"],
 ["18","ESD","Configuración del sistema ESD","Excluida del Contratista; PP entrega un controlador demo para las pruebas.","PP","PP","P-303"],
 ["19","ESD/PSS/USS","Controlador demo para pruebas de integración","Provisto por PP para las pruebas de integración al SCADA.","PP","PP","P-303"],
 # ---------- F&G ----------
 ["20","F&G","Estudio de Fire & Gas (ISA-TR 84.00.07)","Determina cantidad, ubicación y orientación de detectores de mezcla explosiva y llama (votación NooM).","AESA","—","P-171, P-242"],
 ["21","F&G","Detectores de campo (gas/mezcla explosiva, llama, térmico lineal), pulsadores, sirenas-balizas","Sistema de detección de campo F&G.","AESA","AESA","Ad-69 (01.G.08), P-242/250"],
 ["22","F&G","Sistema Fire & Gas — Hardware y Software","Partida 01.G.08 (montaje de alarmas, detectores, pulsadores, sirenas).","AESA","AESA","Ad-69 (01.G.08), P-233"],
 ["23","F&G","Centrales / controladores F&G","Considerados dentro de los instrumentos de campo; parte provistos con el proveedor de shelter (subcontrato).","PP / Prov. shelter","AESA / Prov.","P-232"],
 ["24","F&G","Cajas de conexión F&G + cableado resistente al fuego","Multitriadas (analógicas) y multipares (digitales) desde JB al shelter. ESD-F&G comparten bandeja.","AESA","AESA","Ad-61 (Redef. Rev.1)"],
 ["25","F&G","Habilitación de la red F&G / auditoría entes provinciales (red húmeda)","Se excluye la habilitación de la red de F&G; la auditoría de habilitación está a cargo de PPSA.","PP","PP","P-234"],
 # ---------- OT/IT ----------
 ["26","OT/IT","Sistemas IT/OT (general) — provisión y montaje","Incluye ingeniería constructiva y de detalle para el montaje.","AESA","AESA","P-409"],
 ["27","OT/IT","Servicios de Seguridad y CCTV","Cobertura de seguridad física-patrimonial y operaciones (CCTV Bosch, radares Spotter).","AESA","AESA","P-146, P-409"],
 ["28","OT/IT","Control de Acceso — ingeniería y canalizaciones","Solo desarrollo de ingeniería e instalación de canalizaciones para implementación futura.","AESA","AESA","P-146"],
 ["29","OT/IT","Equipamiento de control de acceso + licencias + integración a seguridad patrimonial","Fuera del alcance de AESA.","PP","PP","P-146"],
 ["30","OT/IT","Servicios WIFI In/Out","Cobertura de áreas exteriores y salas/shelters interiores.","AESA","AESA","P-146/147, P-409"],
 ["31","OT/IT","Rack de Comunicaciones IT","Provisión e instalación del rack; tendidos y fusionado certificado de FO/UTP; alimentación redundante; espejos de FO.","AESA","AESA","P-147, P-409"],
 ["32","OT/IT","Licencias e integración a la infraestructura IT Local de CPF","Fuera del alcance de AESA (aplica a WIFI y Rack IT).","PP","PP","P-147"],
 ["33","OT/IT","Infraestructura OT - FO (anillo de fibra óptica)","Provisión, instalación, fusionado y certificación de FO; continuidad del anillo de 24h con caminos disjuntos; arquitectura PRP + redbox / anillos DLR.","AESA","AESA","P-147, P-305, P-409"],
 ["34","OT/IT","Tendidos y canalizaciones I&C","Bandejas portacables, cañeros de PVC, tritubo PEAD para FO.","AESA","AESA","P-409"],
 ["35","OT/IT","Provisión de equipos IT/OT (servidores, switches, firewalls, radios)","Según lista de vendors aprobados por PPSA (Cisco, Fortinet, Hirschmann, Stratix, MPLS Hitachi-ABB, Dell, etc.).","AESA","AESA","P-148, P-409"],
 # ---------- Ciberseguridad / red (eliminados) ----------
 ["36","Ciberseg./Red","Configuración de routers, switches y firewalls","ELIMINADO del alcance de AESA (según respuestas del Anexo 3.8).","PP","PP","P-409"],
 ["37","Ciberseg./Red","Diagramas de zonas y conductos de ciberseguridad (IEC 62443)","ELIMINADO del alcance de AESA.","PP","—","P-409"],
 ["38","Ciberseg./Red","Cálculo de niveles de disponibilidad · Tabla de IPS","ELIMINADO del alcance de AESA.","PP","—","P-409"],
 ["39","Ciberseguridad","Cláusula contractual 43.15 Ciberseguridad","Marco COBIT / IEC 62443 / RGPD — obligaciones del Contratista ('Ciberseguridad en AESA').","AESA","—","P-118, P-521"],
 # ---------- General / contractual ----------
 ["40","General","Regla por defecto (Adenda cláusula 3.1.4)","AESA suministra todo lo no indicado expresamente como provisión de PP en el Apéndice A, necesario para la obra.","AESA","AESA","Ad-5"],
 ["41","General","Equipos de proceso provistos por PP (Apéndice A)","Equipos suministrados por Pluspetrol; AESA ejecuta su montaje.","PP","AESA","P-408, Ad-5"],
 ["42","Pruebas","FAT / iFAT / SAT / SIT, sintonía de lazos, taller de pantallas","Pruebas y puesta a punto de tableros y sistemas (con participación/insumos de PP).","AESA","AESA","P-145, P-303"],
]

NOTA_INCONSIST=("Nota de inconsistencia documental: en P-408 los gabinetes de seguridad aparecen rotulados con el "
 "código del PCS (CPF2-PCS-SC-07A/B). Según §3.7 (P-144), PP provee los gabinetes de seguridad "
 "CPF-ESD-SC-07/08 y CPF-PSS-SC-07/08, y AESA provee los del PCS. Conviene reconciliar la "
 "codificación de gabinetes en la ingeniería de detalle.")

def color_resp(v):
    v=(v or "").upper()
    if v=="AESA": return VERDE
    if v=="PP": return ROJO
    if v in ("—",""): return None
    return AMAR  # mixto / compartido

# ----------------------------------------------------------- EXCEL
def gen_xlsx(path):
    wb=Workbook(); ws=wb.active; ws.title="Cuadro de Alcance"
    hf=PatternFill("solid",fgColor=AZUL[1:]); hfont=Font(bold=True,color="FFFFFF",size=10)
    thin=Side(style="thin",color="BFBFBF"); border=Border(thin,thin,thin,thin)
    wrap=Alignment(wrap_text=True,vertical="top")
    center=Alignment(wrap_text=True,vertical="center",horizontal="center")
    # título
    ws.merge_cells(start_row=1,start_column=1,end_row=1,end_column=len(COLS))
    ws.cell(row=1,column=1,value=META["titulo"]).font=Font(bold=True,size=13,color=AZUL[1:])
    ws.merge_cells(start_row=2,start_column=1,end_row=2,end_column=len(COLS))
    ws.cell(row=2,column=1,value=META["subtitulo"]+"   |   "+META["fuentes"]).font=Font(italic=True,size=9,color=AZUL2[1:])
    hr=4
    for i,h in enumerate(COLS,1):
        c=ws.cell(row=hr,column=i,value=h); c.fill=hf; c.font=hfont; c.border=border; c.alignment=center
    for r,row in enumerate(ROWS,hr+1):
        for i,v in enumerate(row,1):
            c=ws.cell(row=r,column=i,value=v); c.border=border
            c.alignment=center if i in (1,5,6) else wrap
            if i in (5,6):
                fill=color_resp(v)
                if fill: c.fill=PatternFill("solid",fgColor=fill[1:]); c.font=Font(bold=True,size=10)
    widths=[4,16,30,52,10,14,20]
    for i,w in enumerate(widths,1):
        ws.column_dimensions[get_column_letter(i)].width=w
    ws.freeze_panes="A5"
    ws.auto_filter.ref=f"A{hr}:{get_column_letter(len(COLS))}{hr+len(ROWS)}"
    # nota
    nr=hr+len(ROWS)+2
    ws.merge_cells(start_row=nr,start_column=1,end_row=nr,end_column=len(COLS))
    ws.cell(row=nr,column=1,value=NOTA_INCONSIST).font=Font(italic=True,size=9,color="9C5700")
    ws.cell(row=nr,column=1).alignment=wrap
    # leyenda
    lr=nr+2
    ws.cell(row=lr,column=1,value="Leyenda:").font=Font(bold=True)
    ws.cell(row=lr,column=2,value="AESA").fill=PatternFill("solid",fgColor=VERDE[1:])
    ws.cell(row=lr,column=3,value="PP (Pluspetrol)").fill=PatternFill("solid",fgColor=ROJO[1:])
    ws.cell(row=lr,column=4,value="Mixto / compartido").fill=PatternFill("solid",fgColor=AMAR[1:])
    # hoja resumen por responsable
    ws2=wb.create_sheet("Resumen por responsable")
    def cuenta(idx):
        from collections import Counter
        return Counter((r[idx] or "").strip() for r in ROWS)
    cp=cuenta(4); ci=cuenta(5)
    ws2.cell(row=1,column=1,value="Conteo de ítems por responsable").font=Font(bold=True,size=12,color=AZUL[1:])
    ws2.cell(row=3,column=1,value="PROCURA").font=Font(bold=True)
    r=4
    for k,v in cp.most_common():
        ws2.cell(row=r,column=1,value=k); ws2.cell(row=r,column=2,value=v); r+=1
    r+=1; ws2.cell(row=r,column=1,value="INSTALACIÓN / MONTAJE").font=Font(bold=True); r+=1
    for k,v in ci.most_common():
        ws2.cell(row=r,column=1,value=k); ws2.cell(row=r,column=2,value=v); r+=1
    ws2.column_dimensions["A"].width=28; ws2.column_dimensions["B"].width=10
    wb.save(path); return path

# ----------------------------------------------------------- WORD (apaisado)
def gen_docx(path):
    doc=Document()
    sec=doc.sections[0]; sec.orientation=WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = sec.page_height, sec.page_width
    sec.left_margin=sec.right_margin=Cm(1.2); sec.top_margin=sec.bottom_margin=Cm(1.2)
    doc.styles["Normal"].font.name="Calibri"; doc.styles["Normal"].font.size=Pt(8.5)
    t=doc.add_heading(level=0); rr=t.add_run(META["titulo"]); rr.font.size=Pt(14); rr.font.color.rgb=RGBColor(0x1F,0x38,0x64)
    p=doc.add_paragraph(); r2=p.add_run(META["subtitulo"]); r2.italic=True
    doc.add_paragraph(META["fuentes"]).runs[0].font.size=Pt(8)
    tab=doc.add_table(rows=1,cols=len(COLS)); tab.style="Light Grid Accent 1"
    for i,h in enumerate(COLS):
        cell=tab.rows[0].cells[i]; cell.text=h
        for pp in cell.paragraphs:
            for rn in pp.runs: rn.font.bold=True; rn.font.size=Pt(8.5)
    for row in ROWS:
        cells=tab.add_row().cells
        for i,v in enumerate(row):
            cells[i].text=str(v)
            for pp in cells[i].paragraphs:
                for rn in pp.runs: rn.font.size=Pt(8)
    doc.add_paragraph()
    n=doc.add_paragraph(); rn=n.add_run(NOTA_INCONSIST); rn.italic=True; rn.font.size=Pt(8.5); rn.font.color.rgb=RGBColor(0x9C,0x57,0x00)
    leg=doc.add_paragraph(); leg.add_run("Leyenda: ").bold=True; leg.add_run("AESA = Contratista · PP = Pluspetrol · '—' = no aplica · 'Mixto' = compartido.")
    pie=doc.add_paragraph(f"Documento generado automáticamente — {datetime.now():%Y-%m-%d %H:%M}")
    pie.alignment=WD_ALIGN_PARAGRAPH.CENTER
    for rn in pie.runs: rn.font.size=Pt(7); rn.font.color.rgb=RGBColor(0x80,0x80,0x80)
    doc.save(path); return path

# ----------------------------------------------------------- PDF (apaisado)
def gen_pdf(path):
    doc=SimpleDocTemplate(path,pagesize=landscape(A4),leftMargin=1.0*cm,rightMargin=1.0*cm,topMargin=1.2*cm,bottomMargin=1.0*cm)
    ss=getSampleStyleSheet()
    h0=ParagraphStyle("h0",parent=ss["Title"],textColor=colors.HexColor(AZUL),fontSize=13,leading=15)
    sub=ParagraphStyle("sub",parent=ss["Italic"],fontSize=8.5,textColor=colors.HexColor(AZUL2))
    cell=ParagraphStyle("c",parent=ss["BodyText"],fontSize=7,leading=8.5)
    cellb=ParagraphStyle("cb",parent=cell,alignment=1)
    small=ParagraphStyle("sm",parent=cell,fontSize=7,textColor=colors.HexColor("#9C5700"))
    el=[]; W=doc.width
    el.append(Paragraph(META["titulo"],h0)); el.append(Paragraph(META["subtitulo"]+" — "+META["fuentes"],sub)); el.append(Spacer(1,5))
    cw=[W*0.025,W*0.10,W*0.20,W*0.385,W*0.075,W*0.10,W*0.115]
    data=[[Paragraph(f"<b>{h}</b>",cellb) for h in COLS]]
    for row in ROWS:
        data.append([Paragraph(str(v),cellb if i in(0,4,5) else cell) for i,v in enumerate(row)])
    t=Table(data,colWidths=cw,repeatRows=1)
    style=[("BACKGROUND",(0,0),(-1,0),colors.HexColor(AZUL)),("TEXTCOLOR",(0,0),(-1,0),colors.white),
           ("GRID",(0,0),(-1,-1),0.4,colors.HexColor("#CCCCCC")),("VALIGN",(0,0),(-1,-1),"TOP"),
           ("FONTSIZE",(0,0),(-1,-1),7),("LEFTPADDING",(0,0),(-1,-1),3),("RIGHTPADDING",(0,0),(-1,-1),3),
           ("TOPPADDING",(0,0),(-1,-1),2),("BOTTOMPADDING",(0,0),(-1,-1),2)]
    for r,row in enumerate(ROWS,1):
        for ci,idx in ((4,4),(5,5)):
            f=color_resp(row[idx])
            if f: style.append(("BACKGROUND",(ci,r),(ci,r),colors.HexColor(f)))
    t.setStyle(TableStyle(style)); el.append(t); el.append(Spacer(1,6))
    el.append(Paragraph(NOTA_INCONSIST,small))
    el.append(Paragraph("Leyenda: verde = AESA · rojo = PP (Pluspetrol) · amarillo = mixto/compartido · '—' = no aplica.",cell))
    doc.build(el); return path

def main():
    out=sys.argv[1] if len(sys.argv)>1 else os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    os.makedirs(out,exist_ok=True)
    base=os.path.join(out,"Cuadro_Alcance_Control_OT-IT_CPF2")
    res=[gen_xlsx(base+".xlsx"), gen_docx(base+".docx"), gen_pdf(base+".pdf")]
    print("Generados:")
    for x in res: print("  -",x)

if __name__=="__main__":
    main()
