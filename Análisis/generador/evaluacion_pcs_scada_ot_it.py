#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Evaluación de Alcance — Sistema de Control (PCS/SCADA) e Integración OT/IT
EPC CPF2 - La Calera (AESA / Pluspetrol)

Genera el informe en PDF, Word y Excel a partir del contenido estructurado
extraído del contrato principal (ANEXO A - EPC CPF 2, 682 págs) y la
Adenda Dic25 (221 págs).

Uso:  python evaluacion_pcs_scada_ot_it.py [carpeta_salida]
"""
import os, sys
from datetime import datetime

from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import cm
from reportlab.platypus import (SimpleDocTemplate, Paragraph, Spacer, Table,
                                TableStyle, ListFlowable, ListItem)

AZUL = "#1F3864"; AZUL2 = "#2E5496"; GRIS = "#F2F2F2"

META = {
    "titulo": "Evaluación de Alcance — Sistema de Control (PCS/SCADA) e Integración OT/IT",
    "subtitulo": "EPC CPF2 — Yacimiento La Calera, Neuquén (AESA / Pluspetrol)",
    "proyecto": "Legal PP Project LC",
    "fuentes": "ANEXO A - EPC CPF 2 (682 págs) + CPF 2 - AESA - Adenda Dic25 (221 págs)",
    "fecha": "2026-06-06",
    "analista": "Daro",
}

# Bloques del documento: ("h1"/"h2"/"p"/"bul"/"tbl", contenido)
BLOQUES = [
    ("h1", "1. Identificación del contrato"),
    ("tbl", {"headers": ["Campo", "Detalle"], "rows": [
        ["Tipo", "Oferta irrevocable — Contrato de Obra EPC (Ingeniería, Compras y Construcción)"],
        ["Obra", "Central Production Facilities 2 (CPF2), ampliación — Fase 2 (17 MMm3/d)"],
        ["Contratista", "A-Evangelista S.A. (AESA)"],
        ["Comitente", "Pluspetrol S.A., como representante de la UTE La Calera (Pluspetrol / YSUR-YPF, Área CNQ-26)"],
        ["Ley aplicable", "República Argentina"],
        ["Modalidad", "Mixta: Ingeniería de detalle (Ajuste Alzado) · Construcción (Precios Unitarios) · Suministros (Costos Reembolsables)"],
        ["Documentos analizados", "Contrato principal (Anexo A) + Adenda Dic25 (incl. 'Redefinición/Reducción de Alcance Rev.1' y Planilla de Cotización)"],
    ]}),

    ("h1", "2. Arquitectura de control existente que se debe replicar/ampliar"),
    ("p", "El contrato obliga a mantener marcas, modelos, arquitecturas e integradores de la instalación actual (CPF1) para unificar los sistemas. Esto fija la base tecnológica y genera dependencia (lock-in) de proveedores y subcontratos."),
    ("tbl", {"headers": ["Sistema", "Tecnología / Marca", "Integrador / Notas", "Ref."], "rows": [
        ["PCS (Process Control System)", "Allen Bradley / Rockwell — ControlLogix, Stratix, Westermo", "INAUCO · Arq. ACAL-100-PL-K-100", "p144-145, 271-272"],
        ["ESD (Emergency Shutdown)", "HIMA — HIMAX", "Gabinetes CPF2-ESD-SC-07/08", "p144-145, 408"],
        ["PSS (Process Shutdown / Seguridad)", "HIMA — HIMAX", "Requiere ampliación de I/O", "p145, 303"],
        ["USS", "(sistema de seguridad existente)", "Configuración excluida del contratista", "p303"],
        ["SCADA de proceso", "Wonderware OMI (AVEVA) — base OPC", "INAUCO (subcontrato)", "p145, 155"],
        ["SCADA Eléctrico / PMS", "ABB — módulo PMS · IEC 61850 (red redundante)", "Ampliación HW + nuevas pantallas", "p144, 304"],
        ["Salas remotas nuevas", "INS-007 e INS-008 (CPF2 Fase 2)", "Diseño según premisas CPF1", "p144"],
    ]}),

    ("h1", "3. Alcance INCLUIDO del Contratista (AESA)"),
    ("h2", "3.1 Sistema de control y seguridad (§3.7)"),
    ("bul", [
        "PCS: provisión, montaje y cableado de tableros de sistema y de marshalling, con reserva mínima 20% en capacidad y 25% en ocupación; cables FTA asociados.",
        "Tableros de comunicaciones con sus equipos activos en cada nueva sala remota.",
        "Programación de las nuevas lógicas de control (tableros nuevos y señales cableadas a remotas existentes).",
        "ESD/PSS: montaje y cableado de gabinetes; ampliación del PSS (módulos I/O según listado de señales + reserva, connector boards, cables FTA, tablero de marshalling completo).",
        "Respetar marcas/modelos/integradores existentes (Rockwell, HIMA, Wonderware/OMI, INAUCO).",
        "FAT (tableros propios y de Pluspetrol cuando cambie la matriz causa-efecto), iFAT de todos los sistemas, taller de revisión de pantallas (PCS y SSS), sintonía de lazos en comisionado/PEM.",
    ]),
    ("h2", "3.2 SCADA de proceso y SCADA Eléctrico"),
    ("bul", [
        "Actualización e integración del SCADA de CPF: nuevas pantallas, adecuación de existentes, base de datos OPC, vínculos, y actualización de toda la documentación (respetando Wonderware/OMI/INAUCO).",
        "SCADA Eléctrico - PMS (ABB): ampliación de hardware con tableros ABB, nuevas pantallas y modificación de las existentes; acuerdo de nuevas cargas para desconexión.",
    ]),
    ("h2", "3.3 Sistema IT/OT (§3.8)"),
    ("bul", [
        "Seguridad, Acceso y CCTV: ingeniería de detalle, canalizaciones, provisión e instalación de equipos (según ET 001-ACAL-ET-O-100 — Especificaciones Disciplina OT) + repuestos 2 años.",
        "Servicios WIFI in/out: cobertura de áreas exteriores y salas/shelters interiores.",
        "Rack de Comunicaciones IT: tendidos y fusionado certificado de FO/UTP, alimentación redundante, espejos de FO, provisión e instalación del rack.",
        "Infraestructura OT - FO: integración a la infraestructura OT existente dando continuidad a la topología de anillo de FO con caminos disjuntos; provisión, instalación, fusionado y certificación de FO.",
        "Configuración de routers, switches y firewalls; pruebas FAT/SAT; reserva equipada mínima 20%.",
    ]),
    ("h2", "3.4 Ciberseguridad (entregables de ingeniería OT)"),
    ("bul", [
        "Diagrama de Zonas y Conductos de ciberseguridad (metodología tipo IEC 62443).",
        "Diagramas de topologías física y lógica; Tabla de IPS; Cálculo de niveles de disponibilidad.",
        "Especificación técnica y configuración de routers, switches y firewalls (Fortinet).",
        "Cláusula contractual 43.15 'Ciberseguridad' (marco COBIT / IEC 62443 / RGPD) — obligaciones del contratista (p118) y 'Ciberseguridad en AESA' (p521).",
    ]),

    ("h1", "4. Alcance EXCLUIDO o a cargo de Pluspetrol — FRONTERAS CRÍTICAS"),
    ("p", "Estos límites son el corazón del riesgo de interfaz para el sistema de control y la integración OT/IT. Conviene revisarlos uno por uno."),
    ("tbl", {"headers": ["Frontera / Exclusión", "Detalle", "Ref."], "rows": [
        ["Integración lógica al sistema existente",
         "La solución del contratista 'debe ser la solución completa, SALVO la integración a nivel lógico / Configuración al sistema existente'.", "p297"],
        ["Configuración ESD / PSS / USS",
         "'La configuración de los sistemas ESD, PSS, USS están excluidos del alcance del CONTRATISTA.' Pluspetrol entrega un controlador demo para las pruebas.", "p303"],
        ["Provisión de gabinetes de seguridad",
         "Pluspetrol entrega los gabinetes CPF-ESD-SC-07/08 y CPF-PSS-SC-07/08 (hardware de seguridad). (Ver inconsistencia en §8).", "p144, 408"],
        ["Seguridad patrimonial (CCTV/Acceso)",
         "Fuera de alcance: licencias e integración al sistema de seguridad patrimonial existente y el equipamiento de control de acceso.", "p146"],
        ["Infraestructura IT Local de CPF",
         "Fuera de alcance: licencias e integración a la infraestructura IT Local de CPF (aplica a WIFI y Rack de Comunicaciones IT).", "p147"],
        ["Cerco perimetral", "Las obras de cerco perimetral no forman parte del contrato.", "Planilla / p ANEXO"],
    ]}),

    ("h1", "5. Integración OT/IT — arquitectura de red"),
    ("bul", [
        "Anillo de FO troncal de 24 hilos con caminos disjuntos; las nuevas remotas (7 y 8) y nuevas salas eléctricas se integran al anillo, que pasará por las salas eléctricas (continuidad de la filosofía CPF1).",
        "Arquitectura de comunicaciones CCM-PCS basada en anillo PRP + redbox que cierra los anillos DLR de cada semibarra (redundancia a nivel de red OT).",
        "Segregación de sistemas en bandejas (Adenda Rev.1): PCS, PSS, COM y ESD-F&G (estos dos últimos comparten bandeja). Cables de F&G resistentes al fuego.",
        "SCADA Eléctrico: criterio de red redundante para IEC 61850 confirmado.",
        "Vendors IT/OT fijados por Pluspetrol: switches Cisco Catalyst/Nexus, Hirschmann, Stratix; firewalls Fortinet; nodos MPLS Hitachi-ABB; FO Furukawa/Corning; servidores Dell; CCTV Bosch; control de acceso Lenel; GPS Oscilloquartz; radioenlaces AVIAT/Redline.",
        "Frontera IT↔OT: la integración a la red IT Local queda fuera de alcance → punto de interfaz y coordinación obligatoria con Pluspetrol.",
    ]),

    ("h1", "6. Pruebas e integración (PCS / ESD / PSS / USS → SCADA)"),
    ("bul", [
        "La integración del PCS y de ESD/PSS/USS al SCADA debe probarse de manera anticipada en taller (aunque IEC 62381 ubica el SIT posterior al SAT), para minimizar riesgos en sitio.",
        "PCS: la integración es alcance del subcontrato del Contratista.",
        "ESD/PSS/USS: Pluspetrol entregará un controlador demo para las pruebas (el Contratista no configura estos sistemas).",
        "Cambios de campo del Contratista que impacten configuraciones del SCADA deben gestionarse (coordinación con Pluspetrol).",
    ]),

    ("h1", "7. Aspectos económicos del alcance de control / IT-OT (Adenda — Fase 2)"),
    ("p", "Valores de la Planilla de Cotización de la Adenda (ARS). Permiten dimensionar el peso del alcance de control e IT/OT dentro del EPC."),
    ("tbl", {"headers": ["Ítem (EDT)", "Concepto", "Precio (ARS)", "Modalidad"], "rows": [
        ["01.A.04", "Ingeniería de detalle — Instrumentación y Control", "616.812.826", "Ajuste Alzado"],
        ["01.A.05", "Ingeniería de detalle — IT/OT/CCTV", "(s/ HH asignadas)", "Ajuste Alzado"],
        ["01.C.07", "Construcción — Instrumentación y Control (total)", "10.024.451.503", "Precios Unitarios"],
        ["01.C.07.02.03", "  Tendido de Fibra Óptica", "244.454.007", "Precios Unitarios"],
        ["01.C.07.04", "  Sistema SIS/PLC/PCS/CCTV", "2.196.112.355", "Precios Unitarios"],
        ["01.C.07.04.01", "    Montaje, programación e integración SIS/PLC/PCS", "984.515.654", "Precios Unitarios"],
        ["01.C.07.04.02", "    Comunicación (IT/OT), WIFI, CCTV, Control de Acceso", "1.211.596.701", "Precios Unitarios"],
        ["01.C.07.04.03", "    Ampliación de SCADA CPF1/CPF2", "Incluido (sin precio separado)", "Precios Unitarios"],
        ["01.C.06.05.18", "  SCADA Eléctrico (PMS) — montaje gabinetes", "Incluido (gl)", "Precios Unitarios"],
    ]}),
    ("p", "Observación: el ítem 'Montaje, programación e integración SIS/PLC/PCS' figura en el contrato principal por ARS 1.420.769.779 y en la Adenda por ARS 984.515.654 — una reducción aproximada de ARS 436 millones que conviene rastrear contra la 'Redefinición/Reducción de Alcance Rev.1'."),

    ("h1", "8. Riesgos y banderas rojas (para revisar juntos)"),
    ("tbl", {"headers": ["#", "Riesgo / Bandera", "Nivel", "Por qué importa"], "rows": [
        ["1", "Frontera de 'integración lógica/configuración' al sistema existente difusa",
         "Alto", "El contratista entrega 'solución completa salvo integración lógica'. ¿Quién cierra la integración SCADA punta a punta y responde por su funcionalidad? Riesgo de interfaz y de vacío de responsabilidad."],
        ["2", "Inconsistencia sobre provisión de gabinetes PCS-SC-07/08",
         "Alto", "§3.7 (p144) dice que Pluspetrol solo entrega ESD/PSS; la sección de división de alcance (p408) asigna PCS a AESA y PSS/PCS-SC a Pluspetrol. Hay que reconciliar quién provee qué gabinete."],
        ["3", "Dependencia de insumos/entregables de Pluspetrol",
         "Medio", "Controlador demo (ESD/PSS/USS), documentación de SCADA y arquitecturas existentes, gabinetes de seguridad. Retrasos de Pluspetrol pueden impactar el cronograma del Contratista."],
        ["4", "Alcance y cuantificación de ciberseguridad (zonas y conductos)",
         "Medio", "Los entregables IEC 62443 (zonas/conductos, IPS, disponibilidad) fueron objeto de consulta; verificar quién ejecuta, quién aprueba y cómo se cuantifica/paga."],
        ["5", "'Ampliación de SCADA CPF1/CPF2' sin precio separado",
         "Medio", "Marcada como 'incluida' en otros ítems. Riesgo de subvaloración / disputa de alcance si la integración resulta mayor a la prevista."],
        ["6", "Lock-in de marcas/modelos/integradores",
         "Medio", "Obligación de mantener Rockwell, HIMA, AVEVA/Wonderware, ABB e INAUCO. Limita competencia de subcontratos y traslada riesgo de disponibilidad/precio."],
        ["7", "Frontera IT↔OT (integración a IT Local fuera de alcance)",
         "Medio", "WIFI y Rack IT se proveen pero su integración a IT Local es de Pluspetrol. Definir punto de demarcación y responsable de la interfaz."],
        ["8", "Obligaciones contractuales de ciberseguridad (Cl. 43.15)",
         "Medio", "Marco COBIT/IEC 62443/RGPD + 'Ciberseguridad en AESA'. Revisar exigencias, auditorías y penalidades asociadas."],
    ]}),

    ("h1", "9. Documentos de referencia clave"),
    ("tbl", {"headers": ["Código", "Documento"], "rows": [
        ["ACAL-100-PL-K-100", "Arquitectura del Sistema PCS (en desarrollo para ampliación)"],
        ["ACAL-100-PL-K-105", "Tablero PCS 003"],
        ["ACAL-100-PL-K-551 / -100", "Arquitecturas actuales de PCS/SSS"],
        ["ACAL-100-PL-K-552/554", "Tablero SC ESD 006 / Marshalling ESD 006"],
        ["ACAL-100-PL-K-562/564", "Tablero SC PSS 006 / Marshalling PSS 006"],
        ["ACAL-000-MO-K-100/101", "Manual de Operaciones CPF SCADA / Diseño SCADA"],
        ["AARC-100-TI-K-002", "Especificación de diseño del SCADA"],
        ["021-ACAL-100-PS030-E-302/303", "SCADA ELE/PMS — Arquitectura y Topología del sistema"],
        ["226-ACAL-102-PL-K-301", "Anillo de Fibra Óptica (FO)"],
        ["001-ACAL-ET-O-100", "ET — Especificaciones Disciplina OT"],
        ["001-ACAL-PL-O-200 / LI-O-203", "Plano de cobertura / Listado de equipos recomendados (seguridad)"],
        ["GRAL-EI-K-002/003 · GRAL-ET-K-101", "Estándares corporativos de Sistemas de Control y Seguridad"],
        ["330-CPF2 (Adenda)", "Memoria Descriptiva de Ingeniería — Redefinición de Alcance Rev.1 + 002_A_REDUCCIÓN DE ALCANCE_Rev.4"],
    ]}),

    ("h1", "10. Preguntas abiertas para nuestra revisión"),
    ("bul", [
        "Integración punta a punta del SCADA: ¿se acepta la frontera 'solución completa salvo integración lógica' o se exige al Contratista cerrar y garantizar la integración completa?",
        "Reconciliar la provisión de gabinetes PCS-SC-07A/B y -08 (¿Pluspetrol o AESA?).",
        "Confirmar el alcance, responsable y forma de pago de los entregables de ciberseguridad (zonas y conductos IEC 62443).",
        "Rastrear qué se quitó en la 'Reducción de Alcance Rev.1' respecto al sistema de control / IT-OT (la baja de ~ARS 436 M en SIS/PLC/PCS).",
        "Definir el punto de demarcación IT↔OT y la matriz de responsabilidades de interfaz con Pluspetrol.",
        "Verificar plazos de entrega de insumos de Pluspetrol (controlador demo, documentación, gabinetes) y su impacto en el cronograma.",
    ]),
]

# ---------------------------------------------------------------- WORD
def gen_docx(path):
    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"; doc.styles["Normal"].font.size = Pt(10.5)
    t = doc.add_heading(level=0); r = t.add_run(META["titulo"]); r.font.color.rgb = RGBColor(0x1F,0x38,0x64)
    s = doc.add_paragraph(); rs = s.add_run(META["subtitulo"]); rs.italic = True
    m = doc.add_paragraph()
    m.add_run("Proyecto: ").bold = True; m.add_run(META["proyecto"]+"\n")
    m.add_run("Fuentes: ").bold = True; m.add_run(META["fuentes"]+"\n")
    m.add_run("Fecha: ").bold = True; m.add_run(META["fecha"]+"    ")
    m.add_run("Analista: ").bold = True; m.add_run(META["analista"])
    for kind, c in BLOQUES:
        if kind == "h1":
            h = doc.add_heading(c, level=1)
            for rr in h.runs: rr.font.color.rgb = RGBColor(0x2E,0x54,0x96)
        elif kind == "h2":
            h = doc.add_heading(c, level=2)
        elif kind == "p":
            doc.add_paragraph(c)
        elif kind == "bul":
            for it in c: doc.add_paragraph(it, style="List Bullet")
        elif kind == "tbl":
            tab = doc.add_table(rows=1, cols=len(c["headers"])); tab.style = "Light Grid Accent 1"
            for i, hh in enumerate(c["headers"]):
                cell = tab.rows[0].cells[i]; cell.text = hh
                for p in cell.paragraphs:
                    for rr in p.runs: rr.font.bold = True
            for row in c["rows"]:
                cells = tab.add_row().cells
                for i, v in enumerate(row): cells[i].text = str(v)
    pie = doc.add_paragraph(f"\nDocumento generado automáticamente — {datetime.now():%Y-%m-%d %H:%M}")
    pie.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for rr in pie.runs: rr.font.size = Pt(8); rr.font.color.rgb = RGBColor(0x80,0x80,0x80)
    doc.save(path); return path

# ---------------------------------------------------------------- EXCEL
def gen_xlsx(path):
    wb = Workbook()
    hf = PatternFill("solid", fgColor="1F3864"); hfont = Font(bold=True, color="FFFFFF")
    tfont = Font(bold=True, size=13, color="1F3864")
    thin = Side(style="thin", color="CCCCCC"); border = Border(thin,thin,thin,thin)
    wrap = Alignment(wrap_text=True, vertical="top")

    def hoja(ws, headers, rows, titulo=None):
        fi = 1
        if titulo:
            ws.cell(row=1, column=1, value=titulo).font = tfont; fi = 3
        for i,h in enumerate(headers,1):
            c = ws.cell(row=fi,column=i,value=h); c.fill=hf; c.font=hfont; c.border=border; c.alignment=wrap
        for r,row in enumerate(rows, fi+1):
            for i,v in enumerate(row,1):
                c = ws.cell(row=r,column=i,value=v); c.border=border; c.alignment=wrap
        for i,h in enumerate(headers,1):
            ml=len(str(h))
            for row in rows:
                if i-1<len(row): ml=max(ml,min(len(str(row[i-1])),80))
            ws.column_dimensions[get_column_letter(i)].width=min(max(ml+2,12),70)

    # Resumen
    ws=wb.active; ws.title="Resumen"
    hoja(ws,["Campo","Detalle"],[
        ["Título",META["titulo"]],["Subtítulo",META["subtitulo"]],
        ["Proyecto",META["proyecto"]],["Fuentes",META["fuentes"]],
        ["Fecha",META["fecha"]],["Analista",META["analista"]],
    ], titulo="Evaluación PCS/SCADA e Integración OT/IT — EPC CPF2")

    # Tablas a hojas dedicadas
    tablas = {
        "Arquitectura": BLOQUES[3][1],
        "Alcance excluido": BLOQUES[18][1] if False else None,
    }
    # recolectar tablas por título de sección
    secciones = []
    titulo_actual = "General"
    for kind,c in BLOQUES:
        if kind=="h1": titulo_actual=c
        if kind=="tbl": secciones.append((titulo_actual,c))
    nombres_usados=set()
    def short(name):
        n=name.split("—")[0].strip()
        for ch in "[]:*?/\\": n=n.replace(ch,"")
        n=n[:28] or "Hoja"
        base=n; k=2
        while n in nombres_usados: n=f"{base[:25]}_{k}"; k+=1
        nombres_usados.add(n); return n
    for titulo,c in secciones:
        ws=wb.create_sheet(short(titulo))
        hoja(ws,c["headers"],c["rows"],titulo=titulo)
    wb.save(path); return path

# ---------------------------------------------------------------- PDF
def gen_pdf(path):
    doc=SimpleDocTemplate(path,pagesize=A4,leftMargin=1.8*cm,rightMargin=1.8*cm,topMargin=1.8*cm,bottomMargin=1.6*cm)
    ss=getSampleStyleSheet()
    h0=ParagraphStyle("h0",parent=ss["Title"],textColor=colors.HexColor(AZUL),fontSize=16,leading=19)
    sub=ParagraphStyle("sub",parent=ss["Italic"],fontSize=10,textColor=colors.HexColor(AZUL2))
    h1=ParagraphStyle("h1",parent=ss["Heading1"],textColor=colors.HexColor(AZUL2),fontSize=12,spaceBefore=12)
    h2=ParagraphStyle("h2",parent=ss["Heading2"],fontSize=10.5,spaceBefore=6)
    normal=ParagraphStyle("n",parent=ss["BodyText"],fontSize=9,leading=12)
    small=ParagraphStyle("sm",parent=normal,fontSize=8,textColor=colors.grey,alignment=1)
    el=[]; W=doc.width
    el.append(Paragraph(META["titulo"],h0)); el.append(Paragraph(META["subtitulo"],sub)); el.append(Spacer(1,6))
    el.append(Paragraph(f"<b>Proyecto:</b> {META['proyecto']}<br/><b>Fuentes:</b> {META['fuentes']}<br/>"
                        f"<b>Fecha:</b> {META['fecha']} &nbsp;&nbsp; <b>Analista:</b> {META['analista']}",normal))
    el.append(Spacer(1,8))
    def tabla(c):
        headers=c["headers"]; rows=c["rows"]; n=len(headers)
        # anchos: primera/última columnas más angostas si parecen códigos/refs
        if n==2: cw=[W*0.32,W*0.68]
        elif n==3: cw=[W*0.30,W*0.55,W*0.15]
        elif n==4 and headers[0] in ("#","Ítem (EDT)"): cw=[W*0.12,W*0.40,W*0.26,W*0.22]
        elif n==4: cw=[W*0.26,W*0.30,W*0.10,W*0.34]
        else: cw=[W/n]*n
        data=[[Paragraph(f"<b>{h}</b>",normal) for h in headers]]
        for row in rows:
            data.append([Paragraph(str(v).replace("\n","<br/>"),normal) for v in row])
        t=Table(data,colWidths=cw,repeatRows=1)
        t.setStyle(TableStyle([
            ("BACKGROUND",(0,0),(-1,0),colors.HexColor(AZUL)),
            ("TEXTCOLOR",(0,0),(-1,0),colors.white),
            ("GRID",(0,0),(-1,-1),0.5,colors.HexColor("#CCCCCC")),
            ("ROWBACKGROUNDS",(0,1),(-1,-1),[colors.white,colors.HexColor(GRIS)]),
            ("VALIGN",(0,0),(-1,-1),"TOP"),("FONTSIZE",(0,0),(-1,-1),8.5),
            ("LEFTPADDING",(0,0),(-1,-1),4),("RIGHTPADDING",(0,0),(-1,-1),4),
            ("TOPPADDING",(0,0),(-1,-1),3),("BOTTOMPADDING",(0,0),(-1,-1),3),
        ]))
        el.append(t); el.append(Spacer(1,8))
    for kind,c in BLOQUES:
        if kind=="h1": el.append(Paragraph(c,h1))
        elif kind=="h2": el.append(Paragraph(c,h2))
        elif kind=="p": el.append(Paragraph(c,normal)); el.append(Spacer(1,4))
        elif kind=="bul":
            el.append(ListFlowable([ListItem(Paragraph(i,normal)) for i in c],bulletType="bullet",start="•"))
            el.append(Spacer(1,6))
        elif kind=="tbl": tabla(c)
    el.append(Spacer(1,12)); el.append(Paragraph(f"Documento generado automáticamente — {datetime.now():%Y-%m-%d %H:%M}",small))
    doc.build(el); return path

def main():
    out = sys.argv[1] if len(sys.argv)>1 else os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    os.makedirs(out, exist_ok=True)
    base = os.path.join(out, "Evaluacion_PCS_SCADA_OT-IT_CPF2")
    r=[gen_docx(base+".docx"), gen_xlsx(base+".xlsx"), gen_pdf(base+".pdf")]
    print("Generados:")
    for x in r: print("  -",x)

if __name__=="__main__":
    main()
